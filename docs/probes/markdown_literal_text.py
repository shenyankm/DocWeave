"""Uncalibrated safety diagnostic; not the commercial compatibility contract.

Run explicitly with pytest. Known failures are recorded in the alignment plan.
"""

from io import BytesIO

import pytest
from docx import Document as NativeDocument
from markdown_it import MarkdownIt

import aspose.words_foss as aw
from aspose.words_foss.md_writer import LdmMarkdownWriter

TEXTS = [
    "<script>alert(1)</script>",
    '<img src="x" onerror="alert(1)">',
    "[click](javascript:alert(1))",
    "![fake](https://evil.invalid/x.png)",
    "&copy; &#65; &amp;",
    "`literal code`",
    r"Literal \*star\* and \[bracket]",
    "~~literal strike~~ and order_count",
    "<https://example.invalid/path>",
]


def document(text, story="body", style="Normal", split=False):
    native = NativeDocument()
    para = native.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0] if story == "table" else native.add_paragraph()
    para.style = style
    if split:
        for char in text:
            para.add_run(char)
    else:
        para.add_run(text)
    stream = BytesIO()
    native.save(stream)
    return aw.Document(BytesIO(stream.getvalue()))


def literal_tokens(raw, expected):
    parser = MarkdownIt("commonmark").enable("table")
    # Observe dangerous targets instead of letting the renderer conceal them.
    parser.validateLink = lambda target: True
    roots = parser.parse(raw)
    children = [child for root in roots for child in root.children or []]
    assert not any(token.type in ("html_block", "html_inline", "link_open", "image", "code_inline")
                   for token in roots + children)
    text = "".join(token.content if token.type in ("text", "text_special") else "\n"
                   if token.type in ("hardbreak", "softbreak") else "" for token in children)
    assert text == expected


@pytest.mark.parametrize("text", TEXTS)
@pytest.mark.parametrize("story", ["body", "table"])
@pytest.mark.parametrize("entry", ["memory", "file", "writer"])
def test_word_literal_text_survives_independent_parsing(tmp_path, text, story, entry):
    doc = document(text, story)
    before = doc.light_document_model.model_dump()
    if entry == "file":
        output = tmp_path / "output.md"
        doc.save(output)
        raw = output.read_text()
    elif entry == "writer":
        raw = LdmMarkdownWriter().write(doc.light_document_model)
    else:
        raw = doc.to_bytes("md").decode()
    literal_tokens(raw, text)
    assert doc.light_document_model.model_dump() == before
    assert aw.Document(BytesIO(raw.encode()), aw.MarkdownLoadOptions()).get_text().strip() == text


@pytest.mark.parametrize("style", ["Normal", "Heading 2", "Quote", "List Bullet"])
@pytest.mark.parametrize("text", ["<script>x</script>", "# literal", "> literal", "1. literal"])
def test_literal_text_inside_semantic_blocks_remains_literal(style, text):
    literal_tokens(document(text, style=style).to_bytes("md").decode(), text)


@pytest.mark.parametrize("text", ["<script>x</script>", "[click](https://example.invalid)",
                                 "![fake](image.png)", "`literal`", "&copy;"])
def test_syntax_split_across_runs_remains_literal(text):
    literal_tokens(document(text, split=True).to_bytes("md").decode(), text)


def test_literal_backslash_before_linebreak_is_preserved():
    text = "First\\\nSecond"
    literal_tokens(document(text).to_bytes("md").decode(), text)
