"""Generate independently owned DOCX inputs with conflicting styles and resources."""

import argparse
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from PIL import Image


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    for name, bold, size in (("source", True, 18), ("destination", False, 11)):
        doc = Document()
        doc.styles["Normal"].font.size = Pt(14 if name == "source" else 10)
        conflict = doc.styles.add_style("Import Conflict", WD_STYLE_TYPE.PARAGRAPH)
        conflict.font.bold, conflict.font.size = bold, Pt(size)
        if name == "source":
            derived = doc.styles.add_style("Import Derived", WD_STYLE_TYPE.PARAGRAPH)
            derived.base_style, derived.font.italic = conflict, True
            doc.add_paragraph("IMPORT_CONFLICT", style=conflict)
            doc.add_paragraph("IMPORT_DERIVED", style=derived)
            doc.add_paragraph("IMPORT_NORMAL")
            doc.add_paragraph("IMPORT_NUMBERED", style="List Number")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "IMPORT_CELL"
            image = BytesIO()
            Image.new("RGB", (2, 2), "blue").save(image, format="PNG")
            paragraph = doc.add_paragraph("IMPORT_IMAGE")
            paragraph.add_run().add_picture(BytesIO(image.getvalue()))
        else:
            doc.add_paragraph("DESTINATION", style=conflict)
        doc.save(args.output / (name + ".docx"))


if __name__ == "__main__":
    main()
