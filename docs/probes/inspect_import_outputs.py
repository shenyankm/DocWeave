"""Inspect import outputs independently with python-docx, retaining raw format layers."""

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document


def font_values(font):
    return {"bold": font.bold, "italic": font.italic, "size_pt": None if font.size is None else font.size.pt}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", type=Path)
    args = parser.parse_args()
    report = json.loads((args.folder / "observations.json").read_text())
    records = []
    for row in report["records"]:
        path = args.folder / row["output"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == row["sha256"]
        doc = Document(path)
        paragraphs = []
        for paragraph in doc.paragraphs:
            if not paragraph.text.startswith(("IMPORT_", "DESTINATION")):
                continue
            chain, seen, style = [], set(), paragraph.style
            while style is not None:
                assert style.style_id not in seen, "style inheritance cycle"
                seen.add(style.style_id)
                chain.append({"id": style.style_id, "name": style.name, "font": font_values(style.font)})
                style = style.base_style
            pictures = []
            for drawing in paragraph._p.xpath('.//a:blip'):
                rid = drawing.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                assert rid in doc.part.related_parts, "missing imported image relationship"
                pictures.append(hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest())
            paragraphs.append({"text": paragraph.text, "style_chain": chain,
                               "runs": [{"text": run.text, "font": font_values(run.font)} for run in paragraph.runs],
                               "picture_sha256": pictures})
        imported = any(p["text"].startswith("IMPORT_") or "IMPORT_CONFLICT" in p["text"] for p in paragraphs)
        tables = [[cell.text for line in table.rows for cell in line.cells] for table in doc.tables]
        if row["outcome"]["status"] == "returned" and (row["deep"] or row["subject"] == "run"):
            assert imported or row["subject"] == "table", "imported content missing in saved output"
            if row["subject"] == "table":
                assert any("IMPORT_CELL" in cells for cells in tables), "imported table content missing"
            if row["subject"] == "image":
                assert any(p["picture_sha256"] for p in paragraphs if p["text"] == "IMPORT_IMAGE"), "imported image missing"
        records.append({"output": row["output"], "paragraphs": paragraphs,
                        "table_text": tables})
    (args.folder / "independent.json").write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"checked_outputs": len(records)}))


if __name__ == "__main__":
    main()
