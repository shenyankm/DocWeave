"""Independently check saved DOCX against the structure probe's in-memory state."""

import argparse
import json
from pathlib import Path

from docx import Document


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", type=Path)
    args = parser.parse_args()
    report = json.loads((args.folder / "observations.json").read_text())
    checks = []
    for record in report["records"]:
        document = Document(args.folder / record["saved_docx"])
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        labelled = [text for text in paragraphs if text.startswith(("ALPHA", "BETA", "CLONE"))]
        fonts = [{"text": paragraph.text, "first_run_bold": paragraph.runs[0].bold}
                 for paragraph in document.paragraphs if paragraph.text.startswith(("ALPHA", "CLONE"))]
        cell = document.tables[0].cell(0, 0).text
        matches = labelled == record["after"]["labelled_paragraphs"]
        checks.append({"scenario": record["scenario"], "all_native_paragraphs": paragraphs,
                       "matches_memory": matches, "table_cell": cell, "font_checks": fonts})
        assert cell == "CELL" and all(font["first_run_bold"] is True for font in fonts)
    (args.folder / "independent.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2) + "\n")
    mismatches = [check["scenario"] for check in checks if not check["matches_memory"]]
    print(json.dumps({"checked": len(checks), "memory_save_mismatches": mismatches}))
    assert not mismatches, "saved document differs from the observed DOM"


if __name__ == "__main__":
    main()
