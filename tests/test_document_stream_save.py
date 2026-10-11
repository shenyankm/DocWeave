from io import BytesIO, StringIO

import pytest
import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.saving import SaveOutputParameters

FORMATS = [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC, aw.SaveFormat.FLAT_OPC_MACRO_ENABLED,
           aw.SaveFormat.FLAT_OPC_TEMPLATE, aw.SaveFormat.FLAT_OPC_TEMPLATE_MACRO_ENABLED,
           aw.SaveFormat.TEXT, aw.SaveFormat.MARKDOWN, aw.SaveFormat.PDF]
MIMES = ['application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         'application/vnd.ms-word.document.macroEnabled.12',
         'application/vnd.openxmlformats-officedocument.wordprocessingml.template',
         'application/vnd.ms-word.template.macroEnabled.12', 'text/plain', 'text/plain', 'application/pdf']


def document():
    doc = aw.Document()
    doc.light_document_model.sections[0].body.children = [ldm.Paragraph(children=[ldm.Run(text='OWNED STREAM')])]
    return doc


@pytest.mark.parametrize('fmt,mime', list(zip(FORMATS, MIMES)))
@pytest.mark.parametrize('file_stream', [False, True])
def test_stream_position_tail_and_metadata(tmp_path, fmt, mime, file_stream):
    path = tmp_path/'owned.bin'
    initial = b'Z'*200000
    path.write_bytes(initial)
    stream = path.open('r+b') if file_stream else BytesIO(initial)
    try:
        stream.seek(17)
        result = document().save(stream, fmt)
        end = stream.tell()
        assert not stream.closed
        stream.seek(0)
        data = stream.read()
        start = 0 if fmt == aw.SaveFormat.DOCX else 17
        assert data[:start] == initial[:start]
        assert data[end:] == initial[end:] and len(data) >= len(initial)
        output = data[start:end]
        if fmt == aw.SaveFormat.PDF:
            import fitz
            with fitz.open(stream=output, filetype='pdf') as pdf:
                assert 'OWNED STREAM' in ''.join(page.get_text() for page in pdf)
        elif fmt in [aw.SaveFormat.TEXT, aw.SaveFormat.MARKDOWN]:
            assert 'OWNED STREAM' in output.decode()
        else:
            assert 'OWNED STREAM' in aw.Document(BytesIO(output)).get_text()
        assert isinstance(result, SaveOutputParameters) and result.content_type == mime
        with pytest.raises(AttributeError):result.content_type = 'changed'
    finally:
        stream.close()


@pytest.mark.parametrize('fmt,mime', list(zip(FORMATS, MIMES)))
def test_path_metadata_keeps_existing_save_contract(tmp_path, fmt, mime):
    path = tmp_path/'owned.output'
    assert document().save(path, fmt).content_type == mime
    assert path.is_file()


def test_output_metadata_has_no_public_constructor():
    with pytest.raises(TypeError):SaveOutputParameters()


@pytest.mark.parametrize('fmt', [None, aw.SaveFormat.DOC, 'unknown'])
def test_invalid_format_does_not_change_stream(fmt):
    stream = BytesIO(b'OWNED old bytes');stream.seek(3)
    with pytest.raises(TypeError if fmt is None else ValueError):document().save(stream, fmt)
    assert stream.tell() == 3 and stream.getvalue() == b'OWNED old bytes'


def test_nonbinary_targets_are_rejected_before_writing():
    target = StringIO('OWNED');target.seek(2)
    with pytest.raises(TypeError):document().save(target, aw.SaveFormat.TEXT)
    assert target.getvalue() == 'OWNED' and target.tell() == 2


@pytest.mark.parametrize('closed', [False, True])
def test_unwritable_stream_errors_leave_caller_ownership(tmp_path, closed):
    path = tmp_path/'owned';path.write_bytes(b'OWNED')
    with path.open('rb') as target:
        if closed:target.close()
        with pytest.raises(RuntimeError):document().save(target, aw.SaveFormat.TEXT)
        assert target.closed == closed
    assert path.read_bytes() == b'OWNED'


class Nonseek(BytesIO):
    def seek(self, *args):raise OSError('OWNED seek unavailable')


def test_text_does_not_require_seek_but_docx_does():
    stream = Nonseek()
    document().save(stream, aw.SaveFormat.TEXT)
    assert b'OWNED STREAM' in stream.getvalue()
    before = stream.getvalue()
    with pytest.raises(RuntimeError):document().save(stream, aw.SaveFormat.DOCX)
    assert stream.getvalue() == before and not stream.closed


class Short(BytesIO):
    def write(self, data):return super().write(data[:7])


def test_short_writes_report_failure_without_closing():
    stream = Short()
    with pytest.raises(RuntimeError):document().save(stream, aw.SaveFormat.DOCX)
    assert len(stream.getvalue()) == 7 and not stream.closed


class Stalled(BytesIO):
    def write(self, data):return 0


def test_stalled_stream_reports_failure_without_closing():
    stream = Stalled()
    with pytest.raises(RuntimeError):document().save(stream, aw.SaveFormat.TEXT)
    assert not stream.closed and stream.getvalue() == b''


@pytest.mark.parametrize('option,mime', [(aw.saving.OoxmlSaveOptions(), MIMES[0]), (aw.saving.PdfSaveOptions(), 'application/pdf'), (aw.saving.MarkdownSaveOptions(), 'text/plain')])
def test_options_overload(option, mime):
    stream = BytesIO()
    result = document().save(stream, option)
    assert stream.getvalue() and isinstance(result, SaveOutputParameters)
    assert result.content_type == mime


def image_document():
    from PIL import Image
    from docx import Document
    image = BytesIO()
    Image.new('RGB', (4, 3), 'navy').save(image, format='PNG')
    source = Document()
    source.add_paragraph('OWNED IMAGE')
    source.add_picture(BytesIO(image.getvalue()))
    data = BytesIO();source.save(data)
    return aw.Document(BytesIO(data.getvalue())), image.getvalue()


@pytest.mark.parametrize('failed', [False, True])
def test_markdown_external_images_follow_stream_transaction(tmp_path, failed):
    doc, image = image_document()
    options = aw.saving.MarkdownSaveOptions()
    options.images_folder = str(tmp_path/'images')
    stream = Stalled() if failed else BytesIO()
    if failed:
        with pytest.raises(RuntimeError):doc.save(stream, options)
        assert not list((tmp_path/'images').glob('*.png'))
    else:
        assert doc.save(stream, options).content_type == 'text/plain'
        images = list((tmp_path/'images').glob('*.png'))
        assert len(images) == 1 and images[0].read_bytes() == image
        assert images[0].as_uri() in stream.getvalue().decode()
    assert not stream.closed


@pytest.mark.parametrize('option', [aw.saving.OoxmlSaveOptions(), aw.saving.MarkdownSaveOptions()])
def test_mutated_options_fail_before_touching_stream(option):
    option.save_format = aw.SaveFormat.TEXT
    stream = BytesIO(b'OWNED unchanged');stream.seek(4)
    with pytest.raises(ValueError):document().save(stream, option)
    assert stream.tell() == 4 and stream.getvalue() == b'OWNED unchanged'


@pytest.mark.parametrize('fmt', [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC])
@pytest.mark.parametrize('failed', [False, True])
def test_source_metadata_commits_only_after_successful_write(fmt, failed):
    doc = document()
    row = ldm.Row(row_format=ldm.RowFormat(height=2000, height_rule=1),
                  cells=[ldm.Cell(paragraphs=[ldm.Paragraph(children=[ldm.Run(text='OWNED ROW')])])])
    doc.light_document_model.sections[0].body.children.append(ldm.Table(rows=[row]))
    before = doc.light_document_model.model_dump_json()
    stream = Stalled() if failed else BytesIO()
    if failed:
        with pytest.raises(RuntimeError):doc.save(stream, fmt)
        assert doc.light_document_model.model_dump_json() == before
    else:
        doc.save(stream, fmt)
        assert row.row_format.height == 1584
        assert doc.light_document_model.source_font_table is not None
        cold = aw.Document(BytesIO(stream.getvalue())).light_document_model
        assert cold.sections[0].body.children[1].rows[0].row_format.height == 1584
