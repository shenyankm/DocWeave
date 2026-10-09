"""Unresolved character-unit indents must not silently become successful layout."""

import subprocess
import sys
import warnings
from io import BytesIO

import pytest
from docx import Document as NativeDocx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from aspose.words_foss import Document, DocxDocument
from aspose.words_foss.diagnostics import ContentLossWarning, collect_diagnostics

ATTRIBUTES = ("leftChars", "rightChars", "startChars", "endChars", "firstLineChars", "hangingChars")


def source(attribute, kind, value):
    document = NativeDocx()
    document.add_paragraph("PRIVATE BODY")
    if kind == "body":
        properties = document.paragraphs[0]._p.get_or_add_pPr()
    elif kind == "table":
        properties = document.add_table(1, 1).cell(0, 0).paragraphs[0]._p.get_or_add_pPr()
    elif kind in {"header", "footer"}:
        properties = getattr(document.sections[0], kind).paragraphs[0]._p.get_or_add_pPr()
    elif kind == "style":
        properties = document.styles["Normal"].element.get_or_add_pPr()
    elif kind == "default":
        defaults = document.styles.element.find(qn("w:docDefaults"))
        group = defaults.find(qn("w:pPrDefault"))
        properties = group.find(qn("w:pPr"))
    elif kind in {"footnote", "endnote"}:
        properties = OxmlElement("w:pPr")
    else:
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        properties = OxmlElement("w:pPr")
        level.append(properties)
        for name, item in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1.")):
            setting = OxmlElement("w:" + name)
            setting.set(qn("w:val"), item)
            level.insert(len(level) - 1, setting)
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), "999")
        abstract.append(level)
        document.part.numbering_part.element.append(abstract)
        number = OxmlElement("w:num")
        number.set(qn("w:numId"), "999")
        reference = OxmlElement("w:abstractNumId")
        reference.set(qn("w:val"), "999")
        number.append(reference)
        document.part.numbering_part.element.append(number)
        numbering = OxmlElement("w:numPr")
        for name, item in (("ilvl", "0"), ("numId", "999")):
            setting = OxmlElement("w:" + name)
            setting.set(qn("w:val"), item)
            numbering.append(setting)
        document.paragraphs[0]._p.get_or_add_pPr().append(numbering)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:" + attribute), value)
    properties.append(indent)
    stream = BytesIO()
    document.save(stream)
    raw = stream.getvalue()
    if kind in {"footnote", "endnote"}:
        from zipfile import ZipFile

        with ZipFile(BytesIO(raw)) as package:
            parts = {name: package.read(name) for name in package.namelist()}
        relationship = f'<Relationship Id="notes" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s" Target="{kind}s.xml"/>'
        parts["word/_rels/document.xml.rels"] = parts["word/_rels/document.xml.rels"].replace(b"</Relationships>", relationship.encode() + b"</Relationships>")
        content_type = f'<Override PartName="/word/{kind}s.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml"/>'
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(b"</Types>", content_type.encode() + b"</Types>")
        parts[f"word/{kind}s.xml"] = (f'<w:{kind}s xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:{kind} w:id="1"><w:p>{properties.xml}<w:r><w:t>PRIVATE NOTE</w:t></w:r></w:p></w:{kind}></w:{kind}s>').encode()
        stream = BytesIO()
        with ZipFile(stream, "w") as package:
            for name, data in parts.items():
                package.writestr(name, data)
        raw = stream.getvalue()
    return raw


@pytest.mark.parametrize("attribute", ATTRIBUTES)
@pytest.mark.parametrize("kind", ["body", "table", "header", "footer", "style", "default", "numbering", "footnote", "endnote"])
@pytest.mark.parametrize("value", ["0", "100"])
def test_character_indents_warn_and_strict_snapshot_preserves_original(attribute, kind, value):
    raw = source(attribute, kind, value)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", ContentLossWarning)
        document = Document(BytesIO(raw))
        assert document.get_text()
    observed = [row for row in document.diagnostics if row.code == "load.character_indents_ignored"]
    assert len(observed) == 1
    assert attribute in observed[0].message
    assert "PRIVATE BODY" not in repr(observed)
    assert "PRIVATE NOTE" not in repr(observed)
    editable = DocxDocument(BytesIO(raw))
    assert editable.to_bytes() == raw
    with warnings.catch_warnings():
        warnings.simplefilter("error", ContentLossWarning)
        with pytest.raises(ContentLossWarning, match="character-unit"):
            editable.to_light_document()
    assert editable.to_bytes() == raw


def test_character_indent_warning_state_resets_between_reader_loads():
    from aspose.words_foss.docx_reader import DocumentReader

    reader = DocumentReader()
    diagnostics = []
    with collect_diagnostics(diagnostics), warnings.catch_warnings():
        warnings.simplefilter("ignore", ContentLossWarning)
        reader.load_bytes(source("startChars", "body", "100"))
        reader.load_bytes(source("start", "body", "240"))
    assert len([row for row in diagnostics if row.code == "load.character_indents_ignored"]) == 1


@pytest.mark.skipif(sys.platform == "win32", reason="Bounded CLI requires process groups")
def test_strict_cli_rejects_character_indent_loss_before_replacing_output(tmp_path):
    path = tmp_path / "source.docx"
    path.write_bytes(source("startChars", "body", "100"))
    output = tmp_path / "output.pdf"
    output.write_bytes(b"existing output")
    result = subprocess.run([sys.executable, "-m", "aspose.words_foss.convert", str(path), str(output), "--strict"],
                            capture_output=True, text=True, timeout=20, check=False)
    assert result.returncode != 0
    assert "character-unit" in result.stderr
    assert "PRIVATE BODY" not in result.stderr
    assert output.read_bytes() == b"existing output"
    assert not list(tmp_path.glob(".conversion-*"))


def test_character_indent_warning_aggregates_without_values_or_body_text():
    from zipfile import ZipFile

    raw = source("leftChars", "body", "100")
    with ZipFile(BytesIO(raw)) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    parts["word/document.xml"] = parts["word/document.xml"].replace(b'w:leftChars="100"', b'w:leftChars="100" w:endChars="987654321"')
    output = BytesIO()
    with ZipFile(output, "w") as package:
        for name, data in parts.items():
            package.writestr(name, data)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", ContentLossWarning)
        document = Document(BytesIO(output.getvalue()))
    observed = [row for row in document.diagnostics if row.code == "load.character_indents_ignored"]
    assert len(observed) == 1
    assert "endChars, leftChars" in observed[0].message
    assert "987654321" not in repr(observed)
    assert "PRIVATE BODY" not in repr(observed)


@pytest.mark.parametrize("kind", ["body", "style"])
def test_flat_opc_uses_the_same_character_indent_loss_guard(kind):
    editable = DocxDocument(BytesIO(source("hangingChars", kind, "100")))
    raw = editable.to_flat_opc()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", ContentLossWarning)
        document = Document(BytesIO(raw))
    assert any(row.code == "load.character_indents_ignored" for row in document.diagnostics)
    with warnings.catch_warnings():
        warnings.simplefilter("error", ContentLossWarning)
        with pytest.raises(ContentLossWarning, match="character-unit"):
            Document(BytesIO(raw))
