"""DOCX compatibility settings control flow-table placement and survive conversion."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
import pymupdf
import pytest

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.saving import OoxmlSaveOptions, PdfSaveOptions


WORD_URI = 'http://schemas.microsoft.com/office/word'


def source_docx(mode, *, alignment=0, indent=0, indent_type='dxa', uri=WORD_URI):
    document = Document()
    section = document.sections[0]
    section.page_width = section.page_height = Pt(300)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Pt(20)
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = table.cell(0, 0).width = Pt(160)
    table.alignment = WD_TABLE_ALIGNMENT(alignment)
    width = table._tbl.tblPr.find(qn('w:tblW'))
    width.set(qn('w:w'), '3200')
    width.set(qn('w:type'), 'dxa')
    element = OxmlElement('w:tblInd')
    element.set(qn('w:w'), str(indent * 20) if isinstance(indent, int) else indent)
    element.set(qn('w:type'), indent_type)
    table._tbl.tblPr.append(element)
    cell = table.cell(0, 0)
    margins = OxmlElement('w:tcMar')
    for name in ('left', 'right'):
        element = OxmlElement('w:' + name)
        element.set(qn('w:w'), '160')
        element.set(qn('w:type'), 'dxa')
        margins.append(element)
    cell._tc.get_or_add_tcPr().append(margins)
    cell.paragraphs[0].add_run('CELL').font.size = Pt(12)
    compat = document.settings.element.find(qn('w:compat'))
    setting = next(element for element in compat if element.get(qn('w:name')) == 'compatibilityMode')
    if mode is None:
        compat.remove(setting)
    else:
        setting.set(qn('w:val'), str(mode))
        setting.set(qn('w:uri'), uri)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def cell_x(document, shaping):
    options = PdfSaveOptions()
    options.text_shaping = shaping
    with pymupdf.open(stream=document.to_bytes(options), filetype='pdf') as pdf:
        assert len(pdf) == 1
        words = pdf[0].get_text('words')
        assert [word[4] for word in words] == ['CELL']
        return words[0][0]


@pytest.mark.parametrize('mode', [11, 12, 14, 15])
@pytest.mark.parametrize('alignment', [0, 1, 2])
@pytest.mark.parametrize('indent', [0, 24])
@pytest.mark.parametrize('shaping', [False, True])
def test_source_table_placement_and_docx_roundtrip(mode, alignment, indent, shaping):
    document = aw.Document(BytesIO(source_docx(mode, alignment=alignment, indent=indent)))
    assert document.light_document_model.compatibility_mode == mode
    snapshot = document.light_document_model.model_dump()
    expected = (20 + indent + (8 if mode >= 15 else 0)) if alignment == 0 else 78 if alignment == 1 else 128
    assert cell_x(document, shaping) == pytest.approx(expected, abs=0.05)
    raw = document.to_bytes(OoxmlSaveOptions())
    with ZipFile(BytesIO(raw)) as archive:
        settings = ET.fromstring(archive.read('word/settings.xml'))
        flags = [element for element in settings.findall(qn('w:compat') + '/' + qn('w:compatSetting'))
                 if element.get(qn('w:name')) == 'compatibilityMode']
        assert len(flags) == 1
        assert flags[0].get(qn('w:val')) == str(mode)
        assert flags[0].get(qn('w:uri')) == WORD_URI
    loaded = aw.Document(BytesIO(raw))
    assert loaded.light_document_model.compatibility_mode == mode
    assert cell_x(loaded, shaping) == pytest.approx(expected, abs=0.05)
    assert document.light_document_model.model_dump() == snapshot


@pytest.mark.parametrize('mode', [None, '', '-1', 'bad', '١٥'])
def test_missing_or_invalid_setting_uses_ooxml_default(mode):
    document = aw.Document(BytesIO(source_docx(mode)))
    assert document.light_document_model.compatibility_mode == 12
    assert cell_x(document, False) == pytest.approx(20, abs=0.05)


def test_foreign_compatibility_uri_does_not_change_default():
    document = aw.Document(BytesIO(source_docx(15, uri='https://example.test/word')))
    assert document.light_document_model.compatibility_mode == 12


def test_missing_settings_part_uses_ooxml_default():
    output = BytesIO()
    with ZipFile(BytesIO(source_docx(15))) as source, ZipFile(output, 'w') as target:
        for name in source.namelist():
            if name != 'word/settings.xml':
                target.writestr(name, source.read(name))
    assert aw.Document(BytesIO(output.getvalue())).light_document_model.compatibility_mode == 12


@pytest.mark.parametrize('mode', [11, 12, 14, 15, 16])
def test_full_ldm_json_preserves_compatibility_mode(mode):
    model = ldm.Document(compatibility_mode=mode)
    assert ldm.Document.model_validate_json(model.model_dump_json()).compatibility_mode == mode


def test_new_documents_declare_modern_mode():
    model = ldm.Document()
    assert model.compatibility_mode == 15
    loaded = aw.Document(BytesIO(LdmDocxWriter().write_to_bytes(model)))
    assert loaded.light_document_model.compatibility_mode == 15
    with pytest.raises(ValueError):
        ldm.Document(compatibility_mode=-1)


@pytest.mark.parametrize('indent_type', ['pct', 'auto', 'nil'])
def test_non_point_indent_is_ignored_even_with_non_numeric_width(indent_type):
    document = aw.Document(BytesIO(source_docx(15, indent='ignored', indent_type=indent_type)))
    assert document.light_document_model.sections[0].body.tables[0].left_indent == 0
    assert cell_x(document, False) == pytest.approx(28, abs=0.05)


def test_legacy_table_with_empty_leading_row_keeps_first_real_cell_origin():
    from aspose.words_foss.pdf_writer import LdmPdfWriter

    document = aw.Document(BytesIO(source_docx(14))).light_document_model
    document.sections[0].body.tables[0].rows.insert(0, ldm.Row())
    with pymupdf.open(stream=LdmPdfWriter().write_to_bytes(document), filetype='pdf') as pdf:
        assert pdf[0].get_text('words')[0][0] == pytest.approx(20, abs=0.05)
