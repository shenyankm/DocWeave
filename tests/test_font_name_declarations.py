"""Keep local font declarations separate from the reader's effective names."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document

from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_reader.document_reader import DocumentReader
from aspose.words_foss.docx_writer.writer import LdmDocxWriter
from aspose.words_foss._flat_opc import encode, decode

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
CASES = [None, {}, {'ascii': 'OwnedAscii'}, {'hAnsi': 'OwnedOther'},
         {'asciiTheme': 'minorHAnsi', 'eastAsiaTheme': 'majorEastAsia'},
         {'ascii': 'OwnedAscii', 'hAnsi': 'OwnedOther', 'cs': 'OwnedBi',
          'eastAsia': 'OwnedEast', 'hint': 'eastAsia'}]


def source(attributes, location="body", defaults=None):
    document = Document()
    if location == 'body':
        document.add_paragraph('DECLARATION')
    elif location == 'table':
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'DECLARATION'
    else:
        getattr(document.sections[0], location).paragraphs[0].add_run('DECLARATION')
    stream = BytesIO()
    document.save(stream)
    with ZipFile(BytesIO(stream.getvalue())) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    part = 'word/document.xml' if location in ('body', 'table') else 'word/' + location + '1.xml'
    xml = ET.fromstring(parts[part])
    run = xml.find('.//' + W + 'r')
    if attributes is not None:
        properties = ET.Element(W + 'rPr')
        run.insert(0, properties)
        ET.SubElement(properties, W + 'rFonts', {W + key: value for key, value in attributes.items()})
    parts[part] = ET.tostring(xml)
    if defaults is not None:
        styles = ET.fromstring(parts['word/styles.xml'])
        fonts = styles.find(W + 'docDefaults/' + W + 'rPrDefault/' + W + 'rPr/' + W + 'rFonts')
        fonts.attrib = {W + key: value for key, value in defaults.items()}
        parts['word/styles.xml'] = ET.tostring(styles)
    stream = BytesIO()
    with ZipFile(stream, 'w', ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    return stream.getvalue()


def read(data):
    reader = DocumentReader()
    reader.load_bytes(data)
    return reader.to_light_document()


def names(data, location="body"):
    with ZipFile(BytesIO(data)) as package:
        part = 'word/document.xml' if location in ('body', 'table') else 'word/' + location + '1.xml'
        xml = ET.fromstring(package.read(part))
    fonts = xml.find('.//' + W + 'r/' + W + 'rPr/' + W + 'rFonts')
    return None if fonts is None else {key.removeprefix(W): value for key, value in fonts.attrib.items()}


@pytest.mark.parametrize('attributes', CASES)
@pytest.mark.parametrize('path', ['source', 'json', 'flat'])
@pytest.mark.parametrize('location', ['body', 'table', 'header', 'footer'])
def test_local_font_declarations_survive_two_saves(attributes, path, location):
    model = read(source(attributes, location))
    if path == 'json':
        model = ldm.Document.model_validate_json(model.model_dump_json())
    paragraph = next(p for p in model.get_child_nodes(ldm.NodeType.PARAGRAPH, True) if p.text == 'DECLARATION')
    run = paragraph.runs[0]
    assert run.font.font_names_explicit is (attributes is not None)
    assert (None if run.font.source_font_names is None else run.font.source_font_names.attributes) == attributes
    for _ in range(2):
        data = LdmDocxWriter().write_to_bytes(model)
        if path == 'flat':
            data = decode(encode(data))
        assert names(data, location) == attributes
        model = read(data)


@pytest.mark.parametrize('field,channel', [('name_ascii', 'ascii'), ('name_other', 'hAnsi'),
                                         ('name_bi', 'cs'), ('name_far_east', 'eastAsia')])
def test_editing_one_channel_preserves_the_other_declarations(field, channel):
    attributes = {'asciiTheme': 'minorHAnsi', 'hAnsiTheme': 'majorHAnsi',
                  'cstheme': 'minorBidi', 'eastAsiaTheme': 'majorEastAsia', 'hint': 'eastAsia'}
    model = read(source(attributes))
    run = model.sections[0].body.paragraphs[0].runs[0]
    setattr(run.font, field, 'Edited')
    expected = dict(attributes)
    expected.pop({'ascii': 'asciiTheme', 'hAnsi': 'hAnsiTheme', 'cs': 'cstheme', 'eastAsia': 'eastAsiaTheme'}[channel])
    expected[channel] = 'Edited'
    assert names(LdmDocxWriter().write_to_bytes(model)) == expected


def test_setting_common_name_replaces_all_theme_channels():
    model = read(source({'asciiTheme': 'minorHAnsi', 'hint': 'eastAsia'}))
    model.sections[0].body.paragraphs[0].runs[0].font.name = 'Edited'
    assert names(LdmDocxWriter().write_to_bytes(model)) == dict.fromkeys(['ascii', 'hAnsi', 'cs', 'eastAsia'], 'Edited') | {'hint': 'eastAsia'}


def test_font_channel_metadata_rejects_unknown_attributes():
    with pytest.raises(ValueError):
        ldm.SourceFontNames(attributes={'unknown': 'value'})


@pytest.mark.parametrize('attributes', [None, {}, {'val': 'en-US', 'eastAsia': 'ja-JP', 'bidi': 'ar-SA'}])
def test_theme_languages_keep_absent_and_empty_attributes_through_json(attributes):
    data = source(None)
    with ZipFile(BytesIO(data)) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    xml = ET.fromstring(parts['word/settings.xml'])
    for item in xml.findall(W + 'themeFontLang'):
        xml.remove(item)
    if attributes is not None:
        ET.SubElement(xml, W + 'themeFontLang', {W + key: value for key, value in attributes.items()})
    parts['word/settings.xml'] = ET.tostring(xml)
    stream = BytesIO()
    with ZipFile(stream, 'w', ZIP_DEFLATED) as package:
        for name, content in parts.items():
            package.writestr(name, content)
    model = ldm.Document.model_validate_json(read(stream.getvalue()).model_dump_json())
    with ZipFile(BytesIO(LdmDocxWriter().write_to_bytes(model))) as package:
        xml = ET.fromstring(package.read('word/settings.xml'))
    item = xml.find(W + 'themeFontLang')
    assert (None if item is None else {key.removeprefix(W): value for key, value in item.attrib.items()}) == attributes


@pytest.mark.parametrize('channel,field', [('ascii', 'name_ascii'), ('hAnsi', 'name_other'),
                                         ('cs', 'name_bi'), ('eastAsia', 'name_far_east')])
def test_partial_source_channel_does_not_edit_effective_inherited_channels(channel, field):
    defaults = {'ascii': 'DefaultAscii', 'hAnsi': 'DefaultOther', 'cs': 'DefaultBi', 'eastAsia': 'DefaultEast'}
    model = read(source({channel: 'Direct'}, defaults=defaults))
    font = model.sections[0].body.paragraphs[0].runs[0].font
    fields = {'ascii': 'name_ascii', 'hAnsi': 'name_other', 'cs': 'name_bi', 'eastAsia': 'name_far_east'}
    expected = defaults | {channel: 'Direct'}
    assert {key: getattr(font, name) or font.name for key, name in fields.items()} == expected
