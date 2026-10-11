"""Owned note bodies, ordered marks and resources survive conversion saves."""
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import base64
import warnings

from docx import Document as NativeDocument
import pytest
import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss._flat_opc import encode, decode
from aspose.words_foss._notes import note_id
from aspose.words_foss.docx_writer import LdmDocxWriter

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
PR = '{http://schemas.openxmlformats.org/package/2006/relationships}'
CT = '{http://schemas.openxmlformats.org/package/2006/content-types}'
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aLuoAAAAASUVORK5CYII=')


def package(parts):
    out = BytesIO()
    with ZipFile(out, 'w') as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return out.getvalue()


def owned_notes(kind='footnote', identifier='-1', custom=True):
    d = NativeDocument()
    p = d.add_paragraph('OWNED NOTE BEFORE ')
    p.add_run('OWNED BOLD').bold = True
    p.add_run().add_picture(BytesIO(PNG))
    d.add_table(rows=1, cols=1).cell(0, 0).text = 'OWNED NOTE TABLE'
    d.add_paragraph('OWNED NOTE AFTER')
    out = BytesIO(); d.save(out)
    with ZipFile(BytesIO(out.getvalue())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    doc = ET.fromstring(parts['word/document.xml']);body = doc.find(W + 'body')
    blocks = list(body)[:-1]
    body[:] = []
    main = ET.SubElement(body, W + 'p')
    ET.SubElement(ET.SubElement(main, W + 'r'), W + 't').text = 'OWNED MAIN BEFORE '
    r = ET.SubElement(main, W + 'r');pr = ET.SubElement(r, W + 'rPr');ET.SubElement(pr,W+'b')
    attributes = {W + 'id': identifier}
    if custom: attributes[W+'customMarkFollows'] = '1'
    ET.SubElement(r, W + kind + 'Reference', attributes)
    if custom: ET.SubElement(r, W+'t').text = '*'
    ET.SubElement(ET.SubElement(main, W+'r'), W+'t').text = ' OWNED MAIN AFTER'
    sect = ET.SubElement(body, W+'sectPr')
    secpr = ET.SubElement(sect,W+kind+'Pr');ET.SubElement(secpr,W+'numStart',{W+'val':'3'})
    notes = ET.Element(W+kind+'s')
    note = ET.SubElement(notes, W+kind, {W+'id':identifier})
    note.extend(blocks)
    markrun=ET.Element(W+'r');ET.SubElement(markrun,W+kind+'Ref');note[0].insert(0,markrun)
    hyperlink=ET.SubElement(note[-1],W+'hyperlink',{R+'id':'ownedLink'})
    ET.SubElement(ET.SubElement(hyperlink,W+'r'),W+'t').text='OWNED LINK'
    for id, typ, mark in [('0','separator','separator'),('1','continuationSeparator','continuationSeparator')]:
        entry=ET.SubElement(notes,W+kind,{W+'id':id,W+'type':typ});np=ET.SubElement(entry,W+'p')
        ET.SubElement(ET.SubElement(np,W+'r'),W+mark);ET.SubElement(ET.SubElement(np,W+'r'),W+'t').text='OWNED '+typ
    parts['word/document.xml']=ET.tostring(doc);parts['word/'+kind+'s.xml']=ET.tostring(notes)
    rels=ET.fromstring(parts['word/_rels/document.xml.rels'])
    noterels=ET.Element(PR+'Relationships')
    for relation in rels:
        if relation.get('Type','').endswith('/image'):noterels.append(relation)
    ET.SubElement(noterels,PR+'Relationship',{'Id':'ownedLink','Type':R[1:-1]+'/hyperlink','Target':'https://example.com/owned-note','TargetMode':'External'})
    parts['word/_rels/'+kind+'s.xml.rels']=ET.tostring(noterels)
    ET.SubElement(rels,PR+'Relationship',{'Id':'ownedNotes','Type':R[1:-1]+'/'+kind+'s','Target':kind+'s.xml'})
    parts['word/_rels/document.xml.rels']=ET.tostring(rels)
    types=ET.fromstring(parts['[Content_Types].xml']);ET.SubElement(types,CT+'Override',{'PartName':'/word/'+kind+'s.xml','ContentType':'application/vnd.openxmlformats-officedocument.wordprocessingml.'+kind+'s+xml'})
    parts['[Content_Types].xml']=ET.tostring(types)
    settings=ET.fromstring(parts['word/settings.xml']);pr=ET.SubElement(settings,W+kind+'Pr')
    for tag,val in [('pos','beneathText' if kind=='footnote' else 'sectEnd'),('numFmt','lowerRoman'),('numStart','7'),('numRestart','eachSect')]:ET.SubElement(pr,W+tag,{W+'val':val})
    for id in ['0','1']:ET.SubElement(pr,W+kind,{W+'id':id})
    parts['word/settings.xml']=ET.tostring(settings)
    return package(parts)


@pytest.mark.parametrize('kind',['footnote','endnote'])
@pytest.mark.parametrize('flat_input',[False,True])
@pytest.mark.parametrize('via_json',[False,True])
@pytest.mark.parametrize('flat_output',[False,True])
def test_notes_edit_json_resources_and_two_saves(kind,flat_input,via_json,flat_output):
    raw=owned_notes(kind)
    with warnings.catch_warnings():
        warnings.simplefilter('error', aw.ContentLossWarning)
        model=aw.Document(BytesIO(encode(raw) if flat_input else raw)).light_document_model
        if via_json:model=ldm.Document.model_validate_json(model.model_dump_json())
        story=next(s for s in model.source_stories if s.kind==kind and s.note_type=='normal')
        assert story.identifier=='-1'
        assert [type(n).__name__ for n in story.children]==['Paragraph','Table','Paragraph']
        story.children[0].runs[0].text='EDITED NOTE BEFORE '
        model.sections[0].body.paragraphs[0].runs[0].text='EDITED MAIN BEFORE '
        for _ in range(2):
            before=model.model_dump_json();raw=LdmDocxWriter().write_to_bytes(model)
            assert before==model.model_dump_json()
            with ZipFile(BytesIO(raw)) as z:
                notes=ET.fromstring(z.read('word/'+kind+'s.xml'))
                note=next(n for n in notes if n.get(W+'id')=='-1')
                assert [n.tag for n in note]==[W+'p',W+'tbl',W+'p']
                assert 'EDITED NOTE BEFORE ' in ''.join(note.itertext())
                assert len(note.findall('.//'+W+kind+'Ref'))==1
                assert len(note.findall('.//'+W+'drawing'))==1
                assert len(note.findall('.//'+W+'b'))==1
                assert [''.join(n.itertext()) for n in notes if n.get(W+'type')]==['OWNED separator','OWNED continuationSeparator']
                rels=ET.fromstring(z.read('word/_rels/'+kind+'s.xml.rels'))
                assert any(n.get('Target')=='https://example.com/owned-note' for n in rels)
                assert any(z.read('word/'+n.get('Target'))==PNG for n in rels if n.get('Type').endswith('/image'))
                body=ET.fromstring(z.read('word/document.xml'))
                anchor=body.find('.//'+W+kind+'Reference');assert anchor.get(W+'id')=='-1';assert anchor.get(W+'customMarkFollows')=='1'
                assert 'EDITED MAIN BEFORE ' in ''.join(body.itertext())
                assert body.find('.//'+W+'sectPr/'+W+kind+'Pr/'+W+'numStart').get(W+'val')=='3'
                pr=ET.fromstring(z.read('word/settings.xml')).find(W+kind+'Pr')
                assert pr.find(W+'numFmt').get(W+'val')=='lowerRoman'
                assert [n.get(W+'id') for n in pr.findall(W+kind)]==['0','1']
            model=aw.Document(BytesIO(encode(raw) if flat_output else raw)).light_document_model
            if via_json:model=ldm.Document.model_validate_json(model.model_dump_json())


@pytest.mark.parametrize('value,expected',[
    ('01','1'),('+1','1'),(' 1 ','1'),('-1','-1'),('-0','0'),('1.9','1'),('-1.9','-1'),('1e1','10'),
    ('1.99999999999999999','2'),('bad','0'),('1bad','0'),('١','0'),('\u00a01\u00a0','0'),('','0'),
    ('2147483648','-2147483648'),('4294967297','1'),('9223372036854777856','-1'),
    ('9223372036854774784','-1024'),('-9223372036854774784','1024'),('-9223372036854777856','0')])
def test_native_calibrated_numeric_note_ids(value,expected):
    assert note_id(value)==expected


@pytest.mark.parametrize('failure',['missing-story','duplicate-story','unplaced-anchor','separator-reference','missing-separator','cyclic-story'])
def test_invalid_notes_reject_save_atomically(tmp_path,failure):
    model=aw.Document(BytesIO(owned_notes())).light_document_model
    story=next(s for s in model.source_stories if s.note_type=='normal')
    anchor=next(n for n in model.sections[0].body.paragraphs[0]._children if isinstance(n,ldm.NoteReference))
    if failure=='missing-story':model.source_stories.remove(story)
    elif failure=='duplicate-story':model.source_stories.append(story.model_copy(deep=True))
    elif failure=='unplaced-anchor':model.sections[0].body.paragraphs[0]._children.remove(anchor)
    elif failure=='separator-reference':anchor.identifier='0'
    elif failure=='missing-separator':model.source_stories=[s for s in model.source_stories if s.identifier!='0']
    else:story.children[1].rows[0].cells[0].tables.append(story.children[1])
    destination=tmp_path/'owned.docx';destination.write_bytes(b'ORIGINAL')
    with pytest.raises((ValueError,NotImplementedError)):LdmDocxWriter().write(model,destination)
    assert destination.read_bytes()==b'ORIGINAL'


def test_note_anchor_outside_supported_inline_container_is_rejected():
    from .test_conversion_api import package as simple_package
    with pytest.warns(aw.loading.DocumentLoadWarning, match='footnoteReference'):
        with pytest.raises(NotImplementedError, match='note mark location'):
            aw.Document(BytesIO(simple_package('<w:p><w:footnoteReference w:id="1"/></w:p>')))


def test_reference_only_model_rejects_docx_instead_of_omitting_notes():
    model=ldm.Document(sections=[ldm.Section(body=ldm.Body(children=[ldm.Paragraph(children=[
        ldm.Run(text='body'),ldm.NoteReference(kind='footnote',identifier='7')])]))])
    with pytest.raises(ValueError,match='matching normal story'):LdmDocxWriter().write_to_bytes(model)


def test_custom_note_reference_mark_shares_its_anchor_run():
    model=aw.Document(BytesIO(owned_notes())).light_document_model
    raw=LdmDocxWriter().write_to_bytes(model)
    with ZipFile(BytesIO(raw)) as z:
        doc=ET.fromstring(z.read('word/document.xml'))
        run=next(r for r in doc.iter(W+'r') if r.find(W+'footnoteReference') is not None)
        assert run.find(W+'t').text=='*'


@pytest.mark.parametrize('kind',['footnote','endnote'])
@pytest.mark.parametrize('via_json',[False,True])
@pytest.mark.parametrize('flat_output',[False,True])
def test_continuation_notice_with_matching_property_survives_two_saves(kind,via_json,flat_output):
    model=aw.Document(BytesIO(owned_notes(kind))).light_document_model
    model.source_stories.append(ldm.SourceStory(kind=kind,part_name='word/'+kind+'s.xml',identifier='2',
        note_type='continuationNotice',children=[ldm.Paragraph(children=[ldm.Run(text='OWNED NOTICE')])]))
    model.note_properties[0].children.append(ldm.NoteProperty(kind=kind,value='2'))
    for _ in range(2):
        if via_json:model=ldm.Document.model_validate_json(model.model_dump_json())
        raw=LdmDocxWriter().write_to_bytes(model)
        model=aw.Document(BytesIO(encode(raw) if flat_output else raw)).light_document_model
        notice=next(s for s in model.source_stories if s.note_type=='continuationNotice')
        assert notice.identifier=='2' and notice.children[0].text=='OWNED NOTICE'
        assert model.note_properties[0].children[-1].value=='2'


@pytest.mark.parametrize('bad_root',['footer','comments','footnotes-with-other-entry'])
def test_invalid_note_part_root_is_rejected_even_without_an_anchor(bad_root):
    raw=owned_notes()
    with ZipFile(BytesIO(raw)) as z:parts={n:z.read(n) for n in z.namelist()}
    doc=ET.fromstring(parts['word/document.xml'])
    for run in list(doc.iter(W+'r')):
        for node in list(run):
            if node.tag==W+'footnoteReference':run.remove(node)
    parts['word/document.xml']=ET.tostring(doc)
    name='word/footnotes.xml';root=ET.fromstring(parts[name])
    if bad_root=='footnotes-with-other-entry':ET.SubElement(root,W+'comment')
    else:root.tag=W+bad_root
    parts[name]=ET.tostring(root)
    with pytest.raises(ValueError,match='note story root'):aw.Document(BytesIO(package(parts)))


def test_json_reference_mirror_does_not_override_ordered_mark_edits():
    model=ldm.Document.model_validate_json(aw.Document(BytesIO(owned_notes())).light_document_model.model_dump_json())
    anchor=next(n for n in model.sections[0].body.paragraphs[0]._children if isinstance(n,ldm.NoteReference))
    anchor.reference_mark='EDITED MARK';anchor.font.italic=True
    raw=LdmDocxWriter().write_to_bytes(model)
    restored=aw.Document(BytesIO(raw)).light_document_model
    cold=next(n for n in restored.sections[0].body.paragraphs[0]._children if isinstance(n,ldm.NoteReference))
    assert cold.reference_mark=='EDITED MARK' and cold.font.italic


@pytest.mark.parametrize('kind',['footnote','endnote'])
@pytest.mark.parametrize('via_json',[False,True])
@pytest.mark.parametrize('flat_output',[False,True])
def test_note_body_carriage_return_survives_two_saves(kind,via_json,flat_output):
    model=aw.Document(BytesIO(owned_notes(kind))).light_document_model
    story=next(s for s in model.source_stories if s.note_type=='normal')
    story.children[0].runs[0].text='EDITED NOTE BEFORE\rOWNED CR AFTER '
    for _ in range(2):
        if via_json:model=ldm.Document.model_validate_json(model.model_dump_json())
        raw=LdmDocxWriter().write_to_bytes(model)
        with ZipFile(BytesIO(raw)) as z:
            note=next(n for n in ET.fromstring(z.read('word/'+kind+'s.xml')) if n.get(W+'type') is None)
            assert len(note.findall('.//'+W+'cr'))==1
        model=aw.Document(BytesIO(encode(raw) if flat_output else raw)).light_document_model
        note=next(s for s in model.source_stories if s.note_type=='normal')
        assert note.children[0].runs[0].text=='EDITED NOTE BEFORE\rOWNED CR AFTER '


@pytest.mark.parametrize('kind',['footnote','endnote'])
@pytest.mark.parametrize('mark',['MARK\rTAIL','MARK\nTAIL','MARK\tTAIL','* Ω',''])
@pytest.mark.parametrize('via_json',[False,True])
@pytest.mark.parametrize('flat_output',[False,True])
def test_custom_note_reference_whitespace_survives_two_saves(kind,mark,via_json,flat_output):
    model=aw.Document(BytesIO(owned_notes(kind))).light_document_model
    anchor=next(n for n in model.sections[0].body.paragraphs[0]._children if isinstance(n,ldm.NoteReference))
    anchor.reference_mark=mark
    for _ in range(2):
        if via_json:model=ldm.Document.model_validate_json(model.model_dump_json())
        raw=LdmDocxWriter().write_to_bytes(model)
        model=aw.Document(BytesIO(encode(raw) if flat_output else raw)).light_document_model
        anchor=next(n for n in model.sections[0].body.paragraphs[0]._children if isinstance(n,ldm.NoteReference))
        assert anchor.reference_mark==mark


@pytest.mark.parametrize('source',['<root>CR&#13;LF\nTAB\t&amp;&lt;</root>', '<root><![CDATA[CR]]>&#13;<![CDATA[LF\nTAB\t&<]]></root>'])
def test_shared_xml_serializer_preserves_text_whitespace(source):
    from xml.dom.minidom import parseString
    from aspose.words_foss.utils.xml_helpers import serialize_xml
    original=parseString(source)
    raw=serialize_xml(original)
    assert b'&#13;' in raw
    restored=parseString(raw)
    assert ''.join(n.data for n in original.documentElement.childNodes)==''.join(n.data for n in restored.documentElement.childNodes)


@pytest.mark.parametrize('kind',['footnote','endnote'])
def test_custom_note_reference_invalid_control_rejects_save(kind):
    model=aw.Document(BytesIO(owned_notes(kind))).light_document_model
    anchor=next(n for n in model.sections[0].body.paragraphs[0]._children if isinstance(n,ldm.NoteReference))
    anchor.reference_mark='MARK\x0bTAIL'
    with pytest.raises(ValueError,match='custom note reference mark'):LdmDocxWriter().write_to_bytes(model)


@pytest.mark.parametrize('cdata',[False,True])
@pytest.mark.parametrize('character',['\x00','\x0b','\ud800','\ufffe'])
def test_shared_xml_serializer_rejects_invalid_text_privately(cdata,character):
    from xml.dom.minidom import Document
    from aspose.words_foss.utils.xml_helpers import serialize_xml
    doc=Document();root=doc.createElement('root');doc.appendChild(root)
    root.appendChild(doc.createCDATASection('private'+character+'value') if cdata else doc.createTextNode('private'+character+'value'))
    with pytest.raises(ValueError,match='^Invalid XML text$'):serialize_xml(doc)


@pytest.mark.parametrize('kind',['footnote','endnote'])
@pytest.mark.parametrize('deepcopy',[False,True])
@pytest.mark.parametrize('flat_output',[False,True])
def test_json_public_note_reference_edits_share_ordered_anchor(kind,deepcopy,flat_output):
    model=ldm.Document.model_validate_json(aw.Document(BytesIO(owned_notes(kind))).light_document_model.model_dump_json())
    if deepcopy:model=model.model_copy(deep=True)
    paragraph=model.sections[0].body.paragraphs[0]
    anchor=next(n for n in paragraph._children if isinstance(n,ldm.NoteReference))
    assert paragraph.note_references[0] is anchor
    paragraph.note_references[0].reference_mark='PUBLIC EDIT';paragraph.note_references[0].font.italic=True
    for _ in range(2):
        raw=LdmDocxWriter().write_to_bytes(model)
        model=aw.Document(BytesIO(encode(raw) if flat_output else raw)).light_document_model
        assert model.sections[0].body.paragraphs[0].note_references[0].reference_mark=='PUBLIC EDIT'
        assert model.sections[0].body.paragraphs[0].note_references[0].font.italic


def test_conflicting_note_reference_mirror_metadata_is_rejected():
    model=aw.Document(BytesIO(owned_notes())).light_document_model
    data=model.model_dump();data['sections'][0]['body']['children'][0]['note_references'][0]['reference_mark']='CONFLICT'
    with pytest.raises(ValueError,match='mirror disagrees'):ldm.Document.model_validate(data)
    paragraph=model.sections[0].body.paragraphs[0]
    paragraph.note_references[0]=paragraph.note_references[0].model_copy(deep=True)
    paragraph.note_references[0].reference_mark='DETACHED EDIT'
    with pytest.raises(ValueError,match='metadata disagrees'):LdmDocxWriter().write_to_bytes(model)
    paragraph.note_references.clear()
    with pytest.raises(ValueError,match='no ordered inline anchor'):LdmDocxWriter().write_to_bytes(model)


@pytest.mark.parametrize('kind',['footnote','endnote'])
@pytest.mark.parametrize('deepcopy',[False,True])
@pytest.mark.parametrize('flat_output',[False,True])
def test_public_note_reference_can_switch_back_to_automatic(kind,deepcopy,flat_output):
    model=ldm.Document.model_validate_json(aw.Document(BytesIO(owned_notes(kind))).light_document_model.model_dump_json())
    if deepcopy:model=model.model_copy(deep=True)
    reference=model.sections[0].body.paragraphs[0].note_references[0]
    reference.custom_mark_follows=False;reference.reference_mark=None;reference.font.italic=True
    for _ in range(2):
        raw=LdmDocxWriter().write_to_bytes(model)
        model=aw.Document(BytesIO(encode(raw) if flat_output else raw)).light_document_model
        reference=model.sections[0].body.paragraphs[0].note_references[0]
        assert reference.custom_mark_follows is False and reference.reference_mark is None and reference.font.italic
