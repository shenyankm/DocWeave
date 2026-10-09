"""Logical margins and RTL table geometry across DOCX, JSON and PDF."""

from io import BytesIO

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.saving import (
    MarkdownExportAsHtml,
    MarkdownSaveOptions,
    OoxmlSaveOptions,
    PdfSaveOptions,
)

from .test_table_cell_margins import margins
from .test_table_cell_margins import source_docx as margin_docx


def source_docx(*, bidi=True, scope="table", alignment=0, indent=0):
    document = Document(BytesIO(margin_docx(styled=scope == "style")))
    table = document.tables[0]
    table.columns[0].width = table.cell(0, 0).width = Pt(50)
    table.add_column(Pt(100))
    table.cell(0, 1).width = Pt(100)
    width = table._tbl.tblPr.find(qn("w:tblW"))
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), "3000")
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), ("left", "center", "right")[alignment])
    table._tbl.tblPr.append(jc)
    if indent:
        entry = OxmlElement("w:tblInd")
        entry.set(qn("w:type"), "dxa")
        entry.set(qn("w:w"), str(indent * 20))
        table._tbl.tblPr.append(entry)
    if bidi is not None:
        direction = OxmlElement("w:bidiVisual")
        if bidi != "present":
            direction.set(
                qn("w:val"), "1" if bidi is True else "0" if bidi is False else bidi
            )
        table._tbl.tblPr.append(direction)
    if scope == "style":
        pr = document.styles["Margin Zero"].element.find(qn("w:tblPr"))
        pr.remove(pr.find(qn("w:tblCellMar")))
        margins(pr, "w:tblCellMar", {"start": 6, "end": 12})
    else:
        margins(
            table._tbl.tblPr,
            "w:tblCellMar",
            {
                "left": 3,
                "right": 9,
                "top": 0,
                "bottom": 0,
                **({"start": 6, "end": 12} if scope == "table" else {}),
            },
        )
    for index, label in enumerate(("ONE", "TWO")):
        cell = table.cell(0, index)
        if scope == "cell":
            margins(cell._tc.get_or_add_tcPr(), "w:tcMar", {"start": 6, "end": 12})
        paragraph = cell.paragraphs[0]
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = (
            paragraph.paragraph_format.space_after
        ) = Pt(0)
        run = paragraph.add_run(label)
        run.font.name, run.font.size = "Arial", Pt(12)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def positions(document, shaping=False):
    options = PdfSaveOptions()
    options.text_shaping = shaping
    with pymupdf.open(stream=document.to_bytes(options), filetype="pdf") as pdf:
        assert len(pdf) == 1
        words = pdf[0].get_text("words")
        return {word[4]: word[0] for word in words}


@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("bidi", [False, True])
@pytest.mark.parametrize("scope", ["table", "cell", "style"])
def test_logical_margins_and_unequal_columns_survive_roundtrip(shaping, bidi, scope):
    document = aw.Document(BytesIO(source_docx(bidi=bidi, scope=scope)))
    model = document.light_document_model
    table = model.sections[0].body.tables[0]
    fmt = table.rows[0].cells[0].cell_format if scope == "cell" else table
    assert (fmt.left_padding, fmt.right_padding) == (6, 12)
    assert table.bidi is bidi
    snapshot = model.model_dump()
    expected = {"ONE": 156, "TWO": 56} if bidi else {"ONE": 26, "TWO": 76}
    assert positions(document, shaping) == pytest.approx(expected, abs=0.05)
    loaded = aw.Document(BytesIO(document.to_bytes(OoxmlSaveOptions())))
    assert loaded.light_document_model.sections[0].body.tables[0].bidi is bidi
    assert positions(loaded, shaping) == pytest.approx(expected, abs=0.05)
    assert (
        ldm.Document.model_validate_json(model.model_dump_json())
        .sections[0]
        .body.tables[0]
        .bidi
        is bidi
    )
    assert model.model_dump() == snapshot
    assert document.to_dict()["blocks"][0]["bidi"] is bidi


@pytest.mark.parametrize(
    "value,expected",
    [
        (None, False),
        ("present", True),
        ("1", True),
        ("true", True),
        ("on", True),
        ("0", False),
        ("false", False),
        ("off", False),
    ],
)
def test_source_table_direction_boolean_forms(value, expected):
    document = aw.Document(BytesIO(source_docx(bidi=value)))
    assert document.light_document_model.sections[0].body.tables[0].bidi is expected


@pytest.mark.parametrize("bidi", [False, True])
@pytest.mark.parametrize("alignment", [0, 1, 2])
@pytest.mark.parametrize("indent", [0, 24])
def test_direction_mirrors_alignment_and_leading_indent(bidi, alignment, indent):
    document = aw.Document(
        BytesIO(source_docx(bidi=bidi, alignment=alignment, indent=indent))
    )
    physical_alignment = 2 - alignment if bidi else alignment
    origin = 20 if physical_alignment == 0 else 35 if physical_alignment == 1 else 50
    if alignment == 0:
        origin += -indent if bidi else indent
    expected = {
        "ONE": origin + (100 if bidi else 0) + 6,
        "TWO": origin + (0 if bidi else 50) + 6,
    }
    assert positions(document) == pytest.approx(expected, abs=0.05)


@pytest.mark.parametrize("mode", [12, 14])
@pytest.mark.parametrize("indent", [0, 24])
@pytest.mark.parametrize("shaping", [False, True])
def test_legacy_rtl_table_offset_uses_left_margin(mode, indent, shaping):
    source = Document(BytesIO(source_docx(bidi=True, indent=indent)))
    compat = source.settings.element.find(qn("w:compat"))
    setting = next(e for e in compat if e.get(qn("w:name")) == "compatibilityMode")
    setting.set(qn("w:val"), str(mode))
    output = BytesIO()
    source.save(output)
    document = aw.Document(BytesIO(output.getvalue()))
    assert document.light_document_model.compatibility_mode == mode
    expected = {"ONE": 162 - indent, "TWO": 62 - indent}
    assert positions(document, shaping) == pytest.approx(expected, abs=0.05)
    restored = aw.Document(BytesIO(document.to_bytes(OoxmlSaveOptions())))
    assert positions(restored, shaping) == pytest.approx(expected, abs=0.05)


@pytest.mark.parametrize("direction", [True, False])
def test_style_only_direction_and_explicit_false_override_survive(direction):
    document = Document(BytesIO(source_docx(bidi=None, scope="style")))
    style = document.styles["Margin Base"]
    entry = OxmlElement("w:bidiVisual")
    style.element.find(qn("w:tblPr")).insert(0, entry)
    if not direction:
        entry = OxmlElement("w:bidiVisual")
        entry.set(qn("w:val"), "0")
        document.styles["Margin Zero"].element.find(qn("w:tblPr")).insert(0, entry)
    output = BytesIO()
    document.save(output)
    loaded = aw.Document(BytesIO(output.getvalue()))
    assert loaded.light_document_model.sections[0].body.tables[0].bidi is direction
    assert (
        next(
            s for s in loaded.light_document_model.styles if s.name == "Margin Base"
        ).table_style_format.bidi
        is True
    )
    roundtrip = aw.Document(BytesIO(loaded.to_bytes(OoxmlSaveOptions())))
    assert roundtrip.light_document_model.sections[0].body.tables[0].bidi is direction
    assert (
        next(
            s for s in roundtrip.light_document_model.styles if s.name == "Margin Base"
        ).table_style_format.bidi
        is True
    )


def test_markdown_reports_direction_loss_and_html_preserves_direction():
    document = aw.Document(BytesIO(source_docx()))
    with pytest.warns(aw.ContentLossWarning, match="right-to-left"):
        assert b"ONE" in document.to_bytes(aw.SaveFormat.MARKDOWN)
    assert any(d.code == "markdown.table_direction" for d in document.diagnostics)
    options = MarkdownSaveOptions()
    options.export_as_html = MarkdownExportAsHtml.TABLES
    assert b'<table dir="rtl">' in document.to_bytes(options)


def test_rtl_border_sides_are_mirrored_with_unequal_columns():
    document = Document(BytesIO(source_docx()))
    for cell in document.tables[0].rows[0].cells:
        borders = OxmlElement("w:tcBorders")
        for side, color in (("left", "FF0000"), ("right", "0000FF")):
            edge = OxmlElement("w:" + side)
            for key, value in (("val", "single"), ("sz", "8"), ("color", color)):
                edge.set(qn("w:" + key), value)
            borders.append(edge)
        cell._tc.get_or_add_tcPr().append(borders)
    output = BytesIO()
    document.save(output)
    with pymupdf.open(
        stream=aw.Document(BytesIO(output.getvalue())).to_bytes(PdfSaveOptions()),
        filetype="pdf",
    ) as pdf:
        lines = pdf[0].get_drawings()
        red = sorted(d["rect"].x0 for d in lines if d["color"] == (1, 0, 0))
        blue = sorted(d["rect"].x0 for d in lines if d["color"] == (0, 0, 1))
        assert red == pytest.approx([150, 200], abs=0.05)
        assert blue == pytest.approx([50, 150], abs=0.05)


@pytest.mark.parametrize("bidi", [False, True])
def test_nested_table_direction_is_independent_of_outer_table(bidi):
    document = Document(BytesIO(source_docx(bidi=bidi)))
    cell = document.tables[0].cell(0, 1)
    cell.paragraphs[0].clear()
    nested = cell.add_table(rows=1, cols=2)
    nested.autofit = False
    nested.columns[0].width = nested.columns[1].width = Pt(30)
    direction = OxmlElement("w:bidiVisual")
    direction.set(qn("w:val"), "0" if bidi else "1")
    nested._tbl.tblPr.append(direction)
    for cell, text in zip(nested.rows[0].cells, ("ALPHA", "BETA")):
        cell.text = text
        cell.paragraphs[0].runs[0].font.size = Pt(8)
    output = BytesIO()
    document.save(output)
    result = positions(aw.Document(BytesIO(output.getvalue())))
    assert (result["ALPHA"] < result["BETA"]) is bidi


def test_rtl_spanning_cell_and_repeated_headers_keep_logical_order():
    document = Document(BytesIO(source_docx()))
    table = document.tables[0]
    table.cell(0, 0).merge(table.cell(0, 1)).text = "HEADER"
    heading = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(heading)
    for i in range(20):
        row = table.add_row()
        for cell, width, text in zip(row.cells, (50, 100), (f"A{i}", f"B{i}")):
            cell.width = Pt(width)
            cell.text = text
            cell.paragraphs[0].runs[0].font.size = Pt(8)
    output = BytesIO()
    document.save(output)
    source = aw.Document(BytesIO(output.getvalue()))
    with pymupdf.open(stream=source.to_bytes(PdfSaveOptions()), filetype="pdf") as pdf:
        assert len(pdf) > 1
        labels = []
        for page in pdf:
            words = page.get_text("words")
            coords = {w[4]: w[0] for w in words}
            assert "HEADER" in coords
            assert coords["HEADER"] == pytest.approx(56, abs=0.05)
            for label, x in coords.items():
                if label.startswith("A") and label[1:].isdigit():
                    assert x > coords["B" + label[1:]]
            labels.extend(w[4] for w in words if w[4] != "HEADER")
        assert sorted(labels) == sorted(
            [f"{letter}{i}" for i in range(20) for letter in ("A", "B")]
        )


@pytest.mark.parametrize("bidi", [False, True])
def test_explicit_zero_logical_margins_override_legacy_aliases(bidi):
    document = Document(BytesIO(source_docx(bidi=bidi)))
    for cell in document.tables[0].rows[0].cells:
        margins(
            cell._tc.get_or_add_tcPr(),
            "w:tcMar",
            {"left": 7, "right": 9, "start": 0, "end": 0},
        )
    output = BytesIO()
    document.save(output)
    loaded = aw.Document(BytesIO(output.getvalue()))
    assert positions(loaded) == pytest.approx(
        {"ONE": 150, "TWO": 50} if bidi else {"ONE": 20, "TWO": 70}, abs=0.05
    )
    assert (
        loaded.light_document_model.sections[0]
        .body.tables[0]
        .rows[0]
        .cells[0]
        .cell_format.left_padding
        == 0
    )


@pytest.mark.parametrize("unit", ["pct", "auto", "nil"])
def test_unsupported_logical_margin_units_do_not_override_valid_legacy_margin(unit):
    document = Document(BytesIO(source_docx(bidi=False)))
    margin = document.tables[0]._tbl.tblPr.find(qn("w:tblCellMar")).find(qn("w:start"))
    margin.set(qn("w:type"), unit)
    margin.set(qn("w:w"), "ignored")
    output = BytesIO()
    document.save(output)
    loaded = aw.Document(BytesIO(output.getvalue()))
    assert loaded.light_document_model.sections[0].body.tables[0].left_padding == 3


@pytest.mark.parametrize("scope", ["table", "cell", "style"])
def test_negative_logical_margin_is_rejected(scope):
    document = Document(BytesIO(source_docx(scope=scope)))
    if scope == "table":
        pr = document.tables[0]._tbl.tblPr
        tag = "w:tblCellMar"
    elif scope == "cell":
        pr = document.tables[0].cell(0, 0)._tc.tcPr
        tag = "w:tcMar"
    else:
        pr = document.styles["Margin Zero"].element.find(qn("w:tblPr"))
        tag = "w:tblCellMar"
    pr.find(qn(tag)).find(qn("w:start")).set(qn("w:w"), "-1")
    output = BytesIO()
    document.save(output)
    with pytest.raises(ValueError, match="margins must be nonnegative"):
        aw.Document(BytesIO(output.getvalue()))


@pytest.mark.parametrize("direction", [True, False])
def test_direction_only_table_style_is_retained(direction):
    document = Document(BytesIO(source_docx(bidi=None, scope="style")))
    for name in ("Margin Base", "Margin Zero"):
        pr = document.styles[name].element.find(qn("w:tblPr"))
        for child in list(pr):
            pr.remove(child)
    pr = document.styles["Margin Zero"].element.find(qn("w:tblPr"))
    entry = OxmlElement("w:bidiVisual")
    entry.set(qn("w:val"), "1" if direction else "0")
    pr.append(entry)
    output = BytesIO()
    document.save(output)
    loaded = aw.Document(BytesIO(output.getvalue()))
    assert (
        next(
            s for s in loaded.light_document_model.styles if s.name == "Margin Zero"
        ).table_style_format.bidi
        is direction
    )
    roundtrip = aw.Document(BytesIO(loaded.to_bytes(OoxmlSaveOptions())))
    assert roundtrip.light_document_model.sections[0].body.tables[0].bidi is direction
