"""Conversion-size recovery keeps unknown paragraph references in actual packages."""
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss._comments import _comment_story_nodes
from aspose.words_foss.diagnostics import document_nodes
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.saving import OoxmlSaveOptions
from tests.test_conversion_comments import owned_comment
from tests.test_conversion_notes import W, owned_notes, package


def source(kind, size):
    if kind == 'comment':
        data = owned_comment()
    elif kind in ('footnote', 'endnote'):
        data = owned_notes(kind)
    else:
        doc = Document()
        doc.add_paragraph('OWNED BODY')
        if kind == 'header':
            doc.sections[0].header.paragraphs[0].text = 'OWNED HEADER'
        stream = BytesIO(); doc.save(stream); data = stream.getvalue()
    with ZipFile(BytesIO(data)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    styles = ET.fromstring(parts['word/styles.xml'])
    normal = next(s for s in styles.findall(W+'style') if s.get(W+'styleId') == 'Normal')
    normal.set(W+'default', 'off')
    if size is not None:
        ET.SubElement(ET.SubElement(normal, W+'rPr'), W+'sz', {W+'val': str(size*2)})
    parts['word/styles.xml'] = ET.tostring(styles)
    for name, data in list(parts.items()):
        if name.startswith('word/') and name.endswith('.xml') and any(
                term in name for term in ('document.xml', 'header', 'comments', 'footnotes', 'endnotes')):
            root = ET.fromstring(data)
            for paragraph in root.iter(W+'p'):
                props = paragraph.find(W+'pPr')
                if props is None:
                    props = ET.Element(W+'pPr'); paragraph.insert(0, props)
                for node in list(props):
                    if node.tag == W+'pStyle': props.remove(node)
                ET.SubElement(props, W+'pStyle', {W+'val': 'OwnedMissing'})
            parts[name] = ET.tostring(root)
    return package(parts)


def target(model, kind):
    text = {'body': 'OWNED BODY', 'header': 'OWNED HEADER', 'comment': 'OWNED COMMENT α',
            'footnote': 'OWNED NOTE BEFORE', 'endnote': 'OWNED NOTE BEFORE'}[kind]
    nodes = list(document_nodes(model)) + [n for s in model.source_stories
                                          for n in _comment_story_nodes(s.children)]
    return next(n.font for n in nodes if isinstance(n, ldm.Run) and n.text.strip() == text)


@pytest.mark.parametrize('kind', ['body', 'header', 'comment', 'footnote', 'endnote'])
@pytest.mark.parametrize('size', [None, 0, 16])
@pytest.mark.parametrize('route', ['original', 'json', 'deep'])
def test_unknown_paragraph_size_fallback_edits_and_roundtrips(kind, size, route):
    model = aw.Document(BytesIO(source(kind, size))).light_document_model
    if route == 'json': model = ldm.Document.model_validate_json(model.model_dump_json())
    if route == 'deep': model = model.model_copy(deep=True)
    held = target(model, kind)
    assert held.size == (11 if size is None else size)
    next(s for s in model.styles if s.name == 'Normal').font.size = 28
    assert held.size == held.render_size == 28
    for fmt in (aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC):
        current = model.model_copy(deep=True)
        for _ in range(2):
            current = aw.Document(BytesIO(LdmDocxWriter(OoxmlSaveOptions(fmt)).write_to_bytes(current))).light_document_model
            assert target(current, kind).size == 28
            assert any(p.paragraph_format.style_name == 'OwnedMissing' for p in document_nodes(current)
                       if isinstance(p, ldm.Paragraph))
    assert held.size == 28


@pytest.mark.parametrize('size', [None, 0, 16])
def test_dom_unknown_style_validation_remains_explicit(size):
    document = aw.DocxDocument(BytesIO(source('body', size)))
    before = document.to_bytes()
    with pytest.raises(ValueError, match='Missing style'):
        _ = document.body.paragraphs[0].runs[0].effective_font.size
    assert document.to_bytes() == before


@pytest.mark.parametrize('kind', ['body', 'header'])
@pytest.mark.parametrize('size', [None, 0, 16])
def test_unknown_paragraph_size_reaches_actual_pdf(kind, size):
    model = aw.Document(BytesIO(source(kind, size))).light_document_model
    for expected in (11 if size is None else size, 28):
        if expected == 28: next(s for s in model.styles if s.name == 'Normal').font.size = 28
        with pymupdf.open(stream=LdmPdfWriter().write_to_bytes(model), filetype='pdf') as pdf:
            text = 'OWNED ' + kind.upper()
            spans = [s for p in pdf for s in p.get_texttrace()
                     if text in ''.join(chr(c[0]) for c in s['chars'])]
            assert len(spans) == 1
            assert spans[0]['size'] == pytest.approx(expected or .5, abs=.011)
