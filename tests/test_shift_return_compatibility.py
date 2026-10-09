"""Word manual-break justification survives settings parsing and DOCX round trips."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
import pymupdf
import pytest

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.saving import PdfSaveOptions
from .test_pdf_justification import justified_model, WORDS


def source_docx(value=None, *, style="Code", trailing=False):
    document = Document()
    document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    section = document.sections[0]
    section.page_width, section.page_height = Pt(220), Pt(200)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Pt(20)
    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run("SHORT LINE\n" + ("" if trailing else WORDS))
    run.font.name, run.font.size = "Arial", Pt(12)
    if value is not None:
        setting = OxmlElement("w:doNotExpandShiftReturn")
        if value != "present":
            setting.set(qn("w:val"), value)
        document.settings.element.find(qn("w:compat")).append(setting)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize("value,disabled", [
    (None, False), ("present", True), ("true", True), ("on", True), ("1", True),
    ("false", False), ("off", False), ("0", False)])
@pytest.mark.parametrize("style", ["Normal", "Code"])
@pytest.mark.parametrize("shaping", [False, True])
def test_public_docx_settings_render_and_round_trip(value, disabled, style, shaping):
    source = aw.Document(BytesIO(source_docx(value, style=style)))
    assert source.light_document_model.do_not_expand_shift_return is disabled
    snapshot = source.light_document_model.model_dump()
    options = PdfSaveOptions()
    options.text_shaping = shaping

    def check(document):
        with pymupdf.open(stream=document.to_bytes(options), filetype="pdf") as pdf:
            words = [word for page in pdf for word in page.get_text("words")]
            end = next(word[2] for word in words if word[4] == "LINE")
            if disabled:
                assert end < 120
            else:
                assert end == pytest.approx(200 - 72 / 25.4, abs=0.05)
            assert next(word[2] for word in words if word[4] == "FIVE") == pytest.approx(200 - 72 / 25.4, abs=0.05)
            assert words[-1][2] < 190
            assert [word[4] for word in words] == ["SHORT", "LINE"] + WORDS.split()

    check(source)
    raw = source.to_bytes(aw.saving.OoxmlSaveOptions())
    with ZipFile(BytesIO(raw)) as archive:
        settings = ET.fromstring(archive.read("word/settings.xml"))
        flags = settings.findall(qn("w:compat") + "/" + qn("w:doNotExpandShiftReturn"))
        assert len(flags) == int(disabled)
    loaded = aw.Document(BytesIO(raw))
    assert loaded.light_document_model.do_not_expand_shift_return is disabled
    check(loaded)
    assert source.light_document_model.model_dump() == snapshot


@pytest.mark.parametrize("disabled", [False, True])
@pytest.mark.parametrize("shaping", [False, True])
def test_trailing_manual_break_still_follows_setting(disabled, shaping):
    source = aw.Document(BytesIO(source_docx("on" if disabled else None, trailing=True)))
    options = PdfSaveOptions()
    options.text_shaping = shaping
    with pymupdf.open(stream=source.to_bytes(options), filetype="pdf") as pdf:
        end = next(word[2] for word in pdf[0].get_text("words") if word[4] == "LINE")
        if disabled:
            assert end < 120
        else:
            assert end == pytest.approx(200 - 72 / 25.4, abs=0.05)


@pytest.mark.parametrize("disabled", [False, True])
def test_ldm_json_preserves_manual_break_setting(disabled):
    model, _ = justified_model(text="SHORT LINE\n" + WORDS)
    model.do_not_expand_shift_return = disabled
    copy = ldm.Document.model_validate_json(model.model_dump_json(by_alias=True))
    assert copy.do_not_expand_shift_return is disabled
    assert LdmPdfWriter().write_to_bytes(copy).startswith(b"%PDF")


def test_absent_settings_part_uses_word_default():
    raw = source_docx()
    output = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(output, "w") as target:
        for name in source.namelist():
            if name != "word/settings.xml":
                target.writestr(name, source.read(name))
    assert aw.Document(BytesIO(output.getvalue())).light_document_model.do_not_expand_shift_return is False


def test_off_disables_inherited_run_and_paragraph_flags():
    document = Document()
    document.styles["Normal"].font.bold = True
    document.styles["Normal"].paragraph_format.keep_with_next = True
    paragraph = document.add_paragraph()
    run = paragraph.add_run("OFF")
    for parent, name in [(run._r.get_or_add_rPr(), "b"),
                         (paragraph._p.get_or_add_pPr(), "keepNext")]:
        flag = OxmlElement("w:" + name)
        flag.set(qn("w:val"), "off")
        parent.append(flag)
    raw = BytesIO()
    document.save(raw)
    loaded = aw.Document(BytesIO(raw.getvalue())).light_document_model.sections[0].body.paragraphs[0]
    assert loaded.runs[0].font.bold is False
    assert loaded.paragraph_format.keep_with_next is False
