"""Save switches must not infer output behavior from arbitrary truthiness."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document as NativeDocument
from pypdf import PdfReader

import dockiln as dk

SAVE_FIELDS = [
    (kind, name) for kind in (dk.saving.PdfSaveOptions, dk.saving.OoxmlSaveOptions)
    for name, value in vars(kind()).items() if not name.startswith("_") and isinstance(value, bool)
]


@pytest.mark.parametrize("kind,name", SAVE_FIELDS)
@pytest.mark.parametrize("value", ["false", None, 0, 1, []])
def test_invalid_save_switch_preserves_value_and_explicit_requests(kind, name, value):
    options = kind()
    original = getattr(options, name)
    requested = getattr(options, "_explicit_options", set()).copy()
    with pytest.raises(ValueError, match=name):
        setattr(options, name, value)
    assert getattr(options, name) is original
    assert getattr(options, "_explicit_options", set()) == requested




@pytest.mark.parametrize("kind,name", SAVE_FIELDS)
@pytest.mark.parametrize("value", [False, True])
def test_boolean_assignment_and_copy_keep_public_storage(kind, name, value):
    from copy import deepcopy

    options = kind()
    setattr(options, name, value)
    assert vars(options)[name] is value
    assert getattr(deepcopy(options), name) is value
    if isinstance(options, dk.saving.PdfSaveOptions):
        assert name in options._explicit_options


def source():
    native = NativeDocument()
    p = native.add_paragraph()
    p.add_run("Underlined text").underline = True
    data = BytesIO()
    native.save(data)
    data.seek(0)
    return dk.Document(data)


@pytest.mark.parametrize("value", [False, True])
def test_real_docx_switch_effects_remain_independently_observable(value):
    doc = source()
    pdf = dk.saving.PdfSaveOptions()
    pdf.export_document_structure = value
    reader = PdfReader(BytesIO(doc.to_bytes(pdf)), strict=True)
    assert ("/StructTreeRoot" in reader.trailer["/Root"]) is value
    assert reader.pages[0].extract_text().strip() == "Underlined text"

    ooxml = dk.saving.OoxmlSaveOptions()
    ooxml.pretty_format = value
    with ZipFile(BytesIO(doc.to_bytes(ooxml))) as package:
        xml = package.read("word/document.xml")
        assert (b"<w:body>\r\n" in xml) is value
    reread = NativeDocument(BytesIO(doc.to_bytes(ooxml)))
    assert reread.paragraphs[0].text == "Underlined text"
    assert reread.paragraphs[0].runs[0].underline is True


