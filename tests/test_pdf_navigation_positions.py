"""Navigation destinations must agree with independently extracted content positions."""

from io import BytesIO

import pymupdf
from pypdf import PdfReader
import pytest

from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.saving import PdfSaveOptions


def navigation_model(style, bands=False, columns=1, long=False):
    prefix = ldm.Paragraph(children=[ldm.Run(text="[JUMP](#Target)", is_hyperlink=True,
                                           font=ldm.Font(size=10))],
                           paragraph_format=ldm.ParagraphFormat(space_after=110))
    target = ldm.Paragraph(children=[ldm.BookmarkStart(name="Target"),
                                    ldm.Run(text="TARGET" + ("\n" + "\n".join(f"BODY{i}" for i in range(60)) if long else ""),
                                            font=ldm.Font(size=10 if long else 30))],
                           paragraph_format=ldm.ParagraphFormat(is_heading=style == "heading",
                               outline_level=0, style_name=style, space_before=20))
    section = ldm.Section(page_setup=ldm.PageSetup(page_width=300, page_height=200,
        top_margin=20, bottom_margin=20, left_margin=20, right_margin=20,
        text_columns=ldm.TextColumns(count=columns, spacing=15)),
        body=ldm.Body(children=[prefix, target]))
    if bands:
        section.headers_footers = [ldm.HeaderFooter(header_footer_type=kind, children=[
            ldm.Paragraph(children=[ldm.Run(text=text, font=ldm.Font(size=8))])])
            for kind, text in ((0, "HEADER"), (1, "FOOTER"))]
    return ldm.Document(sections=[section])


def assert_navigation_matches_content(raw):
    parsed = PdfReader(BytesIO(raw))
    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        matches = [(i, word) for i, page in enumerate(pdf) for word in page.get_text("words")
                   if word[4] == "TARGET"]
        assert len(matches) == 1
        target_page, word = matches[0]
        link = pdf[0].get_links()[0]
        assert link["page"] == target_page
        # Destinations use line tops; the DOCX heading's leading shifts its glyph bounds by 3.15 pt.
        assert abs(link["to"].y - word[1]) < 5
        assert 0 <= word[0] - link["to"].x < 5
        assert len(parsed.outline) == 1
        destination = parsed.outline[0]
        assert parsed.get_destination_page_number(destination) == target_page
        assert abs(float(parsed.pages[target_page].mediabox.height) - float(destination.top) - word[1]) < 5
        assert 0 <= word[0] - float(destination.left) < 5
        return target_page, word, len(pdf)


@pytest.mark.parametrize("style", ["heading", "Normal", "Quote", "Code"])
@pytest.mark.parametrize("bands", [False, True])
def test_targets_follow_first_draw_after_spacing_and_page_break(style, bands):
    options = PdfSaveOptions()
    options.outline_options.headings_outline_levels = 6
    raw = LdmPdfWriter(options).write_to_bytes(navigation_model(style, bands=bands))
    page, _, _ = assert_navigation_matches_content(raw)
    assert page == (1 if bands or style in ("heading", "Normal") else 0)


def test_targets_follow_column_position():
    options = PdfSaveOptions()
    options.outline_options.headings_outline_levels = 6
    raw = LdmPdfWriter(options).write_to_bytes(navigation_model("heading", columns=2))
    page, word, _ = assert_navigation_matches_content(raw)
    assert page == 0 and word[0] > 150


def test_multipage_bookmark_stays_at_first_line():
    options = PdfSaveOptions()
    raw = LdmPdfWriter(options).write_to_bytes(navigation_model("Normal", long=True))
    _, _, pages = assert_navigation_matches_content(raw)
    assert pages > 2


def test_public_docx_conversion_navigation_matches_heading(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    import aspose.words_foss as aw

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(300), Pt(200)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(20)
    document.styles['Normal'].font.name = 'DocumentSansSC'
    document.styles['Normal'].font.size = Pt(10)
    document.styles['Heading 1'].font.name = 'DocumentSansSC'
    prefix = document.add_paragraph()
    prefix.paragraph_format.space_after = Pt(110)
    link = OxmlElement('w:hyperlink')
    link.set(qn('w:anchor'), 'Target')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = 'JUMP'
    run.append(text)
    link.append(run)
    prefix._p.append(link)
    heading = document.add_heading('TARGET', level=1)
    heading.runs[0].font.size = Pt(30)
    heading.paragraph_format.space_before = Pt(20)
    heading.paragraph_format.keep_with_next = False
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), '0')
    start.set(qn('w:name'), 'Target')
    heading._p.insert(1, start)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), '0')
    heading._p.append(end)
    source = tmp_path / 'navigation.docx'
    document.save(source)
    options = PdfSaveOptions()
    options.outline_options.headings_outline_levels = 6
    page, _, _ = assert_navigation_matches_content(aw.Document(source).to_bytes(options))
    assert page == 1
