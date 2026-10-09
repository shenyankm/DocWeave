"""Check outline hierarchy using an independent PDF parser."""

from io import BytesIO
import warnings

from pypdf import PdfReader
import pytest

from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter, PdfUnsupportedOptionWarning
from aspose.words_foss.saving import PdfSaveOptions


def model(levels, kind):
    children = []
    for index, level in enumerate(levels):
        title = f"Entry{index}"
        children.append(ldm.Paragraph(
            paragraph_format=ldm.ParagraphFormat(is_heading=kind == "heading", outline_level=level - 1),
            children=([ldm.BookmarkStart(name=title)] if kind == "bookmark" else []) + [ldm.Run(text=title)],
        ))
    return ldm.Document(sections=[ldm.Section(body=ldm.Body(children=children))])


def entries(reader):
    result = []
    def visit(items, depth):
        for item in items:
            if isinstance(item, list):
                visit(item, depth + 1)
            else:
                result.append((depth, item.title))
                assert reader.get_destination_page_number(item) == 0
    visit(reader.outline, 1)
    return result


@pytest.mark.parametrize("kind", ["heading", "bookmark"])
@pytest.mark.parametrize("fill", [False, True])
@pytest.mark.parametrize("levels,compact,filled", [
    ([1, 5, 5], [(1, "Entry0"), (2, "Entry1"), (2, "Entry2")],
     [(1, "Entry0"), (2, ""), (3, ""), (4, ""), (5, "Entry1"), (5, "Entry2")]),
    ([3, 5, 2, 4, 1, 3], [(1, "Entry0"), (2, "Entry1"), (1, "Entry2"), (2, "Entry3"), (1, "Entry4"), (2, "Entry5")],
     [(1, ""), (2, ""), (3, "Entry0"), (4, ""), (5, "Entry1"), (2, "Entry2"), (3, ""), (4, "Entry3"), (1, "Entry4"), (2, ""), (3, "Entry5")]),
])
def test_missing_levels_follow_option(kind, fill, levels, compact, filled):
    options = PdfSaveOptions()
    options.outline_options.headings_outline_levels = 6
    options.outline_options.create_missing_outline_levels = fill
    options.outline_options.bookmarks_outline_levels = {f"Entry{i}": level for i, level in enumerate(levels)}
    reader = PdfReader(BytesIO(LdmPdfWriter(options).write_to_bytes(model(levels, kind))))
    assert entries(reader) == (filled if fill else compact)
    assert all(reader.pages[0].extract_text().count(f"Entry{i}") == 1 for i in range(len(levels)))


def test_reusing_writer_resets_outline_path():
    options = PdfSaveOptions()
    options.outline_options.headings_outline_levels = 6
    writer = LdmPdfWriter(options)
    writer.write_to_bytes(model([1, 3, 5], "heading"))
    reader = PdfReader(BytesIO(writer.write_to_bytes(model([4], "heading"))))
    assert entries(reader) == [(1, "Entry0")]


def test_heading_and_bookmark_share_outline_path():
    source = model([1, 5, 3, 4, 1], "heading")
    for index in (1, 3):
        paragraph = source.sections[0].body.children[index]
        paragraph.paragraph_format.is_heading = False
        paragraph._children.insert(0, ldm.BookmarkStart(name=f"Entry{index}"))
    snapshot = source.model_dump()
    options = PdfSaveOptions()
    options.outline_options.headings_outline_levels = 6
    options.outline_options.bookmarks_outline_levels = {"Entry1": 5, "Entry3": 4}
    reader = PdfReader(BytesIO(LdmPdfWriter(options).write_to_bytes(source)))
    assert entries(reader) == [(1, "Entry0"), (2, "Entry1"), (2, "Entry2"), (3, "Entry3"), (1, "Entry4")]
    assert source.model_dump() == snapshot


def test_public_docx_save_compacts_heading_levels(tmp_path):
    from docx import Document
    import aspose.words_foss as aw

    source = tmp_path / "headings.docx"
    document = Document()
    for index, level in enumerate((1, 5, 5)):
        document.add_heading(f"Entry{index}", level=level)
    document.save(source)
    options = PdfSaveOptions()
    options.outline_options.headings_outline_levels = 5
    reader = PdfReader(BytesIO(aw.Document(source).to_bytes(options)))
    assert entries(reader) == [(1, "Entry0"), (2, "Entry1"), (2, "Entry2")]


def test_compacting_outline_preserves_page_drawing_and_reduces_output():
    outputs = []
    readers = []
    for fill in (False, True):
        options = PdfSaveOptions()
        options.outline_options.headings_outline_levels = 6
        options.outline_options.create_missing_outline_levels = fill
        output = LdmPdfWriter(options).write_to_bytes(model([1, 5, 5], "heading"))
        outputs.append(output)
        readers.append(PdfReader(BytesIO(output)))
    assert len(outputs[0]) < len(outputs[1])
    assert len(readers[0].pages) == len(readers[1].pages)
    assert readers[0].pages[0]["/Contents"].get_object().get_data() == readers[1].pages[0]["/Contents"].get_object().get_data()


@pytest.mark.parametrize("kind,levels", [("heading", [1, 2]), ("bookmark", [1, 2]), ("heading", [])])
def test_expansion_does_not_warn_for_any_document(kind, levels):
    options = PdfSaveOptions()
    options.outline_options.expanded_outline_levels = 2
    writer = LdmPdfWriter(options)
    for _ in range(2):
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            writer.write_to_bytes(model(levels, kind))
        expansion = [item for item in caught if issubclass(item.category, PdfUnsupportedOptionWarning)
                     and "expanded_outline_levels" in str(item.message)]
        assert not expansion


def outline_counts(reader):
    root = reader.trailer["/Root"]["/Outlines"]
    result = []
    def visit(item, depth):
        while item is not None:
            node = item.get_object()
            result.append((depth, str(node["/Title"]), node.get("/Count", 0)))
            if "/First" in node:
                visit(node["/First"], depth + 1)
            item = node.get("/Next")
    visit(root["/First"], 1)
    return root.get("/Count"), result


@pytest.mark.parametrize("kind", ["heading", "bookmark"])
@pytest.mark.parametrize("expanded,root_count,counts", [
    (0, None, [-2, -2, 0, 0, -1, 0, 0]),
    (1, 4, [2, -2, 0, 0, -1, 0, 0]),
    (2, 7, [5, 2, 0, 0, 1, 0, 0]),
    (9, 7, [5, 2, 0, 0, 1, 0, 0]),
])
def test_expansion_counts_visible_descendants(kind, expanded, root_count, counts):
    levels = [1, 2, 3, 3, 2, 3, 1]
    options = PdfSaveOptions()
    options.outline_options.headings_outline_levels = 6
    options.outline_options.bookmarks_outline_levels = {f"Entry{i}": level for i, level in enumerate(levels)}
    options.outline_options.expanded_outline_levels = expanded
    output = LdmPdfWriter(options).write_to_bytes(model(levels, kind))
    reader = PdfReader(BytesIO(output))
    actual_root, nodes = outline_counts(reader)
    assert actual_root == root_count
    assert nodes == [(level, f"Entry{i}", count) for i, (level, count) in enumerate(zip(levels, counts))]
    import pymupdf
    with pymupdf.open(stream=output, filetype="pdf") as independent:
        toc = independent.get_toc(simple=False)
        assert [item[1] for item in toc] == [f"Entry{i}" for i in range(len(levels))]
        for item, count in zip(toc, counts):
            assert item[3].get("collapse") == (count < 0 if count else None)


@pytest.mark.parametrize("value", [-1, 10, "2", 1.5, None])
def test_invalid_expansion_preserves_existing_output(tmp_path, value):
    options = PdfSaveOptions()
    options.outline_options.expanded_outline_levels = value
    target = tmp_path / "existing.pdf"
    target.write_bytes(b"old PDF")
    with pytest.raises(ValueError, match="expanded_outline_levels"):
        LdmPdfWriter(options).write(model([1], "heading"), target)
    assert target.read_bytes() == b"old PDF"


@pytest.mark.parametrize("fill,expected_root,expected_counts", [
    (False, 3, [(1, "Entry0", 2), (2, "Entry1", 0), (2, "Entry2", 0)]),
    (True, 3, [(1, "Entry0", 2), (2, "", 1), (3, "", -1), (4, "", -2), (5, "Entry1", 0), (5, "Entry2", 0)]),
])
def test_expansion_uses_pdf_depth_after_gap_handling(fill, expected_root, expected_counts):
    options = PdfSaveOptions()
    options.outline_options.headings_outline_levels = 6
    options.outline_options.create_missing_outline_levels = fill
    options.outline_options.expanded_outline_levels = 2
    reader = PdfReader(BytesIO(LdmPdfWriter(options).write_to_bytes(model([1, 5, 5], "heading"))))
    assert outline_counts(reader) == (expected_root, expected_counts)


def test_expansion_preserves_page_content_and_destinations():
    readers = []
    for expanded in (0, 1, 2, 9):
        options = PdfSaveOptions()
        options.outline_options.headings_outline_levels = 6
        options.outline_options.expanded_outline_levels = expanded
        readers.append(PdfReader(BytesIO(LdmPdfWriter(options).write_to_bytes(model([1, 2, 3], "heading")))))
    page_stream = readers[0].pages[0]["/Contents"].get_object().get_data()
    for reader in readers:
        assert reader.pages[0]["/Contents"].get_object().get_data() == page_stream
        assert entries(reader) == [(1, "Entry0"), (2, "Entry1"), (3, "Entry2")]
