"""Style saves preserve declarations while public Font getters stay effective."""
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss._flat_opc import decode
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.docx_writer.runs import render_rPr
from aspose.words_foss.saving import OoxmlSaveOptions

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def owned_document(location='body'):
    source = WordDocument()
    source.styles['Normal'].element.set(qn('w:default'), '0')
    first = source.styles.add_style('Owned Base', WD_STYLE_TYPE.PARAGRAPH)
    first.element.set(qn('w:default'), '1')
    first.font.size = Pt(16)
    first.font.name = 'Georgia'
    first.font.color.rgb = RGBColor.from_string('123456')
    last = source.styles.add_style('Owned Default', WD_STYLE_TYPE.PARAGRAPH)
    last.base_style = first
    last.element.set(qn('w:default'), '1')
    last.font.size = Pt(19)
    if location == 'body':
        paragraph = source.add_paragraph()
    elif location == 'table':
        paragraph = source.add_table(1, 1).cell(0, 0).paragraphs[0]
    else:
        paragraph = getattr(source.sections[0], location).paragraphs[0]
    paragraph.style = None
    paragraph.add_run('OWNED')
    stream = BytesIO()
    source.save(stream)
    return stream.getvalue()


def owned_run(model):
    return next(run for run in model.get_child_nodes(ldm.NodeType.RUN, True)
                if run.text == 'OWNED')


def style_properties(raw, name):
    with ZipFile(BytesIO(raw)) as package:
        styles = ET.fromstring(package.read('word/styles.xml'))
    return next(style for style in styles.findall(W + 'style')
                if style.find(W + 'name').get(W + 'val') == name).find(W + 'rPr')


@pytest.mark.parametrize('location', ['body', 'table', 'header', 'footer'])
@pytest.mark.parametrize('through_json', [False, True])
@pytest.mark.parametrize('format', [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC])
def test_sparse_character_default_cannot_override_paragraph_font_after_save(location, through_json, format):
    model = aw.Document(BytesIO(owned_document(location))).light_document_model
    if through_json:
        model = ldm.Document.model_validate_json(model.model_dump_json())
    before = model.model_dump_json()
    held = owned_run(model).font
    assert (held.size, held.name, held.color) == (19, 'Georgia', 'Color [A=255, R=18, G=52, B=86]')
    saved = LdmDocxWriter(OoxmlSaveOptions(format)).write_to_bytes(model)
    package = saved if format == aw.SaveFormat.DOCX else decode(saved)
    default = style_properties(package, 'Default Paragraph Font')
    assert default is None or len(default) == 0
    inherited = style_properties(package, 'Owned Default')
    assert inherited.find(W + 'rFonts') is None
    assert inherited.find(W + 'color') is None
    assert inherited.find(W + 'sz').get(W + 'val') == '38'
    assert model.model_dump_json() == before
    assert held is owned_run(model).font and held.size == 19
    restored = aw.Document(BytesIO(saved)).light_document_model
    assert (owned_run(restored).font.size, owned_run(restored).font.name,
            owned_run(restored).font.color) == (held.size, held.name, held.color)
    second = LdmDocxWriter().write_to_bytes(restored)
    default = style_properties(second, 'Default Paragraph Font')
    assert default is None or len(default) == 0


@pytest.mark.parametrize('field,value,tag,expected', [
    ('size', 0, 'sz', '0'), ('bold', False, 'b', '0'),
    ('italic', True, 'i', '1'), ('color', '123456', 'color', '123456'),
])
def test_bound_style_saves_explicit_local_value_without_mutating_getter(field, value, tag, expected):
    font = ldm.Font()
    setattr(font, field, value)
    font._name_resolver = lambda current: {
        'ascii': 'Inherited', 'hAnsi': 'Inherited', 'cs': 'Inherited', 'eastAsia': 'Inherited',
        'size': 24, 'bold': True, 'italic': False, 'color': 'ABCDEF',
    }
    before = font.model_dump_json()
    xml = ET.fromstring('<root xmlns:w="' + W[1:-1] + '">' + render_rPr(font, for_style=True) + '</root>')
    assert xml.find(W + 'rPr/' + W + tag).get(W + 'val') == expected
    assert xml.find(W + 'rPr/' + W + 'rFonts') is None
    assert font.model_dump_json() == before
    assert font._name_resolver is not None


@pytest.mark.parametrize('target', ['run', 'style'])
@pytest.mark.parametrize('through_json', [False, True])
@pytest.mark.parametrize('format', [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC])
def test_explicit_zero_size_survives_model_edits_and_two_cold_saves(target, through_json, format):
    model = aw.Document(BytesIO(owned_document())).light_document_model
    font = (owned_run(model).font if target == 'run' else
            next(s.font for s in model.styles if s.name == 'Owned Default'))
    font.size = 0
    assert font.size_explicit is True
    if through_json:
        model = ldm.Document.model_validate_json(model.model_dump_json())
    for _ in range(2):
        before = model.model_dump_json()
        raw = LdmDocxWriter(OoxmlSaveOptions(format)).write_to_bytes(model)
        assert model.model_dump_json() == before
        model = aw.Document(BytesIO(raw)).light_document_model
        assert owned_run(model).font.size == 0
