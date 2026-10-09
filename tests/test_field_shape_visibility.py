"""Field instructions must not emit images, text boxes, or occupy layout space."""

import base64
import warnings
from io import BytesIO

import fitz
import pytest
from docx import Document as NativeDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image
from pypdf import PdfReader

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.model.wrap_type import WrapType


def image_shape(color):
    stream = BytesIO()
    Image.new("RGB", (12, 12), color).save(stream, format="PNG")
    return ldm.Shape(has_image=True, is_inline=True, width=24, height=24,
                     image_data=ldm.ImageData(image_bytes=stream.getvalue()))


def document(location, shape_kind, instructions=True):
    secret, visible = image_shape("red"), image_shape("blue")
    if shape_kind == "floating":
        secret.wrap_type = WrapType.NONE
    elif shape_kind == "anchored":
        secret.is_inline = False
    elif shape_kind in ("positioned", "relative"):
        secret._is_positioned = True
        secret.relative_vertical_position = 2 if shape_kind == "relative" else 0
    elif shape_kind == "textbox":
        secret = ldm.Shape(text_box={"paragraphs": [ldm.Paragraph(children=[ldm.Run(text="SECRET_BOX")])]})
    children = [ldm.Run(text="VISIBLE_RESULT"), visible]
    if instructions:
        children = [ldm.FieldStart(), ldm.Run(text="SECRET_CODE"), secret,
                    ldm.FieldSeparator(), *children, ldm.FieldEnd()]
    para = ldm.Paragraph(children=children)
    doc = aw.Document(BytesIO(b"BODY"), aw.MarkdownLoadOptions())
    if location == "body":
        doc.sections[0].body.children = [para]
    elif location == "table":
        doc.sections[0].body.children = [ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[para])])])]
    else:
        doc.sections[0].headers_footers = [ldm.HeaderFooter(
            header_footer_type=0 if location == "header" else 1, children=[para])]
    return doc, secret, visible


def render(doc, options):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", aw.ContentLossWarning)
        return doc.to_bytes(options)


def test_instruction_image_does_not_change_page_break_classification():
    from aspose.words_foss.pdf_writer.text import is_pure_page_break

    para = ldm.Paragraph(children=[ldm.FieldStart(), image_shape("red"),
        ldm.FieldSeparator(), ldm.Run(text="\f"), ldm.FieldEnd()])
    assert is_pure_page_break(para)
    para._children.insert(-1, image_shape("blue"))
    assert not is_pure_page_break(para)


@pytest.mark.parametrize("format", ["md", "pdf"])
def test_instruction_only_table_paragraph_matches_empty_control(format):
    doc, _, _ = document("table", "inline")
    control, _, _ = document("table", "inline", instructions=False)
    paragraph = doc.sections[0].body.children[0].rows[0].cells[0].paragraphs[0]
    paragraph._children = [ldm.FieldStart(), image_shape("red"), ldm.Run(text="SECRET"), ldm.FieldEnd()]
    control.sections[0].body.children[0].rows[0].cells[0].paragraphs[0]._children = []
    raw, expected = render(doc, format), render(control, format)
    if format == "md":
        assert raw == expected
    else:
        with fitz.open(stream=raw, filetype="pdf") as actual, fitz.open(stream=expected, filetype="pdf") as clean:
            assert len(actual) == len(clean)
            for page, reference in zip(actual, clean):
                assert page.get_pixmap().samples == reference.get_pixmap().samples


@pytest.mark.parametrize("location", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("external", [False, True])
def test_markdown_exports_only_result_images(tmp_path, location, external):
    doc, secret, visible = document(location, "inline")
    before = doc.light_document_model.model_dump()
    options = aw.saving.MarkdownSaveOptions()
    if external:
        options.images_folder = str(tmp_path / "images")
        path = tmp_path / "result.md"
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", aw.ContentLossWarning)
            doc.save(path, options)
        result = path.read_text()
    else:
        result = render(doc, options).decode()
    assert "SECRET" not in result
    if external:
        images = list((tmp_path / "images").iterdir())
        assert len(images) == 1
        assert images[0].read_bytes() == visible.image_data.image_bytes
    else:
        assert base64.b64encode(secret.image_data.image_bytes).decode() not in result
        assert base64.b64encode(visible.image_data.image_bytes).decode() in result
    assert doc.light_document_model.model_dump() == before


@pytest.mark.parametrize("location", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("format", ["md", "pdf"])
def test_real_docx_field_images(location, format):
    native = NativeDocument()
    if location == "body":
        paragraph = native.add_paragraph()
    elif location == "table":
        paragraph = native.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    else:
        paragraph = getattr(native.sections[0], location).paragraphs[0]

    def marker(kind):
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), kind)
        paragraph.add_run()._r.append(node)

    secret, visible = image_shape("red"), image_shape("blue")
    marker("begin")
    code = OxmlElement("w:instrText")
    code.text = "SECRET_CODE"
    paragraph.add_run()._r.append(code)
    paragraph.add_run().add_picture(BytesIO(secret.image_data.image_bytes), width=Pt(24))
    marker("separate")
    paragraph.add_run("VISIBLE_RESULT")
    paragraph.add_run().add_picture(BytesIO(visible.image_data.image_bytes), width=Pt(24))
    marker("end")
    source = BytesIO()
    native.save(source)
    doc = aw.Document(BytesIO(source.getvalue()))
    result = render(doc, format)
    if format == "md":
        assert base64.b64encode(secret.image_data.image_bytes) not in result
        assert base64.b64encode(visible.image_data.image_bytes) in result
        assert b"SECRET_CODE" not in result
    else:
        with fitz.open(stream=result, filetype="pdf") as pdf:
            assert "SECRET_CODE" not in "".join(page.get_text() for page in pdf)
            assert "VISIBLE_RESULT" in "".join(page.get_text() for page in pdf)
            images = [pdf.extract_image(item[0])["image"] for page in pdf for item in page.get_images()]
            assert len(images) == 1
            with Image.open(BytesIO(images[0])) as image:
                assert image.convert("RGB").getpixel((0, 0)) == (0, 0, 255)


@pytest.mark.parametrize("location", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("shape_kind", ["inline", "floating", "anchored", "positioned", "relative", "textbox"])
def test_pdf_field_instruction_shapes_match_visible_control(location, shape_kind):
    doc, _, _ = document(location, shape_kind)
    control, _, _ = document(location, shape_kind, instructions=False)
    before = doc.light_document_model.model_dump()
    raw, expected = render(doc, "pdf"), render(control, "pdf")
    text = "".join(page.extract_text() for page in PdfReader(BytesIO(raw)).pages)
    assert "SECRET" not in text
    assert "VISIBLE_RESULT" in text
    with fitz.open(stream=raw, filetype="pdf") as actual, fitz.open(stream=expected, filetype="pdf") as clean:
        assert len(actual) == len(clean)
        for page, reference in zip(actual, clean):
            assert page.get_text("words") == reference.get_text("words")
            assert page.get_pixmap().samples == reference.get_pixmap().samples
            assert len(page.get_images()) == len(reference.get_images()) == 1
    assert doc.light_document_model.model_dump() == before
