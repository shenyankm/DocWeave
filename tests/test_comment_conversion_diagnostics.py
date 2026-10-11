"""Supported comment conversion preserves parts; remaining losses stay diagnostic."""
from io import BytesIO
import warnings
from zipfile import ZipFile

from docx import Document as NativeDocument
import pytest

import dockiln as dk
from dockiln._flat_opc import encode


def commented_document(reference):
    document = NativeDocument()
    run = document.add_paragraph('PRIVATE BODY').runs[0]
    document.add_comment(run, text='PRIVATE COMMENT', author='PRIVATE AUTHOR')
    if not reference:
        for name in ('commentReference', 'commentRangeStart', 'commentRangeEnd'):
            for node in document.element.xpath('.//w:' + name):
                node.getparent().remove(node)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize('reference', [False, True])
@pytest.mark.parametrize('flat', [False, True])
def test_classic_comment_parts_are_preserved_even_when_warnings_are_ignored(reference, flat):
    raw = commented_document(reference)
    source = encode(raw) if flat else raw
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', dk.ContentLossWarning)
        document = dk.Document(BytesIO(source))
        output = document.to_bytes(dk.SaveFormat.DOCX)
    observed = [item for item in document.diagnostics if item.code == 'load.comments_omitted']
    assert observed == []
    assert all(text not in repr(observed) for text in ('PRIVATE BODY', 'PRIVATE COMMENT', 'PRIVATE AUTHOR'))
    with ZipFile(BytesIO(output)) as package:
        assert 'word/comments.xml' in package.namelist()
    assert 'PRIVATE BODY' in document.get_text()


@pytest.mark.parametrize('reference', [False, True])
def test_strict_classic_comment_conversion_keeps_editable_package(reference):
    raw = commented_document(reference)
    editable = dk.DocxDocument(BytesIO(raw))
    with warnings.catch_warnings():
        warnings.simplefilter('error', dk.ContentLossWarning)
        model = editable.to_light_document()
        assert any(story.kind == 'comment' for story in model.source_stories)
    assert editable.to_bytes() == raw


def test_document_without_comments_has_no_comment_loss_diagnostic():
    document = NativeDocument()
    document.add_paragraph('OWNED BODY')
    output = BytesIO()
    document.save(output)
    loaded = dk.Document(BytesIO(output.getvalue()))
    assert not any(item.code == 'load.comments_omitted' for item in loaded.diagnostics)
