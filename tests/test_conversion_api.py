"""Public in-memory output, structured diagnostics and format boundaries."""

from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict
from io import BytesIO
import json
import warnings
from zipfile import ZipFile

import pytest
from pypdf import PdfReader

from ._word_fixtures import word_document

import dockiln as dk
from dockiln import light_document_model as ldm
from dockiln.writers.docx import DocxWriterLossyWarning, LdmDocxWriter
from dockiln.diagnostics import collect_diagnostics, warn
from dockiln.writers.pdf import PdfMissingGlyphWarning
from dockiln.readers.rtf import RtfFileReader

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def package(body):
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("word/document.xml", f'<w:document xmlns:w="{W}" xmlns:r="{R}"><w:body>{body}</w:body></w:document>')
    return stream.getvalue()


def document(text="中文 **粗体** [链接](https://example.com)"):
    return word_document(text)


@pytest.mark.parametrize("mode", ["path", "stream", "bytes", "reader_path", "reader_stream", "reader_bytes"])
def test_standard_text_rtf_has_actionable_error(tmp_path, mode):
    raw = b"{\\rtf1\\ansi ordinary text}"
    path = tmp_path / "standard.rtf"
    path.write_bytes(raw)
    reader = RtfFileReader()
    with pytest.raises(ValueError, match="Standard text RTF.*libreoffice"):
        if mode == "path":
            dk.Document(path)
        elif mode == "stream":
            dk.Document(BytesIO(raw))
        elif mode == "bytes":
            with pytest.warns(DeprecationWarning):
                dk.Document(data=raw)
        else:
            getattr(reader, "load_file" if mode == "reader_path" else "load_" + mode.removeprefix("reader_"))(
                path if mode == "reader_path" else BytesIO(raw) if mode == "reader_stream" else raw)


@pytest.mark.parametrize("tag", ["fldSimple", "ins", "del", "sdt", "oMath"])
def test_load_loss_is_recorded_even_when_warning_is_ignored(tag):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", dk.loading.DocumentLoadWarning)
        doc = dk.Document(BytesIO(package(f"<w:p><w:{tag}/></w:p>")))
    assert len(doc.diagnostics) == 1
    item = asdict(doc.diagnostics[0])
    assert item["code"] == "load.content_loss" and tag in item["message"]
    assert item["location"] and item["severity"] == "warning"


@pytest.mark.parametrize("kind", ["headerReference", "footerReference"])
@pytest.mark.parametrize("variant", ["different_sections", "first", "even"])
def test_missing_header_footer_resource_is_reported(kind, variant):
    if variant == "different_sections":
        body = (f'<w:p><w:pPr><w:sectPr><w:{kind} w:type="default" r:id="r1"/></w:sectPr></w:pPr></w:p>'
                f'<w:sectPr><w:{kind} w:type="default" r:id="r2"/></w:sectPr>')
    else:
        body = f'<w:sectPr><w:{kind} w:type="{variant}" r:id="r1"/></w:sectPr>'
    with pytest.warns(dk.loading.DocumentLoadWarning, match="Missing header/footer resource"):
        doc = dk.Document(BytesIO(package(body)))
    assert doc.diagnostics[0].code == "load.missing_header_footer"


def test_distinct_model_headers_are_saved_and_unknown_nodes_warn():
    doc = document()
    doc.sections[0].headers_footers = [ldm.HeaderFooter(header_footer_type=0, children=[ldm.Paragraph(children=[ldm.Run(text="A")])])]
    doc.sections.append(ldm.Section(headers_footers=[ldm.HeaderFooter(header_footer_type=0, children=[ldm.Paragraph(children=[ldm.Run(text="B")])])]))
    cold = dk.Document(BytesIO(doc.to_bytes(dk.SaveFormat.DOCX))).light_document_model
    assert [[p.text for h in s.headers_footers for p in h.paragraphs] for s in cold.sections] == [["A"], ["B"]]
    assert not any(d.code == "docx.header_footer_flattened" for d in doc.diagnostics)
    doc.sections[0].body.children.append(ldm.UnknownNode(_type="not_supported"))
    with pytest.warns(DocxWriterLossyWarning):
        doc.to_bytes(dk.SaveFormat.DOCX)
    assert any(item.code == "docx.unknown_node" for item in doc.diagnostics)


def test_complex_fields_are_retained_not_misreported():
    model = ldm.Document(sections=[ldm.Section(body=ldm.Body(children=[ldm.Paragraph(children=[
        ldm.FieldStart(), ldm.Run(text=" DATE "), ldm.FieldSeparator(),
        ldm.Run(text="2026-01-01"), ldm.FieldEnd(),
    ])]))])
    raw = LdmDocxWriter().write_to_bytes(model)
    with ZipFile(BytesIO(raw)) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "instrText" in xml and " DATE " in xml and "2026-01-01" in xml
    assert all(f'w:fldCharType="{part}"' in xml for part in ("begin", "separate", "end"))
    doc = dk.Document(BytesIO(raw))
    assert not doc.diagnostics
    assert "instrText" in ZipFile(BytesIO(doc.to_bytes("docx"))).read("word/document.xml").decode()


@pytest.mark.parametrize("fmt", [dk.SaveFormat.DOCX, dk.SaveFormat.PDF])
def test_memory_output_matches_path_content(tmp_path, fmt):
    doc = document()
    raw = doc.to_bytes(fmt)
    output = tmp_path / f"report.{fmt.name.lower()}"
    doc.save(output, fmt)
    if fmt == dk.SaveFormat.DOCX:
        with ZipFile(BytesIO(raw)) as memory, ZipFile(output) as disk:
            assert {n: memory.read(n) for n in memory.namelist()} == {n: disk.read(n) for n in disk.namelist()}
        assert "中文" in dk.Document(BytesIO(raw)).get_text()
    elif fmt == dk.SaveFormat.PDF:
        assert "中文" in "".join(p.extract_text() for p in PdfReader(BytesIO(raw)).pages)
        assert "".join(p.extract_text() for p in PdfReader(BytesIO(raw)).pages) == "".join(p.extract_text() for p in PdfReader(output).pages)
    else:
        assert raw == output.read_bytes()
    assert sorted(p.name for p in tmp_path.iterdir()) == [output.name]


def test_memory_options_and_validation(tmp_path):
    for fmt in (None, "unknown", "text", "md", dk.SaveFormat.DOC):
        with pytest.raises(ValueError, match="Unsupported"):
            document().to_bytes(fmt)
    assert not list(tmp_path.iterdir())


def test_error_filtered_warning_is_recorded_and_output_preserved(tmp_path):
    doc = document("missing " + chr(0x1FAE0))
    output = tmp_path / "report.pdf"
    output.write_bytes(b"old")
    with warnings.catch_warnings():
        warnings.simplefilter("error", PdfMissingGlyphWarning)
        with pytest.raises(PdfMissingGlyphWarning):
            doc.save(output)
    assert output.read_bytes() == b"old"
    assert doc.diagnostics[-1].code == "pdf.missing_glyph"


def test_diagnostic_collection_is_context_local():
    def collect(number):
        target = []
        with collect_diagnostics(target), warnings.catch_warnings():
            warnings.simplefilter("ignore", dk.loading.DocumentLoadWarning)
            warn(str(number), dk.loading.DocumentLoadWarning)
        return [item.message for item in target]
    with ThreadPoolExecutor(max_workers=4) as pool:
        assert list(pool.map(collect, range(12))) == [[str(n)] for n in range(12)]


def test_structure_export_marks_links_and_excludes_hidden_text():
    doc = document("公开链接")
    from dockiln._links import format_link
    run = doc.sections[0].body.children[0].runs[0]
    run.text = format_link("公开链接", "https://example.com")
    run.is_hyperlink = True
    doc.sections[0].body.children[0]._children.append(ldm.Run(text="hidden", font=ldm.Font(hidden=True)))
    block = doc.to_dict()["blocks"][0]
    assert block["text"] == "公开链接"
    assert block["runs"][0]["link"] == "https://example.com"
    assert "hidden" not in json.dumps(block)


def test_structure_export_keeps_body_order_and_locations():
    from docx import Document
    source = Document()
    source.add_heading("标题", 1)
    source.add_paragraph("正文")
    source.add_paragraph()
    table = source.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("A", "B"), ("1", "2"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    source.add_paragraph("结尾")
    stream = BytesIO()
    source.save(stream)
    doc = dk.Document(BytesIO(stream.getvalue()))
    result = json.loads(json.dumps(doc.to_dict(), ensure_ascii=False))
    assert result["schema_version"] == 1
    assert [block["type"] for block in result["blocks"]] == ["paragraph", "paragraph", "paragraph", "table", "paragraph"]
    assert result["blocks"][0]["heading_level"] == 1
    assert result["blocks"][3]["rows"][1]["cells"][0]["paragraphs"][0]["text"] == "1"
    assert result["blocks"][-1]["location"] == "sections[0].body.children[4]"
