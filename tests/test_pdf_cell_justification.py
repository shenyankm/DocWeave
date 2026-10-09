"""Cell justification uses the measured grid width and retains rich-text geometry."""

from io import BytesIO

import pymupdf
from pypdf import PdfReader
import pytest

from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.saving import PdfSaveOptions
from .test_pdf_justification import WORDS, lines


def cell_model(*, text=WORDS, mixed=False, disabled=False, wrap=True, listed=False, merged=False):
    runs = ([ldm.Run(text=word + (" " if index < 9 else ""),
                    font=ldm.Font(size=12 if index % 2 else 16, bold=index % 2 == 0))
             for index, word in enumerate(text.split())] if mixed else
            [ldm.Run(text=text, font=ldm.Font(size=12))])
    paragraph = ldm.Paragraph(children=runs, paragraph_format=ldm.ParagraphFormat(
        alignment=3, left_indent=24 if listed else 0, first_line_indent=-12 if listed else 0,
        right_indent=6 if listed else 0))
    if listed:
        paragraph.list_format = ldm.ListFormat(is_list_item=True, list_id=1)
        paragraph.list_label = ldm.ListLabel(label_string="1.")
    cell = ldm.Cell(paragraphs=[paragraph], cell_format=ldm.CellFormat(
        left_padding=8, right_padding=8, top_padding=4, bottom_padding=4,
        grid_span=2 if merged else 1, wrap_text=wrap))
    table = ldm.Table(rows=[ldm.Row(cells=[cell])])
    model = ldm.Document(do_not_expand_shift_return=disabled, sections=[ldm.Section(
        page_setup=ldm.PageSetup(page_width=220, page_height=180,
            left_margin=20, right_margin=20, top_margin=20, bottom_margin=20),
        body=ldm.Body(children=[table]))])
    return model, paragraph


def render(model, shaping=False, tagged=False):
    options = PdfSaveOptions()
    options.text_shaping = shaping
    options.export_document_structure = tagged
    return LdmPdfWriter(options).write_to_bytes(model)


@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("first_indent", [-12, 12])
def test_cell_paragraph_indents_apply_to_first_and_continuation_rows(shaping, first_indent):
    model, paragraph = cell_model()
    paragraph.paragraph_format.left_indent = 24
    paragraph.paragraph_format.right_indent = 6
    paragraph.paragraph_format.first_line_indent = first_indent
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        words = pdf[0].get_text("words")
        rows = {}
        for word in words:
            rows.setdefault(round(word[1], 2), []).append(word)
        rows = list(rows.values())
        assert len(rows) > 2
        assert rows[0][0][0] == pytest.approx(52 + first_indent, abs=0.05)
        assert all(row[0][0] == pytest.approx(52, abs=0.05) for row in rows[1:])
        assert all(row[-1][2] == pytest.approx(186, abs=0.05) for row in rows[:-1])
        assert [word[4] for word in words] == WORDS.split()


def test_cell_paragraph_indents_reject_exhausted_first_row():
    model, paragraph = cell_model()
    paragraph.paragraph_format.first_line_indent = 164
    with pytest.raises(ValueError, match="cell paragraph indents"):
        render(model)


@pytest.mark.parametrize("shaping", [False, True])
def test_cell_picture_uses_first_indent_and_text_continues_at_body_indent(shaping):
    from PIL import Image

    image = BytesIO()
    Image.new("RGB", (16, 8), "red").save(image, format="PNG")
    model, paragraph = cell_model(text="AFTER IMAGE")
    paragraph.paragraph_format.left_indent = 24
    paragraph.paragraph_format.right_indent = 6
    paragraph.paragraph_format.first_line_indent = 12
    paragraph._children.insert(0, ldm.Shape(width=240, height=120, has_image=True, image_data=ldm.ImageData(
        image_bytes=image.getvalue())))
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        picture = pdf[0].get_image_rects(pdf[0].get_images()[0][0])[0]
        assert picture.x0 == pytest.approx(64, abs=0.05)
        assert picture.x1 == pytest.approx(186, abs=0.05)
        assert lines(pdf[0])[0]["bbox"][0] == pytest.approx(52, abs=0.05)


@pytest.mark.parametrize("shaping", [False, True])
def test_nested_cell_justifies_inside_both_padding_layers(shaping):
    model, _ = cell_model()
    table = model.sections[0].body.children[0]
    inner = table.rows[0].cells[0]
    table.rows[0].cells[0] = ldm.Cell(children=[ldm.Table(rows=[ldm.Row(cells=[inner])])],
        cell_format=ldm.CellFormat(left_padding=8, right_padding=8))
    raw = render(model, shaping, tagged=True)
    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        words = pdf[0].get_text("words")
        assert words[0][0] == pytest.approx(36, abs=0.05)
        assert [word[4] for word in words] == WORDS.split()
        assert max(word[2] for word in words) == pytest.approx(184, abs=0.05)
    assert PdfReader(BytesIO(raw)).trailer["/Root"]["/StructTreeRoot"]


@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("alignment", [1, 2])
def test_center_and_right_cell_alignment_use_paragraph_indents(shaping, alignment):
    model, paragraph = cell_model(text="SHORT LINE")
    paragraph.paragraph_format.alignment = alignment
    paragraph.paragraph_format.left_indent = 24
    paragraph.paragraph_format.right_indent = 6
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        words = pdf[0].get_text("words")
        start, end = words[0][0], words[-1][2]
        if alignment == 1:
            assert (start + end) / 2 == pytest.approx((52 + 186) / 2, abs=0.05)
        else:
            assert end == pytest.approx(186, abs=0.05)


@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("mixed", [False, True])
@pytest.mark.parametrize("merged", [False, True])
def test_cell_soft_rows_reach_inner_margin_and_final_row_stays_natural(shaping, mixed, merged):
    model, _ = cell_model(mixed=mixed, merged=merged)
    snapshot = model.model_dump()
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        rows = lines(pdf[0])
        assert len(rows) > 1
        for row in rows[:-1]:
            assert row["bbox"][0] == pytest.approx(28, abs=0.05)
            assert row["bbox"][2] == pytest.approx(192, abs=0.05)
        assert rows[-1]["bbox"][2] < 185
        assert [word[4] for page in pdf for word in page.get_text("words")] == WORDS.split()
    assert model.model_dump() == snapshot


@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("disabled", [False, True])
@pytest.mark.parametrize("wrap", [False, True])
def test_cell_manual_break_respects_compatibility_even_without_auto_wrap(shaping, disabled, wrap):
    model, _ = cell_model(text="SHORT LINE\nLAST", disabled=disabled, wrap=wrap)
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        words = pdf[0].get_text("words")
        assert [word[4] for word in words] == ["SHORT", "LINE", "LAST"]
        short, last = words[1], words[-1]
        if disabled:
            assert short[2] < 120
        else:
            assert short[2] == pytest.approx(192, abs=0.05)
        assert last[1] > short[1]
        assert last[2] < 100


@pytest.mark.parametrize("shaping", [False, True])
def test_justified_cell_list_marker_is_not_stretched(shaping):
    model, _ = cell_model(listed=True)
    raw = render(model, shaping, tagged=True)
    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        words = pdf[0].get_text("words")
        marker = next(word for word in words if word[4] == "1.")
        assert marker[0] == pytest.approx(40, abs=0.05)
        assert marker[2] - marker[0] < 15
        rows = lines(pdf[0])
        for row in rows[:-1]:
            assert row["bbox"][2] == pytest.approx(186, abs=0.05)
        assert [word[4] for word in words if word[4] != "1."] == WORDS.split()
    assert PdfReader(BytesIO(raw)).trailer["/Root"]["/StructTreeRoot"]


@pytest.mark.parametrize("shaping", [False, True])
def test_cell_links_highlights_and_decorations_follow_expanded_width(shaping):
    model, paragraph = cell_model()
    run = paragraph.runs[0]
    run.text = "[" + WORDS + "](https://example.test/cell)"
    run.font.underline = run.font.strike_through = True
    run.font.highlight_color = "FFFF00"
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        page = pdf[0]
        for row in lines(page):
            end = row["bbox"][2]
            assert any(abs(link["from"].x1 - end) < 0.05 for link in page.get_links())
            drawings = page.get_drawings()
            assert any(d["fill"] == (1, 1, 0) and abs(d["rect"].x1 - end) < 0.05 for d in drawings)
            assert sum(d["rect"].height < 2 and abs(d["rect"].x1 - end) < 0.05 for d in drawings) >= 2


@pytest.mark.parametrize("shaping", [False, True])
def test_justified_long_cell_and_repeated_header_keep_all_body_tokens(shaping):
    text = " ".join(f"WORD{i:03}" for i in range(100))
    model, _ = cell_model(text=text)
    table = model.sections[0].body.children[0]
    header = ldm.Paragraph(children=[ldm.Run(text="HEAD ONE\nHEAD END", font=ldm.Font(size=12))],
                           paragraph_format=ldm.ParagraphFormat(alignment=3))
    table.rows.insert(0, ldm.Row(row_format=ldm.RowFormat(heading_format=True), cells=[ldm.Cell(
        paragraphs=[header], cell_format=ldm.CellFormat(left_padding=8, right_padding=8))]))
    raw = render(model, shaping, tagged=True)
    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        assert len(pdf) > 2
        words = [word[4] for page in pdf for word in page.get_text("words")]
        assert [word for word in words if word.startswith("WORD")] == text.split()
        for page in pdf:
            assert next(word[2] for word in page.get_text("words") if word[4] == "ONE") == pytest.approx(192, abs=0.05)


@pytest.mark.parametrize("shaping", [False, True])
def test_no_wrap_retains_overwide_line_without_automatic_breaks(shaping):
    model, _ = cell_model(wrap=False)
    model.sections[0].page_setup.page_width = 500
    model.sections[0].body.children[0].preferred_width = ldm.PreferredWidth.from_points(180)
    raw = render(model, shaping)
    assert PdfReader(BytesIO(raw)).pages[0].extract_text().split() == WORDS.split()
    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        words = pdf[0].get_text("words")
        assert len(words) < len(WORDS.split())  # The fixed cell clips the overwide visual line.
        assert max(word[1] for word in words) - min(word[1] for word in words) < 0.05


@pytest.mark.parametrize("shaping", [False, True])
def test_trimming_large_trailing_space_keeps_original_line_metrics(shaping):
    model, paragraph = cell_model()
    paragraph._children = [ldm.Run(text="ONE TWO THREE", font=ldm.Font(size=12)),
                           ldm.Run(text=" ", font=ldm.Font(size=40)),
                           ldm.Run(text="LONGWORD FOUR FIVE SIX", font=ldm.Font(size=12))]
    control = model.model_copy(deep=True)
    control.sections[0].body.children[0].rows[0].cells[0].paragraphs[0].paragraph_format.alignment = 0

    def origin(document):
        with pymupdf.open(stream=render(document, shaping), filetype="pdf") as pdf:
            return next(span["origin"][1] for block in pdf[0].get_text("dict")["blocks"]
                        for line in block.get("lines", []) for span in line["spans"]
                        if "ONE" in span["text"])

    assert origin(model) == pytest.approx(origin(control), abs=0.05)


@pytest.mark.parametrize("shaping", [False, True])
def test_cell_superscript_keeps_reduced_glyphs_and_stretched_links(shaping):
    model, paragraph = cell_model(text="SHORT LINE\nLAST")
    run = paragraph.runs[0]
    run.text = "[SHORT LINE\nLAST](https://example.test/script)"
    run.font.superscript = True
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        spans = [span for block in pdf[0].get_text("dict")["blocks"]
                 for line in block.get("lines", []) for span in line["spans"]]
        assert all(span["size"] == pytest.approx(8.4, abs=0.05) for span in spans)
        assert next(word[2] for word in pdf[0].get_text("words") if word[4] == "LINE") == pytest.approx(192, abs=0.05)
        assert any(link["from"].x1 == pytest.approx(192, abs=0.05) for link in pdf[0].get_links())


def test_public_docx_cell_justification_uses_cell_padding(tmp_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    import aspose.words_foss as aw

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(220), Pt(180)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Pt(20)
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Pt(180)
    cell = table.cell(0, 0)
    margins = OxmlElement("w:tcMar")
    for name in ["left", "right"]:
        element = OxmlElement("w:" + name)
        element.set(qn("w:w"), "160")
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    cell._tc.get_or_add_tcPr().append(margins)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(WORDS).font.size = Pt(12)
    source, target = tmp_path / "cell.docx", tmp_path / "cell.pdf"
    document.save(source)
    loaded = aw.Document(source)
    assert loaded.light_document_model.compatibility_mode == 14
    assert loaded.light_document_model.sections[0].body.tables[0].rows[0].cells[0].cell_format.left_padding == 8
    loaded.save(target)
    with pymupdf.open(target) as pdf:
        rows = lines(pdf[0])
        assert rows[0]["bbox"][0] == pytest.approx(20, abs=0.05)
        assert rows[0]["bbox"][2] == pytest.approx(184, abs=0.05)
        assert rows[-1]["bbox"][2] < 177
