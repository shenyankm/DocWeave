"""Ordinary size views keep table declarations sparse across edits and saves."""
from io import BytesIO
from xml.etree import ElementTree as E
from zipfile import ZipFile

import pytest
import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.saving import OoxmlSaveOptions
from .test_docx_dom import package

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'


def owned_table(path, normal=12, mode=15):
    normal_size = '' if normal is None else f'<w:rPr><w:sz w:val="{normal * 2}"/></w:rPr>'
    styles = (f'<w:styles xmlns:w="{W}"><w:style w:type="paragraph" w:default="1" w:styleId="Normal">'
              f'<w:name w:val="Normal"/>{normal_size}</w:style>'
              '<w:style w:type="table" w:styleId="Owned"><w:name w:val="Owned"/>'
              '<w:tblStylePr w:type="firstRow"><w:rPr><w:sz w:val="44"/></w:rPr></w:tblStylePr>'
              '</w:style></w:styles>')
    settings = (f'<w:settings xmlns:w="{W}"><w:compat><w:compatSetting w:name="compatibilityMode" '
                f'w:uri="http://schemas.microsoft.com/office/word" w:val="{mode}"/></w:compat></w:settings>')
    rels = (f'<Relationships xmlns="{REL}"><Relationship Id="s" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
            '<Relationship Id="c" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings" '
            'Target="settings.xml"/></Relationships>')
    package(path, '<w:tbl><w:tblPr><w:tblStyle w:val="Owned"/><w:tblLook w:firstRow="1"/></w:tblPr>'
            '<w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid><w:tr><w:tc><w:p>'
            '<w:r><w:t>OWNED_CELL</w:t></w:r></w:p></w:tc></w:tr></w:tbl>',
            {'word/styles.xml': styles.encode(), 'word/settings.xml': settings.encode(),
             'word/_rels/document.xml.rels': rels.encode()})
    return path


def target(model):
    return next(r for r in model.get_child_nodes(ldm.NodeType.RUN, True) if r.text == 'OWNED_CELL')


@pytest.mark.parametrize('normal,expected', [(12, 11), (None, 22), (13, 13)])
@pytest.mark.parametrize('copy', ['direct', 'deep', 'json'])
def test_sparse_table_size_survives_copies_and_two_saves(tmp_path, normal, expected, copy):
    path = owned_table(tmp_path / 'owned.docx', normal)
    model = aw.Document(path).light_document_model
    run = target(model)
    assert run.font.size == expected
    assert run.font.size_explicit is False
    assert aw.DocxDocument(path).body.tables[0].rows[0].cells[0].paragraphs[0].runs[0].effective_font.size == expected
    before = model.model_dump_json()
    value = (model if copy == 'direct' else model.model_copy(deep=True) if copy == 'deep'
             else ldm.Document.model_validate_json(before))
    cold_expected = 12 if normal == 12 else expected
    for fmt in (aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC):
        current = value
        for _ in range(2):
            data = LdmDocxWriter(OoxmlSaveOptions(fmt)).write_to_bytes(current)
            current = aw.Document(BytesIO(data)).light_document_model
            assert target(current).font.size == cold_expected
            assert target(current).font.size_explicit is False
            dom = aw.DocxDocument(BytesIO(data))
            assert dom.body.tables[0].rows[0].cells[0].paragraphs[0].runs[0].effective_font.size == cold_expected
            if fmt == aw.SaveFormat.DOCX:
                with ZipFile(BytesIO(data)) as archive:
                    tree = E.fromstring(archive.read('word/document.xml'))
                assert tree.find('.//{' + W + '}rPr/{' + W + '}sz') is None
    assert model.model_dump_json() == before


def test_held_size_follows_region_edit_and_direct_zero(tmp_path):
    model = aw.Document(owned_table(tmp_path / 'owned.docx', None)).light_document_model
    font = target(model).font
    region = next(s for s in model.styles if s.type == 3).table_style_properties[0]
    assert font.size == 22
    region.font.size = 36
    assert font.size == 36
    font.size = 0
    assert font.size == 0
    assert font.size_explicit is True
    cold = aw.Document(BytesIO(LdmDocxWriter().write_to_bytes(model)))
    assert target(cold.light_document_model).font.size == 0


def test_size_copy_does_not_retain_original_document_binding(tmp_path):
    model = aw.Document(owned_table(tmp_path / 'owned.docx', None)).light_document_model
    clone = model.model_copy(deep=True)
    next(s for s in clone.styles if s.type == 3).table_style_properties[0].font.size = 36
    assert target(clone).font.size == 36
    assert target(model).font.size == 22


def test_inherited_look_only_style_survives_json_and_cold_saves(tmp_path):
    path = owned_table(tmp_path / 'owned.docx', None)
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    document = E.fromstring(parts['word/document.xml'])
    table_properties = document.find('.//{' + W + '}tblPr')
    look = table_properties.find('{' + W + '}tblLook')
    table_properties.remove(look)
    styles = E.fromstring(parts['word/styles.xml'])
    table_style = next(s for s in styles if s.get('{' + W + '}type') == 'table')
    E.SubElement(table_style, '{' + W + '}tblPr').append(look)
    parts['word/document.xml'], parts['word/styles.xml'] = E.tostring(document), E.tostring(styles)
    with ZipFile(path, 'w') as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    model = aw.Document(path).light_document_model
    assert target(model).font.size == 22
    assert model.get_child_nodes(ldm.NodeType.TABLE, True)[0].style_look is None
    for fmt in (aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC):
        current = ldm.Document.model_validate_json(model.model_dump_json())
        for _ in range(2):
            data = LdmDocxWriter(OoxmlSaveOptions(fmt)).write_to_bytes(current)
            current = aw.Document(BytesIO(data)).light_document_model
            assert target(current).font.size == 22
            assert aw.DocxDocument(BytesIO(data)).body.tables[0].rows[0].cells[0].paragraphs[0].runs[0].effective_font.size == 22


@pytest.mark.parametrize('route', ['direct', 'deep', 'json'])
@pytest.mark.parametrize('fmt', [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC])
def test_conditional_size_with_classic_comment_image(tmp_path, route, fmt):
    from docx import Document as NativeDocument
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from PIL import Image

    native = NativeDocument()
    native.styles.add_style('OwnedTable', WD_STYLE_TYPE.TABLE)
    table = native.add_table(rows=1, cols=1)
    table.style = 'OwnedTable'
    anchor = table.cell(0, 0).paragraphs[0].add_run('OWNED_CELL')
    comment = native.add_comment([anchor], text='OWNED COMMENT', author='Owned')
    picture = BytesIO()
    Image.new('RGB', (2, 2), (70, 110, 140)).save(picture, format='PNG')
    comment.paragraphs[0].add_run().add_picture(BytesIO(picture.getvalue()))
    style = native.styles['OwnedTable'].element
    region = OxmlElement('w:tblStylePr'); region.set(qn('w:type'), 'firstRow')
    rpr = OxmlElement('w:rPr'); size = OxmlElement('w:sz'); size.set(qn('w:val'), '44')
    rpr.append(size); region.append(rpr); style.append(region)
    output = BytesIO(); native.save(output)
    model = aw.Document(BytesIO(output.getvalue())).light_document_model
    held = target(model).font
    assert held.size == 22
    next(s for s in model.styles if s.name == 'OwnedTable').table_style_properties[0].font.size = 36
    assert held.size == 36
    current = (model if route == 'direct' else model.model_copy(deep=True) if route == 'deep'
               else ldm.Document.model_validate_json(model.model_dump_json()))
    for _ in range(2):
        data = LdmDocxWriter(OoxmlSaveOptions(fmt)).write_to_bytes(current)
        current = aw.Document(BytesIO(data)).light_document_model
        assert target(current).font.size == 36
        assert target(current).font.size_explicit is False
        story = next(s for s in current.source_stories if s.kind == 'comment')
        assert story.comment_metadata.author == 'Owned'
        assert any('OWNED COMMENT' in p.text for p in story.children if isinstance(p, ldm.Paragraph))
        raw = data
        if fmt == aw.SaveFormat.FLAT_OPC:
            from aspose.words_foss._flat_opc import decode
            raw = decode(raw)
        with ZipFile(BytesIO(raw)) as archive:
            images = [n for n in archive.namelist() if n.startswith('word/media/')]
            assert len(images) == 1 and archive.read(images[0]) == picture.getvalue()


@pytest.mark.parametrize('fmt', [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC])
@pytest.mark.parametrize('destination', ['path', 'stream', 'writer'])
def test_successful_save_updates_held_size_without_replacing_source(tmp_path, fmt, destination):
    document = aw.Document(owned_table(tmp_path / 'source.docx'))
    model = document.light_document_model
    source = model.source_settings
    source_bytes = source.data
    held = target(model).font
    assert held.size == 11
    assert model.current_compatibility_settings is None
    document.to_bytes(fmt)
    assert held.size == 11 and model.current_compatibility_settings is None
    path = tmp_path / 'result'
    if destination == 'writer':
        LdmDocxWriter(OoxmlSaveOptions(fmt)).write(model, path)
        data = path.read_bytes()
    elif destination == 'stream':
        stream = BytesIO()
        document.save(stream, fmt)
        data = stream.getvalue()
    else:
        document.save(path, fmt)
        data = path.read_bytes()
    assert held.size == 12
    assert target(aw.Document(BytesIO(data)).light_document_model).font.size == 12
    assert model.source_settings is source and source.data == source_bytes
    assert model.current_compatibility_settings is not None
    for copied in (model.model_copy(deep=True), ldm.Document.model_validate_json(model.model_dump_json())):
        assert target(copied).font.size == 12


@pytest.mark.parametrize('fmt', [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC])
def test_failed_save_does_not_commit_compatibility_view(tmp_path, fmt):
    class FailedStream(BytesIO):
        def write(self, data):
            raise OSError('Owned failure')

    document = aw.Document(owned_table(tmp_path / 'source.docx'))
    model = document.light_document_model
    source = model.source_settings
    for destination in (FailedStream(), tmp_path):
        with pytest.raises((RuntimeError, OSError, ValueError)):
            document.save(destination, fmt)
        assert model.source_settings is source
        assert model.current_compatibility_settings is None
        assert target(model).font.size == 11


def test_programmatic_save_preserves_unset_source_and_public_preferences(tmp_path):
    document = aw.Document()
    model = document.light_document_model
    model.source_settings = None
    document.compatibility_options.do_not_expand_shift_return = True
    before = dict(model.compatibility_options_state.flags)
    document.save(BytesIO(), aw.SaveFormat.DOCX)
    assert model.source_settings is None
    assert model.current_compatibility_settings is not None
    assert model.compatibility_options_state.flags == before
    assert document.compatibility_options.do_not_expand_shift_return is True



def test_typed_document_inputs_do_not_rebind_caller_or_previous_document():
    run = ldm.Run(text='OWNED')
    original_font = run.font
    paragraph = ldm.Paragraph(children=[run])
    section = ldm.Section(body=ldm.Body(children=[paragraph]))
    first = ldm.Document(sections=[section], doc_defaults_font=ldm.Font(size=20))
    held = first.sections[0].body.paragraphs[0].runs[0].font
    assert held.size == 20
    assert original_font.size == 20
    assert first.sections[0] is section
    second = ldm.Document(sections=first.sections, doc_defaults_font=ldm.Font(size=30))
    assert held.size == 20
    assert first.sections[0].body.paragraphs[0].runs[0].font.size == 20
    assert second.sections[0].body.paragraphs[0].runs[0].font.size == 30
    assert second.sections[0] is not first.sections[0]
    assert original_font.size == 20


@pytest.mark.parametrize('version', range(8))
@pytest.mark.parametrize('preload', [False, True])
def test_optimize_profile_updates_warm_size_then_actual_saved_view(tmp_path, version, preload):
    from aspose.words_foss.settings import MsWordVersion
    document = aw.Document(owned_table(tmp_path / 'source.docx'))
    model = document.light_document_model
    snapshot = model.source_settings
    held = target(model).font
    if preload:
        document.save(BytesIO(), aw.SaveFormat.DOCX)
        assert held.size == 12
    document.compatibility_options.optimize_for(MsWordVersion(version))
    assert held.size == (11 if version < 4 else 12)
    for copied in (model.model_copy(deep=True), ldm.Document.model_validate_json(model.model_dump_json())):
        assert target(copied).font.size == held.size
    document.to_bytes(aw.SaveFormat.DOCX)
    assert held.size == (11 if version < 4 else 12)
    stream = BytesIO()
    document.save(stream, aw.SaveFormat.DOCX)
    expected = 11 if version == 3 else 12
    assert held.size == expected
    assert target(aw.Document(BytesIO(stream.getvalue())).light_document_model).font.size == expected
    assert model.source_settings is snapshot


def test_shallow_document_copy_keeps_shared_font_binding():
    run = ldm.Run(text="OWNED")
    first = ldm.Document(sections=[ldm.Section(body=ldm.Body(
        children=[ldm.Paragraph(children=[run])]))], doc_defaults_font=ldm.Font(size=20))
    second = first.model_copy(update={"doc_defaults_font": ldm.Font(size=30)})
    assert second.sections[0] is first.sections[0]
    assert run.font.size == 20
    first.doc_defaults_font = ldm.Font(size=22)
    assert run.font.size == 22
    independent = first.model_copy(deep=True, update={"sections": first.sections,
                                                     "doc_defaults_font": ldm.Font(size=30)})
    assert independent.sections[0].body.paragraphs[0].runs[0].font.size == 30
    assert run.font.size == 22


def test_reference_style_projection_does_not_rebind_source_runs():
    from aspose.words_foss.docx_writer.styles_part import apply_reference_styles, render_styles_xml
    run = ldm.Run(text="OWNED")
    source = ldm.Document(sections=[ldm.Section(body=ldm.Body(
        children=[ldm.Paragraph(children=[run])]))], doc_defaults_font=ldm.Font(size=20))
    apply_reference_styles(render_styles_xml(source), source)
    source.doc_defaults_font = ldm.Font(size=30)
    assert run.font.size == 30
