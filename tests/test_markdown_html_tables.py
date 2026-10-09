"""HTML table output must contain HTML geometry and inline semantics."""

import warnings
from io import BytesIO
from xml.etree import ElementTree as ET

import pytest
from docx import Document as NativeDocument
from PIL import Image

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm


def html_output(doc):
    options = aw.saving.MarkdownSaveOptions()
    options.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
    return doc.to_bytes(options).decode().strip()


def merged_source(merge):
    native = NativeDocument()
    table = native.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f"R{r}C{c}"
    table.cell(0, 0).merge(table.cell(0, 1) if merge == "horizontal" else
                         table.cell(1, 0) if merge == "vertical" else table.cell(1, 1))
    raw = BytesIO()
    native.save(raw)
    return raw.getvalue()


@pytest.mark.parametrize("merge", ["horizontal", "vertical", "rectangle"])
def test_real_docx_merge_has_html_spans(merge):
    doc = aw.Document(BytesIO(merged_source(merge)))
    before = doc.light_document_model.model_dump()
    root = ET.fromstring(html_output(doc))
    first = root.find("tr/th")
    assert first.get("colspan", "1") == ("2" if merge != "vertical" else "1")
    assert first.get("rowspan", "1") == ("2" if merge != "horizontal" else "1")
    text = " ".join(root.itertext())
    assert all(text.count(f"R{r}C{c}") == 1 for r in range(3) for c in range(3))
    assert doc.light_document_model.model_dump() == before


def test_independent_mammoth_agrees_on_real_docx_grid():
    mammoth = pytest.importorskip("mammoth")

    def records(root):
        return [[(cell.get("colspan", "1"), cell.get("rowspan", "1"),
                  " ".join(" ".join(cell.itertext()).split())) for cell in row]
                for row in root.findall("tr")]

    for merge in ("horizontal", "vertical", "rectangle"):
        raw = merged_source(merge)
        own = ET.fromstring(html_output(aw.Document(BytesIO(raw))))
        independent = ET.fromstring(mammoth.convert_to_html(BytesIO(raw)).value)
        assert records(own) == records(independent)


@pytest.mark.parametrize("attribute,tag", [("bold", "strong"), ("italic", "em"),
                                        ("strike_through", "del"), ("superscript", "sup"),
                                        ("subscript", "sub")])
def test_html_uses_format_tags_without_interpreting_literal_markdown(attribute, tag):
    doc = aw.Document(BytesIO(b"BODY"), aw.MarkdownLoadOptions())
    doc.sections[0].body.children = [ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[
        ldm.Paragraph(children=[ldm.Run(text="LITERAL ** & <x>", font=ldm.Font(**{attribute: True}))])])])])]
    root = ET.fromstring(html_output(doc))
    assert root.find(".//" + tag).text == "LITERAL ** & <x>"


def test_html_preserves_nested_table_and_paragraph_order():
    def paragraph(text):
        return ldm.Paragraph(children=[ldm.Run(text=text)])
    nested = ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[paragraph("NESTED")])])])
    doc = aw.Document(BytesIO(b"BODY"), aw.MarkdownLoadOptions())
    doc.sections[0].body.children = [ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(
        children=[paragraph("BEFORE"), nested, paragraph("AFTER")])])])]
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        root = ET.fromstring(html_output(doc))
    assert len(root.findall(".//table")) == 1
    assert [e.tag for e in root.find("tr/th")] == ["p", "table", "p"]
    assert " ".join(root.itertext()).index("BEFORE") < " ".join(root.itertext()).index("NESTED") < " ".join(root.itertext()).index("AFTER")
    assert not caught


def model_document(paragraph):
    doc = aw.Document(BytesIO(b"BODY"), aw.MarkdownLoadOptions())
    doc.sections[0].body.children = [ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[paragraph])])])]
    return doc


@pytest.mark.parametrize("target", ['https://example.com/a_(b)?x=1&y=2', '#anchor', 'mailto:a@example.com'])
def test_encoded_links_have_html_targets_and_formatting(target):
    from aspose.words_foss._links import format_link

    doc = model_document(ldm.Paragraph(children=[ldm.Run(
        text=format_link('Label [x] & "quoted"', target), font=ldm.Font(style_name="Hyperlink", bold=True))]))
    root = ET.fromstring(html_output(doc))
    link = root.find(".//strong/a")
    assert link.get("href") == target and link.text == 'Label [x] & "quoted"'


@pytest.mark.parametrize("target", ['javascript:alert(1)', 'JaVaScRiPt:alert(1)', 'data:text/html,<script>x</script>'])
def test_dangerous_link_targets_are_not_activated(target):
    from aspose.words_foss._links import format_link

    doc = model_document(ldm.Paragraph(children=[ldm.Run(
        text=format_link("LABEL", target), font=ldm.Font(style_name="Hyperlink"))]))
    with pytest.warns(aw.ContentLossWarning, match="Unsafe HTML"):
        root = ET.fromstring(html_output(doc))
    assert root.find(".//a") is None and "LABEL" in "".join(root.itertext())


def test_html_images_preserve_bytes_and_escape_attributes(tmp_path):
    import base64

    stream = BytesIO()
    Image.new("RGB", (2, 2), "red").save(stream, format="PNG")
    shape = ldm.Shape(has_image=True, alternative_text='Alt ](" & <x>', image_data=ldm.ImageData(
        image_type=ldm.ImageData.from_mime("image/png"), image_bytes=stream.getvalue()))
    doc = model_document(ldm.Paragraph(children=[shape]))
    root = ET.fromstring(html_output(doc))
    image = root.find(".//img")
    assert image.get("alt") == shape.alternative_text
    assert base64.b64decode(image.get("src").split(",", 1)[1]) == stream.getvalue()
    options = aw.saving.MarkdownSaveOptions()
    options.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
    options.images_folder = str(tmp_path / "images")
    options.images_folder_alias = 'img&"quoted'
    output = tmp_path / "table.md"
    doc.save(output, options)
    image = ET.fromstring(output.read_text()).find(".//img")
    assert image.get("src") == 'img&"quoted/image1.png'
    assert (tmp_path / "images/image1.png").read_bytes() == stream.getvalue()


def test_html_fields_hide_instructions_and_hidden_runs_with_explicit_loss():
    doc = model_document(ldm.Paragraph(children=[ldm.FieldStart(), ldm.Run(text="SECRET_CODE"),
        ldm.FieldSeparator(), ldm.Run(text="RESULT", font=ldm.Font(bold=True)), ldm.FieldEnd(),
        ldm.Run(text="SECRET_HIDDEN", font=ldm.Font(hidden=True))]))
    with pytest.warns(aw.ContentLossWarning, match="field actions"):
        root = ET.fromstring(html_output(doc))
    assert root.find(".//strong").text == "RESULT"
    assert "SECRET" not in "".join(root.itertext())


def test_orphan_vertical_merge_retains_content_and_warns():
    doc = model_document(ldm.Paragraph(children=[ldm.Run(text="ORPHAN")]))
    doc.sections[0].body.tables[0].rows[0].cells[0].cell_format.vertical_merge = 2
    with pytest.warns(aw.ContentLossWarning, match="Orphan"):
        root = ET.fromstring(html_output(doc))
    assert "ORPHAN" in "".join(root.itertext())


def test_legacy_horizontal_merge_retains_continuation_content():
    doc = model_document(ldm.Paragraph(children=[ldm.Run(text="START")]))
    row = doc.sections[0].body.tables[0].rows[0]
    row.cells[0].cell_format.horizontal_merge = 1
    row.cells.append(ldm.Cell(cell_format=ldm.CellFormat(horizontal_merge=2),
        paragraphs=[ldm.Paragraph(children=[ldm.Run(text="CONTINUATION")])]))
    root = ET.fromstring(html_output(doc))
    assert root.find("tr/th").get("colspan") == "2"
    assert len(root.findall("tr/th")) == 1
    assert "START" in "".join(root.itertext()) and "CONTINUATION" in "".join(root.itertext())


def test_html_underline_code_and_line_breaks_are_markup():
    doc = model_document(ldm.Paragraph(children=[
        ldm.Run(text="UNDER\nLINE", font=ldm.Font(underline=1)),
        ldm.Run(text="**literal**", font=ldm.Font(style_name="InlineCode"))]))
    options = aw.saving.MarkdownSaveOptions()
    options.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
    options.export_underline_formatting = True
    root = ET.fromstring(doc.to_bytes(options).decode())
    assert root.find(".//u/br") is not None
    assert root.find(".//code").text == "**literal**"


def test_html_note_anchor_loss_is_explicit_and_hidden_notes_do_not_leak():
    doc = model_document(ldm.Paragraph(children=[ldm.Run(text="BODY"),
        ldm.NoteReference(kind="footnote", identifier="1"),
        ldm.NoteReference(kind="footnote", identifier="2", hidden=True)]))
    doc.light_document_model.source_stories = [ldm.SourceStory(kind="footnote",
        part_name="word/footnotes.xml", identifier=identifier,
        children=[ldm.Paragraph(children=[ldm.Run(text=text)])])
        for identifier, text in [("1", "VISIBLE_NOTE"), ("2", "SECRET_NOTE")]]
    options = aw.saving.MarkdownSaveOptions()
    options.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
    options.export_notes = True
    with pytest.warns(aw.ContentLossWarning, match="semantic note anchors"):
        result = doc.to_bytes(options).decode()
    assert "VISIBLE_NOTE" in result and "SECRET_NOTE" not in result


def test_html_unknown_shape_has_loss_diagnostic_before_publish(tmp_path):
    doc = model_document(ldm.Paragraph(children=[ldm.Shape(text_box={"paragraphs": [
        ldm.Paragraph(children=[ldm.Run(text="UNSUPPORTED_BOX")]) ]})]))
    output = tmp_path / "original.md"
    output.write_bytes(b"ORIGINAL")
    options = aw.saving.MarkdownSaveOptions()
    options.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
    with warnings.catch_warnings():
        warnings.simplefilter("error", aw.ContentLossWarning)
        with pytest.raises(aw.ContentLossWarning, match="non-image shape"):
            doc.save(output, options)
    assert output.read_bytes() == b"ORIGINAL"
