"""Model JSON keeps sparse font declarations instead of creating direct off flags."""

import json
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document as IndependentDocument
from docx.enum.style import WD_STYLE_TYPE

from aspose.words_foss import Document, SaveFormat
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.saving import OoxmlSaveOptions

BENCHMARKS = Path(__file__).parents[1] / 'docs/benchmarks'
TOGGLES = json.loads((BENCHMARKS / 'style-toggles-26.9.json').read_text())['records']


@pytest.mark.parametrize('row', TOGGLES, ids=[row['input'] for row in TOGGLES])
def test_conversion_run_flags_match_native_category_inheritance(row):
    with ZipFile(BENCHMARKS / 'corpus/style-toggles-26.9.zip') as archive:
        source = archive.read(row['input'])
    model = Document(BytesIO(source)).light_document_model
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    def direct(data):
        with ZipFile(BytesIO(data)) as archive:
            document = ET.fromstring(archive.read('word/document.xml'))
        run = next(document.iter(w + 'r'))
        return tuple(None if (element := run.find(w + 'rPr/' + w + tag)) is None
                     else element.get(w + 'val', '1') for tag in ('b', 'i'))
    for current in (model, ldm.Document.model_validate_json(model.model_dump_json())):
        font = current.sections[0].body.paragraphs[0].runs[0].font
        assert font.bold is row['bold'] and font.italic is row['italic']
        assert direct(LdmDocxWriter().write_to_bytes(current)) == direct(source)


@pytest.mark.parametrize('attribute', ['bold', 'italic', 'hidden', 'all_caps', 'small_caps', 'strike'])
@pytest.mark.parametrize('declared', [None, False, True])
@pytest.mark.parametrize('format', [SaveFormat.DOCX, SaveFormat.FLAT_OPC])
def test_sparse_character_flags_survive_model_json(attribute, declared, format):
    source = IndependentDocument()
    base = source.styles.add_style('Base', WD_STYLE_TYPE.CHARACTER)
    setattr(base.font, attribute, True)
    child = source.styles.add_style('Derived', WD_STYLE_TYPE.CHARACTER)
    child.base_style = base
    setattr(child.font, attribute, declared)
    source.add_paragraph().add_run('IMPORT').style = child
    stream = BytesIO()
    source.save(stream)
    model = Document(BytesIO(stream.getvalue())).light_document_model
    model = ldm.Document.model_validate_json(model.model_dump_json())
    data = LdmDocxWriter(OoxmlSaveOptions(format)).write_to_bytes(model)
    from aspose.words_foss import DocxDocument
    reopened = DocxDocument(BytesIO(data)).to_bytes()
    actual = IndependentDocument(BytesIO(reopened))
    assert getattr(actual.styles['Derived'].font, attribute) is declared
    assert actual.styles['Derived'].base_style.name == 'Base'


def test_font_field_origin_survives_json_and_later_assignments():
    font = ldm.Font(bold=False, size=12)
    restored = ldm.Font.model_validate_json(font.model_dump_json())
    assert restored.model_fields_set == font.model_fields_set
    restored.italic = False
    again = ldm.Font.model_validate_json(restored.model_dump_json())
    assert again.model_fields_set == {'bold', 'size', 'italic', 'italic_explicit'}
    assert again.italic_explicit is True
    assert again.bold is again.italic is False and again.size == 12


def test_font_json_keeps_defaults_and_legacy_input_policy():
    font = ldm.Font()
    data = font.model_dump()
    assert data['bold'] is False and data['size'] == 0
    assert ldm.Font.model_validate(data).model_fields_set == set()
    legacy = {key: value for key, value in data.items() if key != '_fields_set'}
    assert ldm.Font.model_validate(legacy).model_fields_set == set(legacy)


@pytest.mark.parametrize('metadata', [None, 'bold', [1], ['unknown'], ['italic']])
def test_invalid_font_origin_metadata_is_rejected(metadata):
    with pytest.raises(ValueError, match='Invalid font field-origin metadata'):
        ldm.Font.model_validate({'bold': False, '_fields_set': metadata})


def test_font_origin_validation_does_not_mutate_input_or_invent_excluded_fields():
    data = {'bold': False, '_fields_set': ['bold']}
    ldm.Font.model_validate(data)
    assert data == {'bold': False, '_fields_set': ['bold']}
    font = ldm.Font(bold=False, italic=True)
    selected = font.model_dump(include={'bold'})
    assert selected == {'bold': False, '_fields_set': ['bold']}
    assert ldm.Font.model_validate(selected).model_fields_set == {'bold'}
