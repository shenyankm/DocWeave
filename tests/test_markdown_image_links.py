"""Image labels and external destinations survive independent Markdown parsing."""

import base64
import os
import warnings
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET

import pytest
from docx import Document as NativeDocument
from markdown_it import MarkdownIt
from PIL import Image

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.md_writer import LdmMarkdownWriter
from aspose.words_foss.model.enums.image import ImageType
from aspose.words_foss.models import ConversionOptions


def picture():
    stream = BytesIO()
    Image.new("RGB", (3, 3), "blue").save(stream, format="PNG")
    return stream.getvalue()


def source(name, alt, story):
    native = NativeDocument()
    container = native.add_table(rows=1, cols=1).cell(0, 0) if story == "table" else (
        native.sections[0].header if story == "header" else native.sections[0].footer if story == "footer" else native)
    para = container.paragraphs[0] if story != "body" else container.add_paragraph()
    inline = para.add_run().add_picture(BytesIO(picture()))._inline
    inline.docPr.set("descr", alt)
    stream = BytesIO()
    native.save(stream)
    stream.seek(0)
    doc = aw.Document(stream)
    shape = next(item for item in doc.get_child_nodes(aw.NodeType.SHAPE, True))
    shape.image_data.source_full_name = name
    return doc


def images(markdown):
    parser = MarkdownIt("commonmark").enable("table")
    # Parse grammar without the renderer concealing unsafe URL schemes.
    parser.validateLink = lambda target: True
    roots = parser.parse(markdown)
    tokens = roots + [child for token in roots for child in (token.children or [])]
    assert not any(token.type in ("html_inline", "html_block", "link_open") for token in tokens)
    return [token for token in tokens if token.type == "image"]


@pytest.mark.parametrize("name", ["space name.png", "image#1?.png", "literal%20.png", "图[片](蓝).png"])
@pytest.mark.parametrize("story", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("entry", ["file", "writer"])
def test_external_paths_and_labels_resolve_to_original_assets(tmp_path, name, story, entry):
    name = name.replace("?", "!") if os.name == "nt" else name
    alt = 'Alt ](https://evil.invalid) <script>x</script> ![label & "quoted"'
    doc = source(name, alt, story)
    folder = tmp_path / "assets # folder"
    output = tmp_path / "out.md"
    if entry == "writer":
        raw = LdmMarkdownWriter(ConversionOptions(images_folder=str(folder))).write(doc.light_document_model, output)
    else:
        options = aw.saving.MarkdownSaveOptions()
        options.images_folder = str(folder)
        doc.save(output, options)
        raw = output.read_text()
    parsed = images(raw)
    assert len(parsed) == 1
    assert "".join(child.content for child in parsed[0].children) == alt
    uri = urlsplit(parsed[0].attrGet("src"))
    assert not uri.query and not uri.fragment
    linked = Path(unquote(uri.path))
    if not linked.is_absolute():
        linked = output.parent / linked
    assert linked == folder / name and linked.read_bytes() == picture()
    assert sorted(p.name for p in folder.iterdir()) == [name]


@pytest.mark.parametrize("html", [False, True])
@pytest.mark.parametrize("alias", ["public images", "https://cdn.invalid/assets%20v1?token=ab&mode=1#logo"])
def test_alias_appends_encoded_filename_to_path_and_retains_url_components(tmp_path, html, alias):
    name = "image#1!.png" if os.name == "nt" else "image#1?.png"
    doc = source(name, "Plain alt", "table" if html else "body")
    options = aw.saving.MarkdownSaveOptions()
    options.images_folder = str(tmp_path / "images")
    options.images_folder_alias = alias
    if html:
        options.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
    output = tmp_path / "out.md"
    doc.save(output, options)
    raw = output.read_text()
    target = ET.fromstring(raw).find(".//img").get("src") if html else images(raw)[0].attrGet("src")
    uri, prefix = urlsplit(target), urlsplit(alias)
    assert (uri.scheme, uri.netloc, uri.query, uri.fragment) == (prefix.scheme, prefix.netloc, prefix.query, prefix.fragment)
    assert unquote(uri.path) == unquote(prefix.path.rstrip("/")) + "/" + name
    assert (tmp_path / "images" / name).read_bytes() == picture()


@pytest.mark.parametrize("target", ["javascript:alert(1)", "data:text/html,<script>x</script>"])
@pytest.mark.parametrize("html", [False, True])
def test_unsafe_linked_image_is_omitted_with_diagnostic(target, html):
    doc = aw.Document(BytesIO(b"BODY"), aw.MarkdownLoadOptions())
    shape = ldm.Shape(has_image=True, alternative_text="Plain alt", image_data=ldm.ImageData(source_full_name=target))
    para = ldm.Paragraph(children=[shape])
    doc.sections[0].body.children = [ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[para])])])] if html else [para]
    options = aw.saving.MarkdownSaveOptions()
    if html:
        options.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
    with pytest.warns(aw.ContentLossWarning, match="Unsafe"):
        raw = doc.to_bytes(options).decode()
    assert "Plain alt" in raw
    assert ET.fromstring(raw).find(".//img") is None if html else not images(raw)
    assert any(d.code.endswith("unsafe_image") for d in doc.diagnostics)


@pytest.mark.parametrize("story", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("direct", [False, True])
@pytest.mark.parametrize("alt", ['Alt ](https://evil.invalid) <script>x</script> ![label',
                                 'First\r\n\nSecond\t[label]', r'Literal \[x] &amp; *text*',
                                 r'Pipe | with literal \ slash'])
def test_inline_memory_images_keep_literal_alternative_text(story, direct, alt):
    doc = source("plain.png", alt, story)
    raw = LdmMarkdownWriter().write(doc.light_document_model) if direct else doc.to_bytes("md").decode()
    parsed = images(raw)
    assert len(parsed) == 1
    assert "".join(child.content for child in parsed[0].children) == alt
    assert base64.b64decode(parsed[0].attrGet("src").split(",", 1)[1]) == picture()
    restored = aw.Document(BytesIO(raw.encode()), aw.MarkdownLoadOptions())
    shapes = list(restored.get_child_nodes(aw.NodeType.SHAPE, True))
    assert len(shapes) == 1
    assert shapes[0].image_data.source_full_name == alt
    assert shapes[0].image_data.image_bytes == picture()


@pytest.mark.parametrize("story", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("image_type", [ImageType.NO_IMAGE, ImageType.UNKNOWN])
def test_inline_images_without_format_metadata_use_actual_png_mime(story, image_type):
    doc = source("plain.png", "Plain alt", story)
    shape = next(iter(doc.get_child_nodes(aw.NodeType.SHAPE, True)))
    shape.image_data.image_type = image_type
    before = doc.light_document_model.model_dump()
    parsed = images(doc.to_bytes("md").decode())
    assert len(parsed) == 1
    target = parsed[0].attrGet("src")
    assert target.startswith("data:image/png;base64,")
    assert base64.b64decode(target.split(",", 1)[1]) == picture()
    assert doc.light_document_model.model_dump() == before


def test_html_table_image_without_format_metadata_uses_actual_png_mime():
    doc = source("plain.png", "Plain alt", "table")
    shape = next(iter(doc.get_child_nodes(aw.NodeType.SHAPE, True)))
    shape.image_data.image_type = ImageType.UNKNOWN
    options = aw.saving.MarkdownSaveOptions()
    options.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
    raw = doc.to_bytes(options).decode()
    target = ET.fromstring(raw).find(".//img").get("src")
    assert target.startswith("data:image/png;base64,")
    assert base64.b64decode(target.split(",", 1)[1]) == picture()


def test_external_memory_output_remains_rejected_without_side_effects(tmp_path):
    options = aw.saving.MarkdownSaveOptions()
    options.images_folder = str(tmp_path / "images")
    with pytest.raises(ValueError, match="External Markdown images"):
        source("plain.png", "Plain", "body").to_bytes(options)
    assert not (tmp_path / "images").exists()


@pytest.mark.parametrize("alias", [None, 1, "bad\nfolder", "bad\x00folder", "https://[invalid"])
def test_bad_active_alias_fails_before_assets_are_created(tmp_path, alias):
    options = aw.saving.MarkdownSaveOptions()
    options.images_folder = str(tmp_path / "images")
    options.images_folder_alias = alias
    output = tmp_path / "out.md"
    output.write_bytes(b"ORIGINAL")
    with pytest.raises(ValueError):
        source("plain.png", "Plain", "body").save(output, options)
    assert output.read_bytes() == b"ORIGINAL" and not (tmp_path / "images").exists()


@pytest.mark.parametrize("html", [False, True])
def test_strict_unsafe_alias_rolls_back_new_assets_and_retains_existing(tmp_path, html):
    folder = tmp_path / "images"
    folder.mkdir()
    (folder / "keep.txt").write_bytes(b"KEEP")
    options = aw.saving.MarkdownSaveOptions()
    options.images_folder = str(folder)
    options.images_folder_alias = "javascript:alert(1)"
    if html:
        options.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
    output = tmp_path / "out.md"
    output.write_bytes(b"ORIGINAL")
    with warnings.catch_warnings():
        warnings.simplefilter("error", aw.ContentLossWarning)
        with pytest.raises(aw.ContentLossWarning, match="Unsafe"):
            source("plain.png", "Plain", "table" if html else "body").save(output, options)
    assert output.read_bytes() == b"ORIGINAL"
    assert list(folder.iterdir()) == [folder / "keep.txt"]


def test_direct_writer_without_output_path_uses_absolute_file_uri(tmp_path):
    folder = tmp_path / "images # folder"
    doc = source("image%20.png", "Plain alt", "body")
    raw = LdmMarkdownWriter(ConversionOptions(images_folder=str(folder))).write(doc.light_document_model)
    assert images(raw)[0].attrGet("src") == (folder / "image%20.png").as_uri()
