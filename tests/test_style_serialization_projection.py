"""Projection preserves live state and opaque metadata, including failure paths."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest

import aspose.words_foss as aw

from .test_docx_dom import payloads

CORPUS = Path(__file__).parents[1] / "docs/benchmarks/corpus/style-import-conflicts-26.9.zip"


def documents():
    key = "character-b-True-None-True-False-False-None"
    with ZipFile(CORPUS) as archive:
        return tuple(aw.DocxDocument(BytesIO(archive.read(key + "/" + phase + ".docx"))) for phase in ("source", "destination"))


@pytest.mark.parametrize("decoration", [b'w:future="keep"', b'><!--keep--></w:b>'])
def test_decorated_used_root_toggle_is_refused_before_any_commit(decoration):
    source, target = documents()
    part = "word/styles.xml"
    original = target._package.payload(part)
    if decoration.startswith(b">"):
        data = original.replace(b'<w:b w:val="0"/>', b'<w:b w:val="0"' + decoration, 1)
    else:
        data = original.replace(b'<w:b w:val="0"/>', b'<w:b w:val="0" ' + decoration + b'/>', 1)
    target._package.set_parts({part: data})
    before_source, before_target = source.to_bytes(), target.to_bytes()
    with pytest.raises(NotImplementedError, match="decorated root"):
        target.import_node(source.body.paragraphs[0], True)
    assert payloads(source.to_bytes()) == payloads(before_source)
    assert payloads(target.to_bytes()) == payloads(before_target)
    assert target._package._style_projection_part is None


def test_style_part_outer_comments_and_processing_instructions_survive_import_and_projection():
    source, target = documents()
    part = "word/styles.xml"
    original = target._package.payload(part)
    prefix, suffix = b'<!--before--><?opaque before?>', b'<!--after--><?opaque after?>'
    target._package.set_parts({part: prefix + original + suffix})
    copied = target.import_node(source.body.paragraphs[0], True)
    target.body.append_child(copied)
    live = target.part_xml(part)
    data = target.to_bytes()
    saved = payloads(data)[part]
    assert prefix in saved and suffix in saved
    assert target.part_xml(part) == live
    assert target.to_bytes() == data


def test_failed_projection_save_keeps_existing_destination(tmp_path):
    source, target = documents()
    target.body.append_child(target.import_node(source.body.paragraphs[0], True))
    part = "word/styles.xml"
    raw = target._package.tree(part)
    for element in raw.getElementsByTagNameNS("http://schemas.openxmlformats.org/wordprocessingml/2006/main", "b"):
        if element.getAttributeNS("http://schemas.openxmlformats.org/wordprocessingml/2006/main", "val") == "0":
            element.setAttribute("future", "keep")
            break
    path = tmp_path / "existing.docx"
    path.write_bytes(b"unchanged")
    with pytest.raises(NotImplementedError, match="decorated root"):
        target.save(path)
    assert path.read_bytes() == b"unchanged"


def test_custom_styles_part_receives_projection():
    source, target = documents()
    original = payloads(target.to_bytes())
    name = "word/custom/styles.xml"
    original[name] = original.pop("word/styles.xml")
    original["word/_rels/document.xml.rels"] = original["word/_rels/document.xml.rels"].replace(b'Target="styles.xml"', b'Target="custom/styles.xml"')
    original["[Content_Types].xml"] = original["[Content_Types].xml"].replace(b'/word/styles.xml', b'/word/custom/styles.xml')
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        for part, data in original.items():
            archive.writestr(part, data)
    target = aw.DocxDocument(BytesIO(stream.getvalue()))
    target.body.append_child(target.import_node(source.body.paragraphs[0], True))
    assert target._package._style_projection_part == name
    assert payloads(target.to_bytes())[name] != original[name]
    assert "word/styles.xml" not in target.part_names


def test_unknown_relationship_namespace_does_not_select_an_opaque_story():
    from docx import Document

    source, target = documents()
    for run in target.body.paragraphs[0].runs:
        run.font.style_id = None
    rels = "word/_rels/document.xml.rels"
    target._package.set_parts({
        rels: target._package.payload(rels).replace(b'</Relationships>', (
            b'<x:Relationship xmlns:x="urn:opaque" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" '
            b'Target="opaque.xml"/></Relationships>')),
        "word/opaque.xml": b'<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:pPr><w:pStyle w:val="P"/></w:pPr><w:r><w:rPr><w:rStyle w:val="Base"/></w:rPr></w:r></w:p></w:hdr>',
    })
    copied = target.import_node(source.body.paragraphs[0], True)
    assert copied.parent_node is None
    # The imported node is detached; no real story uses a character style yet.
    data = target.to_bytes()
    saved = payloads(data)
    assert Document(BytesIO(data)).styles["P"].font.bold is False
    assert saved["word/opaque.xml"] == target._package.payload("word/opaque.xml")
