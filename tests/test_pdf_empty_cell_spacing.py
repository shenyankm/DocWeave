"""Empty cell paragraphs retain their source line-spacing constraints."""

from io import BytesIO

import pymupdf
import pytest
from docx import Document as WordDocument
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

from aspose.words_foss import Document
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.pdf_writer.constants import POST_TABLE_SPACING_MM, PT_TO_MM


def empty_table(paragraphs, padding=0):
    table = ldm.Table(
        top_padding=padding,
        bottom_padding=padding,
        rows=[
            ldm.Row(
                row_format=ldm.RowFormat(
                    borders=[ldm.Border(line_style=1, line_width=0.5) for _ in range(6)]
                ),
                cells=[ldm.Cell(paragraphs=paragraphs)],
            )
        ],
    )
    model = ldm.Document(sections=[ldm.Section(body=ldm.Body(children=[table]))])
    return model, table


def rendered_height(writer, model):
    with pymupdf.open(stream=writer.write_to_bytes(model), filetype="pdf") as pdf:
        assert len(pdf) == 1
        assert not pdf[0].get_text().strip()
        rects = [drawing["rect"] for drawing in pdf[0].get_drawings()]
        return max(rect.y1 for rect in rects) - min(rect.y0 for rect in rects)


@pytest.mark.parametrize(
    "rule,spacing,expected",
    [
        (1, 8, 8),
        (1, 30, 30),
        (0, 30, 30),
        (2, 24, 30.8),
        (2, 0, 15.4),
    ],
)
@pytest.mark.parametrize("padding", [0, 2])
def test_empty_cell_applies_line_spacing_and_shared_measurement(
    rule, spacing, expected, padding
):
    para = ldm.Paragraph(
        paragraph_format=ldm.ParagraphFormat(
            line_spacing_rule=rule,
            line_spacing=spacing,
            space_before=3,
            space_after=4,
        )
    )
    model, table = empty_table([para], padding)
    original = model.model_dump_json()
    writer = LdmPdfWriter()
    height = rendered_height(writer, model)
    assert height == pytest.approx(expected + 7 + padding * 2, abs=0.02)
    assert writer._estimate_table_height(table, 180 * PT_TO_MM) == pytest.approx(
        height * PT_TO_MM + POST_TABLE_SPACING_MM, abs=0.02
    )
    assert model.model_dump_json() == original


@pytest.mark.parametrize("rule", [WD_LINE_SPACING.EXACTLY, WD_LINE_SPACING.AT_LEAST])
def test_public_docx_empty_paragraph_preserves_fixed_spacing(rule):
    source = WordDocument()
    para = source.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = Pt(30)
    para.paragraph_format.line_spacing_rule = rule
    stream = BytesIO()
    source.save(stream)
    stream.seek(0)
    loaded = Document(stream).light_document_model
    table = loaded.sections[0].body.tables[0]
    assert not table.rows[0].cells[0].paragraphs[0]._children
    # Use an independent bordered model to expose the measured empty line.
    model, _ = empty_table(table.rows[0].cells[0].paragraphs)
    assert rendered_height(LdmPdfWriter(), model) == pytest.approx(30, abs=0.02)


def test_multiple_empty_paragraphs_each_reserve_their_fixed_height():
    model, _ = empty_table(
        [
            ldm.Paragraph(
                paragraph_format=ldm.ParagraphFormat(
                    line_spacing_rule=1, line_spacing=size
                )
            )
            for size in [8, 20, 30]
        ]
    )
    assert rendered_height(LdmPdfWriter(), model) == pytest.approx(58, abs=0.02)
