"""Independent consumers resolve the same character style in every output story."""

from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
import pytest

from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.docx_reader.constants import PAGE_FIELD_SENTINEL

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


@pytest.mark.parametrize('location', ['body', 'table', 'header', 'footer'])
def test_colliding_character_styles_resolve_in_all_stories(location):
    runs = [ldm.Run(text='FIRST', font=ldm.Font(style_name='Brand Mark')),
            ldm.Run(text='SECOND', font=ldm.Font(style_name='BrandMark'))]
    para = ldm.Paragraph(children=runs)
    section = ldm.Section()
    part = 'word/document.xml'
    if location == 'body':
        section.body.children = [para]
    elif location == 'table':
        section.body.children = [ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[para])])])]
    else:
        kind = 0 if location == 'header' else 1
        section.headers_footers = [ldm.HeaderFooter(header_footer_type=kind, children=[para])]
        part = f'word/{location}1.xml'
    doc = ldm.Document(sections=[section], styles=[
        ldm.Style(name='Brand Mark', type=2, font=ldm.Font(bold=True)),
        ldm.Style(name='BrandMark', type=2, font=ldm.Font(italic=True))])
    raw = LdmDocxWriter().write_to_bytes(doc)
    independent = Document(BytesIO(raw))
    if location == 'body':
        actual = independent.paragraphs[0].runs
    elif location == 'table':
        actual = independent.tables[0].cell(0, 0).paragraphs[0].runs
    else:
        story = getattr(independent.sections[0], location)
        actual = next(p for p in story.paragraphs if p.text).runs
    assert [r.style.name for r in actual] == ['Brand Mark', 'BrandMark']
    with ZipFile(BytesIO(raw)) as package:
        xml = ET.fromstring(package.read(part))
    references = [x.get(W+'val') for x in xml.iter(W+'rStyle')]
    assert len(set(references)) == 2


@pytest.mark.parametrize('text', ['[LABEL](https://example.com)\tTAIL', PAGE_FIELD_SENTINEL])
def test_link_suffix_and_page_field_use_resolved_character_style(text):
    doc = ldm.Document(styles=[ldm.Style(name='Brand Mark', type=2),
                              ldm.Style(name='BrandMark', type=2)],
        sections=[ldm.Section(body=ldm.Body(children=[ldm.Paragraph(children=[
            ldm.Run(text=text, font=ldm.Font(style_name='BrandMark'))])]))])
    raw = LdmDocxWriter().write_to_bytes(doc)
    with ZipFile(BytesIO(raw)) as package:
        styles = ET.fromstring(package.read('word/styles.xml'))
        body = ET.fromstring(package.read('word/document.xml'))
    wanted = next(s.get(W+'styleId') for s in styles.iter(W+'style')
                  if s.find(W+'name').get(W+'val') == 'BrandMark')
    references = [r.get(W+'val') for r in body.iter(W+'rStyle')]
    assert references and set(references) == {wanted}
    if text != PAGE_FIELD_SENTINEL:
        assert len(references) == 2


@pytest.mark.parametrize('text', ['[LABEL](https://example.com)\tTAIL', PAGE_FIELD_SENTINEL])
def test_link_suffix_and_page_field_do_not_freeze_inherited_font(text):
    doc = ldm.Document(styles=[
        ldm.Style(name='Brand Base', type=2, font=ldm.Font(size=16, color='Blue')),
        ldm.Style(name='Brand Child', type=2, base_style_name='Brand Base')],
        sections=[ldm.Section(body=ldm.Body(children=[ldm.Paragraph(children=[
            ldm.Run(text=text, font=ldm.Font(style_name='Brand Child', size=16, color='Blue'))])]))])
    raw = LdmDocxWriter().write_to_bytes(doc)
    with ZipFile(BytesIO(raw)) as package:
        body = ET.fromstring(package.read('word/document.xml'))
    properties = [rpr for rpr in body.iter(W+'rPr') if rpr.find(W+'rStyle') is not None]
    assert properties
    assert all(rpr.find(W+'sz') is None and rpr.find(W+'color') is None for rpr in properties)
