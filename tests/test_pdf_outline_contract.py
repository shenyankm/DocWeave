"""Reject malformed navigation requests before rendering or allocating gaps."""

from io import BytesIO
from types import MappingProxyType

import pytest
from docx import Document as NativeDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader

import aspose.words_foss as aw
from aspose.words_foss.pdf_writer import LdmPdfWriter

INVALID = [
    (field, value)
    for field in ("headings_outline_levels", "expanded_outline_levels", "default_bookmarks_outline_level")
    for value in (-1, True, False, None, "3", 1.5, float("nan"))
] + [
    ("expanded_outline_levels", 10), ("default_bookmarks_outline_level", 10**9),
    ("bookmarks_outline_levels", None), ("bookmarks_outline_levels", []),
    ("bookmarks_outline_levels", {1: 1}),
    *[("bookmarks_outline_levels", {"Unused": value}) for value in (-1, True, "3", 1.5, 10**9)],
    *[(field, value) for field in ("create_missing_outline_levels", "create_outlines_for_headings_in_tables")
      for value in (None, "false", 1)],
]


@pytest.mark.parametrize("field,value", INVALID)
@pytest.mark.parametrize("entry", ["memory", "path", "writer"])
def test_invalid_outline_options_fail_before_rendering(tmp_path, monkeypatch, field, value, entry):
    doc = aw.Document(BytesIO(b"Untouched body"), aw.MarkdownLoadOptions())
    options = aw.saving.PdfSaveOptions()
    writer = LdmPdfWriter(options)
    options.outline_options.create_missing_outline_levels = True
    setattr(options.outline_options, field, value)
    target = tmp_path / "original.pdf"
    target.write_bytes(b"ORIGINAL")

    def unexpected_render(*args, **kwargs):
        pytest.fail("Invalid navigation options reached the renderer")

    # Guard also makes the billion-level baseline safe to reproduce.
    monkeypatch.setattr(LdmPdfWriter, "_render_pdf", unexpected_render)
    with pytest.raises(ValueError, match=field):
        if entry == "memory":
            doc.to_bytes(options)
        elif entry == "path":
            doc.save(target, options)
        else:
            writer.write_to_bytes(doc.light_document_model)
    assert target.read_bytes() == b"ORIGINAL"
    assert list(tmp_path.iterdir()) == [target]


def native_document():
    doc = NativeDocument()
    doc.add_heading("Top", level=1)
    doc.add_heading("Deep", level=6)
    for index, name in enumerate(("Visible", "Skip", "_Hidden")):
        p = doc.add_paragraph()
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(index))
        start.set(qn("w:name"), name)
        p._p.append(start)
        p.add_run("Body " + name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(index))
        p._p.append(end)
    source = BytesIO()
    doc.save(source)
    source.seek(0)
    return aw.Document(source)


def outline_entries(reader):
    def visit(items, depth=1):
        for item in items:
            if isinstance(item, list):
                yield from visit(item, depth + 1)
            else:
                assert reader.get_destination_page_number(item) == 0
                yield depth, item.title
    return list(visit(reader.outline))


@pytest.mark.parametrize("fill", [False, True])
@pytest.mark.parametrize("headings", [6, 10**100])
def test_valid_real_docx_navigation_keeps_body_and_bounded_depth(fill, headings):
    doc = native_document()
    options = aw.saving.PdfSaveOptions()
    oo = options.outline_options
    oo.headings_outline_levels = headings
    oo.default_bookmarks_outline_level = 9
    oo.bookmarks_outline_levels = MappingProxyType({"Skip": 0})
    oo.create_missing_outline_levels = fill
    oo.expanded_outline_levels = 9
    reader = PdfReader(BytesIO(doc.to_bytes(options)), strict=True)
    entries = outline_entries(reader)
    assert [(depth, title) for depth, title in entries if title] == (
        [(1, "Top"), (6, "Deep"), (9, "Visible")] if fill else [(1, "Top"), (2, "Deep"), (3, "Visible")]
    )
    assert max(depth for depth, title in entries) <= 9
    text = reader.pages[0].extract_text()
    for body in ("Top", "Deep", "Body Visible", "Body Skip", "Body _Hidden"):
        assert text.count(body) == 1


def test_mapping_mutation_is_validated_on_each_save_and_writer_recovers():
    options = aw.saving.PdfSaveOptions()
    writer = LdmPdfWriter(options)
    doc = native_document()
    levels = options.outline_options.bookmarks_outline_levels
    levels["Visible"] = 10**9
    with pytest.raises(ValueError, match="bookmarks_outline_levels"):
        writer.write_to_bytes(doc.light_document_model)
    levels["Visible"] = 1
    reader = PdfReader(BytesIO(writer.write_to_bytes(doc.light_document_model)), strict=True)
    assert outline_entries(reader) == [(1, "Visible")]
