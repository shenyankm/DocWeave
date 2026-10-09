"""Failed Markdown exports roll back newly created external image files."""

import os
import re
import threading
import warnings
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as NativeDocument
from PIL import Image

import aspose.words_foss as aw
from aspose.words_foss.md_writer import LdmMarkdownWriter
from aspose.words_foss.models import ConversionOptions


def document(two_images=False, color="blue"):
    native = NativeDocument()
    image = BytesIO()
    Image.new("RGB", (8, 8), color).save(image, format="PNG")
    native.add_picture(BytesIO(image.getvalue()))
    if two_images:
        second = BytesIO()
        Image.new("RGB", (8, 8), "red").save(second, format="PNG")
        native.add_picture(BytesIO(second.getvalue()))
    native.add_paragraph("中文")
    table = native.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "MERGED"
    source = BytesIO()
    native.save(source)
    return aw.Document(BytesIO(source.getvalue()))


@pytest.mark.parametrize("failure", ["encoding", "unknown_encoding", "strict", "main_write"])
@pytest.mark.parametrize("existing", [False, True])
def test_failed_save_removes_only_new_assets(tmp_path, monkeypatch, failure, existing):
    images = tmp_path / "assets"
    if existing:
        images.mkdir()
        (images / "image.png").write_bytes(b"KEEP")
    output = tmp_path / "out.md"
    output.write_bytes(b"ORIGINAL")
    opts = aw.saving.MarkdownSaveOptions()
    opts.images_folder = str(images)
    if failure == "encoding":
        opts.encoding = "ascii"
    elif failure == "unknown_encoding":
        opts.encoding = "not-an-encoding"
    elif failure == "main_write":
        original = os.replace

        def replace(source, destination):
            if destination == output:
                raise OSError("main publication failed")
            return original(source, destination)

        monkeypatch.setattr(os, "replace", replace)
    with warnings.catch_warnings():
        warnings.simplefilter("error" if failure == "strict" else "ignore", aw.ContentLossWarning)
        with pytest.raises((UnicodeError, LookupError, aw.ContentLossWarning, OSError)):
            document().save(output, opts)
    assert output.read_bytes() == b"ORIGINAL"
    assert {p.name: p.read_bytes() for p in images.glob("*")} == ({"image.png": b"KEEP"} if existing else {})
    assert not list(tmp_path.rglob(".*"))


def test_direct_writer_failure_rolls_back_and_writer_can_be_reused(tmp_path):
    opts = ConversionOptions(images_folder=str(tmp_path / "assets"))
    writer = LdmMarkdownWriter(opts)
    doc = document()
    with warnings.catch_warnings():
        warnings.simplefilter("error", aw.ContentLossWarning)
        with pytest.raises(aw.ContentLossWarning):
            writer.write(doc.light_document_model, tmp_path / "out.md")
    assert not list((tmp_path / "assets").glob("*"))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", aw.ContentLossWarning)
        result = writer.write(doc.light_document_model, tmp_path / "out.md")
    assert "assets/image.png" in result
    assert list((tmp_path / "assets").glob("*.png"))


def test_reused_existing_image_is_preserved_on_failure(tmp_path):
    images = tmp_path / "assets"
    images.mkdir()
    image = BytesIO()
    Image.new("RGB", (8, 8), "blue").save(image, format="PNG")
    existing = images / "image.png"
    existing.write_bytes(image.getvalue())
    opts = aw.saving.MarkdownSaveOptions()
    opts.images_folder = str(images)
    with warnings.catch_warnings():
        warnings.simplefilter("error", aw.ContentLossWarning)
        with pytest.raises(aw.ContentLossWarning):
            document().save(tmp_path / "out.md", opts)
    assert existing.read_bytes() == image.getvalue()
    assert list(images.iterdir()) == [existing]


def test_later_image_write_failure_rolls_back_prior_image(tmp_path, monkeypatch):
    second = BytesIO()
    Image.new("RGB", (8, 8), "red").save(second, format="PNG")
    original = Path.write_bytes

    def write_bytes(path, data):
        if data == second.getvalue():
            raise OSError("second image write failed")
        return original(path, data)

    monkeypatch.setattr(Path, "write_bytes", write_bytes)
    opts = aw.saving.MarkdownSaveOptions()
    opts.images_folder = str(tmp_path / "assets")
    with pytest.raises(OSError, match="second image"):
        document(two_images=True).save(tmp_path / "out.md", opts)
    assert not list((tmp_path / "assets").glob("*"))
    assert not (tmp_path / "out.md").exists()


def test_other_export_does_not_reuse_an_image_that_can_roll_back(tmp_path, monkeypatch):
    failed, successful = document(), document()
    options = aw.saving.MarkdownSaveOptions()
    options.images_folder = str(tmp_path / "assets")
    options.encoding = "ascii"
    other_options = aw.saving.MarkdownSaveOptions()
    other_options.images_folder = options.images_folder
    render = failed._render_markdown

    def interleaved(*args, **kwargs):
        result = render(*args, **kwargs)
        successful.save(tmp_path / "success.md", other_options)
        return result

    monkeypatch.setattr(failed, "_render_markdown", interleaved)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", aw.ContentLossWarning)
        with pytest.raises(UnicodeError):
            failed.save(tmp_path / "failure.md", options)
    result = (tmp_path / "success.md").read_text()
    assert "assets/image_2.png" in result
    assert (tmp_path / "assets" / "image_2.png").is_file()
    assert not (tmp_path / "assets" / "image.png").exists()
    assert not list(tmp_path.rglob(".*"))


def test_simultaneous_exports_claim_distinct_assets(tmp_path, monkeypatch):
    barrier = threading.Barrier(2)
    state = threading.local()
    original = os.open

    def open_file(path, *args, **kwargs):
        if str(path).endswith(".pending") and not getattr(state, "waited", False):
            state.waited = True
            barrier.wait(timeout=10)
        return original(path, *args, **kwargs)

    monkeypatch.setattr(os, "open", open_file)

    def export(color):
        opts = aw.saving.MarkdownSaveOptions()
        opts.images_folder = str(tmp_path / "assets")
        doc = document(color=color)
        path = tmp_path / f"{color}.md"
        doc.save(path, opts)
        image_path = tmp_path / re.search(r"\]\(([^)]+)\)", path.read_text()).group(1)
        return image_path

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", aw.ContentLossWarning)
        with ThreadPoolExecutor(max_workers=2) as pool:
            paths = list(pool.map(export, ("red", "blue")))
    colors = []
    for path in paths:
        with Image.open(path) as image:
            colors.append(image.convert("RGB").getpixel((0, 0)))
    assert colors == [(255, 0, 0), (0, 0, 255)]
    assert not list(tmp_path.rglob(".*"))


def test_claim_close_error_removes_pending_marker(tmp_path, monkeypatch):
    open_file, close_file = os.open, os.close
    claims = set()

    def opened(path, *args, **kwargs):
        fd = open_file(path, *args, **kwargs)
        if str(path).endswith(".pending"):
            claims.add(fd)
        return fd

    def closed(fd):
        close_file(fd)
        if fd in claims:
            claims.remove(fd)
            raise OSError("claim close failed")

    monkeypatch.setattr(os, "open", opened)
    monkeypatch.setattr(os, "close", closed)
    opts = aw.saving.MarkdownSaveOptions()
    opts.images_folder = str(tmp_path / "assets")
    with pytest.raises(OSError, match="claim close"):
        document().save(tmp_path / "out.md", opts)
    assert not list((tmp_path / "assets").iterdir())
