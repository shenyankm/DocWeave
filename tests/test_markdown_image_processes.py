"""Independent installed Python processes must not share uncommitted assets."""

import json
import os
import subprocess
import sys
import time
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as NativeDocument
from PIL import Image

import aspose.words_foss as aw

WORKER = r'''
import json, os, sys, time
from pathlib import Path
import aspose.words_foss as aw
mode, source, output, images, ready, release = sys.argv[1:]
ready, release = Path(ready), Path(release)
doc = aw.Document(source)
opts = aw.saving.MarkdownSaveOptions()
opts.images_folder = images
if mode == "rollback":
    opts.encoding = "ascii"
def pause():
    ready.write_text(str(os.getpid()))
    deadline = time.monotonic() + 20
    while not release.exists():
        if time.monotonic() > deadline:
            raise TimeoutError("process release missing")
        time.sleep(0.01)
if mode == "claim":
    opened = os.open
    waited = False
    def open_file(path, *args, **kwargs):
        global waited
        fd = opened(path, *args, **kwargs)
        if str(path).endswith(".pending") and not waited:
            waited = True
            pause()
        return fd
    os.open = open_file
elif mode in ("rollback", "interrupted"):
    render = doc._render_markdown
    def held_render(*args, **kwargs):
        result = render(*args, **kwargs)
        pause()
        return result
    doc._render_markdown = held_render
try:
    doc.save(output, opts)
except UnicodeError:
    print(json.dumps({"pid": os.getpid(), "module": aw.__file__, "result": "encoding_error"}))
    sys.exit(3)
print(json.dumps({"pid": os.getpid(), "module": aw.__file__, "result": "success"}))
'''


def input_docx(path, color):
    native = NativeDocument()
    image = BytesIO()
    Image.new("RGB", (8, 8), color).save(image, format="PNG")
    native.add_picture(BytesIO(image.getvalue()))
    native.add_paragraph("中文")
    native.save(path)


@pytest.mark.parametrize("mode", ["claim", "rollback", "interrupted"])
def test_independent_processes_preserve_shared_directory_outputs(tmp_path, mode):
    first_source, second_source = tmp_path / "first.docx", tmp_path / "second.docx"
    input_docx(first_source, "red" if mode == "claim" else "blue")
    input_docx(second_source, "blue")
    first_output, second_output = tmp_path / "first.md", tmp_path / "second.md"
    first_output.write_bytes(b"ORIGINAL")
    images, ready, release = tmp_path / "assets", tmp_path / "ready", tmp_path / "release"
    env = os.environ.copy()
    env["PYTHONPATH"] = os.pathsep.join([str(Path(aw.__file__).parents[2]), env.get("PYTHONPATH", "")])
    first = subprocess.Popen([sys.executable, "-c", WORKER, mode, str(first_source),
        str(first_output), str(images), str(ready), str(release)], cwd=tmp_path,
        env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        deadline = time.monotonic() + 15
        while not ready.exists() or not ready.read_text().strip():
            assert first.poll() is None, first.communicate(timeout=1)
            assert time.monotonic() < deadline, "first process did not reach the held operation"
            time.sleep(0.01)
        if mode == "interrupted":
            first.kill()
            first.communicate(timeout=10)
        second = subprocess.run([sys.executable, "-c", WORKER, "normal", str(second_source),
            str(second_output), str(images), str(ready), str(release)], cwd=tmp_path,
            env=env, capture_output=True, text=True, timeout=20, check=False)
        assert second.returncode == 0, second.stderr
        report = json.loads(second.stdout)
        assert report["pid"] != int(ready.read_text())
        assert Path(report["module"]).resolve() == Path(aw.__file__).resolve()
        assert "assets/image_2.png" in second_output.read_text()
        with Image.open(images / "image_2.png") as image:
            assert image.convert("RGB").getpixel((0, 0)) == (0, 0, 255)
        release.write_text("continue")
        if mode != "interrupted":
            stdout, stderr = first.communicate(timeout=20)
            assert first.returncode == (3 if mode == "rollback" else 0), stderr
            assert json.loads(stdout)["result"] == ("encoding_error" if mode == "rollback" else "success")
        if mode == "claim":
            assert "assets/image.png" in first_output.read_text()
            with Image.open(images / "image.png") as image:
                assert image.convert("RGB").getpixel((0, 0)) == (255, 0, 0)
        else:
            assert first_output.read_bytes() == b"ORIGINAL"
        if mode == "rollback":
            assert not (images / "image.png").exists()
        if mode == "interrupted":
            assert (images / ".image.png.pending").exists()
            assert (images / "image.png").exists()
        else:
            assert not list(images.glob(".*"))
    finally:
        if first.poll() is None:
            first.kill()
        first.communicate(timeout=10)
