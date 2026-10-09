"""Optional live competitors validate our output; no mocked engines or core dependencies."""

from html.parser import HTMLParser
from io import BytesIO
from zipfile import ZipFile
import warnings

from docx import Document
import pytest

import aspose.words_foss as aw
from .test_content_integrity import source_package
from .test_docx_markdown_corpus import corpus
from .test_markdown_notes import options
from .test_reference_styles import reference

mammoth = pytest.importorskip('mammoth', reason='Install mammoth for independent ecosystem checks')
pypandoc = pytest.importorskip('pypandoc', reason='Install pypandoc and Pandoc for independent checks')
try:
    pypandoc.get_pandoc_version()
except OSError:
    pytest.skip('A runnable Pandoc executable is required', allow_module_level=True)


class HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data)

    @classmethod
    def read(cls, html):
        parser = cls()
        parser.feed(html)
        return ' '.join(' '.join(parser.parts).split())


@pytest.mark.parametrize('case,tokens', [
    ('text_and_emphasis', ['中文报告', '普通', '粗体', '粗斜体', 'Literal *stars* _under_ [brackets]']),
    ('business_style', ['业务章节', '第一项', '第二项']),
    ('merged_table', ['表前', 'Header left', 'Header right', 'A | B', '中文表格', '表后']),
    ('links', ['链接 [说明]']),
])
def test_competitors_agree_on_common_docx_content(tmp_path, case, tokens):
    raw = corpus(case)
    source = tmp_path / 'source.docx'
    source.write_bytes(raw)
    own = aw.Document(BytesIO(raw)).get_text()
    other = mammoth.convert_to_html(BytesIO(raw), style_map="p[style-name='业务标题'] => h2:fresh").value
    pandoc = pypandoc.convert_file(str(source), 'html')
    assert all(token in own and token in HTMLText.read(other) and token in HTMLText.read(pandoc)
               for token in tokens)
    if case == 'links':
        assert 'https://example.com/a_(b)?x=1&amp;y=2' in other
        assert 'https://example.com/a_(b)?x=1&amp;y=2' in pandoc


def standard_note_package():
    with ZipFile(BytesIO(source_package())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts['word/footnotes.xml'] = parts.pop('notes/fn text.xml')
    parts['word/_rels/document.xml.rels'] = parts['word/_rels/document.xml.rels'].replace(
        b'../notes/fn%20text.xml', b'footnotes.xml')
    parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace(
        b'/notes/fn text.xml', b'/word/footnotes.xml')
    parts['word/document.xml'] = parts['word/document.xml'].replace(
        b'<w:t>body</w:t><w:footnoteReference', b'<w:t>body</w:t></w:r><w:r><w:footnoteReference').replace(
        b'w:id="7"/><w:endnoteReference', b'w:id="7"/></w:r><w:r><w:endnoteReference')
    output = BytesIO()
    with ZipFile(output, 'w') as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return output.getvalue()


def test_note_content_and_references_survive_independent_markdown_parser(tmp_path):
    raw = standard_note_package()
    source = tmp_path / 'source.docx'
    source.write_bytes(raw)
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', aw.ContentLossWarning)
        own = aw.Document(BytesIO(raw)).to_bytes(options()).decode()
    mammoth_html = mammoth.convert_to_html(BytesIO(raw)).value
    pandoc_md = pypandoc.convert_file(str(source), 'markdown')
    own_html = pypandoc.convert_text(own, 'html', format='markdown')
    for value in (own, mammoth_html, pandoc_md, own_html):
        assert all(token in value for token in ('body', 'second', 'footnote text', 'endnote text'))
    assert own.index('body[^note1][^note2]') < own.index('second') < own.index('[^note1]: footnote text')
    assert own.index('[^note1]: footnote text') < own.index('[^note2]: endnote text')
    assert own_html.count('role="doc-noteref"') == 2
    assert own_html.count('role="doc-backlink"') == 2


def test_reference_heading_formatting_agrees_with_pandoc(tmp_path):
    brand = reference(tmp_path)
    before = brand.read_bytes()
    markdown = '# Report\n\nBody\n'
    opts = aw.saving.OoxmlSaveOptions()
    opts.reference_docx = brand
    own = tmp_path / 'own.docx'
    aw.Document(BytesIO(markdown.encode()), aw.MarkdownLoadOptions()).save(own, opts)
    other = tmp_path / 'pandoc.docx'
    pypandoc.convert_text(markdown, 'docx', format='markdown', outputfile=str(other),
                         extra_args=['--reference-doc=' + str(brand)])
    records = []
    for path in (own, other):
        independent = Document(path)
        effective_size = aw.Document(path).light_document_model.sections[0].body.paragraphs[0].runs[0].font.size
        records.append((effective_size, str(independent.paragraphs[0].style.font.color.rgb),
                        independent.styles['Normal'].font.name,
                        ' '.join(p.text for p in independent.paragraphs)))
    assert records[0] == records[1] == (30, '008800', 'Courier New', 'Report Body')
    assert brand.read_bytes() == before
