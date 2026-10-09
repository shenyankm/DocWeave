"""PDF compression preserves AUTO transparency and validates JPEG quality."""

import warnings
from io import BytesIO

import pymupdf
import pytest
from docx import Document as NativeDocument
from PIL import Image
from pypdf import PdfReader

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter


def picture(mode):
    if mode == "RGBA":
        image = Image.new(mode, (20, 20), (255, 0, 0, 0))
        for x in range(10, 20):
            for y in range(20):
                image.putpixel((x, y), (0, 0, 255, 128))
    elif mode == "LA":
        image = Image.new(mode, (20, 20), (0, 0))
        for x in range(10, 20):
            for y in range(20):
                image.putpixel((x, y), (0, 128))
    else:
        image = Image.new("P", (20, 20), 0)
        image.putpalette([255, 0, 0, 0, 0, 255] + [0] * 762)
        image.info["transparency"] = bytes([0, 128])
        for x in range(10, 20):
            for y in range(20):
                image.putpixel((x, y), 1)
    data = BytesIO()
    image.save(data, format="PNG")
    return data.getvalue()


def document(mode="RGBA", location="body"):
    native = NativeDocument()
    native.add_picture(BytesIO(picture(mode)))
    source = BytesIO()
    native.save(source)
    doc = aw.Document(BytesIO(source.getvalue()))
    paragraph = doc.sections[0].body.children[0]
    if location == "table":
        doc.sections[0].body.children = [ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[paragraph])])])]
    elif location in ("header", "footer"):
        doc.sections[0].body.children = []
        doc.sections[0].headers_footers = [ldm.HeaderFooter(
            header_footer_type=0 if location == "header" else 1, children=[paragraph])]
    elif location == "anchored":
        next(child for child in paragraph._children if isinstance(child, ldm.Shape)).is_inline = False
    return doc


@pytest.mark.parametrize("mode", ["RGBA", "LA", "P"])
@pytest.mark.parametrize("location", ["body", "table", "header", "footer", "anchored"])
def test_auto_low_quality_preserves_transparency(mode, location):
    doc = document(mode, location)
    before = doc.light_document_model.model_dump()
    options = aw.saving.PdfSaveOptions()
    options.jpeg_quality = 80
    actual = doc.to_bytes(options)
    reference = doc.to_bytes("pdf")
    with pymupdf.open(stream=actual, filetype="pdf") as pdf, pymupdf.open(stream=reference, filetype="pdf") as clean:
        assert len(pdf) == len(clean)
        for page, other in zip(pdf, clean):
            assert page.get_pixmap().samples == other.get_pixmap().samples
            assert any(item[1] for item in page.get_images())
    assert not any(item.code == "pdf.image_transparency_lost" for item in doc.diagnostics)
    assert doc.light_document_model.model_dump() == before


@pytest.mark.parametrize("mode", ["RGBA", "LA", "P"])
@pytest.mark.parametrize("location", ["body", "table", "header", "footer", "anchored"])
def test_explicit_jpeg_flattens_against_white_and_diagnoses_loss(mode, location):
    doc = document(mode, location)
    options = aw.saving.PdfSaveOptions()
    options.image_compression = aw.saving.PdfImageCompression.JPEG
    with pytest.warns(aw.ContentLossWarning, match="transparency"):
        raw = doc.to_bytes(options)
    assert any(item.code == "pdf.image_transparency_lost" for item in doc.diagnostics)
    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        images = pdf[0].get_images()
        assert images and all(item[1] == 0 for item in images)
        with Image.open(BytesIO(pdf.extract_image(images[0][0])["image"])) as image:
            assert min(image.convert("RGB").getpixel((2, 2))) >= 240


def test_strict_jpeg_loss_preserves_original_output(tmp_path):
    target = tmp_path / "original.pdf"
    target.write_bytes(b"ORIGINAL")
    doc = document()
    options = aw.saving.PdfSaveOptions()
    options.image_compression = aw.saving.PdfImageCompression.JPEG
    with warnings.catch_warnings():
        warnings.simplefilter("error", aw.ContentLossWarning)
        with pytest.raises(aw.ContentLossWarning):
            doc.save(target, options)
    assert target.read_bytes() == b"ORIGINAL"
    assert list(tmp_path.iterdir()) == [target]


@pytest.mark.parametrize("quality", [-1, 101, 1.5, "80", None, True, float("nan"), float("inf")])
@pytest.mark.parametrize("entry", ["memory", "path", "writer"])
def test_invalid_quality_rejected_before_output(tmp_path, quality, entry):
    doc = aw.Document(BytesIO(b"BODY"), aw.MarkdownLoadOptions())
    options = aw.saving.PdfSaveOptions()
    options.jpeg_quality = quality
    target = tmp_path / "original.pdf"
    target.write_bytes(b"ORIGINAL")
    with pytest.raises(ValueError, match="jpeg_quality"):
        if entry == "memory":
            doc.to_bytes(options)
        elif entry == "path":
            doc.save(target, options)
        else:
            LdmPdfWriter(options).write_to_bytes(doc.light_document_model)
    assert target.read_bytes() == b"ORIGINAL"
    assert list(tmp_path.iterdir()) == [target]


@pytest.mark.parametrize("quality", [0, 1, 80, 100])
def test_valid_jpeg_quality_embeds_jpeg(quality):
    doc = document()
    options = aw.saving.PdfSaveOptions()
    options.image_compression = aw.saving.PdfImageCompression.JPEG
    options.jpeg_quality = quality
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", aw.ContentLossWarning)
        raw = doc.to_bytes(options)
    reader = PdfReader(BytesIO(raw))
    resources = reader.pages[0]["/Resources"]["/XObject"]
    assert any(obj.get_object().get("/Filter") == "/DCTDecode" for obj in resources.values())


@pytest.mark.parametrize("compression", [aw.saving.PdfImageCompression.AUTO, aw.saving.PdfImageCompression.JPEG])
def test_opaque_alpha_image_still_uses_jpeg_without_loss_diagnostic(compression):
    doc = document()
    data = BytesIO()
    Image.new("RGBA", (20, 20), (0, 0, 255, 255)).save(data, format="PNG")
    shape = next(child for child in doc.sections[0].body.children[0]._children if isinstance(child, ldm.Shape))
    shape.image_data.image_bytes = data.getvalue()
    options = aw.saving.PdfSaveOptions()
    options.image_compression = compression
    options.jpeg_quality = 80
    raw = doc.to_bytes(options)
    assert not any(item.code == "pdf.image_transparency_lost" for item in doc.diagnostics)
    resources = PdfReader(BytesIO(raw)).pages[0]["/Resources"]["/XObject"]
    assert any(obj.get_object().get("/Filter") == "/DCTDecode" for obj in resources.values())
