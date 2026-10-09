"""Code paragraphs retain source formatting and paint each row after pagination."""

import pymupdf
import pytest

from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.pdf_writer.constants import CODE_BLOCK_BG_RGB


def code_model(runs, *, columns=1, align=0):
    paragraph = ldm.Paragraph(children=runs, paragraph_format=ldm.ParagraphFormat(
        style_name="Code", alignment=align, widow_control=False))
    return ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
        page_width=260, page_height=180, left_margin=20, right_margin=20,
        top_margin=20, bottom_margin=20,
        text_columns=ldm.TextColumns(count=columns, spacing=12)),
        body=ldm.Body(children=[paragraph]))])


def backgrounds(page):
    expected = tuple(value / 255 for value in CODE_BLOCK_BG_RGB)
    return [drawing["rect"] for drawing in page.get_drawings()
            if drawing["fill"] is not None and drawing["fill"] == pytest.approx(expected, abs=0.001)]


def test_code_preserves_sizes_colors_styles_and_source():
    model = code_model([ldm.Run(text="Q", font=ldm.Font(size=24, bold=True, color="FF0000")),
                        ldm.Run(text="W", font=ldm.Font(size=12, italic=True, color="0000FF")),
                        ldm.Run(text="E", font=ldm.Font(size=0))])
    snapshot = model.model_dump()
    with pymupdf.open(stream=LdmPdfWriter().write_to_bytes(model), filetype="pdf") as pdf:
        spans = [span for block in pdf[0].get_text("dict")["blocks"]
                 for line in block.get("lines", []) for span in line["spans"]]
        chars = {span["text"]: span for span in spans}
        assert chars["Q"]["size"] == pytest.approx(24)
        assert chars["W"]["size"] == pytest.approx(12)
        assert chars["E"]["size"] == pytest.approx(10)
        assert chars["Q"]["color"] == 0xFF0000
        assert chars["W"]["color"] == 0x0000FF
        assert "Bold" in chars["Q"]["font"]
        assert "Oblique" in chars["W"]["font"]
        assert len(backgrounds(pdf[0])) == 1
    assert model.model_dump() == snapshot


@pytest.mark.parametrize("columns", [1, 2])
def test_code_backgrounds_follow_rows_across_pages_and_columns(columns):
    model = code_model([ldm.Run(text="\n".join(f"ROW{i:03}" for i in range(50)),
                               font=ldm.Font(size=12, highlight_color="FFFF00"))], columns=columns)
    with pymupdf.open(stream=LdmPdfWriter().write_to_bytes(model), filetype="pdf") as pdf:
        assert len(pdf) > 1
        words = [word[4] for page in pdf for word in page.get_text("words")]
        assert words == [f"ROW{i:03}" for i in range(50)]
        assert sum(len(backgrounds(page)) for page in pdf) == 50
        for page in pdf:
            rectangles = backgrounds(page)
            for rectangle in rectangles:
                assert rectangle.y0 >= 20 - 0.05
                assert rectangle.y1 <= page.rect.height - 20 + 0.05
            for word in page.get_text("words"):
                assert any(rectangle.contains(pymupdf.Point(word[0], word[1])) for rectangle in rectangles)


@pytest.mark.parametrize("runs", [[], [ldm.Run(text="")], [ldm.Run(text="\n\n")]])
def test_empty_code_rows_have_background(runs):
    with pymupdf.open(stream=LdmPdfWriter().write_to_bytes(code_model(runs)), filetype="pdf") as pdf:
        assert len(backgrounds(pdf[0])) == (3 if runs and runs[0].text else 1)


def test_public_docx_code_keeps_source_formatting(tmp_path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt, RGBColor
    import aspose.words_foss as aw

    document = Document()
    document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph(style="Code")
    run = paragraph.add_run("SOURCE")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(255, 0, 0)
    script = paragraph.add_run("SUP")
    script.font.size = Pt(24)
    script.font.superscript = True
    source, target = tmp_path / "code.docx", tmp_path / "code.pdf"
    document.save(source)
    aw.Document(source).save(target)
    with pymupdf.open(target) as pdf:
        spans = [span for block in pdf[0].get_text("dict")["blocks"]
                 for line in block.get("lines", []) for span in line["spans"]]
        regular = next(span for span in spans if span["text"] == "SOURCE")
        raised = next(span for span in spans if span["text"] == "SUP")
        assert regular["size"] == pytest.approx(24)
        assert regular["color"] == 0xFF0000
        assert raised["size"] == pytest.approx(16.8, abs=0.05)
        assert raised["origin"][1] == pytest.approx(regular["origin"][1] - 9.6, abs=0.05)
        assert len(backgrounds(pdf[0])) == 1
