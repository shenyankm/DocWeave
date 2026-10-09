"""Directory and padded two-column lines obey the native text inset."""

import pymupdf
import pytest

from aspose.words_foss import Document
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.saving import PdfSaveOptions


def split_line_model(kind, columns, location="body"):
    paragraph = ldm.Paragraph(
        paragraph_format=ldm.ParagraphFormat(
            style_name="TOC 1" if kind == "toc" else "",
            alignment=0 if kind == "toc" else 2,
            left_indent=24,
            right_indent=6,
        ),
        children=[
            ldm.Run(text="TITLE", font=ldm.Font(size=12)),
            ldm.Run(text="\t" if kind == "toc" else " " * 150),
            ldm.Run(text="123", font=ldm.Font(size=16, bold=True)),
        ],
    )
    return ldm.Document(
        sections=[
            ldm.Section(
                page_setup=ldm.PageSetup(
                    page_width=300,
                    page_height=180,
                    left_margin=20,
                    right_margin=20,
                    top_margin=20,
                    bottom_margin=20,
                    text_columns=ldm.TextColumns(count=columns, spacing=15),
                ),
                body=ldm.Body(
                    children=[paragraph]
                    if location == "body"
                    else [ldm.Paragraph(children=[ldm.Run(text="BODY")])]
                ),
                headers_footers=[]
                if location == "body"
                else [
                    ldm.HeaderFooter(
                        header_footer_type=0 if location == "header" else 1,
                        children=[paragraph],
                    )
                ],
            )
        ]
    )


@pytest.mark.parametrize("kind", ["toc", "gap"])
@pytest.mark.parametrize("columns", [1, 2])
@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("location", ["body", "header", "footer"])
def test_split_line_ends_at_effective_right_text_edge(kind, columns, shaping, location):
    model = split_line_model(kind, columns, location)
    original = model.model_dump()
    options = PdfSaveOptions()
    options.text_shaping = shaping
    raw = LdmPdfWriter(options).write_to_bytes(model)
    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        assert len(pdf) == 1
        words = {word[4]: word for word in pdf[0].get_text("words")}
        assert set(words) == (
            {"TITLE", "123"} if location == "body" else {"TITLE", "123", "BODY"}
        )
        assert words["TITLE"][0] == pytest.approx(20 + 24 + 72 / 25.4, abs=0.05)
        width = (260 - 15 * (columns - 1)) / columns if location == "body" else 260
        assert words["123"][2] == pytest.approx(20 + width - 6 - 72 / 25.4, abs=0.05)
        assert words["123"][0] > words["TITLE"][2]
        if location == "body":
            assert words["123"][1] == pytest.approx(20, abs=0.05)
    assert model.model_dump() == original


@pytest.mark.parametrize("kind", ["toc", "gap"])
@pytest.mark.parametrize("shaping", [False, True])
def test_public_docx_split_line_keeps_right_text_inset(tmp_path, kind, shaping):
    from docx import Document as DocxDocument
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    source = DocxDocument()
    section = source.sections[0]
    section.page_width = Pt(300)
    section.page_height = Pt(180)
    section.left_margin = section.right_margin = Pt(20)
    section.top_margin = section.bottom_margin = Pt(20)
    paragraph = source.add_paragraph()
    if kind == "toc":
        paragraph.style = source.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.paragraph_format.right_indent = Pt(6)
    paragraph.add_run("TITLE").font.size = Pt(12)
    paragraph.add_run("\t" if kind == "toc" else " " * 150)
    number = paragraph.add_run("123")
    number.font.size = Pt(16)
    number.bold = True
    path = tmp_path / "split.docx"
    source.save(path)
    options = PdfSaveOptions()
    options.text_shaping = shaping
    raw = Document(path).to_bytes(options)
    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        assert len(pdf) == 1
        words = {word[4]: word for word in pdf[0].get_text("words")}
        assert set(words) == {"TITLE", "123"}
        assert words["123"][2] == pytest.approx(300 - 20 - 6 - 72 / 25.4, abs=0.05)
        assert words["123"][0] > words["TITLE"][2]
