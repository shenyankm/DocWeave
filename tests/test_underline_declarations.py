"""Underline declarations retain inheritance and explicitly cancel it on save."""
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document as SourceDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from dockiln import Document, SaveFormat, light_document_model as ldm
from dockiln.readers.docx import DocumentReader
from dockiln.writers.docx import LdmDocxWriter


def source(scope, value, inherited=False):
    document = SourceDocument()
    paragraph = document.add_paragraph()
    run = paragraph.add_run('Owned underline')
    if scope == 'style':
        style = document.styles.add_style('Owned underline style', WD_STYLE_TYPE.PARAGRAPH)
        paragraph.style = style
    if inherited:
        if scope == 'style':
            base = document.styles.add_style('Owned underline base', WD_STYLE_TYPE.PARAGRAPH)
            base.font.underline = True
            style.base_style = base
        else:
            document.styles['Normal'].font.underline = True
    if scope == 'run':
        properties = run._r.get_or_add_rPr()
    elif scope == 'style':
        properties = style.element.get_or_add_rPr()
    else:
        defaults = document.styles.element.find(qn('w:docDefaults'))
        properties = defaults.find(qn('w:rPrDefault')).find(qn('w:rPr'))
    for child in list(properties):
        if child.tag == qn('w:u'):
            properties.remove(child)
    if value != 'absent':
        underline = OxmlElement('w:u')
        if value is not None:
            underline.set(qn('w:val'), value)
        properties.append(underline)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def font(document):
    return document.light_document_model.sections[0].body.paragraphs[0].runs[0].font


@pytest.mark.parametrize('scope,inherited', [('run', False), ('style', False),
    ('default', False), ('run', True), ('style', True)])
@pytest.mark.parametrize('value,explicit', [(None, None), ('', 0), ('none', 0),
    ('single', 1), ('double', 3), ('absent', None)])
def test_loaded_underline_and_two_format_cold_roundtrip(tmp_path, scope, inherited, value, explicit):
    expected = int(inherited) if explicit is None else explicit
    document = Document(BytesIO(source(scope, value, inherited)))
    assert font(document).underline == expected
    restored = ldm.Document.model_validate_json(document.light_document_model.model_dump_json())
    owner = Document(BytesIO(LdmDocxWriter().write_to_bytes(restored)))
    for format_name in ('DOCX', 'FLAT_OPC'):
        path = tmp_path / format_name
        owner.save(path, getattr(SaveFormat, format_name))
        assert font(Document(path)).underline == expected


@pytest.mark.parametrize('value,expected', [(None, False), ('', False), ('none', False),
    ('single', True), ('double', True), ('absent', False)])
def test_legacy_direct_underline_flag(value, expected):
    reader = DocumentReader()
    with ZipFile(BytesIO(source('run', value))) as package:
        run = ET.fromstring(package.read('word/document.xml')).find('.//' + qn('w:r'))
    assert reader._parse_run(run).underline is expected


def test_writer_explicitly_cancels_inherited_underline():
    model = Document(BytesIO(source('run', 'none', True))).light_document_model
    data = LdmDocxWriter().write_to_bytes(model)
    with ZipFile(BytesIO(data)) as package:
        underline = ET.fromstring(package.read('word/document.xml')).find('.//' + qn('w:rPr') + '/' + qn('w:u'))
    assert underline is not None
    assert underline.get(qn('w:val')) == 'none'
