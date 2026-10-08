"""Independent DOCX producers/extractors check the supported semantic round trip."""

from io import BytesIO
import warnings

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx2python import docx2python
import pytest
from pypdf import PdfReader

import aspose.words_foss as aw
from .test_content_integrity import interleaved_document
from aspose.words_foss.docx_writer import LdmDocxWriter


def corpus(case):
    doc = Document()
    if case == "text_and_emphasis":
        doc.add_heading("中文报告", 1)
        p = doc.add_paragraph("普通 ")
        p.add_run("粗体").bold = True
        p.add_run(" 和 ")
        run = p.add_run("粗斜体")
        run.bold = run.italic = True
        doc.add_paragraph(r"Literal *stars* _under_ [brackets] \ slashes | pipe # hash")
    elif case == "business_style":
        style = doc.styles.add_style("业务标题", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Heading 2"]
        doc.add_paragraph("业务章节", style="业务标题")
        doc.add_paragraph("第一项", style="List Bullet")
        doc.add_paragraph("第二项", style="List Bullet")
    elif case == "merged_table":
        doc.add_paragraph("表前")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header left"
        table.cell(0, 1).text = "Header right"
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(1, 0).text = "A | B"
        table.cell(1, 1).text = "中文表格"
        doc.add_paragraph("表后")
    elif case == "links":
        doc.add_paragraph("链接：")
    else:
        raise AssertionError(case)
    stream = BytesIO()
    doc.save(stream)
    raw = stream.getvalue()
    if case == "links":
        editable = aw.DocxDocument(BytesIO(raw))
        editable.body.paragraphs[0].add_hyperlink("链接 [说明]", "https://example.com/a_(b)?x=1&y=2")
        raw = editable.to_bytes()
    return raw


def semantics(doc):
    def normalize(value):
        return " ".join(value.split())
    def paragraph(block):
        return (normalize(block["text"]), block["heading_level"],
                block["list"]["list_level_number"] if block["list"] else None)
    result = []
    for block in doc.to_dict()["blocks"]:
        if block["type"] == "paragraph" and block["text"].strip():
            result.append(("paragraph", paragraph(block)))
        elif block["type"] == "table":
            grid = []
            for row in block["rows"]:
                cells = []
                for cell in row["cells"]:
                    cells.append(normalize(" ".join(p["text"] for p in cell["paragraphs"])))
                    cells.extend("" for _ in range(cell["grid_span"] - 1))
                grid.append(cells)
            result.append(("table", grid))
    return result


@pytest.mark.parametrize("case", ["text_and_emphasis", "business_style", "merged_table", "links"])
def test_python_docx_corpus_round_trip_preserves_supported_semantics(case):
    raw = corpus(case)
    original = aw.Document(BytesIO(raw))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", aw.ContentLossWarning)
        markdown = original.to_bytes("md")
    reread = aw.Document(BytesIO(markdown), aw.MarkdownLoadOptions())
    assert semantics(reread) == semantics(original)
    final = aw.Document(BytesIO(reread.to_bytes("docx")))
    assert semantics(final) == semantics(original)
    if case == "text_and_emphasis":
        for doc in (original, reread, final):
            runs = doc.to_dict()["blocks"][1]["runs"]
            assert next(run for run in runs if "粗体" in run["text"])["bold"]
            nested = next(run for run in runs if "粗斜体" in run["text"])
            assert nested["bold"] and nested["italic"]
    if case == "links":
        for doc in (original, reread, final):
            link = next(run for run in doc.to_dict()["blocks"][0]["runs"] if run["link"])
            assert link["text"] == "链接 [说明]"
            assert link["link"] == "https://example.com/a_(b)?x=1&y=2"
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", aw.ConversionWarning)
            pdf = PdfReader(BytesIO(original.to_bytes("pdf")))
        annotation = next(item.get_object() for item in pdf.pages[0]["/Annots"])
        assert annotation["/A"]["/URI"] == "https://example.com/a_(b)?x=1&y=2"


@pytest.mark.parametrize("case", ["text_and_emphasis", "business_style", "merged_table", "links"])
def test_independent_extractor_agrees_on_content(case):
    raw = corpus(case)
    doc = aw.Document(BytesIO(raw))
    with docx2python(BytesIO(raw), duplicate_merged_cells=False) as extracted:
        # The extractor includes list labels, so compare content tokens, not renderer formatting.
        for block in doc.to_dict()["blocks"]:
            if block["type"] == "paragraph":
                assert all(run["text"] in extracted.text for run in block["runs"] if run["text"])
            else:
                for row in block["rows"]:
                    for cell in row["cells"]:
                        for paragraph in cell["paragraphs"]:
                            assert paragraph["text"] in extracted.text


@pytest.mark.parametrize("kind", ["text", "docx"])
def test_get_text_preserves_literal_link_syntax(kind):
    text = "Literal [label](https://example.com) and [1](note)"
    if kind == "text":
        raw = text.encode()
    else:
        original = Document()
        original.add_paragraph(text)
        stream = BytesIO()
        original.save(stream)
        raw = stream.getvalue()
    doc = aw.Document(BytesIO(raw))
    assert doc.get_text() == doc.to_dict()["blocks"][0]["text"] == text
    assert not any(run["link"] for run in doc.to_dict()["blocks"][0]["runs"])


def test_independent_extractor_sees_every_interleaved_cell_paragraph():
    raw = LdmDocxWriter().write_to_bytes(interleaved_document())
    with docx2python(BytesIO(raw)) as extracted:
        assert all(value in extracted.text for value in ("before", "nested", "after"))
    assert aw.Document(BytesIO(raw)).get_text().splitlines() == ["before", "nested", "after"]
