"""Explicit zero cell margins survive OOXML and share measured/rendered geometry."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image
import pymupdf
import pytest

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.pdf_writer.constants import POST_TABLE_SPACING_MM, PT_TO_MM
from aspose.words_foss.saving import OoxmlSaveOptions, PdfSaveOptions


def margins(parent, tag, values):
    element = OxmlElement(tag)
    for side, value in values.items():
        child = OxmlElement('w:' + side)
        child.set(qn('w:type'), 'dxa')
        child.set(qn('w:w'), str(round(value * 20)))
        element.append(child)
    parent.append(element)
    return element


def source_docx(*, table=None, cell=None, styled=False):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(220), Pt(180)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Pt(20)
    compat = document.settings.element.find(qn('w:compat'))
    next(e for e in compat if e.get(qn('w:name')) == 'compatibilityMode').set(qn('w:val'), '15')
    grid = document.add_table(rows=1, cols=1)
    grid.autofit = False
    grid.columns[0].width = grid.cell(0, 0).width = Pt(180)
    if styled:
        base = document.styles.add_style('Margin Base', WD_STYLE_TYPE.TABLE)
        pr = OxmlElement('w:tblPr')
        margins(pr, 'w:tblCellMar', {'left': 8, 'right': 8, 'top': 0, 'bottom': 0})
        base.element.append(pr)
        derived = document.styles.add_style('Margin Zero', WD_STYLE_TYPE.TABLE)
        derived.base_style = base
        pr = OxmlElement('w:tblPr')
        margins(pr, 'w:tblCellMar', {'left': 0})
        derived.element.append(pr)
        grid.style = derived
    if table is not None:
        margins(grid._tbl.tblPr, 'w:tblCellMar', table)
    if cell is not None:
        margins(grid.cell(0, 0)._tc.get_or_add_tcPr(), 'w:tcMar', cell)
    paragraph = grid.cell(0, 0).paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run('SHORT LINE\nLAST')
    run.font.name, run.font.size = 'Arial', Pt(12)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def geometry(document, shaping=False):
    options = PdfSaveOptions()
    options.text_shaping = shaping
    with pymupdf.open(stream=document.to_bytes(options), filetype='pdf') as pdf:
        words = pdf[0].get_text('words')
        assert [word[4] for word in words] == ['SHORT', 'LINE', 'LAST']
        assert len(pdf) == 1
        return words[0][0], words[0][1], words[1][2]


@pytest.mark.parametrize('shaping', [False, True])
@pytest.mark.parametrize('override', [None, 0, 4])
def test_cell_zero_overrides_table_and_survives_docx_roundtrip(shaping, override):
    source = source_docx(table={'left': 8, 'right': 8, 'top': 6, 'bottom': 6},
        cell=None if override is None else {side: override for side in ('left','right','top','bottom')})
    document = aw.Document(BytesIO(source))
    model = document.light_document_model
    snapshot = model.model_dump()
    expected = (28, 26, 192) if override is None else (20 + override, 20 + override, 200 - override)
    assert geometry(document, shaping) == pytest.approx(expected, abs=0.05)
    raw = document.to_bytes(OoxmlSaveOptions())
    loaded = aw.Document(BytesIO(raw))
    assert geometry(loaded, shaping) == pytest.approx(expected, abs=0.05)
    cell_format = loaded.light_document_model.sections[0].body.tables[0].rows[0].cells[0].cell_format
    assert cell_format.left_padding == override
    assert model.model_dump() == snapshot


@pytest.mark.parametrize('shaping', [False, True])
def test_zero_table_margin_overrides_style_chain_and_zero_style_is_retained(shaping):
    document = aw.Document(BytesIO(source_docx(styled=True, table={'right': 0})))
    assert geometry(document, shaping) == pytest.approx((20, 20, 200), abs=0.05)
    style = next(style for style in document.light_document_model.styles if style.name == 'Margin Zero')
    assert style.table_style_format.left_padding == 0
    assert style.table_style_format.right_padding is None
    loaded = aw.Document(BytesIO(document.to_bytes(OoxmlSaveOptions())))
    assert geometry(loaded, shaping) == pytest.approx((20, 20, 200), abs=0.05)
    assert next(style for style in loaded.light_document_model.styles if style.name == 'Margin Zero').table_style_format.left_padding == 0


@pytest.mark.parametrize('shaping', [False, True])
def test_implicit_default_table_style_is_used(shaping):
    document = aw.Document(BytesIO(source_docx()))
    table = document.light_document_model.sections[0].body.tables[0]
    assert table.left_padding == pytest.approx(5.4)
    assert table.top_padding == 0
    assert geometry(document, shaping) == pytest.approx((25.4, 20, 194.6), abs=0.05)


@pytest.mark.parametrize('tag', ['w:tcPr', 'w:tblPr'])
def test_missing_property_block_still_inherits_default_table_style(tag):
    output = BytesIO()
    with ZipFile(BytesIO(source_docx())) as source, ZipFile(output, 'w') as target:
        for name in source.namelist():
            data = source.read(name)
            if name == 'word/document.xml':
                root = ET.fromstring(data)
                for parent in root.iter():
                    child = parent.find(qn(tag))
                    if child is not None:
                        parent.remove(child)
                data = ET.tostring(root)
            target.writestr(name, data)
    document = aw.Document(BytesIO(output.getvalue()))
    assert geometry(document) == pytest.approx((25.4,20,194.6),abs=0.05)


@pytest.mark.parametrize('kind', [ldm.CellFormat, ldm.Table, ldm.TableStyleFormat])
def test_json_distinguishes_unset_and_zero_margins(kind):
    unset, zero = kind(), kind(left_padding=0, top_padding=0)
    assert kind.model_validate_json(unset.model_dump_json()).left_padding is None
    assert kind.model_validate_json(zero.model_dump_json()).left_padding == 0
    assert kind.model_validate_json(zero.model_dump_json()).right_padding is None


@pytest.mark.parametrize('kind', [ldm.CellFormat, ldm.Table, ldm.TableStyleFormat])
@pytest.mark.parametrize('value', [-1, float('inf'), float('nan')])
def test_model_rejects_invalid_margin_geometry(kind, value):
    with pytest.raises(ValueError):
        kind(left_padding=value)


@pytest.mark.parametrize('scope', ['table','cell','style'])
def test_source_rejects_negative_margins(scope):
    kwargs = {'table': {'left': -1}} if scope == 'table' else {'cell': {'left': -1}}
    if scope == 'style':
        raw = source_docx(styled=True)
        output = BytesIO()
        with ZipFile(BytesIO(raw)) as source, ZipFile(output, 'w') as target:
            for name in source.namelist():
                data = source.read(name)
                if name == 'word/styles.xml':
                    data = data.replace(b'w:w="160"', b'w:w="-20"')
                target.writestr(name, data)
        raw = output.getvalue()
    else:
        raw = source_docx(**kwargs)
    with pytest.raises(ValueError, match='margins must be nonnegative'):
        aw.Document(BytesIO(raw))


@pytest.mark.parametrize('margin_type', ['pct','auto','nil'])
def test_unsupported_cell_margin_units_do_not_override_parent(margin_type):
    raw = source_docx(table={'left':8,'right':8,'top':0,'bottom':0}, cell={'left':0})
    output = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(output, 'w') as target:
        for name in source.namelist():
            data = source.read(name)
            if name == 'word/document.xml':
                data = data.replace(b'<w:tcMar><w:left w:type="dxa" w:w="0"',
                    ('<w:tcMar><w:left w:type="'+margin_type+'" w:w="ignored"').encode())
            target.writestr(name, data)
    document = aw.Document(BytesIO(output.getvalue()))
    assert geometry(document) == pytest.approx((28,20,192),abs=0.05)


def test_ldm_nested_tables_inherit_their_own_margins_without_mutating_source():
    paragraph = ldm.Paragraph(children=[ldm.Run(text='INNER', font=ldm.Font(size=12))])
    inner = ldm.Table(left_padding=0, top_padding=0, right_padding=0, bottom_padding=0,
        rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[paragraph])])])
    outer = ldm.Table(left_padding=8, top_padding=4, right_padding=8, bottom_padding=4,
        rows=[ldm.Row(cells=[ldm.Cell(children=[inner])])])
    model = ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
        page_width=220,page_height=180,left_margin=20,right_margin=20,top_margin=20,bottom_margin=20),
        body=ldm.Body(children=[outer]))])
    snapshot = model.model_dump()
    with pymupdf.open(stream=LdmPdfWriter().write_to_bytes(model), filetype='pdf') as pdf:
        word = pdf[0].get_text('words')[0]
        assert word[0:2] == pytest.approx((28,24),abs=0.05)
    assert model.model_dump() == snapshot


def test_zero_margin_picture_uses_full_width_after_json_and_docx_roundtrip():
    image = BytesIO()
    Image.new('RGB',(16,8),'red').save(image,format='PNG')
    picture = ldm.Shape(has_image=True,width=240,height=120,image_data=ldm.ImageData(image_bytes=image.getvalue()))
    table = ldm.Table(left_padding=0,right_padding=0,top_padding=0,bottom_padding=0,
        rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[ldm.Paragraph(children=[picture])])])])
    model = ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
        page_width=220,page_height=180,left_margin=20,right_margin=20,top_margin=20,bottom_margin=20),
        body=ldm.Body(children=[table]))])
    for document in (model,ldm.Document.model_validate_json(model.model_dump_json()),
                     aw.Document(BytesIO(LdmDocxWriter().write_to_bytes(model))).light_document_model):
        with pymupdf.open(stream=LdmPdfWriter().write_to_bytes(document),filetype='pdf') as pdf:
            rect = pdf[0].get_image_rects(pdf[0].get_images()[0][0])[0]
            assert tuple(rect) == pytest.approx((20,20,200,110),abs=0.05)


@pytest.mark.parametrize('shaping', [False, True])
@pytest.mark.parametrize('padding', [0, 10])
def test_table_height_measurement_matches_painted_border_with_inherited_margins(shaping, padding):
    paragraph = ldm.Paragraph(children=[ldm.Run(text='MEASURE',font=ldm.Font(size=12))])
    row = ldm.Row(row_format=ldm.RowFormat(borders=[ldm.Border(line_style=1,line_width=0.5) for _ in range(6)]),
                  cells=[ldm.Cell(paragraphs=[paragraph])])
    table = ldm.Table(top_padding=padding,bottom_padding=padding,rows=[row])
    model = ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
        page_width=220,page_height=180,left_margin=20,right_margin=20,top_margin=20,bottom_margin=20),
        body=ldm.Body(children=[table]))])
    options=PdfSaveOptions()
    options.text_shaping=shaping
    writer=LdmPdfWriter(options)
    with pymupdf.open(stream=writer.write_to_bytes(model),filetype='pdf') as pdf:
        borders=[drawing['rect'] for drawing in pdf[0].get_drawings()]
        actual_height=(max(rect.y1 for rect in borders)-min(rect.y0 for rect in borders))*PT_TO_MM
    measured=writer._estimate_table_height(table,180*PT_TO_MM)
    assert measured == pytest.approx(actual_height+POST_TABLE_SPACING_MM,abs=0.02)
