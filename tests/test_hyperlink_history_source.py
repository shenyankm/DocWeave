from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from xml.etree.ElementTree import fromstring

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

import dockiln as dk
from dockiln import light_document_model as ldm
from dockiln._flat_opc import encode, decode
from dockiln.writers.docx import LdmDocxWriter
from dockiln.saving import OoxmlSaveOptions

HISTORY = [None, "0", "1", "false", "true", "off", "on"]


def source(history, target="https://owned.test/path", story="body"):
    doc = Document()
    p = doc.add_paragraph() if story == "body" else doc.sections[0].header.paragraphs[0]
    link = OxmlElement("w:hyperlink")
    if target.startswith("#"):
        link.set(qn("w:anchor"), target[1:])
    else:
        link.set(
            qn("r:id"),
            p.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True),
        )
    if history is not None:
        link.set(qn("w:history"), history)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "OWNED LINK"
    run.append(text)
    link.append(run)
    p._p.append(link)
    data = BytesIO()
    doc.save(data)
    return data.getvalue()


def linked_run(model, story):
    para = (
        model.sections[0].body.children[0]
        if story == "body"
        else model.sections[0].headers_footers[0].children[0]
    )
    return next(r for r in para.runs if r.is_hyperlink)


@pytest.mark.parametrize("history", HISTORY)
@pytest.mark.parametrize("target", ["https://owned.test/path", "#OwnedBookmark"])
@pytest.mark.parametrize("fmt", [dk.SaveFormat.DOCX, dk.SaveFormat.FLAT_OPC])
@pytest.mark.parametrize("copy_mode", ["direct", "deep", "json"])
def test_history_survives_read_copy_save_cold_without_mutating_source(
    history, target, fmt, copy_mode
):
    model = dk.Document(BytesIO(source(history, target))).light_document_model
    before = model.model_dump_json()
    current = (
        deepcopy(model)
        if copy_mode == "deep"
        else ldm.Document.model_validate_json(before)
        if copy_mode == "json"
        else model
    )
    for _ in range(2):
        assert linked_run(current, "body").source_hyperlink_history == history
        raw = LdmDocxWriter(OoxmlSaveOptions(fmt)).write_to_bytes(current)
        with ZipFile(
            BytesIO(decode(raw) if fmt == dk.SaveFormat.FLAT_OPC else raw)
        ) as z:
            link = fromstring(z.read("word/document.xml")).find(
                ".//" + qn("w:hyperlink")
            )
            assert link.get(qn("w:history")) == history
            assert link.find(".//" + qn("w:t")).text == "OWNED LINK"
        current = dk.Document(BytesIO(raw)).light_document_model
    assert model.model_dump_json() == before


@pytest.mark.parametrize("history", HISTORY)
def test_flat_input_and_header_share_history_mapping(history):
    model = dk.Document(
        BytesIO(encode(source(history, story="header")))
    ).light_document_model
    assert linked_run(model, "header").source_hyperlink_history == history
    raw = LdmDocxWriter().write_to_bytes(
        ldm.Document.model_validate_json(model.model_dump_json())
    )
    assert (
        linked_run(
            dk.Document(BytesIO(raw)).light_document_model, "header"
        ).source_hyperlink_history
        == history
    )


@pytest.mark.parametrize(
    "invalid", ["", "2", "invalid", "TRUE", " false ", "OWNED SECRET"]
)
def test_invalid_history_is_rejected_without_source_value_in_error(invalid):
    with pytest.raises(RuntimeError) as error:
        dk.Document(BytesIO(source(invalid)))
    assert "OWNED SECRET" not in str(error.value)


@pytest.mark.parametrize("invalid", ["", "2", "OWNED SECRET", 1, object()])
def test_mutated_history_fails_before_path_or_stream_publication(tmp_path, invalid):
    doc = dk.Document(BytesIO(source("1")))
    linked_run(doc.light_document_model, "body").source_hyperlink_history = invalid
    path = tmp_path / "existing.docx"
    path.write_bytes(b"OWNED EXISTING")
    with pytest.raises(RuntimeError):
        doc.save(path, dk.SaveFormat.DOCX)
    assert path.read_bytes() == b"OWNED EXISTING"
    stream = BytesIO(b"OWNED EXISTING")
    stream.seek(3)
    with pytest.raises(RuntimeError):
        doc.save(stream, dk.SaveFormat.DOCX)
    assert stream.tell() == 3 and stream.getvalue() == b"OWNED EXISTING"
