"""Repeated style resolution is bounded and never shares editable font models."""

import gc
from io import BytesIO
import weakref

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from aspose.words_foss.docx_reader.document_reader import DocumentReader
from aspose.words_foss.docx_writer.runs import color_to_hex

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def reader(size=16):
    source = Document()
    para = source.styles.add_style('Brand Paragraph', WD_STYLE_TYPE.PARAGRAPH)
    para.font.size = Pt(size)
    char = source.styles.add_style('Brand Character', WD_STYLE_TYPE.CHARACTER)
    char.font.italic = True
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), 'CCDDFF')
    char.element.get_or_add_rPr().append(shading)
    large = source.styles.add_style('Large Character', WD_STYLE_TYPE.CHARACTER)
    large.base_style = char
    large.font.size = Pt(22)
    for name, color in [('Blue Table', '0000FF'), ('Red Table', 'FF0000')]:
        table = source.styles.add_style(name, WD_STYLE_TYPE.TABLE)
        table.font.color.rgb = RGBColor.from_string(color)
    paragraph = source.add_paragraph(style=para)
    for text in ['FIRST', 'SECOND']:
        paragraph.add_run(text).style = char
    stream = BytesIO()
    source.save(stream)
    result = DocumentReader()
    result.load_bytes(stream.getvalue())
    result._build_name_to_style_id_map()
    result._ensure_builders_installed()
    return result, stream.getvalue()


def test_inherited_styles_are_not_reparsed_for_every_run(monkeypatch):
    doc, _ = reader()
    resolver = doc._font_resolver
    rpr = doc._document_xml.find('.//'+W+'rPr')
    direct = doc._fonts_builder.build
    counts = {}

    def track(element):
        counts[element] = counts.get(element, 0) + 1
        return direct(element)

    monkeypatch.setattr(doc._fonts_builder, 'build', track)
    for _ in range(50):
        assert resolver.resolve(rpr, 'BrandParagraph').size == 16
    assert counts[doc._doc_default_rPr] == 1
    assert counts[doc._style_elem_cache['BrandParagraph'].find(W+'rPr')] == 1
    assert counts[doc._style_elem_cache['BrandCharacter'].find(W+'rPr')] == 1
    assert counts[rpr] == 50


def test_cached_fonts_and_nested_shading_remain_independent():
    doc, _ = reader()
    resolver = doc._font_resolver
    rpr = doc._document_xml.find('.//'+W+'rPr')
    first = resolver.resolve(rpr, 'BrandParagraph')
    second = resolver.resolve(rpr, 'BrandParagraph')
    before = second.model_dump()
    first.size = 99
    first.shading.background_pattern_color = 'Yellow'
    assert second.model_dump() == before
    assert resolver.resolve(rpr, 'BrandParagraph').model_dump() == before
    model = doc.to_light_document()
    runs = model.sections[0].body.paragraphs[0].runs
    runs[0].font.shading.background_pattern_color = 'Yellow'
    assert runs[1].font.shading.background_pattern_color != 'Yellow'
    style = next(s for s in model.styles if s.name == 'Brand Character')
    assert style.font.shading.background_pattern_color != 'Yellow'


def test_font_cache_includes_table_and_character_context():
    doc, _ = reader()
    resolver = doc._font_resolver
    rpr = doc._document_xml.find('.//'+W+'rPr')
    doc._current_table_style_id = 'BlueTable'
    assert color_to_hex(resolver.resolve(rpr, 'BrandParagraph').color) == '0000FF'
    doc._current_table_style_id = 'RedTable'
    assert color_to_hex(resolver.resolve(rpr, 'BrandParagraph').color) == 'FF0000'
    large = OxmlElement('w:rPr')
    reference = OxmlElement('w:rStyle')
    reference.set(qn('w:val'), 'LargeCharacter')
    large.append(reference)
    assert resolver.resolve(large, 'BrandParagraph').size == 22
    doc._current_table_style_id = ''
    assert resolver.resolve(rpr, 'BrandParagraph').size == 16
    assert color_to_hex(resolver.resolve(rpr, 'BrandParagraph').color) != 'FF0000'


def test_reloading_reader_does_not_reuse_previous_fonts():
    doc, _ = reader(16)
    assert doc.to_light_document().sections[0].body.paragraphs[0].runs[0].font.size == 16
    _, replacement = reader(24)
    doc.load_bytes(replacement)
    assert doc.to_light_document().sections[0].body.paragraphs[0].runs[0].font.size == 24


def test_font_cache_is_bounded_per_reader():
    doc, _ = reader()
    resolver = doc._font_resolver
    rpr = doc._document_xml.find('.//'+W+'rPr')
    before = resolver.resolve(rpr, 'BrandParagraph').model_dump()
    for index in range(300):
        resolver.resolve(None, f'Unknown{index}')
    assert resolver._inherited_font.cache_info().currsize == 128
    misses = resolver._inherited_font.cache_info().misses
    assert resolver.resolve(rpr, 'BrandParagraph').model_dump() == before
    assert resolver._inherited_font.cache_info().misses == misses + 1


def test_cached_reader_and_resolver_can_be_collected():
    doc, _ = reader()
    doc.to_light_document()
    refs = [weakref.ref(doc), weakref.ref(doc._font_resolver)]
    del doc
    gc.collect()
    assert all(ref() is None for ref in refs)
