"""Explicit off settings in character styles survive OOXML regeneration."""

from io import BytesIO
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
import pytest

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer.styles_part import render_styles_xml


@pytest.mark.parametrize('attribute', ['bold', 'italic', 'hidden', 'all_caps', 'small_caps', 'strike'])
def test_character_style_explicit_off_is_not_dropped(attribute):
    source = Document()
    base = source.styles.add_style('Brand Base', WD_STYLE_TYPE.CHARACTER)
    setattr(base.font, attribute, True)
    child = source.styles.add_style('Brand Child', WD_STYLE_TYPE.CHARACTER)
    child.base_style = base
    setattr(child.font, attribute, False)
    source.add_paragraph().add_run('TEXT').style = child
    stream = BytesIO()
    source.save(stream)
    doc = aw.Document(BytesIO(stream.getvalue()))
    font = doc.light_document_model.sections[0].body.paragraphs[0].runs[0].font
    model_attribute = 'strike_through' if attribute == 'strike' else attribute
    assert getattr(font, model_attribute) is False
    result = Document(BytesIO(doc.to_bytes('docx')))
    assert getattr(result.styles['Brand Child'].font, attribute) is False
    assert result.styles['Brand Child'].base_style.name == 'Brand Base'


def test_unset_character_style_bold_stays_unset():
    source = Document()
    source.styles.add_style('Brand Child', WD_STYLE_TYPE.CHARACTER)
    source.add_paragraph().add_run('TEXT').style = source.styles['Brand Child']
    stream = BytesIO()
    source.save(stream)
    result = Document(BytesIO(aw.Document(BytesIO(stream.getvalue())).to_bytes('docx')))
    assert result.styles['Brand Child'].font.bold is None


@pytest.mark.parametrize('paragraph_style_name', ['Normal', 'Brand Paragraph'])
def test_sparse_character_style_keeps_paragraph_font_and_run_off_override(paragraph_style_name):
    source = Document()
    paragraph_style = source.styles['Normal'] if paragraph_style_name == 'Normal' else source.styles.add_style(
        paragraph_style_name, WD_STYLE_TYPE.PARAGRAPH)
    paragraph_style.font.bold = True
    child = source.styles.add_style('Brand Child', WD_STYLE_TYPE.CHARACTER)
    child.font.italic = True
    para = source.add_paragraph(style=paragraph_style)
    para.add_run('INHERITED').style = child
    direct = para.add_run('OFF')
    direct.style = child
    direct.font.bold = False
    stream = BytesIO()
    source.save(stream)
    doc = aw.Document(BytesIO(stream.getvalue()))
    actual = doc.light_document_model.sections[0].body.paragraphs[0].runs
    assert [run.font.bold for run in actual] == [True, False]
    assert all(run.font.italic for run in actual)
    result = Document(BytesIO(doc.to_bytes('docx')))
    assert result.paragraphs[0].runs[0].font.bold is None
    assert result.paragraphs[0].runs[1].font.bold is False
    assert all(run.font.italic is None for run in result.paragraphs[0].runs)


def test_conditional_table_style_keeps_explicit_off():
    model = ldm.Document(styles=[ldm.Style(name='Brand Table', type=3,
        table_style_properties=[ldm.TableStyleProperty(type='firstRow', font=ldm.Font(bold=False))])])
    xml = ET.fromstring(render_styles_xml(model))
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    bold = xml.find('.//' + w + 'tblStylePr/' + w + 'rPr/' + w + 'b')
    assert bold is not None and bold.get(w + 'val') == '0'
