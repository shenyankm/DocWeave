"""Optional native renderer smoke test, not a Word visual-fidelity assertion."""

from io import BytesIO
from pathlib import Path
import os
import shutil
import subprocess

from docx import Document
from PIL import Image
from fontTools.ttLib import TTFont
from pypdf import PdfReader
import pytest

import aspose.words_foss as aw
from aspose.words_foss.libreoffice import convert_to_pdf


AVAILABLE = bool(shutil.which("soffice") or shutil.which("libreoffice") or
                 Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").is_file())


@pytest.mark.skipif(not AVAILABLE, reason="LibreOffice is not installed")
def test_native_renderer_accepts_dom_picture_link_and_horizontal_merge(tmp_path, monkeypatch):
    # Native rendering needs installed fonts; ActualText can hide missing visual glyphs.
    font = TTFont(Path(aw.__file__).parent / "pdf_writer/fonts/DocumentSansSC-Regular.woff")
    font.flavor = None
    font.save(tmp_path / "DocumentSansSC.ttf")
    from xml.sax.saxutils import escape
    config = tmp_path / "fonts.conf"
    config.write_text(f'<fontconfig><dir>{escape(str(tmp_path))}</dir>'
                      f'<cachedir>{escape(str(tmp_path / "cache"))}</cachedir></fontconfig>')
    monkeypatch.setenv("FONTCONFIG_FILE", str(config))
    original = Document()
    original.add_paragraph("DOM 中文 smoke test")
    table = original.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left cell"
    table.cell(0, 1).text = "right cell"
    raw = BytesIO()
    original.save(raw)
    doc = aw.DocxDocument(BytesIO(raw.getvalue()))
    paragraph = doc.body.paragraphs[0]
    paragraph.add_hyperlink("链接 [说明]", "https://example.com/a_(b)?x=1&y=2")
    image = BytesIO()
    Image.new("RGB", (100, 50), "navy").save(image, format="PNG")
    paragraph.add_picture(image.getvalue(), width=72, alternative_text="中文图片")
    doc.body.tables[0].merge_cells(0, 0, 2)
    source, output = tmp_path / "edited.docx", tmp_path / "native.pdf"
    doc.save(source)
    convert_to_pdf(source, output)
    rendered = PdfReader(output)
    text = "".join(page.extract_text() for page in rendered.pages)
    assert all(token in text for token in ("DOM", "中文", "链接", "说明", "left cell", "right cell"))
    cmaps = []
    for ref in rendered.pages[0]["/Resources"]["/Font"].values():
        descriptor = ref.get_object().get("/FontDescriptor")
        if descriptor and "/FontFile2" in descriptor:
            embedded = TTFont(BytesIO(descriptor["/FontFile2"].get_data()))
            cmaps.extend(embedded.getGlyphOrder())
    assert cmaps, "No actual embedded font outlines"
    assert any(page.images for page in rendered.pages)
    links = [item.get_object()["/A"].get("/URI") for page in rendered.pages
             for item in page.get("/Annots", []) if "/A" in item.get_object()]
    assert "https://example.com/a_(b)?x=1&y=2" in links


@pytest.mark.skipif(not AVAILABLE or os.name != "posix", reason="Requires native LibreOffice and POSIX cleanup")
def test_native_timeout_then_changed_input_succeeds_with_fresh_profiles(tmp_path, monkeypatch):
    from aspose.words_foss import libreoffice

    real_process = libreoffice.run_process
    profiles = []

    def monitored(command, *args, **kwargs):
        profiles.append(next(arg for arg in command if arg.startswith("-env:UserInstallation=")))
        return real_process(command, *args, **kwargs)

    monkeypatch.setattr(libreoffice, "run_process", monitored)
    source, output = tmp_path / "source.docx", tmp_path / "output.pdf"
    doc = Document()
    doc.add_paragraph("FIRST REQUEST")
    doc.save(source)
    output.write_bytes(b"existing output")
    with pytest.raises(subprocess.TimeoutExpired):
        convert_to_pdf(source, output, timeout=0.001)
    assert output.read_bytes() == b"existing output"
    assert not list(tmp_path.glob(".libreoffice-*"))
    for label in ("FIRST REQUEST", "SECOND REQUEST"):
        doc.paragraphs[0].text = label
        doc.save(source)
        original = source.read_bytes()
        convert_to_pdf(source, output, timeout=30)
        text = "".join(page.extract_text() for page in PdfReader(output).pages)
        assert label in text
        if label == "SECOND REQUEST":
            assert "FIRST REQUEST" not in text
        assert source.read_bytes() == original
        assert not list(tmp_path.glob(".libreoffice-*"))
    assert len(profiles) == len(set(profiles)) == 3
