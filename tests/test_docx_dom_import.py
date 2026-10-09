"""Ownership and transactional destination-style import, including saved packages."""

from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.shared import Pt

import aspose.words_foss as aw
from aspose.words_foss.dom.nodes import W


def document(name="Conflict", style_id="Conflict", *, bold=True, base=None):
    doc = Document()
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.element.set(f"{{{W}}}styleId", style_id)
    style.font.bold, style.font.size = bold, Pt(18 if bold else 11)
    if base is not None:
        style.base_style = doc.styles[base]
    doc.add_paragraph("IMPORT", style=style)
    stream = BytesIO()
    doc.save(stream)
    return aw.DocxDocument(BytesIO(stream.getvalue()))


@pytest.mark.parametrize("deep", [False, True])
def test_foreign_import_maps_style_name_and_preserves_ownership_and_source(deep):
    source = document(style_id="SourceConflict")
    target = document(style_id="DestinationConflict", bold=False)
    before_source, before_target = source.to_bytes(), target.to_bytes()
    original = source.body.paragraphs[0]
    copied = target.import_node(original, deep)
    assert copied.owner_document is target and copied.parent_node is None
    assert original.parent_node is source.body
    assert copied.text == ("IMPORT" if deep else "")
    assert copied.paragraph_format.style_id == "DestinationConflict"
    assert source.to_bytes() == before_source and target.to_bytes() == before_target
    target.body.append_child(copied)
    saved = Document(BytesIO(target.to_bytes())).paragraphs[-1]
    assert saved.text == copied.text and saved.style.style_id == "DestinationConflict"
    assert saved.style.font.bold is False and saved.style.font.size.pt == 11


def test_import_new_style_renames_id_collision_and_reuses_mapping():
    source = document(name="New Style", style_id="Clash")
    target = document(name="Existing Style", style_id="Clash", bold=False)
    original = source.to_bytes()
    first = target.import_node(source.body.paragraphs[0], True)
    second = target.import_node(source.body.paragraphs[0], True)
    assert first.paragraph_format.style_id == second.paragraph_format.style_id == "Clash_0"
    target.body.append_child(first)
    target.body.append_child(second)
    saved = Document(BytesIO(target.to_bytes()))
    assert saved.paragraphs[-1].style.name == "New Style"
    assert saved.paragraphs[-1].style.font.bold is True
    assert len([style for style in saved.styles if style.name == "New Style"]) == 1
    assert source.to_bytes() == original


@pytest.mark.parametrize("mode", list(aw.ImportFormatMode))
@pytest.mark.parametrize("deep", [False, True])
def test_same_document_import_respects_depth_and_detaches(mode, deep):
    doc = document()
    original = doc.to_bytes()
    copied = doc.import_node(doc.body.paragraphs[0], deep, mode)
    assert copied.parent_node is None and copied.owner_document is doc
    assert copied.text == ("IMPORT" if deep else "")
    assert doc.to_bytes() == original


@pytest.mark.parametrize("mode", [aw.ImportFormatMode.KEEP_SOURCE_FORMATTING, aw.ImportFormatMode.KEEP_DIFFERENT_STYLES])
def test_unimplemented_format_modes_fail_before_destination_changes(mode):
    source, target = document(), document(bold=False)
    before = target.to_bytes()
    with pytest.raises(NotImplementedError):
        target.import_node(source.body.paragraphs[0], True, mode)
    assert target.to_bytes() == before


@pytest.mark.parametrize("name", ["b", "i"])
@pytest.mark.parametrize("child", [None, False, True])
def test_default_on_character_base_projects_saved_story_without_mutating_live_format(name, child):
    import runpy
    from pathlib import Path

    helpers = runpy.run_path(str(Path(__file__).parents[1] / "docs" / "probes" / "import_style_conflicts.py"))
    style, build = helpers["style"], helpers["document"]
    common = (f'<w:docDefaults><w:rPrDefault><w:rPr><w:{name}/><w:sz w:val="22"/>' +
              '</w:rPr></w:rPrDefault></w:docDefaults>' + style("P", "paragraph", {name: False}) +
              style("Base", "character", {name: True}))
    source_data = build(common + style("Derived", "character", {name: child}, "Base"),
                        '<w:pStyle w:val="P"/>', '<w:rStyle w:val="Derived"/>')
    target_data = build(common, '<w:pStyle w:val="P"/>', '<w:rStyle w:val="Base"/>')
    source, target = aw.DocxDocument(BytesIO(source_data)), aw.DocxDocument(BytesIO(target_data))
    existing = target.body.paragraphs[0].runs[0]
    copied = target.import_node(source.body.paragraphs[0], True)
    target.body.append_child(copied)
    attribute = "bold" if name == "b" else "italic"
    assert getattr(existing.effective_font, attribute) is True
    saved = target.to_bytes()
    assert getattr(existing.effective_font, attribute) is True
    assert getattr(aw.DocxDocument(BytesIO(saved)).body.paragraphs[0].runs[0].effective_font, attribute) is False
    assert target.to_bytes() == saved and source.to_bytes() == source_data


def test_implicit_source_font_size_is_refused_before_destination_changes():
    import runpy
    from pathlib import Path

    helpers = runpy.run_path(str(Path(__file__).parents[1] / "docs" / "probes" / "import_style_conflicts.py"))
    style, build = helpers["style"], helpers["document"]
    source_data = build(style("Base", "paragraph", {}) + style("Derived", "paragraph", {}, "Base"),
                        '<w:pStyle w:val="Derived"/>', "")
    target_data = build(style("Base", "paragraph", {"sz": 36}), '<w:pStyle w:val="Base"/>', "")
    source, target = aw.DocxDocument(BytesIO(source_data)), aw.DocxDocument(BytesIO(target_data))
    before = target.to_bytes()
    with pytest.raises(NotImplementedError, match="implicit application font size"):
        target.import_node(source.body.paragraphs[0], True)
    assert target.to_bytes() == before and source.to_bytes() == source_data


def test_new_paragraph_style_preserves_source_font_over_conflicting_base():
    source, target = document(), document(bold=False)
    raw = Document(BytesIO(source.to_bytes()))
    style = raw.styles.add_style("Derived", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = raw.styles["Conflict"]
    raw.paragraphs[0].style = style
    data = BytesIO()
    raw.save(data)
    source = aw.DocxDocument(BytesIO(data.getvalue()))
    before_source = source.to_bytes()
    copied = target.import_node(source.body.paragraphs[0], True)
    target.body.append_child(copied)
    assert copied.runs[0].effective_font == source.body.paragraphs[0].runs[0].effective_font
    reopened = aw.DocxDocument(BytesIO(target.to_bytes()))
    assert reopened.body.paragraphs[-1].runs[0].effective_font == copied.runs[0].effective_font
    saved = Document(BytesIO(target.to_bytes()))
    assert saved.styles["Conflict"].font.bold is False and saved.styles["Conflict"].font.size.pt == 11
    assert saved.styles["Derived"].font.bold is True and saved.styles["Derived"].font.size.pt == 18
    assert source.to_bytes() == before_source


@pytest.mark.parametrize("kind", [WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER])
def test_import_resolves_conflicting_ancestor_with_identical_immediate_base(kind):
    def make(bold):
        raw = Document()
        ancestor = raw.styles.add_style("Ancestor", kind)
        ancestor.font.bold = bold
        base = raw.styles.add_style("Base", kind)
        base.base_style = ancestor
        style = raw.styles.add_style("Derived", kind) if bold else base
        if bold:
            style.base_style = base
        paragraph = raw.add_paragraph("IMPORT")
        if kind == WD_STYLE_TYPE.PARAGRAPH:
            paragraph.style = style
        else:
            paragraph.runs[0].style = style
        stream = BytesIO()
        raw.save(stream)
        return aw.DocxDocument(BytesIO(stream.getvalue()))

    source, target = make(True), make(False)
    before_source = source.to_bytes()
    copied = target.import_node(source.body.paragraphs[0], True)
    target.body.append_child(copied)
    assert copied.runs[0].effective_font.bold is True
    assert aw.DocxDocument(BytesIO(target.to_bytes())).body.paragraphs[-1].runs[0].effective_font.bold is True
    assert Document(BytesIO(target.to_bytes())).styles["Ancestor"].font.bold is (None if kind == WD_STYLE_TYPE.CHARACTER else False)
    assert source.to_bytes() == before_source


def test_conflict_in_later_conditional_table_style_is_not_ignored():
    def make(color, derived):
        raw = Document()
        ancestor = raw.styles.add_style("Ancestor", WD_STYLE_TYPE.TABLE)
        for region, fill in (("firstRow", "0000FF"), ("lastRow", color)):
            conditional = OxmlElement("w:tblStylePr")
            conditional.set(f"{{{W}}}type", region)
            cell = OxmlElement("w:tcPr")
            shading = OxmlElement("w:shd")
            shading.set(f"{{{W}}}fill", fill)
            cell.append(shading)
            conditional.append(cell)
            ancestor.element.append(conditional)
        base = raw.styles.add_style("Base", WD_STYLE_TYPE.TABLE)
        base.base_style = ancestor
        style = raw.styles.add_style("Derived", WD_STYLE_TYPE.TABLE) if derived else base
        if derived:
            style.base_style = base
        table = raw.add_table(rows=1, cols=1)
        table.style = style
        table.cell(0, 0).text = "CELL"
        stream = BytesIO()
        raw.save(stream)
        return aw.DocxDocument(BytesIO(stream.getvalue()))

    source, target = make("FF0000", True), make("00FF00", False)
    before_source, before_target = source.to_bytes(), target.to_bytes()
    with pytest.raises(NotImplementedError, match="conflicting base"):
        target.import_node(source.body.tables[0], True)
    assert source.to_bytes() == before_source and target.to_bytes() == before_target


@pytest.mark.parametrize("decoration", ["color", "duplicate", "comment", "attribute", "nested_style"])
def test_unsupported_style_translation_is_atomic(decoration):
    raw_source = Document(BytesIO(document().to_bytes()))
    raw_target = Document(BytesIO(document(bold=False).to_bytes()))
    derived = raw_source.styles.add_style("Derived", WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = raw_source.styles["Conflict"]
    raw_source.paragraphs[0].style = derived
    if decoration == "color":
        color = OxmlElement("w:color")
        color.set(f"{{{W}}}val", "FF0000")
        raw_source.styles["Conflict"].element.get_or_add_rPr().append(color)
    elif decoration == "nested_style":
        reference = OxmlElement("w:rStyle")
        reference.set(f"{{{W}}}val", "DefaultParagraphFont")
        derived.element.get_or_add_rPr().append(reference)
    else:
        bold = OxmlElement("w:b")
        derived.element.get_or_add_rPr().append(bold)
        if decoration == "duplicate":
            derived.element.get_or_add_rPr().append(OxmlElement("w:b"))
        elif decoration == "comment":
            from lxml.etree import Comment

            bold.append(Comment("retained metadata"))
        else:
            bold.set(f"{{{W}}}future", "retained metadata")
    def load(raw):
        stream = BytesIO()
        raw.save(stream)
        return aw.DocxDocument(BytesIO(stream.getvalue()))
    source, target = load(raw_source), load(raw_target)
    before_source, before_target = source.to_bytes(), target.to_bytes()
    with pytest.raises(ValueError if decoration == "duplicate" else NotImplementedError):
        target.import_node(source.body.paragraphs[0], True)
    assert source.to_bytes() == before_source and target.to_bytes() == before_target


def test_new_style_dependency_is_resolved_from_planned_destination_styles():
    raw = Document()
    parent = raw.styles.add_style("New Parent", WD_STYLE_TYPE.PARAGRAPH)
    parent.font.bold, parent.font.size = True, Pt(16)
    child = raw.styles.add_style("New Child", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = parent
    raw.add_paragraph("IMPORT", style=child)
    stream = BytesIO()
    raw.save(stream)
    source, target = aw.DocxDocument(BytesIO(stream.getvalue())), document(bold=False)
    copied = target.import_node(source.body.paragraphs[0], True)
    target.body.append_child(copied)
    saved = Document(BytesIO(target.to_bytes()))
    assert saved.styles["New Child"].base_style.name == "New Parent"
    assert copied.runs[0].effective_font.bold and copied.runs[0].effective_font.size == 16


@pytest.mark.parametrize("value", [0, 1, None, "USE_DESTINATION_STYLES"])
def test_mode_requires_enum(value):
    source, target = document(), document()
    with pytest.raises(TypeError):
        target.import_node(source.body.paragraphs[0], True, value)


def test_invalid_run_style_does_not_commit_an_already_planned_paragraph_style():
    source = document(name="New Style")
    raw = Document(BytesIO(source.to_bytes()))
    reference = OxmlElement("w:rStyle")
    reference.set(f"{{{W}}}val", "MissingStyle")
    raw.paragraphs[0].runs[0]._r.get_or_add_rPr().append(reference)
    data = BytesIO()
    raw.save(data)
    source, target = aw.DocxDocument(BytesIO(data.getvalue())), document()
    before = target.to_bytes()
    with pytest.raises(ValueError, match="source style"):
        target.import_node(source.body.paragraphs[0], True)
    assert target.to_bytes() == before


def test_numbered_style_requires_reference_translation_without_mutation():
    raw = Document()
    raw.add_paragraph("LIST", style="List Number")
    data = BytesIO()
    raw.save(data)
    source, target = aw.DocxDocument(BytesIO(data.getvalue())), document()
    before = target.to_bytes()
    with pytest.raises(NotImplementedError, match="numbering"):
        target.import_node(source.body.paragraphs[0], True)
    assert target.to_bytes() == before


def test_different_theme_does_not_silently_change_an_imported_style():
    source, target = document(name="New Style"), document()
    part = "word/theme/theme1.xml"
    source._package.set_parts({part: source._package.payload(part).replace(b'val="000000"', b'val="112233"', 1)})
    before = target.to_bytes()
    with pytest.raises(NotImplementedError, match="themes"):
        target.import_node(source.body.paragraphs[0], True)
    assert target.to_bytes() == before


@pytest.mark.parametrize("mode", list(aw.ImportFormatMode))
def test_same_document_header_import_rebinds_part_and_preserves_header(mode):
    raw = Document()
    raw.add_paragraph("BODY")
    raw.sections[0].header.paragraphs[0].text = "HEADER"
    stream = BytesIO()
    raw.save(stream)
    doc = aw.DocxDocument(BytesIO(stream.getvalue()))
    header = doc.story("word/header1.xml").paragraphs[0]
    original_header = doc._package.payload("word/header1.xml")
    copied = doc.import_node(header, True, mode)
    assert copied.part_name == "word/document.xml" and copied.parent_node is None
    doc.body.append_child(copied)
    assert doc._package.payload("word/header1.xml") == original_header
    saved = Document(BytesIO(doc.to_bytes()))
    assert saved.paragraphs[-1].text == saved.sections[0].header.paragraphs[0].text == "HEADER"


def test_shallow_imported_table_remains_in_dom_but_is_omitted_from_saved_output():
    raw = Document()
    raw.add_table(rows=1, cols=1).cell(0, 0).text = "CELL"
    stream = BytesIO()
    raw.save(stream)
    source, target = aw.DocxDocument(BytesIO(stream.getvalue())), document()
    copied = target.import_node(source.body.tables[0], False)
    target.body.append_child(copied)
    assert copied.rows == () and copied.parent_node is target.body
    assert len(target.body.tables) == 1
    saved = target.to_bytes()
    assert len(Document(BytesIO(saved)).tables) == 0
    assert len(target.body.tables) == 1
    assert source.body.tables[0].rows[0].cells[0].paragraphs[0].text == "CELL"


@pytest.mark.parametrize("metadata", ["comment", "revision", "unknown", "description"])
def test_empty_table_with_opaque_or_revision_metadata_is_retained(metadata):
    raw = Document()
    raw.add_paragraph("BODY")
    table = raw.add_table(rows=1, cols=1)
    table._tbl.remove(table.rows[0]._tr)
    if metadata != "comment":
        tag = {"revision": "tblPrChange", "unknown": "futureMetadata", "description": "tblDescription"}[metadata]
        table._tbl.tblPr.append(OxmlElement("w:" + tag))
    else:
        from lxml.etree import Comment

        table._tbl.tblPr.append(Comment("retained metadata"))
    stream = BytesIO()
    raw.save(stream)
    doc = aw.DocxDocument(BytesIO(stream.getvalue()))
    doc.body.paragraphs[0].runs[0].font.bold = True
    saved = Document(BytesIO(doc.to_bytes()))
    assert len(saved.tables) == 1
    expected = {"revision": "tblPrChange", "unknown": "futureMetadata", "description": "tblDescription",
                "comment": "retained metadata"}[metadata]
    assert expected in saved.tables[0]._tbl.xml
