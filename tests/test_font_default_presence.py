"""Implicit ordinary defaults and explicit size origin survive conversion and JSON."""

import json
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest

from aspose.words_foss import Document, DocxDocument, SaveFormat
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.saving import OoxmlSaveOptions

ROOT = Path(__file__).parents[1]
BENCHMARKS = ROOT / 'docs/benchmarks'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
CASES = [(report, row) for name in ('font-defaults-26.9.json', 'font-default-matrix-26.9.json')
         for report in [json.loads((BENCHMARKS / name).read_text())] for row in report['records']]


def present(data):
    with ZipFile(BytesIO(data)) as archive:
        root = ET.fromstring(archive.read('word/styles.xml'))
    return root.find(W + 'docDefaults/' + W + 'rPrDefault') is not None


@pytest.mark.parametrize('report,row', CASES, ids=[report['corpus'] + row['input'] for report, row in CASES])
@pytest.mark.parametrize('format', [SaveFormat.DOCX, SaveFormat.FLAT_OPC])
@pytest.mark.parametrize('json_roundtrip', [False, True])
def test_default_group_presence_and_inherited_size_survive_conversion(report, row, format, json_roundtrip):
    with ZipFile(BENCHMARKS / report['corpus']) as archive:
        source = archive.read(row['input'])
    document = Document(BytesIO(source))
    model = document.light_document_model
    expected = present(source)
    assert model.doc_defaults_rpr_present is expected
    assert model.sections[0].body.paragraphs[0].runs[0].font.size == row['run_size']
    assert model.sections[0].body.paragraphs[0].runs[0].font.size_explicit is False
    if json_roundtrip:
        model = ldm.Document.model_validate_json(model.model_dump_json())
    before = model.model_dump_json()
    data = LdmDocxWriter(OoxmlSaveOptions(format)).write_to_bytes(model)
    assert model.model_dump_json() == before
    reopened = DocxDocument(BytesIO(data))
    assert present(reopened.to_bytes()) is expected
    assert reopened.body.paragraphs[0].runs[0].effective_font.size == row['run_size']
    assert reopened.styles.get_by_name('Derived').font.size == row['style_size']
    assert reopened.body.paragraphs[0].runs[0].font.size is None
    assert reopened.styles.get_by_name('Derived').direct_font.size is None


@pytest.mark.parametrize('report,row', CASES, ids=[report['corpus'] + row['input'] for report, row in CASES])
def test_implicit_rendered_size_matches_native_ordinary_getter(report, row):
    import pymupdf

    with ZipFile(BENCHMARKS / report['corpus']) as archive:
        source = archive.read(row['input'])
    data = Document(BytesIO(source)).to_bytes('pdf')
    with pymupdf.open(stream=data, filetype='pdf') as pdf:
        spans = [span for page in pdf for block in page.get_text('dict')['blocks']
                 for line in block.get('lines', []) for span in line['spans'] if span['text'] == 'IMPORT']
    assert len(spans) == 1 and spans[0]['size'] == pytest.approx(row['run_size'], abs=.001)


@pytest.mark.parametrize('declared', [False, True])
@pytest.mark.parametrize('json_roundtrip', [False, True])
def test_equal_explicit_and_inherited_size_remain_distinct(declared, json_roundtrip):
    from docx import Document as IndependentDocument
    from docx.shared import Pt

    source = IndependentDocument()
    source.styles['Normal'].font.size = Pt(11)
    run = source.add_paragraph().add_run('IMPORT')
    if declared:
        run.font.size = Pt(11)
    stream = BytesIO()
    source.save(stream)
    model = Document(BytesIO(stream.getvalue())).light_document_model
    assert model.sections[0].body.paragraphs[0].runs[0].font.size_explicit is declared
    if json_roundtrip:
        model = ldm.Document.model_validate_json(model.model_dump_json())
    data = LdmDocxWriter().write_to_bytes(model)
    reopened = DocxDocument(BytesIO(data))
    actual = reopened.body.paragraphs[0].runs[0]
    assert actual.effective_font.size == 11
    assert actual.font.size == (11 if declared else None)


@pytest.mark.parametrize('state,expected', [(None, True), (False, False), (True, True)])
def test_new_and_legacy_models_keep_the_default_group_policy(state, expected):
    model = ldm.Document(doc_defaults_rpr_present=state)
    assert present(LdmDocxWriter().write_to_bytes(model)) is expected
    model.doc_defaults_font = ldm.Font(size=14)
    data = LdmDocxWriter().write_to_bytes(model)
    assert present(data)
    with ZipFile(BytesIO(data)) as archive:
        root = ET.fromstring(archive.read('word/styles.xml'))
    assert root.find(W + 'docDefaults/' + W + 'rPrDefault/' + W + 'rPr/' + W + 'sz').get(W + 'val') == '28'
