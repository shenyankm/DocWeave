"""Declared relative hyperlinks survive conversion without rewriting literal text."""
from io import BytesIO
from xml.etree.ElementTree import fromstring
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss._flat_opc import decode, encode
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.saving import OoxmlSaveOptions


def linked_source(target):
    document = Document()
    paragraph = document.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = 'LINK'
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize('target', ['relative/path#', 'note', 'folder/report.docx?x=1&y=2', 'résumé.docx'])
@pytest.mark.parametrize('flat_input', [False, True])
@pytest.mark.parametrize('json_copy', [False, True])
@pytest.mark.parametrize('save_format', [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC])
def test_declared_relative_link_retains_relationship_and_display_text(target, flat_input, json_copy, save_format):
    source = linked_source(target)
    model = aw.Document(BytesIO(encode(source) if flat_input else source)).light_document_model
    if json_copy:
        model = ldm.Document.model_validate_json(model.model_dump_json())
    held = model.get_child_nodes(ldm.NodeType.RUN, True)[0]
    assert held.is_hyperlink
    original_text = held.text
    for _ in range(2):
        output = LdmDocxWriter(OoxmlSaveOptions(save_format)).write_to_bytes(model)
        package = decode(output) if save_format == aw.SaveFormat.FLAT_OPC else output
        with ZipFile(BytesIO(package)) as archive:
            body = fromstring(archive.read('word/document.xml'))
            links = body.findall('.//' + qn('w:hyperlink'))
            assert len(links) == 1
            assert links[0].find('.//' + qn('w:t')).text == 'LINK'
            relationships = fromstring(archive.read('word/_rels/document.xml.rels'))
            relationship = next(node for node in relationships if node.get('Id') == links[0].get(qn('r:id')))
            assert relationship.get('Target') == target
            assert relationship.get('TargetMode') == 'External'
        cold = aw.Document(BytesIO(output)).light_document_model.get_child_nodes(ldm.NodeType.RUN, True)[0]
        assert cold.is_hyperlink and cold.text == original_text
        assert held.text == original_text


@pytest.mark.parametrize('save_format', [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC])
@pytest.mark.parametrize('text, marked', [('[1](note)', False), ('[1]()', True), ('[1](<>)', True)])
def test_unmarked_relative_or_empty_target_pattern_remains_literal(save_format, text, marked):
    model = ldm.Document(sections=[ldm.Section(body=ldm.Body(children=[
        ldm.Paragraph(children=[ldm.Run(text=text, is_hyperlink=marked)])]))])
    output = LdmDocxWriter(OoxmlSaveOptions(save_format)).write_to_bytes(model)
    cold = aw.Document(BytesIO(output)).light_document_model.get_child_nodes(ldm.NodeType.RUN, True)[0]
    assert cold.text == text and not cold.is_hyperlink
