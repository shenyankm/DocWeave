"""Preserve character units independently of unresolved point/layout conversion."""

import json
import runpy
import warnings
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document as IndependentDocument
from docx.oxml.ns import qn

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.diagnostics import ContentLossWarning, collect_diagnostics
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.pdf_writer import LdmPdfWriter

ROOT = Path(__file__).parents[1]
BENCHMARKS = ROOT / "docs/benchmarks"
REPORT = json.loads((BENCHMARKS / "paragraph-character-indents-26.9.json").read_text())
FIELDS = ("character_unit_left_indent", "character_unit_right_indent", "character_unit_first_line_indent")
ATTRIBUTES = ("leftChars", "rightChars", "firstLineChars")
source = runpy.run_path(str(ROOT / "tests/test_character_indent_diagnostics.py"))["source"]


def load(raw):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", ContentLossWarning)
        return aw.Document(BytesIO(raw))


def units(document):
    paragraph = next(p for p in document.light_document_model.sections[0].body.paragraphs if "IMPORT" in p.text)
    return {name: getattr(paragraph.paragraph_format, name) or 0.0 for name in FIELDS}


@pytest.mark.parametrize("row", REPORT["records"], ids=lambda row: f'{row["input"]}-{row["value"]}-{row["layout"]}')
@pytest.mark.parametrize("format", [aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC])
def test_native_character_units_survive_model_json_and_ooxml_output(row, format):
    with ZipFile(BENCHMARKS / REPORT["outputs"]) as archive:
        document = load(archive.read(row["output"]))
    expected = {name: row["after_reopen"][name] for name in FIELDS}
    assert units(document) == expected
    model = document.light_document_model
    rebuilt = ldm.Document.model_validate_json(model.model_dump_json())
    assert rebuilt.model_dump() == model.model_dump()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", ContentLossWarning)
        saved = LdmDocxWriter(aw.saving.OoxmlSaveOptions(format)).write_to_bytes(rebuilt)
    assert units(load(saved)) == expected
    if format == aw.SaveFormat.DOCX:
        independent = IndependentDocument(BytesIO(saved))
        paragraph = next(p for p in independent.paragraphs if "IMPORT" in p.text)
        ind = paragraph._p.find(qn("w:pPr") + "/" + qn("w:ind"))
        for name, attr in zip(FIELDS, ATTRIBUTES):
            value = expected[name]
            if attr == "firstLineChars" and value < 0:
                attr, value = "hangingChars", -value
            if value:
                assert int(ind.get(qn("w:" + attr))) / 100 == value


@pytest.mark.parametrize("attribute", ["leftChars", "rightChars", "startChars", "endChars", "firstLineChars", "hangingChars"])
@pytest.mark.parametrize("kind", ["body", "table", "header", "footer", "style", "default", "numbering"])
@pytest.mark.parametrize("value", ["0", "113"])
def test_alias_units_and_explicit_zero_survive_supported_contexts(attribute, kind, value):
    document = load(source(attribute, kind, value))
    target = "character_unit_first_line_indent" if attribute in {"firstLineChars", "hangingChars"} else "character_unit_left_indent" if attribute in {"leftChars", "startChars"} else "character_unit_right_indent"
    expected = int(value) / (-100 if attribute == "hangingChars" else 100)
    def found(data):
        if isinstance(data, dict):
            return any(k == target and v == expected for k, v in data.items()) or any(found(v) for v in data.values())
        return isinstance(data, list) and any(found(v) for v in data)
    assert found(document.light_document_model.model_dump())
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", ContentLossWarning)
        saved = document.to_bytes(aw.SaveFormat.DOCX)
    assert found(load(saved).light_document_model.model_dump())
    with ZipFile(BytesIO(saved)) as archive:
        assert any((('Chars="' + value + '"').encode() in archive.read(name)) for name in archive.namelist() if name.endswith('.xml'))


@pytest.mark.parametrize("field", FIELDS)
@pytest.mark.parametrize("value", [0, 1.13, -1.13])
def test_model_edit_serialization_uses_hundredths_without_float_truncation(field, value):
    document = load(b"plain")
    model = document.light_document_model
    setattr(model.sections[0].body.paragraphs[0].paragraph_format, field, value)
    rebuilt = ldm.Document.model_validate_json(model.model_dump_json())
    saved = LdmDocxWriter().write_to_bytes(rebuilt)
    actual = load(saved).light_document_model.sections[0].body.paragraphs[0].paragraph_format
    assert getattr(actual, field) == value


@pytest.mark.parametrize("field", FIELDS)
@pytest.mark.parametrize("value,error", [(float("nan"), ValueError), (float("inf"), ValueError),
                                         (1e100, ValueError), (10**1000, ValueError), (True, TypeError), ("bad", TypeError)])
def test_invalid_model_character_edit_preserves_existing_output(tmp_path, field, value, error):
    model = load(b"plain").light_document_model
    setattr(model.sections[0].body.paragraphs[0].paragraph_format, field, value)
    destination = tmp_path / "existing.docx"
    destination.write_bytes(b"EXISTING")
    with pytest.raises(error):
        LdmDocxWriter().write(model, destination)
    assert destination.read_bytes() == b"EXISTING"


def test_new_model_character_edits_have_private_pdf_loss_diagnostics_and_strict_refusal(tmp_path):
    model = load(b"PRIVATE BODY").light_document_model
    model.sections[0].body.paragraphs[0].paragraph_format.character_unit_left_indent = 1.13
    diagnostics = []
    with collect_diagnostics(diagnostics), pytest.warns(ContentLossWarning, match="Character-unit"):
        assert LdmPdfWriter().write_to_bytes(model).startswith(b"%PDF")
    assert len([d for d in diagnostics if d.code == "pdf.character_indents_ignored"]) == 1
    assert "PRIVATE BODY" not in repr(diagnostics)
    destination = tmp_path / "existing.pdf"
    destination.write_bytes(b"EXISTING")
    with warnings.catch_warnings():
        warnings.simplefilter("error", ContentLossWarning)
        with pytest.raises(ContentLossWarning, match="Character-unit"):
            LdmPdfWriter().write(model, destination)
    assert destination.read_bytes() == b"EXISTING"


@pytest.mark.parametrize("attrs,field,expected", [
    ('w:leftChars="100" w:startChars="200"', FIELDS[0], 2.0),
    ('w:startChars="200" w:leftChars="100"', FIELDS[0], 1.0),
    ('w:rightChars="100" w:endChars="200"', FIELDS[1], 2.0),
    ('w:endChars="200" w:rightChars="100"', FIELDS[1], 1.0),
    ('w:firstLineChars="50" w:hangingChars="100"', FIELDS[2], -1.0),
    ('w:hangingChars="100" w:firstLineChars="50"', FIELDS[2], 0.5),
])
def test_native_character_alias_order_is_retained_in_model(attrs, field, expected):
    helper = runpy.run_path(str(ROOT / "docs/probes/style_toggles.py"))["document"]
    document = load(helper("", f"<w:ind {attrs}/>", ""))
    assert units(document)[field] == expected
    assert units(load(document.to_bytes(aw.SaveFormat.DOCX)))[field] == expected


def test_character_style_clear_and_zero_have_distinct_inheritance():
    helper = runpy.run_path(str(ROOT / "docs/probes/style_toggles.py"))["document"]
    styles = '<w:style w:type="paragraph" w:styleId="P"><w:name w:val="P"/><w:pPr><w:ind w:leftChars="113"/></w:pPr></w:style>'
    inherited = load(helper(styles, '<w:pStyle w:val="P"/>', ""))
    zero = load(helper(styles, '<w:pStyle w:val="P"/><w:ind w:leftChars="0"/>', ""))
    assert units(inherited)[FIELDS[0]] == 1.13
    assert units(zero)[FIELDS[0]] == 0
    assert units(load(inherited.to_bytes(aw.SaveFormat.DOCX)))[FIELDS[0]] == 1.13
    assert units(load(zero.to_bytes(aw.SaveFormat.DOCX)))[FIELDS[0]] == 0


@pytest.mark.parametrize("field", FIELDS)
def test_numbering_override_character_units_and_pdf_refusal(tmp_path, field):
    model = load(b"plain").light_document_model
    level = ldm.ListLevel()
    setattr(level, field, 1.13)
    model.lists.append(ldm.DocList(list_id=9, overrides=[ldm.ListLevelOverride(list_level=level)]))
    saved = LdmDocxWriter().write_to_bytes(model)
    actual = load(saved).light_document_model.lists[0].overrides[0].list_level
    assert getattr(actual, field) == 1.13
    destination = tmp_path / "existing.pdf"
    destination.write_bytes(b"EXISTING")
    with warnings.catch_warnings():
        warnings.simplefilter("error", ContentLossWarning)
        with pytest.raises(ContentLossWarning, match="Character-unit"):
            LdmPdfWriter().write(model, destination)
    assert destination.read_bytes() == b"EXISTING"


@pytest.mark.parametrize("field", FIELDS)
@pytest.mark.parametrize("value", [-21474836.48, 21474836.47])
def test_character_hundredths_signed_boundaries_roundtrip(field, value):
    model = load(b"plain").light_document_model
    setattr(model.sections[0].body.paragraphs[0].paragraph_format, field, value)
    actual = load(LdmDocxWriter().write_to_bytes(model)).light_document_model.sections[0].body.paragraphs[0].paragraph_format
    assert getattr(actual, field) == value


@pytest.mark.parametrize("value", ["1.13", "nan", "PRIVATE VALUE"])
def test_malformed_character_hundredths_fail_without_exposing_value(value):
    import traceback

    with pytest.raises(ValueError, match="integer hundredths") as error:
        load(source("leftChars", "body", value))
    assert value not in str(error.value)
    assert error.value.__suppress_context__
    assert value not in "".join(traceback.format_exception(error.value))


def test_direct_model_style_character_edit_has_pdf_loss_diagnostic(tmp_path):
    model = load(b"plain").light_document_model
    model.styles.append(ldm.Style(name="PRIVATE STYLE", paragraph_format=ldm.ParagraphFormat(character_unit_left_indent=1.13)))
    destination = tmp_path / "existing.pdf"
    destination.write_bytes(b"EXISTING")
    with warnings.catch_warnings():
        warnings.simplefilter("error", ContentLossWarning)
        with pytest.raises(ContentLossWarning, match="Character-unit") as error:
            LdmPdfWriter().write(model, destination)
    assert "PRIVATE STYLE" not in str(error.value)
    assert destination.read_bytes() == b"EXISTING"


@pytest.mark.parametrize("field", FIELDS)
def test_character_hundredths_do_not_depend_on_caller_decimal_precision(field):
    from decimal import localcontext

    from aspose.words_foss.docx_writer.constants import character_indent_attrs

    for value in (1.13, -1.13):
        format = ldm.ParagraphFormat(**{field: value})
        expected = character_indent_attrs(format)
        with localcontext() as context:
            context.prec = 1
            assert character_indent_attrs(format) == expected
