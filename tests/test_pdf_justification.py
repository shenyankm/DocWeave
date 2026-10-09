"""Justification stretches only soft-wrapped rows, including rich-text geometry."""

import pymupdf
import pytest

from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.saving import PdfSaveOptions


WORDS = "ONE TWO THREE FOUR FIVE SIX SEVEN EIGHT NINE TEN"


def justified_model(style="Normal", *, text=WORDS, mixed=False, widow=True, columns=1):
    runs = ([ldm.Run(text=word + (" " if index < 9 else ""),
                    font=ldm.Font(size=12 if index % 2 else 16, bold=index % 2 == 0))
             for index, word in enumerate(text.split())] if mixed else
            [ldm.Run(text=text, font=ldm.Font(size=12))])
    paragraph = ldm.Paragraph(children=runs, paragraph_format=ldm.ParagraphFormat(
        alignment=3, style_name=style, widow_control=widow))
    model = ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
        page_width=220, page_height=160, left_margin=20, right_margin=20,
        top_margin=20, bottom_margin=20,
        text_columns=ldm.TextColumns(count=columns, spacing=12)),
        body=ldm.Body(children=[paragraph]))])
    return model, paragraph


def lines(page, columns=1):
    rows = {}
    # PDF extraction can split a visual row at wide justified gaps or font boundaries.
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                column = int(columns == 2 and span["origin"][0] >= page.rect.width / 2)
                rows.setdefault((column, round(span["origin"][1], 2)), []).append(span)
    return [{"spans": spans, "bbox": (min(s["bbox"][0] for s in spans), min(s["bbox"][1] for s in spans),
                                       max(s["bbox"][2] for s in spans), max(s["bbox"][3] for s in spans))}
            for _, spans in sorted(rows.items())]


def render(model, shaping=False):
    options = PdfSaveOptions()
    options.text_shaping = shaping
    return LdmPdfWriter(options).write_to_bytes(model)


@pytest.mark.parametrize("style", ["Normal", "Code"])
@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("mixed", [False, True])
@pytest.mark.parametrize("widow", [False, True])
def test_justified_soft_rows_reach_margin_and_final_row_stays_natural(style, shaping, mixed, widow):
    model, _ = justified_model(style, mixed=mixed, widow=widow)
    snapshot = model.model_dump()
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        rows = lines(pdf[0])
        assert len(rows) >= 2
        assert [word[4] for page in pdf for word in page.get_text("words")] == WORDS.split()
        for row in rows[:-1]:
            assert row["bbox"][2] == pytest.approx(200, abs=0.05)
        assert rows[-1]["bbox"][2] < 190
        if mixed:
            assert {span["size"] for row in rows for span in row["spans"]} == {12, 16}
    assert model.model_dump() == snapshot


@pytest.mark.parametrize("style", ["Normal", "Code"])
@pytest.mark.parametrize("shaping", [False, True])
def test_explicit_breaks_keep_short_rows_natural(style, shaping):
    model, _ = justified_model(style, text="SHORT LINE\n" + WORDS)
    model.do_not_expand_shift_return = True
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        rows = lines(pdf[0])
        assert rows[0]["bbox"][2] < 120
        assert rows[1]["bbox"][2] == pytest.approx(200, abs=0.05)


@pytest.mark.parametrize("shaping", [False, True])
def test_no_space_rows_do_not_overflow_or_divide_by_zero(shaping):
    model, _ = justified_model(text="X" * 100)
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        assert sum(page.get_text().count("X") for page in pdf) == 100
        assert all(line["bbox"][2] <= 200 for page in pdf for line in lines(page))


@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("scripted", [False, True])
def test_justified_link_highlight_and_decorations_cover_stretched_text(shaping, scripted):
    model, paragraph = justified_model("Code")
    run = paragraph.runs[0]
    run.text = "[" + WORDS + "](https://example.test/code)"
    run.font.underline = run.font.strike_through = True
    run.font.highlight_color = "FFFF00"
    run.font.superscript = scripted
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        for page in pdf:
            links = page.get_links()
            for row in lines(page):
                rect = pymupdf.Rect(row["bbox"])
                assert any(abs(link["from"].x0 - rect.x0) < 0.05 and
                           abs(link["from"].x1 - rect.x1) < 0.05 for link in links)
                drawings = page.get_drawings()
                assert any(d["fill"] == (1, 1, 0) and abs(d["rect"].x1 - rect.x1) < 0.05
                           for d in drawings)
                assert sum(d["rect"].height < 2 and abs(d["rect"].x1 - rect.x1) < 0.05
                           for d in drawings) >= 2


@pytest.mark.parametrize("columns", [1, 2])
@pytest.mark.parametrize("shaping", [False, True])
def test_justification_survives_page_and_column_breaks(columns, shaping):
    text = " ".join(f"WORD{i:03}" for i in range(100))
    model, _ = justified_model("Code", text=text, columns=columns)
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        assert len(pdf) > 1
        assert [word[4] for page in pdf for word in page.get_text("words")] == text.split()
        width = (180 - (columns - 1) * 12) / columns
        rows = [line for page in pdf for line in lines(page, columns)]
        for row in rows[:-1]:
            if " " in "".join(span["text"] for span in row["spans"]):
                assert row["bbox"][2] - row["bbox"][0] == pytest.approx(width, abs=0.05)
            else:
                assert row["bbox"][2] - row["bbox"][0] <= width + 0.05


@pytest.mark.parametrize("shaping", [False, True])
def test_first_line_indent_preserves_justified_right_edge(shaping):
    model, paragraph = justified_model()
    paragraph.paragraph_format.first_line_indent = 12
    with pymupdf.open(stream=render(model, shaping), filetype="pdf") as pdf:
        rows = lines(pdf[0])
        assert rows[0]["bbox"][0] == pytest.approx(20 + 12, abs=0.05)
        for row in rows[:-1]:
            assert row["bbox"][2] == pytest.approx(200, abs=0.05)


@pytest.mark.parametrize("linked", [False, True])
def test_justified_bidi_fallback_format_boundaries_preserve_glyph_positions(linked):
    from .test_pdf_shaping import arabic_font
    from .test_pdf_bidi_formatting import glyphs

    options = PdfSaveOptions()
    options.text_shaping = True
    options.fallback_fonts = [str(arabic_font())]
    parts = ("سلام ", "(123) ", "عالم ", "ABC ", "نهاية ") * 3

    def output(fragmented):
        model, paragraph = justified_model(text="".join(parts))
        paragraph._children = [ldm.Run(
            text=f"[{part}](https://example.test/{index})" if linked and fragmented else part,
            font=ldm.Font(size=12, color="FF0000" if fragmented and index % 2 else ""))
            for index, part in enumerate(parts if fragmented else ("".join(parts),))]
        return LdmPdfWriter(options).write_to_bytes(model)

    actual, expected = glyphs(output(True)), glyphs(output(False))
    assert [item[0] for item in actual] == [item[0] for item in expected]
    for item, target in zip(actual, expected, strict=True):
        assert item[1:] == pytest.approx(target[1:], abs=0.05)


def test_public_docx_justified_code_reaches_margin(tmp_path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    import aspose.words_foss as aw

    document = Document()
    document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    section = document.sections[0]
    section.page_width, section.page_height = Pt(220), Pt(160)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Pt(20)
    paragraph = document.add_paragraph(style="Code")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(WORDS)
    run.font.size = Pt(12)
    source, target = tmp_path / "justified.docx", tmp_path / "justified.pdf"
    document.save(source)
    aw.Document(source).save(target)
    with pymupdf.open(target) as pdf:
        rows = lines(pdf[0])
        assert rows[0]["bbox"][2] == pytest.approx(200, abs=0.05)
        assert rows[-1]["bbox"][2] < 190
