"""Explicit export semantics and independently verified package preservation."""

from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
import pytest

import dockiln as dk
from .test_docx_dom import package, payloads
from .test_docx_dom_resources import image_bytes










def assert_independent_report(doc, original):
    current = payloads(doc.to_bytes())
    report = doc.preservation_report()
    assert report == {
        "added": sorted(current.keys() - original.keys()),
        "removed": sorted(original.keys() - current.keys()),
        "modified": sorted(name for name in current.keys() & original.keys()
                           if current[name] != original[name]),
        "unchanged": sorted(name for name in current.keys() & original.keys()
                            if current[name] == original[name]),
    }
    return report


def test_preservation_report_includes_resources_and_does_not_reset_on_save(tmp_path):
    source = tmp_path / "original.docx"
    original = package(source)
    doc = dk.DocxDocument(source)
    doc.part_xml("word/styles.xml")
    doc.create_paragraph("detached")
    assert assert_independent_report(doc, original)["modified"] == []
    doc.body.paragraphs[0].runs[0].text = "新客户"
    doc.body.paragraphs[0].add_picture(image_bytes(), width=30)
    report = assert_independent_report(doc, original)
    assert "word/document.xml" in report["modified"]
    assert "[Content_Types].xml" in report["modified"]
    assert any(name.startswith("word/media/") for name in report["added"])
    assert "customXml/item.xml" in report["unchanged"]
    doc.save(tmp_path / "edited.docx")
    assert doc.preservation_report() == report
    reopened = dk.DocxDocument(tmp_path / "edited.docx")
    assert reopened.preservation_report()["modified"] == []


def test_resource_replacement_is_compared_with_original_bytes(tmp_path):
    source = tmp_path / "original.docx"
    original = package(source)
    doc = dk.DocxDocument(source)
    name = "customXml/item.xml"
    doc._package.set_parts({name: b"changed"})
    assert assert_independent_report(doc, original)["modified"] == [name]
    doc._package.set_parts({name: original[name]})
    assert assert_independent_report(doc, original)["modified"] == []
