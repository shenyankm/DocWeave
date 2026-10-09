"""Default styles apply without pStyle; explicit formatting still wins."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
import pytest

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer.styles_part import apply_reference_styles, render_styles_xml

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def load(source):
    stream = BytesIO()
    source.save(stream)
    return aw.Document(BytesIO(stream.getvalue()))


@pytest.mark.parametrize('location', ['body', 'table', 'header', 'footer'])
@pytest.mark.parametrize('flag', ['1', 'true', 'on'])
def test_implicit_default_style_reaches_all_stories(location, flag):
    source = Document()
    normal = source.styles['Normal']
    normal.element.set(qn('w:default'), flag)
    normal.font.size = Pt(16)
    normal.font.bold = True
    normal.paragraph_format.left_indent = Pt(23)
    normal.paragraph_format.space_after = Pt(19)
    if location == 'body':
        para = source.add_paragraph('DEFAULT')
    elif location == 'table':
        para = source.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        para.add_run('DEFAULT')
    else:
        para = getattr(source.sections[0], location).paragraphs[0]
        para.add_run('DEFAULT')
    para.style = None
    assert para._p.find(W+'pPr/'+W+'pStyle') is None
    doc = load(source)
    section = doc.light_document_model.sections[0]
    if location == 'body':
        actual = section.body.paragraphs[0]
    elif location == 'table':
        actual = section.body.tables[0].rows[0].cells[0].paragraphs[0]
    else:
        actual = next(p for hf in section.headers_footers for p in hf.paragraphs if p.text)
    assert actual.runs[0].font.size == 16
    assert actual.runs[0].font.bold is True
    assert actual.paragraph_format.left_indent == 23
    assert actual.paragraph_format.space_after == 19
    assert actual.paragraph_format.style_name == 'Normal'


def test_custom_default_based_on_chain_and_default_marker_survive_output():
    source = Document()
    normal = source.styles['Normal']
    normal.element.set(qn('w:default'), '0')
    normal.font.italic = True
    custom = source.styles.add_style('Brand Default', WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = normal
    custom.element.set(qn('w:default'), '1')
    custom.font.size = Pt(17)
    custom.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.add_paragraph('CUSTOM')
    doc = load(source)
    para = doc.light_document_model.sections[0].body.paragraphs[0]
    assert para.paragraph_format.style_name == 'Brand Default'
    assert para.runs[0].font.size == 17
    assert para.runs[0].font.italic is True
    raw = doc.to_bytes('docx')
    with ZipFile(BytesIO(raw)) as package:
        styles = ET.fromstring(package.read('word/styles.xml'))
    defaults = [s.find(W+'name').get(W+'val') for s in styles.findall(W+'style')
                if s.get(W+'type') == 'paragraph' and s.get(W+'default') == '1']
    assert defaults == ['Brand Default']
    result = Document(BytesIO(raw))
    assert result.paragraphs[0].style.name == 'Brand Default'
    assert result.paragraphs[0].style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_explicit_style_and_direct_overrides_take_precedence():
    source = Document()
    normal = source.styles['Normal']
    normal.font.size = Pt(16)
    normal.font.bold = True
    normal.paragraph_format.left_indent = Pt(23)
    explicit = source.styles.add_style('Explicit', WD_STYLE_TYPE.PARAGRAPH)
    explicit.font.size = Pt(14)
    explicit.font.bold = False
    source.add_paragraph('STYLE', style=explicit)
    para = source.add_paragraph('DIRECT')
    para.runs[0].font.bold = False
    para.runs[0].font.size = Pt(21)
    para.paragraph_format.left_indent = Pt(7)
    doc = load(source)
    actual = doc.light_document_model.sections[0].body.paragraphs
    assert [p.runs[0].font.size for p in actual] == [14, 21]
    assert [p.runs[0].font.bold for p in actual] == [False, False]
    assert [p.paragraph_format.left_indent for p in actual] == [0, 7]
    result = Document(BytesIO(doc.to_bytes('docx')))
    assert result.paragraphs[1].runs[0].font.bold is False
    assert result.paragraphs[1].runs[0].font.size.pt == 21
    assert result.paragraphs[1].paragraph_format.left_indent.pt == 7


@pytest.mark.parametrize('flag', ['0', 'false', 'off'])
def test_normal_name_does_not_imply_default(flag):
    source = Document()
    normal = source.styles['Normal']
    normal.element.set(qn('w:default'), flag)
    normal.font.size = Pt(16)
    source.add_paragraph('NO DEFAULT')
    doc = load(source)
    assert doc.light_document_model.sections[0].body.paragraphs[0].runs[0].font.size == 11
    with ZipFile(BytesIO(doc.to_bytes('docx'))) as package:
        styles = ET.fromstring(package.read('word/styles.xml'))
    normal_xml = next(s for s in styles.findall(W+'style')
                      if s.find(W+'name').get(W+'val') == 'Normal')
    assert normal_xml.get(W+'default') != '1'


def test_default_style_font_reaches_actual_pdf():
    import fitz

    source = Document()
    source.styles['Normal'].font.size = Pt(16)
    source.add_paragraph('DEFAULT')
    doc = load(source)
    with fitz.open(stream=doc.to_bytes('pdf'), filetype='pdf') as pdf:
        spans = [s for b in pdf[0].get_text('dict')['blocks']
                 for line in b.get('lines', []) for s in line['spans']]
        assert next(s for s in spans if s['text'] == 'DEFAULT')['size'] == pytest.approx(16)


def test_reference_default_replaces_source_default_without_ambiguity():
    source = ldm.Document(styles=[ldm.Style(name='Source Default', type=1, is_default=True)])
    reference = ldm.Document(styles=[ldm.Style(name='Reference Default', type=1, is_default=True)])
    styles = ET.fromstring(apply_reference_styles(render_styles_xml(source), reference))
    defaults = [s.find(W+'name').get(W+'val') for s in styles.findall(W+'style')
                if s.get(W+'type') == 'paragraph' and s.get(W+'default') == '1']
    assert defaults == ['Reference Default']


def test_default_style_inherits_numbering_from_its_base():
    source = Document()
    source.styles['Normal'].element.set(qn('w:default'), '0')
    default = source.styles.add_style('Default List', WD_STYLE_TYPE.PARAGRAPH)
    default.base_style = source.styles['List Number']
    default.element.set(qn('w:default'), '1')
    source.add_paragraph('FIRST')
    source.add_paragraph('SECOND')
    doc = load(source)
    paras = doc.light_document_model.sections[0].body.paragraphs
    assert all(p.list_format.is_list_item for p in paras)
    assert paras[0].list_format.list_id == paras[1].list_format.list_id > 0
    markdown = doc.to_bytes('md').decode()
    assert markdown.splitlines() == ['1. FIRST', '1. SECOND']
    result = aw.Document(BytesIO(doc.to_bytes('docx')))
    assert all(p.list_format.is_list_item for p in result.light_document_model.sections[0].body.paragraphs)


def test_old_model_without_default_field_keeps_normal_default():
    old = ldm.Document.model_validate({'styles': [
        {'name': 'Normal', 'type': 1, 'built_in': True, 'font': {'size': 17}}]})
    assert old.styles[0].is_default is None
    styles = ET.fromstring(render_styles_xml(old))
    normal = next(s for s in styles.findall(W+'style')
                  if s.find(W+'name').get(W+'val') == 'Normal')
    assert normal.get(W+'default') == '1'


@pytest.mark.parametrize('location', ['body', 'table', 'header', 'footer'])
@pytest.mark.parametrize('flag', ['1', 'true', 'on', '0', 'false', 'off'])
def test_last_enabled_default_of_matching_type_wins(location, flag):
    source = Document()
    source.styles['Normal'].element.set(qn('w:default'), '0')
    first = source.styles.add_style('First Default', WD_STYLE_TYPE.PARAGRAPH)
    first.element.set(qn('w:default'), '1')
    first.font.size = Pt(16)
    last = source.styles.add_style('Last Default', WD_STYLE_TYPE.PARAGRAPH)
    last.element.set(qn('w:default'), flag)
    last.base_style = first
    last.font.size = Pt(19)
    for name, kind in [('Character Default', WD_STYLE_TYPE.CHARACTER),
                       ('Table Default', WD_STYLE_TYPE.TABLE)]:
        other = source.styles.add_style(name, kind)
        other.element.set(qn('w:default'), '1')
        other.font.size = Pt(32)
    if location == 'body':
        container = source
        implicit = source.add_paragraph()
    elif location == 'table':
        container = source.add_table(1, 1).cell(0, 0)
        implicit = container.paragraphs[0]
    else:
        container = getattr(source.sections[0], location)
        implicit = container.paragraphs[0]
    implicit.style = None
    assert implicit._p.pPr.pStyle is None
    implicit.add_run('IMPLICIT')
    container.add_paragraph('EXPLICIT', style=first)
    direct = container.add_paragraph().add_run('DIRECT')
    direct.font.size = Pt(21)
    document = load(source)
    expected_size = 19 if flag in ('1', 'true', 'on') else 16
    expected_name = 'Last Default' if expected_size == 19 else 'First Default'

    def check(model):
        paragraphs = model.get_child_nodes(ldm.NodeType.PARAGRAPH, True)
        by_text = {p.text: p for p in paragraphs}
        assert by_text['IMPLICIT'].paragraph_format.style_name == expected_name
        assert by_text['IMPLICIT'].runs[0].font.size == expected_size
        assert by_text['EXPLICIT'].runs[0].font.size == 16
        assert by_text['DIRECT'].runs[0].font.size == 21

    check(document.light_document_model)
    output = document.to_bytes('docx')
    check(aw.Document(BytesIO(output)).light_document_model)
    with ZipFile(BytesIO(output)) as archive:
        styles = ET.fromstring(archive.read('word/styles.xml'))
    defaults = [style.find(W+'name').get(W+'val') for style in styles.findall(W+'style')
                if style.get(W+'type') == 'paragraph' and style.get(W+'default') == '1']
    assert defaults[-1] == expected_name


def test_last_default_style_reaches_actual_pdf():
    import fitz

    source = Document()
    source.styles['Normal'].font.size = Pt(12)
    last = source.styles.add_style('Last Default', WD_STYLE_TYPE.PARAGRAPH)
    last.element.set(qn('w:default'), 'on')
    last.font.size = Pt(19)
    source.add_paragraph('LAST DEFAULT')
    document = load(source)
    with fitz.open(stream=document.to_bytes('pdf'), filetype='pdf') as pdf:
        spans = [s for block in pdf[0].get_text('dict')['blocks']
                 for line in block.get('lines', []) for s in line['spans']]
        assert next(s for s in spans if s['text'] == 'LAST DEFAULT')['size'] == pytest.approx(19)
