"""Public declarations can be inventoried without installing the commercial runtime."""

import runpy
from pathlib import Path

import pytest


@pytest.mark.parametrize("name", ["verify_paragraph_pagination", "verify_first_paragraph_trial", "verify_pagination_rendering", "verify_paragraph_dimensions", "verify_paragraph_character_indents", "verify_paragraph_character_setters", "verify_paragraph_character_reads", "verify_paragraph_character_inheritance_edits", "verify_character_indent_roundtrips", "verify_font_size_loading"])
def test_generated_packages_accept_windows_zip_creator_metadata(monkeypatch, name):
    from zipfile import ZipInfo

    original = ZipInfo.__init__

    def windows_info(self, *args, **kwargs):
        original(self, *args, **kwargs)
        self.create_system = 0  # Windows central-directory creator; payloads are identical.

    monkeypatch.setattr(ZipInfo, "__init__", windows_info)
    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))[name]
    assert check(root / "docs/benchmarks") in {10, 18, 24, 61, 122, 126, 144, 216, 360, 1296}



@pytest.mark.parametrize("metric,delta,accepted", [("advance", 0.00001, True), ("advance", 0.05, False),
                                                   ("origin", 0.00001, True), ("origin", 0.05, False),
                                                   ("page", 1, False), ("size", 0.05, False)])
def test_pagination_snapshot_uses_declared_tolerance(monkeypatch, metric, delta, accepted):
    root = Path(__file__).parents[1]
    original = runpy.run_path
    check = original(str(root / "scripts/verify_commercial_baseline.py"))["verify_pagination_rendering"]

    def measured(path, *args, **kwargs):
        functions = original(path, *args, **kwargs)
        if Path(path).name == "pagination_rendering.py":
            snapshot = functions["pdf_snapshot"]

            def perturbed(raw):
                result = snapshot(raw)
                line = result["lines"]["ANCHOR"]
                if metric == "origin":
                    line[metric][0] += delta
                else:
                    line[metric] += delta
                return result

            functions["pdf_snapshot"] = perturbed
        return functions

    monkeypatch.setattr(runpy, "run_path", measured)
    if accepted:
        assert check(root / "docs/benchmarks") == 10
    else:
        with pytest.raises(AssertionError):
            check(root / "docs/benchmarks")


def test_pagination_evidence_rejects_forged_setter_result(tmp_path):
    import json
    import shutil

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_paragraph_pagination"]
    (tmp_path / "corpus").mkdir()
    for name in ("paragraph-pagination-26.9.zip", "paragraph-pagination-26.9-outputs.zip"):
        shutil.copyfile(root / "docs/benchmarks/corpus" / name, tmp_path / "corpus" / name)
    report = json.loads((root / "docs/benchmarks/paragraph-pagination-26.9.json").read_text())
    report["records"][0]["edited"]["derived"] = not report["records"][0]["value"]
    (tmp_path / "paragraph-pagination-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


def test_style_font_edit_evidence_rejects_forged_setter_result(tmp_path):
    import json
    import shutil

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_style_font_edits"]
    benchmarks = root / "docs/benchmarks"
    (tmp_path / "corpus").mkdir()
    for name in ("style-import-conflicts-26.9.zip", "style-font-edits-26.9.zip"):
        shutil.copyfile(benchmarks / "corpus" / name, tmp_path / "corpus" / name)
    report = json.loads((benchmarks / "style-font-edits-26.9.json").read_text())
    row = report["records"][0]
    row["after_edit"]["styles"]["Derived"][row["property"]] = not row["value"]
    (tmp_path / "style-font-edits-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


def test_implicit_size_evidence_rejects_forged_observation(tmp_path):
    import json
    import shutil

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_font_default_matrix"]
    (tmp_path / "corpus").mkdir()
    shutil.copyfile(root / "docs/benchmarks/corpus/font-default-matrix-26.9.zip",
                    tmp_path / "corpus/font-default-matrix-26.9.zip")
    report = json.loads((root / "docs/benchmarks/font-default-matrix-26.9.json").read_text())
    report["records"][0]["run_size"] = 99
    (tmp_path / "font-default-matrix-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


def test_style_paragraph_evidence_rejects_forged_setter_result(tmp_path):
    import json
    import shutil

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_style_paragraph_formats"]
    (tmp_path / "corpus").mkdir()
    for name in ("style-paragraph-format-26.9.zip", "style-paragraph-format-26.9-outputs.zip"):
        shutil.copyfile(root / "docs/benchmarks/corpus" / name, tmp_path / "corpus" / name)
    report = json.loads((root / "docs/benchmarks/style-paragraph-format-26.9.json").read_text())
    report["records"][0]["after_edit"]["style_alignment"]["Derived"] = "RIGHT"
    (tmp_path / "style-paragraph-format-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


def test_inventory_preserves_overloads_setters_and_enum_values(tmp_path):
    inventory = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "inventory_commercial_api.py")
    )["inventory"]
    (tmp_path / "__init__.pyi").write_text('''
from enum import Enum
class Format(Enum):
    DOCX: int = 20
class Document:
    def __getitem__(self, index: int) -> str: ...
    def save(self, filename: str) -> None: ...
    def save(self, stream: object, format: Format) -> None: ...
    @property
    def text(self) -> str:
        """Malformed published prose - "[Red]""""
        ...
    @text.setter
    def text(self, value: str): ...
    def _private(self): ...
''')
    hooks = tmp_path / "__nuitka"
    hooks.mkdir()
    (hooks / "__init__.pyi").write_text("class PackagingHook: ...")
    modules = inventory(tmp_path)
    assert len(modules) == 1
    symbols = {symbol["id"]: symbol for symbol in modules[0]["symbols"]}
    assert symbols["aspose.words.Format"]["kind"] == "enum"
    assert symbols["aspose.words.Format.DOCX"]["value"] == "20"
    assert len(symbols["aspose.words.Document.save"]["declarations"]) == 2
    assert symbols["aspose.words.Document.__getitem__"]["kind"] == "method"
    text = symbols["aspose.words.Document.text"]
    assert text["kind"] == "property"
    assert [item["decorators"] for item in text["declarations"]] == [["property"], ["text.setter"]]
    assert "aspose.words.Document._private" not in symbols


def test_runtime_inventory_keeps_inherited_owner_and_unstubbed_module():
    from types import ModuleType

    observe = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "inventory_runtime_surface.py")
    )["observe"]
    root = ModuleType("aspose.words")
    child_module = ModuleType("aspose.words.lowcode")

    class Parent:
        def visit(self):
            return None

    class Child(Parent):
        pass

    Parent.__module__ = Child.__module__ = root.__name__
    root.Parent = Parent
    root.Child = Child
    root.lowcode = child_module
    # A module cycle must terminate without dropping either surface.
    child_module.parent = root
    declarations = {"modules": [{"module": root.__name__, "symbols": [{"id": "aspose.words.Parent.visit"}]}]}
    modules = observe(root, declarations)
    assert len(modules) == 2
    assert modules[1]["declaration_module_present"] is False
    child = next(item for item in modules[0]["exports"] if item["id"] == "aspose.words.Child")
    visit = next(item for item in child["members"] if item["name"] == "visit")
    assert visit == {"name": "visit", "owner": "aspose.words.Parent.visit", "declaration_present": True}


def test_ledger_does_not_promote_matching_names_to_compatibility():
    build = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "build_alignment_ledger.py")
    )["build"]
    declarations = {"version": "26.9.0", "capability_defaults": {"baseline_behavior": "not_measured"},
                    "modules": [{"symbols": [{"id": "aspose.words.Document.save"}]}]}
    baseline = {"root_module": "aspose.words", "modules": [{"exports": [
        {"id": "aspose.words.Document", "kind": "type", "members": [
            {"name": "save", "owner": "aspose.words.Document.save"},
            {"name": "dynamic", "owner": "aspose.words.Document.dynamic"},
        ]},
    ]}]}
    current = {"root_module": "aspose.words_foss", "runtime_version": "26.7.0.post2",
               "modules": [{"exports": [{"id": "aspose.words_foss.Document", "kind": "type",
                                         "members": [{"name": "save", "owner": "aspose.words_foss.Document.save"}]}]}]}
    ledger = build(declarations, baseline, current)
    records = {record["id"]: record for record in ledger["records"]}
    assert records["aspose.words.Document.save"]["implementation"] == "name_present_behavior_unverified"
    assert records["aspose.words.Document.dynamic"]["current_entrypoint"] is None
    assert records["aspose.words.Document.dynamic"]["scope_confirmation"] == "runtime_only_requires_public_support_confirmation"
    assert ledger["record_defaults"]["baseline_behavior"] == "not_measured"


def test_runtime_enum_inventory_preserves_aliases():
    from enum import IntEnum

    members = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "inventory_runtime_enums.py")
    )["members"]

    class Format(IntEnum):
        PRIMARY = 1
        ALIAS = 1
        OTHER = 2

    assert members(Format) == [
        {"name": "PRIMARY", "value": 1, "canonical_name": "PRIMARY"},
        {"name": "ALIAS", "value": 1, "canonical_name": "PRIMARY"},
        {"name": "OTHER", "value": 2, "canonical_name": "OTHER"},
    ]


def test_enum_comparison_keeps_missing_types_and_protocol_gaps():
    from types import SimpleNamespace

    compare = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "compare_enum_values.py")
    )["compare"]
    root = SimpleNamespace(NodeType=type("NodeType", (), {"RUN": 21, "BODY": 99}))
    enums = [{"id": "aspose.words.NodeType", "values": [
        {"name": "RUN", "value": 21}, {"name": "BODY", "value": 3}, {"name": "SECTION", "value": 2}]},
        {"id": "aspose.words.lowcode.SplitCriteria", "values": [{"name": "PAGE", "value": 0}]}]
    rows = compare(enums, root)
    assert rows[0]["current_is_enum"] is False
    assert [value["status"] for value in rows[0]["values"]] == ["integer_matches", "integer_differs", "absent_or_not_integer"]
    assert rows[1]["current_type_present"] is False


def test_format_ledger_keeps_missing_samples_and_literal_markup_failures():
    format_records = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "build_alignment_ledger.py")
    )["format_records"]
    saved = [{"name": "UNKNOWN", "value": 0, "outcome": {"status": "raised"}}]
    loaded = [{"name": "CHM", "value": 60, "status": "missing_legal_sample_not_verified"},
              {"name": "HTML", "value": 50, "status": "returned", "loaded": True}]
    reports = {"save_commercial": {"records": saved}, "save_current": {"records": saved},
               "load_commercial": {"records": loaded}, "load_current": {"records": loaded}}
    rows = format_records({"reports": reports, "independent_docx_checks": {"load_current": [
        {"format": "HTML", "literal_markup_present": True}]}})
    assert len(rows) == 3
    assert rows[1]["gap"] == "legal input sample missing; capability remains unverified"
    assert rows[2]["gap"] == "markup accepted as literal text; semantic structure not preserved"
    assert all(row["delivery"] == "baseline_observation_only" for row in rows)


@pytest.mark.parametrize("tamper", ["archive_digest", "output_digest", "missing_output"])
def test_format_output_verifier_rejects_tampered_or_missing_evidence(tmp_path, tamper):
    from hashlib import sha256
    from zipfile import ZipFile

    check = runpy.run_path(str(Path(__file__).parents[1] / "scripts" / "verify_commercial_baseline.py"))["verify_format_outputs"]
    archive = tmp_path / "outputs.zip"
    with ZipFile(archive, "w") as zipfile:
        zipfile.writestr("save_current/TEXT/output.bin", b"CONTENT")
    report = {"outputs": {"sha256": sha256(archive.read_bytes()).hexdigest()},
              "reports": {"save_current": {"records": [{"name": "TEXT", "files": [
                  {"name": "output.bin", "size": 7, "sha256": sha256(b"CONTENT").hexdigest()}]}]}},
              "independent_docx_checks": {}}
    assert check(report, archive) == 1
    if tamper == "archive_digest":
        report["outputs"]["sha256"] = "incorrect"
    elif tamper == "output_digest":
        report["reports"]["save_current"]["records"][0]["files"][0]["sha256"] = "incorrect"
    else:
        report["reports"]["save_current"]["records"][0]["files"][0]["name"] = "missing.bin"
    with pytest.raises((AssertionError, KeyError)):
        check(report, archive)


def test_frozen_evidence_verifies_after_windows_text_checkout_and_detects_binary_change(tmp_path):
    import json
    from shutil import copyfile

    root = Path(__file__).parents[1]
    verify = runpy.run_path(str(root / "scripts" / "verify_commercial_baseline.py"))["verify"]
    source = root / "docs" / "benchmarks"
    names = ["commercial-26.9-api.json", "commercial-26.9-capabilities.json", "commercial-26.9-all-enums.json",
             "commercial-26.9-baseline.json", "commercial-26.9-doc-registry.json", "commercial-26.9-format-behavior.json",
             "commercial-26.9-literal-text.json", "corpus/commercial-26.9-dom.docx",
             "corpus/commercial-26.9-format-outputs.zip", "corpus/commercial-26.9-literal-text.zip",
             "corpus/commercial-26.9-literal-outputs.zip", "import-node-26.9.json", "import-node-current.json",
             "corpus/import-node-26.9.zip", "corpus/import-node-current.zip",
             "style-toggles-26.9.json", "style-toggles-current.json", "corpus/style-toggles-26.9.zip",
             "style-import-conflicts-26.9.json", "corpus/style-import-conflicts-26.9.zip",
             "corpus/style-import-conflict-outputs.zip", "style-import-translated.json",
             "corpus/style-import-translated.zip", "font-defaults-26.9.json", "corpus/font-defaults-26.9.zip",
             "style-import-default-on.json", "corpus/style-import-default-on.zip",
             "style-import-character-default-on.json", "corpus/style-import-character-default-on.zip",
             "paragraph-style-defaults-26.9.json", "corpus/paragraph-style-defaults-26.9.zip",
             "corpus/paragraph-style-default-outputs.zip", "corpus/paragraph-default-repeat-outputs.zip",
             "style-save-roundtrips-26.9.json", "corpus/style-save-roundtrips-26.9.zip",
             "style-save-state-26.9.json", "corpus/style-save-state-26.9.zip"]
    names += ["character-style-save-state-26.9.json", "corpus/character-style-save-state-26.9.zip",
              "corpus/character-style-defaults-26.9.zip"]
    names += ["paragraph-character-indents-current.json", "corpus/paragraph-character-indents-current.zip"]
    names += ["paragraph-character-setters-26.9.json", "corpus/paragraph-character-setters-26.9.zip"]
    names += ["paragraph-character-reads-26.9.json", "corpus/paragraph-character-reads-26.9.zip"]
    names += ["paragraph-character-inheritance-edits-26.9.json", "corpus/paragraph-character-inheritance-edits-26.9.zip"]
    names += ["paragraph-character-indents-26.9.json", "corpus/paragraph-character-indents-26.9.zip",
              "corpus/paragraph-character-indents-26.9-outputs.zip"]
    names += ["paragraph-logical-indents-26.9.json", "corpus/paragraph-logical-indents-26.9.zip",
              "corpus/paragraph-logical-indents-26.9-outputs.zip", "paragraph-indent-limits-26.9.json", "corpus/paragraph-indent-limits-26.9.zip",
              "corpus/paragraph-indent-limits-26.9-outputs.zip", "corpus/paragraph-indent-limits-26.9-normalized.zip",
              "paragraph-spacing-limits-26.9.json", "corpus/paragraph-spacing-limits-26.9.zip",
              "corpus/paragraph-spacing-limits-26.9-outputs.zip", "paragraph-dimensions-26.9.json", "corpus/paragraph-dimensions-26.9.zip",
              "corpus/paragraph-dimensions-26.9-outputs.zip", "paragraph-dimensions-rendering-26.9.json",
              "corpus/paragraph-dimensions-rendering-26.9.zip", "corpus/paragraph-dimensions-rendering-26.9-commercial.zip",
              "corpus/paragraph-dimensions-rendering-26.9-current.zip"]
    names += ["pagination-rendering-26.9.json", "corpus/pagination-rendering-26.9.zip",
              "corpus/pagination-rendering-26.9-commercial.zip", "corpus/pagination-rendering-26.9-current.zip",
              "corpus/pagination-rendering-before-padding-fix.zip",
              "first-paragraph-trial-26.9.json", "corpus/first-paragraph-trial-26.9.zip", "corpus/first-paragraph-trial-26.9-outputs.zip",
              "first-paragraph-page-break-26.9.json", "corpus/first-paragraph-page-break-26.9.zip", "corpus/first-paragraph-page-break-26.9-outputs.zip",
              "paragraph-pagination-26.9.json", "corpus/paragraph-pagination-26.9.zip", "corpus/paragraph-pagination-26.9-outputs.zip",
              "style-paragraph-format-26.9.json", "corpus/style-paragraph-format-26.9.zip", "corpus/style-paragraph-format-26.9-outputs.zip",
              "font-default-matrix-26.9.json", "corpus/font-default-matrix-26.9.zip",
              "style-font-edits-26.9.json", "corpus/style-font-edits-26.9.zip",
              "style-normalization-contexts-26.9.json", "corpus/style-normalization-contexts-26.9.zip",
              "corpus/style-normalization-context-outputs-26.9.zip", "style-import-projections.json",
              "corpus/style-import-projections.zip"]
    names += ["font-size-loading-26.9.json", "corpus/font-size-loading-26.9.zip",
              "corpus/font-size-loading-26.9-outputs.zip", "corpus/font-size-loading-current.zip"]
    names += ["font-default-presence-26.9.json", "corpus/font-default-presence-current.zip"]
    for name in names:
        target = tmp_path / name
        target.parent.mkdir(exist_ok=True)
        if target.suffix == ".json":
            target.write_bytes((source / name).read_bytes().replace(b"\r\n", b"\n").replace(b"\n", b"\r\n"))
        else:
            copyfile(source / name, target)
    result = verify(tmp_path)
    assert result["checked_format_outputs"] == 100 and result["behavioral_acceptance"] is False
    assert result["checked_import_outputs"] == 168
    assert result["checked_style_inputs"] == 745
    assert result["checked_style_import_outputs"] == 1878
    assert result["checked_style_roundtrip_outputs"] == 494
    assert result["checked_style_save_states"] == 523
    assert result["checked_style_projection_outputs"] == 523
    assert result["checked_style_font_edit_outputs"] == 1482
    assert result["checked_font_default_matrix_inputs"] == 16
    assert result["checked_style_paragraph_edit_outputs"] == 108
    assert result["checked_paragraph_pagination_outputs"] == 1296
    assert result["checked_unresolved_first_paragraph_outputs"] == 4
    assert result["checked_trial_first_paragraph_outputs"] == 18
    assert result["checked_pagination_rendering_pairs"] == 10
    assert result["checked_paragraph_dimension_edits"] == 360
    assert result["checked_paragraph_spacing_limits"] == 32
    assert result["checked_paragraph_indent_limits"] == 54
    assert result["checked_paragraph_logical_indents"] == 168
    assert result["checked_paragraph_character_indents"] == 61
    assert result["checked_paragraph_character_setters"] == 144
    assert result["checked_paragraph_character_reads"] == 24
    assert result["checked_paragraph_character_inheritance_edits"] == 216
    assert result["checked_character_indent_roundtrips"] == 122
    assert result["checked_paragraph_dimension_rendering_pairs"] == 10
    assert result["checked_font_size_loading_inputs"] == 126
    assert result["checked_font_default_inputs"] == 5
    font_rows = [row for row in json.loads((tmp_path / "commercial-26.9-capabilities.json").read_text())["records"]
                 if row["id"] in {"aspose.words.Font.bold", "aspose.words.Font.italic"}]
    assert len(font_rows) == 2
    assert all(row["behavior_evidence"]["file"] == "style-toggles-26.9.json" for row in font_rows)
    assert all("canonical Font getter alignment remains incomplete" in row["baseline_behavior"] for row in font_rows)
    import_rows = [row for row in json.loads((tmp_path / "commercial-26.9-capabilities.json").read_text())["records"]
                   if row["id"] in {"aspose.words.Document.import_node", "aspose.words.DocumentBase.import_node"}]
    assert len(import_rows) == 2
    assert all(row["style_conflict_evidence"]["file"] == "style-import-projections.json" for row in import_rows)
    assert all("523 cases returned" in row["style_conflict_evidence"]["delivery"] for row in import_rows)
    archive = tmp_path / "corpus" / "commercial-26.9-literal-outputs.zip"
    archive.write_bytes(archive.read_bytes() + b"corruption")
    with pytest.raises(AssertionError):
        verify(tmp_path)


def test_import_evidence_rejects_fabricated_observation_and_changed_archive(tmp_path):
    import json
    from shutil import copyfile

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts" / "verify_commercial_baseline.py"))["verify_import_outputs"]
    source = root / "docs" / "benchmarks"
    (tmp_path / "corpus").mkdir()
    for name in ("import-node-26.9.json", "import-node-current.json",
                 "corpus/import-node-26.9.zip", "corpus/import-node-current.zip"):
        copyfile(source / name, tmp_path / name)
    assert check(tmp_path) == 168
    path = tmp_path / "import-node-current.json"
    original = path.read_bytes()
    data = json.loads(original)
    data["independent_checks"][0]["paragraphs"][0]["style_chain"][0]["font"]["bold"] = True
    path.write_text(json.dumps(data))
    with pytest.raises(AssertionError, match="independent observation"):
        check(tmp_path)
    path.write_bytes(original)
    archive = tmp_path / "corpus" / "import-node-current.zip"
    archive.write_bytes(archive.read_bytes() + b"corruption")
    with pytest.raises(AssertionError, match="archive digest"):
        check(tmp_path)


def test_style_evidence_rejects_changed_digest_missing_case_and_fabricated_match(tmp_path):
    import json
    from shutil import copyfile

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts" / "verify_commercial_baseline.py"))["verify_style_toggles"]
    source = root / "docs" / "benchmarks"
    (tmp_path / "corpus").mkdir()
    for name in ("style-toggles-26.9.json", "style-toggles-current.json", "corpus/style-toggles-26.9.zip"):
        copyfile(source / name, tmp_path / name)
    assert check(tmp_path) == 745
    path = tmp_path / "style-toggles-26.9.json"
    original = path.read_bytes()
    report = json.loads(original)
    report["records"][0]["sha256"] = "incorrect"
    path.write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)
    path.write_bytes(original)
    candidate_path = tmp_path / "style-toggles-current.json"
    original_candidate = candidate_path.read_bytes()
    candidate = json.loads(original_candidate)
    row = candidate["reports"]["candidate"]["records"][0]
    row["bold"] = not row["bold"]
    candidate_path.write_text(json.dumps(candidate))
    with pytest.raises(AssertionError):
        check(tmp_path)
    candidate_path.write_bytes(original_candidate)
    report = json.loads(original)
    report["records"].pop()
    path.write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("translated", [False, True])
def test_style_import_evidence_rejects_forged_layers_and_changed_outputs(tmp_path, translated):
    import json
    from shutil import copyfile

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts" / "verify_commercial_baseline.py"))["verify_style_imports"]
    source = root / "docs" / "benchmarks"
    (tmp_path / "corpus").mkdir()
    for name in ("style-import-conflicts-26.9.json", "corpus/style-import-conflicts-26.9.zip",
                 "corpus/style-import-conflict-outputs.zip"):
        copyfile(source / name, tmp_path / name)
    name = "style-import-translated.json" if translated else "style-import-conflicts-26.9.json"
    if translated:
        for file in (name, "corpus/style-import-translated.zip"):
            copyfile(source / file, tmp_path / file)
    assert check(tmp_path, name) == (247 if translated else 741)
    path = tmp_path / name
    original = path.read_bytes()
    report = json.loads(original)
    row = report["independent_checks"]["candidate" if translated else "commercial"][0]["styles"][0]
    row["font"]["bold"] = not row["font"]["bold"]
    path.write_text(json.dumps(report))
    with pytest.raises(AssertionError, match="independent observation"):
        check(tmp_path, name)
    path.write_bytes(original)
    archive = tmp_path / report["outputs"]["archive"]
    archive.write_bytes(archive.read_bytes() + b"corruption")
    with pytest.raises(AssertionError, match="outputs digest"):
        check(tmp_path, name)


@pytest.mark.parametrize("name", ["style-import-default-on.json", "paragraph-style-defaults-26.9.json"])
def test_story_evidence_rejects_forged_destination_format(tmp_path, name):
    import json
    from shutil import copyfile

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts" / "verify_commercial_baseline.py"))["verify_style_imports"]
    source = root / "docs" / "benchmarks"
    report = json.loads((source / name).read_text())
    files = {name, report["corpus"]["archive"], report["outputs"]["archive"]}
    files.update(report[key]["file"] for key in ("baseline", "before_baseline") if key in report)
    if "repeat_outputs" in report:
        files.add(report["repeat_outputs"]["archive"])
    for file in files:
        path = tmp_path / file
        path.parent.mkdir(exist_ok=True)
        copyfile(source / file, path)
    assert check(tmp_path, name) == report["outputs"]["files"]
    row = report["story_rereads"]["candidate"]["records"][0]["formats"]["DESTINATION"]
    row["bold"] = not row["bold"]
    (tmp_path / name).write_text(json.dumps(report))
    with pytest.raises(AssertionError, match="saved story observation"):
        check(tmp_path, name)


def test_roundtrip_evidence_rejects_forged_after_format(tmp_path):
    import json
    from shutil import copyfile

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts" / "verify_commercial_baseline.py"))["verify_style_roundtrips"]
    source = root / "docs" / "benchmarks"
    name = "style-save-roundtrips-26.9.json"
    report = json.loads((source / name).read_text())
    for file in (name, report["corpus"]["archive"], report["outputs"]["archive"]):
        path = tmp_path / file
        path.parent.mkdir(exist_ok=True)
        copyfile(source / file, path)
    assert check(tmp_path) == 494
    report["report"]["records"][0]["after"]["bold"] = not report["report"]["records"][0]["after"]["bold"]
    (tmp_path / name).write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("phase,category", [("after_save_live", "styles"), ("after_reopen", "styles"),
                                           ("after_reopen", "paragraphs")])
def test_save_state_evidence_rejects_forged_getters(tmp_path, phase, category):
    import json
    from shutil import copyfile

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts" / "verify_commercial_baseline.py"))["verify_style_save_state"]
    source = root / "docs" / "benchmarks"
    name = "style-save-state-26.9.json"
    report = json.loads((source / name).read_text())
    for file in (name, report["outputs"]["archive"], *(item["corpus"] for item in report["reports"])):
        path = tmp_path / file
        path.parent.mkdir(exist_ok=True)
        copyfile(source / file, path)
    assert check(tmp_path) == 379
    fonts = report["reports"][0]["records"][0][phase][category]
    font = next(iter(fonts.values()))
    font["bold"] = not font["bold"]
    (tmp_path / name).write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("phase", ["cold", "warm"])
def test_projection_evidence_rejects_forged_formats(tmp_path, phase):
    import json
    from shutil import copyfile

    root = Path(__file__).parents[1]
    source = root / "docs" / "benchmarks"
    check = runpy.run_path(str(root / "scripts" / "verify_commercial_baseline.py"))["verify_style_projections"]
    name = "style-import-projections.json"
    report = json.loads((source / name).read_text())
    files = {name, report["outputs"]["archive"]}
    files.update(matrix["native"]["file"] for matrix in report["reports"])
    files.update(matrix["corpus"]["archive"] for matrix in report["reports"])
    for file in files:
        path = tmp_path / file
        path.parent.mkdir(exist_ok=True)
        copyfile(source / file, path)
    assert check(tmp_path) == 523
    matrix = report["reports"][0]
    font = matrix["official_reread"]["records"][0]["formats"]["DESTINATION"] if phase == "cold" else matrix["candidate"]["records"][0]["imported_format"]
    font["bold"] = not font["bold"]
    (tmp_path / name).write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


def test_font_default_evidence_rejects_fabricated_value_and_missing_input(tmp_path):
    import json
    from shutil import copyfile

    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / "scripts" / "verify_commercial_baseline.py"))["verify_font_defaults"]
    source = root / "docs" / "benchmarks"
    (tmp_path / "corpus").mkdir()
    for name in ("font-defaults-26.9.json", "corpus/font-defaults-26.9.zip"):
        copyfile(source / name, tmp_path / name)
    assert check(tmp_path) == 5
    path = tmp_path / "font-defaults-26.9.json"
    original = path.read_bytes()
    report = json.loads(original)
    report["records"][0]["run_size"] = 12.0
    path.write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)
    report = json.loads(original)
    report["records"].pop()
    path.write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


def test_format_ledger_identifies_refusal_return_contract_and_load_error_gaps():
    format_records = runpy.run_path(str(Path(__file__).parents[1] / "scripts" / "build_alignment_ledger.py"))["format_records"]
    baseline_saved = [{"name": name, "value": value, "outcome": {"status": "returned", "return_type": "SaveOutputParameters"}}
                      for name, value in [("DOC", 10), ("DOCX", 20)]]
    current_saved = [{"name": "DOC", "value": 10, "outcome": {"status": "raised", "exception_type": "ValueError"}},
                     {"name": "DOCX", "value": 20, "outcome": {"status": "returned", "return_type": "NoneType"}}]
    baseline_loaded = [{"name": "PDF", "value": 40, "status": "returned", "loaded": True}]
    current_loaded = [{"name": "PDF", "value": 40, "status": "raised", "loaded": False,
                       "phase": "load", "exception_type": "UnicodeDecodeError"}]
    reports = {key: {"records": rows} for key, rows in [("save_commercial", baseline_saved),
               ("save_current", current_saved), ("load_commercial", baseline_loaded), ("load_current", current_loaded)]}
    rows = format_records({"reports": reports, "independent_docx_checks": {"load_current": []}})
    assert "rejects or fails" in rows[0]["gap"]
    assert "return contract differs" in rows[1]["gap"]
    assert rows[2]["implementation"]["exception_type"] == "UnicodeDecodeError"
    assert rows[2]["implementation"]["phase"] == "load" and "roundtrip fails" in rows[2]["gap"]


def test_indent_normalization_rejects_unrelated_package_edits(tmp_path):
    import hashlib
    import json
    from io import BytesIO
    from shutil import copyfile
    from zipfile import ZipFile

    root = Path(__file__).parents[1]
    source = root / "docs/benchmarks"
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_paragraph_dimensions"]
    name = "paragraph-indent-limits-26.9.json"
    report = json.loads((source / name).read_text())
    for file in (name, report["corpus"], report["outputs"], report["normalized_outputs"]):
        destination = tmp_path / file
        destination.parent.mkdir(exist_ok=True)
        copyfile(source / file, destination)
    assert check(tmp_path, name) == 54
    row = next(row for row in report["records"] if row.get("native_xml_issue"))
    path = tmp_path / report["normalized_outputs"]
    with ZipFile(path) as archive:
        entries = {part: archive.read(part) for part in archive.namelist()}
    with ZipFile(BytesIO(entries[row["output"]])) as archive:
        parts = {part: archive.read(part) for part in archive.namelist()}
    parts["[Content_Types].xml"] += b" "
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        for part, data in parts.items():
            archive.writestr(part, data)
    entries[row["output"]] = stream.getvalue()
    row["normalized_sha256"] = hashlib.sha256(entries[row["output"]]).hexdigest()
    with ZipFile(path, "w") as archive:
        for part, data in entries.items():
            archive.writestr(part, data)
    report["normalized_outputs_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    (tmp_path / name).write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path, name)


@pytest.mark.parametrize("tamper", ["loaded", "after_edit", "before_save", "after_save_live", "after_reopen", "coverage", "setter_error", "acceptance"])
def test_character_indent_evidence_rejects_forged_state(tmp_path, tamper):
    import json
    import shutil

    root = Path(__file__).parents[1]
    benchmarks = root / "docs/benchmarks"
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_paragraph_character_indents"]
    report = json.loads((benchmarks / "paragraph-character-indents-26.9.json").read_text())
    (tmp_path / "corpus").mkdir()
    for key in ("corpus", "outputs"):
        shutil.copyfile(benchmarks / report[key], tmp_path / report[key])
    if tamper == "coverage":
        report["records"][-1] = report["records"][-2]
    elif tamper == "setter_error":
        report["setter_errors"][0]["error"] = "returned"
    elif tamper == "acceptance":
        report["sdk_acceptance"] = True
    else:
        report["records"][0][tamper]["left_indent"] += 1
    (tmp_path / "paragraph-character-indents-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("attribute", ["left", "leftChars"])
def test_character_indent_evidence_checks_xml_after_digests_are_updated(tmp_path, attribute):
    import json
    import shutil
    from hashlib import sha256
    from io import BytesIO
    from zipfile import ZipFile

    from docx import Document
    from docx.oxml.ns import qn

    root = Path(__file__).parents[1]
    benchmarks = root / "docs/benchmarks"
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_paragraph_character_indents"]
    report = json.loads((benchmarks / "paragraph-character-indents-26.9.json").read_text())
    (tmp_path / "corpus").mkdir()
    shutil.copyfile(benchmarks / report["corpus"], tmp_path / report["corpus"])
    row = report["records"][0]
    with ZipFile(benchmarks / report["outputs"]) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    document = Document(BytesIO(parts[row["output"]]))
    paragraph = next(p for p in document.paragraphs if "IMPORT" in p.text)
    paragraph._p.find(qn("w:pPr") + "/" + qn("w:ind")).set(qn("w:" + attribute), "12345")
    stream = BytesIO()
    document.save(stream)
    parts[row["output"]] = stream.getvalue()
    row["output_sha256"] = sha256(stream.getvalue()).hexdigest()
    archive_path = tmp_path / report["outputs"]
    with ZipFile(archive_path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    report["outputs_sha256"] = sha256(archive_path.read_bytes()).hexdigest()
    (tmp_path / "paragraph-character-indents-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("tamper", ["loaded", "after_edit", "after_save_live", "after_reopen", "coverage", "setter_error", "acceptance", "corpus_digest", "property"])
def test_character_inheritance_edits_reject_forged_evidence(tmp_path, tamper):
    import json
    import shutil

    root = Path(__file__).parents[1]
    source = root / "docs/benchmarks"
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_paragraph_character_inheritance_edits"]
    report = json.loads((source / "paragraph-character-inheritance-edits-26.9.json").read_text())
    (tmp_path / "corpus").mkdir()
    for key in ("corpus", "outputs"):
        shutil.copyfile(source / report[key], tmp_path / report[key])
    shutil.copyfile(source / "paragraph-character-reads-26.9.json", tmp_path / "paragraph-character-reads-26.9.json")
    if tamper == "coverage":
        report["records"][-1] = report["records"][-2]
    elif tamper == "setter_error":
        report["setter_errors"][0]["error"] = "returned"
    elif tamper == "acceptance":
        report["sdk_acceptance"] = True
    elif tamper == "corpus_digest":
        report["corpus_sha256"] = "forged"
    elif tamper == "property":
        report["records"][0]["property"] = "other"
    else:
        report["records"][0][tamper]["derived"]["left_indent"] += 1
    (tmp_path / "paragraph-character-inheritance-edits-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("target,attribute", [(target, attribute) for target in ("base", "derived", "paragraph")
                                             for attribute in ("left", "leftChars")] +
                         [("base", "basedOn"), ("derived", "basedOn"), ("paragraph", "pStyle")])
def test_character_inheritance_edits_verify_xml_with_updated_digests(tmp_path, target, attribute):
    import json
    import shutil
    from hashlib import sha256
    from io import BytesIO
    from zipfile import ZipFile

    from docx import Document
    from docx.oxml.ns import qn

    root = Path(__file__).parents[1]
    source = root / "docs/benchmarks"
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_paragraph_character_inheritance_edits"]
    report = json.loads((source / "paragraph-character-inheritance-edits-26.9.json").read_text())
    (tmp_path / "corpus").mkdir()
    shutil.copyfile(source / report["corpus"], tmp_path / report["corpus"])
    shutil.copyfile(source / "paragraph-character-reads-26.9.json", tmp_path / "paragraph-character-reads-26.9.json")
    row = next(row for row in report["records"] if row["target"] == target and row["value"] == 1.235)
    with ZipFile(source / report["outputs"]) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    document = Document(BytesIO(entries[row["output"]]))
    element = (next(p for p in document.paragraphs if "IMPORT" in p.text)._p
               if target == "paragraph" else document.styles[target.capitalize()].element)
    if attribute in {"basedOn", "pStyle"}:
        from docx.oxml import OxmlElement

        parent = element.find(qn("w:pPr")) if attribute == "pStyle" else element
        setting = parent.find(qn("w:" + attribute))
        if setting is None:
            setting = OxmlElement("w:" + attribute)
            parent.append(setting)
        setting.set(qn("w:val"), "Derived" if target == "base" else "Base" if target == "paragraph" else "Normal")
    else:
        element.find(qn("w:pPr") + "/" + qn("w:ind")).set(qn("w:" + attribute), "12345")
    stream = BytesIO()
    document.save(stream)
    entries[row["output"]] = stream.getvalue()
    row["output_sha256"] = sha256(stream.getvalue()).hexdigest()
    path = tmp_path / report["outputs"]
    with ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    report["outputs_sha256"] = sha256(path.read_bytes()).hexdigest()
    (tmp_path / "paragraph-character-inheritance-edits-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("tamper", ["value", "coverage", "digest", "acceptance"])
def test_character_read_evidence_rejects_forged_inheritance(tmp_path, tamper):
    import json
    import shutil

    root = Path(__file__).parents[1]
    source = root / "docs/benchmarks"
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_paragraph_character_reads"]
    report = json.loads((source / "paragraph-character-reads-26.9.json").read_text())
    (tmp_path / "corpus").mkdir()
    shutil.copyfile(source / report["corpus"], tmp_path / report["corpus"])
    if tamper == "coverage":
        report["records"][-1] = report["records"][-2]
    elif tamper == "digest":
        report["corpus_sha256"] = "forged"
    elif tamper == "acceptance":
        report["sdk_acceptance"] = True
    else:
        report["records"][0]["values"]["derived"]["character_unit_left_indent"] = 1
    (tmp_path / "paragraph-character-reads-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("tamper", ["loaded", "after_edit", "after_save_live", "after_reopen", "coverage", "setter_error", "acceptance"])
def test_character_setter_context_evidence_rejects_forged_state(tmp_path, tamper):
    import json
    import shutil

    root = Path(__file__).parents[1]
    source = root / "docs/benchmarks"
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_paragraph_character_setters"]
    report = json.loads((source / "paragraph-character-setters-26.9.json").read_text())
    (tmp_path / "corpus").mkdir()
    for key in ("corpus", "outputs"):
        shutil.copyfile(source / report[key], tmp_path / report[key])
    shutil.copyfile(source / "paragraph-character-indents-26.9.json", tmp_path / "paragraph-character-indents-26.9.json")
    if tamper == "coverage":
        report["records"][-1] = report["records"][-2]
    elif tamper == "setter_error":
        report["setter_errors"][0]["error"] = "returned"
    elif tamper == "acceptance":
        report["sdk_acceptance"] = True
    else:
        report["records"][0][tamper]["target"]["left_indent"] += 1
    (tmp_path / "paragraph-character-setters-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("target,attribute", [("paragraph", "left"), ("paragraph", "leftChars"),
                                             ("style", "left"), ("style", "leftChars")])
def test_character_setter_context_evidence_checks_xml_after_updated_digests(tmp_path, target, attribute):
    import json
    import shutil
    from hashlib import sha256
    from io import BytesIO
    from zipfile import ZipFile

    from docx import Document
    from docx.oxml.ns import qn

    root = Path(__file__).parents[1]
    source = root / "docs/benchmarks"
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_paragraph_character_setters"]
    report = json.loads((source / "paragraph-character-setters-26.9.json").read_text())
    (tmp_path / "corpus").mkdir()
    shutil.copyfile(source / report["corpus"], tmp_path / report["corpus"])
    shutil.copyfile(source / "paragraph-character-indents-26.9.json", tmp_path / "paragraph-character-indents-26.9.json")
    row = next(row for row in report["records"] if row["target"] == target and row["value"] == 1.235)
    with ZipFile(source / report["outputs"]) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    document = Document(BytesIO(entries[row["output"]]))
    element = (next(p for p in document.paragraphs if "IMPORT" in p.text)._p
               if target == "paragraph" else document.styles["P"].element)
    element.find(qn("w:pPr") + "/" + qn("w:ind")).set(qn("w:" + attribute), "12345")
    stream = BytesIO()
    document.save(stream)
    entries[row["output"]] = stream.getvalue()
    row["output_sha256"] = sha256(stream.getvalue()).hexdigest()
    archive_path = tmp_path / report["outputs"]
    with ZipFile(archive_path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    report["outputs_sha256"] = sha256(archive_path.read_bytes()).hexdigest()
    (tmp_path / "paragraph-character-setters-26.9.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize("tamper", ["character_units", "coverage", "native_digest", "acceptance", "official_reread"])
def test_character_roundtrip_evidence_rejects_forged_observation(tmp_path, tamper):
    import json
    import shutil

    root = Path(__file__).parents[1]
    benchmarks = root / "docs/benchmarks"
    check = runpy.run_path(str(root / "scripts/verify_commercial_baseline.py"))["verify_character_indent_roundtrips"]
    report = json.loads((benchmarks / "paragraph-character-indents-current.json").read_text())
    (tmp_path / "corpus").mkdir()
    shutil.copyfile(benchmarks / report["outputs"], tmp_path / report["outputs"])
    shutil.copyfile(benchmarks / report["native_report"], tmp_path / report["native_report"])
    if tamper == "coverage":
        report["records"][-1] = report["records"][-2]
    elif tamper == "native_digest":
        report["native_report_sha256"] = "forged"
    elif tamper == "acceptance":
        report["full_format_acceptance"] = True
    elif tamper == "official_reread":
        report["official_reread"]["records"][0]["character_units"]["character_unit_left_indent"] += 1
    else:
        report["records"][0]["character_units"]["character_unit_left_indent"] += 1
    (tmp_path / "paragraph-character-indents-current.json").write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        check(tmp_path)


@pytest.mark.parametrize('phase', ['loaded', 'after_save', 'cold', 'current_run', 'current_hash', 'current_coverage'])
def test_font_size_loading_evidence_rejects_forged_measurement(tmp_path, phase):
    import json
    import shutil

    root = Path(__file__).parents[1]
    benchmarks = root / 'docs/benchmarks'
    report = json.loads((benchmarks / 'font-size-loading-26.9.json').read_text())
    (tmp_path / 'corpus').mkdir()
    for key in ('corpus', 'outputs'):
        shutil.copyfile(benchmarks / report[key], tmp_path / report[key])
    shutil.copyfile(benchmarks / report['docweave_roundtrip']['outputs'],
                    tmp_path / report['docweave_roundtrip']['outputs'])
    if phase == 'current_run':
        report['docweave_roundtrip']['native_reread']['records'][0]['run_size'] += .5
    elif phase == 'current_hash':
        report['docweave_roundtrip']['records'][0]['output_sha256'] = 'forged'
    elif phase == 'current_coverage':
        rows = report['docweave_roundtrip']['native_reread']['records']
        rows[0] = rows[1]
    else:
        report['records'][0][phase]['run_size'] += .5
    (tmp_path / 'font-size-loading-26.9.json').write_text(json.dumps(report))
    check = runpy.run_path(str(root / 'scripts/verify_commercial_baseline.py'))['verify_font_size_loading']
    with pytest.raises(AssertionError):
        check(tmp_path)


def test_font_size_loading_evidence_accepts_frozen_outputs():
    root = Path(__file__).parents[1]
    check = runpy.run_path(str(root / 'scripts/verify_commercial_baseline.py'))['verify_font_size_loading']
    assert check(root / 'docs/benchmarks') == 126


@pytest.mark.parametrize('change', [None, 'native_size', 'source_hash', 'presence', 'coverage', 'xml_direct_size'])
def test_font_default_presence_evidence_checks_independent_outputs(tmp_path, change):
    import hashlib
    import json
    import shutil
    from io import BytesIO
    from xml.etree import ElementTree as ET
    from zipfile import ZipFile

    root = Path(__file__).parents[1]
    source = root / 'docs/benchmarks'
    report = json.loads((source / 'font-default-presence-26.9.json').read_text())
    (tmp_path / 'corpus').mkdir()
    for filename in report['sources']:
        shutil.copyfile(source / filename, tmp_path / filename)
        corpus = json.loads((source / filename).read_text())['corpus']
        shutil.copyfile(source / corpus, tmp_path / corpus)
    output_path = tmp_path / report['outputs']
    shutil.copyfile(source / report['outputs'], output_path)
    row = report['records'][0]
    if change == 'native_size':
        report['native_reread']['records'][0]['run_size'] += .5
    elif change == 'source_hash':
        row['source_sha256'] = 'forged'
    elif change == 'presence':
        row['rpr_default_present'] = not row['rpr_default_present']
    elif change == 'coverage':
        report['records'][0] = report['records'][1]
    elif change == 'xml_direct_size':
        w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        with ZipFile(output_path) as archive:
            outputs = {name: archive.read(name) for name in archive.namelist()}
        with ZipFile(BytesIO(outputs[row['output']])) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        document = ET.fromstring(parts['word/document.xml'])
        run = next(r for r in document.iter(w + 'r') if ''.join(t.text or '' for t in r.iter(w + 't')) == 'IMPORT')
        properties = run.find(w + 'rPr')
        if properties is None:
            properties = ET.SubElement(run, w + 'rPr')
        ET.SubElement(properties, w + 'sz', {w + 'val': str(int(row['expected_run_size'] * 2))})
        parts['word/document.xml'] = ET.tostring(document)
        stream = BytesIO()
        with ZipFile(stream, 'w') as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        outputs[row['output']] = stream.getvalue()
        with ZipFile(output_path, 'w') as archive:
            for name, data in outputs.items():
                archive.writestr(name, data)
        row['output_sha256'] = hashlib.sha256(outputs[row['output']]).hexdigest()
        report['outputs_sha256'] = hashlib.sha256(output_path.read_bytes()).hexdigest()
    (tmp_path / 'font-default-presence-26.9.json').write_text(json.dumps(report))
    check = runpy.run_path(str(root / 'scripts/verify_commercial_baseline.py'))['verify_font_default_presence']
    if change is None:
        assert check(tmp_path) == 348
    else:
        with pytest.raises(AssertionError):
            check(tmp_path)
