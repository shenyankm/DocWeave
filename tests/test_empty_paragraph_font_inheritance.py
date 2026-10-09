"""Blank paragraphs use the same font cascade as visible text."""

from io import BytesIO

import pymupdf
import pytest
from docx import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from aspose.words_foss import Document
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_reader.ldm_builder.cascading import ParagraphFormatResolver
from aspose.words_foss.pdf_writer import LdmPdfWriter


def source_document(kind, mark):
    source = WordDocument()
    paragraph = source.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if kind == "normal":
        source.styles["Normal"].font.size = Pt(8)
        expected = 8
    elif kind == "chain":
        parent = source.styles.add_style("Blank Parent", WD_STYLE_TYPE.PARAGRAPH)
        parent.font.size = Pt(30)
        parent.font.bold = True
        child = source.styles.add_style("Blank Child", WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = parent
        paragraph.style = child
        expected = 30
    else:
        rpr = (
            source.styles.element.find(qn("w:docDefaults"))
            .find(qn("w:rPrDefault"))
            .find(qn("w:rPr"))
        )
        size = rpr.find(qn("w:sz"))
        size.set(qn("w:val"), "12")
        source.styles["Normal"].font.size = None
        expected = 6
    if mark is not None:
        properties = OxmlElement("w:rPr")
        child = OxmlElement("w:b" if mark == "partial" else "w:sz")
        child.set(qn("w:val"), "0" if mark == "partial" else "32")
        properties.append(child)
        paragraph._p.get_or_add_pPr().append(properties)
        if mark != "partial":
            expected = 16
    return source, expected


@pytest.mark.parametrize("kind", ["normal", "chain", "defaults"])
@pytest.mark.parametrize("mark", [None, "partial", "size"])
def test_blank_cell_inherits_font_and_direct_mark_overrides(kind, mark):
    source, expected = source_document(kind, mark)
    stream = BytesIO()
    source.save(stream)
    loaded = Document(BytesIO(stream.getvalue()))
    para = (
        loaded.light_document_model.sections[0]
        .body.tables[0]
        .rows[0]
        .cells[0]
        .paragraphs[0]
    )
    assert not para._children
    assert para.paragraph_break_font.size == expected
    if mark == "partial":
        assert para.paragraph_break_font.bold is False
    table = ldm.Table(
        top_padding=0,
        bottom_padding=0,
        rows=[
            ldm.Row(
                row_format=ldm.RowFormat(
                    borders=[ldm.Border(line_style=1, line_width=0.5) for _ in range(6)]
                ),
                cells=[ldm.Cell(paragraphs=[para])],
            )
        ],
    )
    model = ldm.Document(sections=[ldm.Section(body=ldm.Body(children=[table]))])
    original = model.model_dump_json()
    with pymupdf.open(
        stream=LdmPdfWriter().write_to_bytes(model), filetype="pdf"
    ) as pdf:
        assert not pdf[0].get_text().strip()
        rects = [drawing["rect"] for drawing in pdf[0].get_drawings()]
        assert max(rect.y1 for rect in rects) - min(
            rect.y0 for rect in rects
        ) == pytest.approx(expected * 1.4, abs=0.02)
    assert model.model_dump_json() == original
    # Effective empty-line size survives DOCX rebuilding and an independent reload.
    rewritten = Document(BytesIO(loaded.to_bytes("docx")))
    actual = (
        rewritten.light_document_model.sections[0]
        .body.tables[0]
        .rows[0]
        .cells[0]
        .paragraphs[0]
    )
    assert actual.paragraph_break_font.size == expected


def test_partial_mark_merge_preserves_inherited_fields_without_aliasing():
    parent = ldm.Font(size=8, bold=True)
    base = ldm.ParagraphFormat(paragraph_break_font=parent)
    override = ldm.ParagraphFormat(
        paragraph_break_font=ldm.Font(bold=False, italic=True)
    )
    ParagraphFormatResolver.merge(base, override)
    assert base.paragraph_break_font.size == 8
    assert base.paragraph_break_font.bold is False
    assert base.paragraph_break_font.italic is True
    base.paragraph_break_font.size = 20
    assert parent.size == 8 and parent.bold is True
    assert override.paragraph_break_font.size == 0
