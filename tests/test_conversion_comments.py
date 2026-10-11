"""Owned classic comments survive edited conversion and JSON save cycles."""
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import warnings

from defusedxml import DefusedXmlException
from docx import Document as NativeDocument
import pytest
import dockiln as dk
from dockiln import light_document_model as ldm
from dockiln._flat_opc import encode, decode
from dockiln.writers.docx import LdmDocxWriter

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def owned_comment(case='range'):
    d = NativeDocument()
    p = d.add_paragraph()
    p.add_run('OWNED BEFORE ')
    anchor = p.add_run('OWNED ANCHOR')
    p.add_run(' OWNED AFTER')
    runs = [anchor]
    if case == 'cross-paragraph':
        runs.append(d.add_paragraph('OWNED END').runs[0])
    elif case == 'table':
        runs = [d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run('OWNED TABLE')]
    d.add_comment(runs, text='OWNED COMMENT α', author='Owned Author Ω', initials='OA')
    if case in {'point', 'unanchored'}:
        for name in ['commentRangeStart', 'commentRangeEnd'] + (['commentReference'] if case == 'unanchored' else []):
            for node in d.element.xpath('.//w:' + name):
                node.getparent().remove(node)
    output = BytesIO()
    d.save(output)
    with ZipFile(BytesIO(output.getvalue())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    comments = ET.fromstring(parts['word/comments.xml'])
    comments.find(W + 'comment').set(W + 'date', '2026-10-11T12:34:56Z')
    parts['word/comments.xml'] = ET.tostring(comments)
    return package(parts)


def package(parts):
    output = BytesIO()
    with ZipFile(output, 'w') as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return output.getvalue()


@pytest.mark.parametrize('case', ['point', 'range', 'cross-paragraph', 'table', 'unanchored'])
@pytest.mark.parametrize('flat_input', [False, True])
@pytest.mark.parametrize('via_json', [False, True])
@pytest.mark.parametrize('flat_output', [False, True])
def test_classic_comment_edit_json_and_consecutive_save(case, flat_input, via_json, flat_output):
    raw = owned_comment(case)
    with warnings.catch_warnings():
        warnings.simplefilter('error', dk.ContentLossWarning)
        loaded = dk.Document(BytesIO(encode(raw) if flat_input else raw))
    model = loaded.light_document_model
    if via_json:
        model = ldm.Document.model_validate_json(model.model_dump_json())
    story = next(s for s in model.source_stories if s.kind == 'comment')
    assert story.comment_metadata.author == 'Owned Author Ω'
    assert story.comment_metadata.initial == 'OA'
    assert story.comment_metadata.date == '2026-10-11T12:34:56Z'
    model.sections[0].body.paragraphs[0].runs[0].text = 'EDITED BODY '
    story.children[0].runs[0].text = 'EDITED COMMENT α'
    for _ in range(2):
        initial = model.model_dump_json()
        raw = LdmDocxWriter().write_to_bytes(model)
        assert model.model_dump_json() == initial
        if flat_output:
            raw = encode(raw)
        with ZipFile(BytesIO(decode(raw) if flat_output else raw)) as z:
            comments = ET.fromstring(z.read('word/comments.xml'))
            comment = comments.find(W + 'comment')
            assert comment.get(W + 'author') == 'Owned Author Ω'
            assert comment.get(W + 'initials') == 'OA'
            assert comment.get(W + 'date') == '2026-10-11T12:34:56Z'
            assert ''.join(comment.itertext()) == 'EDITED COMMENT α'
            assert len(comment.findall('.//' + W + 'annotationRef')) == 1
            body = ET.fromstring(z.read('word/document.xml'))
            assert 'EDITED BODY ' in ''.join(body.itertext())
            assert len(body.findall('.//' + W + 'commentReference')) == (0 if case == 'unanchored' else 1)
            assert len(body.findall('.//' + W + 'commentRangeStart')) == (0 if case in {'point', 'unanchored'} else 1)
            assert len(body.findall('.//' + W + 'commentRangeEnd')) == (0 if case in {'point', 'unanchored'} else 1)
        model = dk.Document(BytesIO(raw)).light_document_model
        if via_json:
            model = ldm.Document.model_validate_json(model.model_dump_json())


@pytest.mark.parametrize('failure', ['missing-story', 'duplicate-reference', 'missing-end', 'reversed', 'duplicate-story', 'invalid-metadata'])
def test_invalid_comment_model_rejects_save_atomically(tmp_path, failure):
    model = dk.Document(BytesIO(owned_comment())).light_document_model
    p = model.sections[0].body.paragraphs[0]
    markers = [c for c in p._children if isinstance(c, ldm.CommentMarker)]
    story = model.source_stories[0]
    if failure == 'missing-story':
        model.source_stories.clear()
    elif failure == 'duplicate-reference':
        p._children.append(markers[-1].model_copy(deep=True))
    elif failure == 'missing-end':
        p._children.remove(next(m for m in markers if m.kind == 'end'))
    elif failure == 'reversed':
        start, end = next(m for m in markers if m.kind == 'start'), next(m for m in markers if m.kind == 'end')
        a, b = p._children.index(start), p._children.index(end)
        p._children[a], p._children[b] = p._children[b], p._children[a]
    elif failure == 'duplicate-story':
        model.source_stories.append(story.model_copy(deep=True))
    else:
        story.comment_metadata.author = 'PRIVATE INVALID\x00'
    initial = model.model_dump_json()
    path = tmp_path / 'output.docx'
    path.write_bytes(b'ORIGINAL')
    with pytest.raises(ValueError) as error:
        LdmDocxWriter().write(model, path)
    assert 'PRIVATE' not in str(error.value)
    assert path.read_bytes() == b'ORIGINAL'
    assert model.model_dump_json() == initial


@pytest.mark.parametrize('failure', ['missing-part', 'wrong-root', 'wrong-content-type', 'duplicate-entry', 'missing-story', 'dtd'])
@pytest.mark.parametrize('flat', [False, True])
def test_invalid_comment_source_graph_is_rejected(failure, flat):
    with ZipFile(BytesIO(owned_comment())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    if failure == 'missing-part':
        parts.pop('word/comments.xml')
    elif failure == 'wrong-root':
        parts['word/comments.xml'] = b'<wrong/>'
    elif failure == 'dtd':
        parts['word/comments.xml'] = b'<!DOCTYPE comments><w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
    elif failure == 'wrong-content-type':
        parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace(
            b'application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml', b'application/xml')
    elif failure == 'duplicate-entry':
        root = ET.fromstring(parts['word/comments.xml'])
        root.append(ET.fromstring(ET.tostring(root.find(W + 'comment'))))
        parts['word/comments.xml'] = ET.tostring(root)
    else:
        root = ET.fromstring(parts['word/comments.xml'])
        root.find(W + 'comment').set(W + 'id', '9')
        parts['word/comments.xml'] = ET.tostring(root)
    raw = package(parts)
    with pytest.raises((ValueError, DefusedXmlException)) as error:
        dk.Document(BytesIO(encode(raw) if flat else raw))
    assert 'OWNED COMMENT' not in str(error.value)
    assert 'Owned Author' not in str(error.value)


def test_other_outputs_still_report_comment_loss_without_body_metadata():
    loaded = dk.Document(BytesIO(owned_comment()))
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', dk.ContentLossWarning)
        loaded.to_bytes(dk.SaveFormat.PDF)
    observed = [d for d in loaded.diagnostics if d.code == 'pdf.comments_omitted']
    assert observed
    assert all(s not in repr(observed) for s in ('OWNED COMMENT', 'Owned Author', 'OWNED ANCHOR'))


@pytest.mark.parametrize('flat', [False, True])
def test_comment_anchors_inside_hyperlink_are_not_dropped(flat):
    with ZipFile(BytesIO(owned_comment())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    paragraph = root.find('.//' + W + 'p')
    anchor = next(r for r in paragraph.findall(W + 'r') if 'OWNED ANCHOR' in ''.join(r.itertext()))
    nodes = [paragraph.find(W + 'commentRangeStart'), anchor, paragraph.find(W + 'commentRangeEnd')]
    link = ET.Element(W + 'hyperlink', {'{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id': 'rIdOwnedLink'})
    position = list(paragraph).index(nodes[0])
    for node in nodes:
        paragraph.remove(node)
        link.append(node)
    paragraph.insert(position, link)
    parts['word/document.xml'] = ET.tostring(root)
    relationships = ET.fromstring(parts['word/_rels/document.xml.rels'])
    relationships.append(ET.Element('{http://schemas.openxmlformats.org/package/2006/relationships}Relationship', {
        'Id': 'rIdOwnedLink', 'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        'Target': 'https://owned.test/path', 'TargetMode': 'External'}))
    parts['word/_rels/document.xml.rels'] = ET.tostring(relationships)
    raw = package(parts)
    loaded = dk.Document(BytesIO(encode(raw) if flat else raw))
    output = loaded.to_bytes(dk.SaveFormat.DOCX)
    cold = dk.Document(BytesIO(output))
    markers = [node for node in cold.light_document_model.sections[0].body.paragraphs[0]._children
               if isinstance(node, ldm.CommentMarker)]
    assert [n.kind for n in markers] == ['start', 'end', 'reference']
    assert any(run.is_hyperlink and 'OWNED ANCHOR' in run.text for run in cold.sections[0].body.paragraphs[0].runs)


def test_unsupported_anchor_container_fails_without_losing_content():
    with ZipFile(BytesIO(owned_comment())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    paragraph = root.find('.//' + W + 'p')
    insertion = ET.Element(W + 'ins', {W + 'id': '1', W + 'author': 'PRIVATE AUTHOR'})
    for node in list(paragraph):
        if node.tag != W + 'pPr':
            paragraph.remove(node)
            insertion.append(node)
    paragraph.append(insertion)
    parts['word/document.xml'] = ET.tostring(root)
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', dk.ContentLossWarning)
        with pytest.raises(NotImplementedError, match='Unsupported comment anchor location') as error:
            dk.Document(BytesIO(package(parts)))
    assert 'PRIVATE AUTHOR' not in str(error.value)


def test_malformed_source_marker_identifier_does_not_disclose_its_value():
    with ZipFile(BytesIO(owned_comment())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    root.find('.//' + W + 'commentRangeStart').set(W + 'id', 'PRIVATE INVALID ID')
    parts['word/document.xml'] = ET.tostring(root)
    with pytest.raises(ValueError, match='Invalid annotation identifier') as error:
        dk.Document(BytesIO(package(parts)))
    assert 'PRIVATE INVALID ID' not in str(error.value)


def test_comment_marker_character_style_survives_json_and_save():
    loaded = dk.Document(BytesIO(owned_comment())).light_document_model
    loaded = ldm.Document.model_validate_json(loaded.model_dump_json())
    raw = LdmDocxWriter().write_to_bytes(loaded)
    with ZipFile(BytesIO(raw)) as z:
        for part, tag in [('word/document.xml', 'commentReference'), ('word/comments.xml', 'annotationRef')]:
            root = ET.fromstring(z.read(part))
            run = next(run for run in root.iter(W + 'r') if run.find(W + tag) is not None)
            assert run.find(W + 'rPr/' + W + 'rStyle').get(W + 'val') == 'CommentReference'


@pytest.mark.parametrize('location', ['root', 'story'])
def test_unknown_comment_container_is_rejected_without_disclosing_content(location):
    with ZipFile(BytesIO(owned_comment())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    root = ET.fromstring(parts['word/comments.xml'])
    parent = root if location == 'root' else root.find(W + 'comment')
    child = ET.SubElement(parent, W + 'customXml')
    child.text = 'PRIVATE CONTENT'
    parts['word/comments.xml'] = ET.tostring(root)
    with pytest.raises((ValueError, NotImplementedError)) as error:
        dk.Document(BytesIO(package(parts)))
    assert 'PRIVATE CONTENT' not in str(error.value)


@pytest.mark.parametrize('flat_input', [False, True])
@pytest.mark.parametrize('flat_output', [False, True])
def test_comment_images_and_links_have_their_own_saved_relationships(flat_input, flat_output):
    from PIL import Image
    d = NativeDocument(BytesIO(owned_comment()))
    comment = next(iter(d.comments))
    picture = BytesIO()
    Image.new('RGB', (2, 2), (70, 110, 140)).save(picture, format='PNG')
    comment.paragraphs[0].add_run().add_picture(BytesIO(picture.getvalue()))
    source = BytesIO()
    d.save(source)
    with ZipFile(BytesIO(source.getvalue())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    root = ET.fromstring(parts['word/comments.xml'])
    rel_ns = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    office_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    link = ET.SubElement(root.find(W + 'comment/' + W + 'p'), W + 'hyperlink', {office_ns + 'id': 'rIdOwnedCommentLink'})
    ET.SubElement(ET.SubElement(link, W + 'r'), W + 't').text = 'OWNED LINK'
    parts['word/comments.xml'] = ET.tostring(root)
    relationships = ET.fromstring(parts['word/_rels/comments.xml.rels'])
    relationships.append(ET.Element(rel_ns + 'Relationship', {'Id': 'rIdOwnedCommentLink',
        'Type': office_ns[1:-1] + '/hyperlink', 'Target': 'https://owned.test/comment', 'TargetMode': 'External'}))
    parts['word/_rels/comments.xml.rels'] = ET.tostring(relationships)
    source = package(parts)
    model = dk.Document(BytesIO(encode(source) if flat_input else source)).light_document_model
    model = ldm.Document.model_validate_json(model.model_dump_json())
    for _ in range(2):
        raw = LdmDocxWriter().write_to_bytes(model)
        if flat_output:
            raw = encode(raw)
        with ZipFile(BytesIO(decode(raw) if flat_output else raw)) as z:
            relationships = ET.fromstring(z.read('word/_rels/comments.xml.rels'))
            image = next(r for r in relationships if r.get('Type').endswith('/image'))
            assert z.read('word/' + image.get('Target')) == picture.getvalue()
            assert any(r.get('Target') == 'https://owned.test/comment' for r in relationships)
        model = dk.Document(BytesIO(raw)).light_document_model
        children = model.source_stories[0].children[0]._children
        assert sum(isinstance(node, ldm.Shape) for node in children) == 1
        assert any(isinstance(node, ldm.Run) and node.is_hyperlink and 'OWNED LINK' in node.text for node in children)


def comment_settings_position_package():
    with ZipFile(BytesIO(owned_comment('table'))) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    settings = ET.fromstring(parts['word/settings.xml'])
    compat = settings.find(W + 'compat')
    ET.SubElement(compat, W + 'compatSetting', {W + 'name': 'ownedCommentSetting', W + 'uri': 'owned', W + 'val': 'keep'})
    parts['word/settings.xml'] = ET.tostring(settings)
    body = ET.fromstring(parts['word/document.xml'])
    ET.SubElement(body.find('.//' + W + 'tblPr'), W + 'tblpPr', {W + 'tblpX': '0', W + 'horzAnchor': 'page'})
    parts['word/document.xml'] = ET.tostring(body)
    return package(parts)


@pytest.mark.parametrize('flat_input', [False, True])
@pytest.mark.parametrize('flat_output', [False, True])
def test_comment_conversion_preserves_source_settings_and_table_position(flat_input, flat_output):
    source = comment_settings_position_package()
    model = dk.Document(BytesIO(encode(source) if flat_input else source)).light_document_model
    model = ldm.Document.model_validate_json(model.model_dump_json())
    for _ in range(2):
        initial = model.model_dump_json()
        raw = LdmDocxWriter().write_to_bytes(model)
        assert model.model_dump_json() == initial
        if flat_output:
            raw = encode(raw)
        with ZipFile(BytesIO(decode(raw) if flat_output else raw)) as z:
            settings = ET.fromstring(z.read('word/settings.xml'))
            assert any(n.get(W + 'name') == 'ownedCommentSetting' and n.get(W + 'val') == 'keep'
                       for n in settings.iter(W + 'compatSetting'))
            assert ET.fromstring(z.read('word/document.xml')).find('.//' + W + 'tblpPr').attrib == {
                W + 'tblpX': '0', W + 'horzAnchor': 'page'}
            assert 'OWNED COMMENT α' in ''.join(ET.fromstring(z.read('word/comments.xml')).itertext())
        model = dk.Document(BytesIO(raw)).light_document_model
        table = model.sections[0].body.tables[0]
        assert dict(table.source_position.attributes) == {'tblpX': '0', 'horzAnchor': 'page'}
        assert model.source_stories[0].comment_metadata.author == 'Owned Author Ω'


@pytest.mark.parametrize('cycle', ['table', 'textbox'])
def test_cyclic_comment_story_is_rejected_before_output_and_source_changes(tmp_path, cycle):
    model = dk.Document(BytesIO(owned_comment())).light_document_model
    story = model.source_stories[0]
    if cycle == 'table':
        node = ldm.Table(rows=[ldm.Row(cells=[ldm.Cell()])])
        children = node.rows[0].cells[0].tables
        children.append(node)
    else:
        node = ldm.Paragraph()
        shape = ldm.Shape(text_box={'paragraphs': []})
        node._children.append(shape)
        children = shape.text_box['paragraphs']
        children.append(node)
    story.children = [node]
    destination = tmp_path / 'owned.docx'
    destination.write_bytes(b'ORIGINAL')
    with pytest.raises(ValueError, match='Cyclic comment story') as error:
        LdmDocxWriter().write(model, destination)
    assert destination.read_bytes() == b'ORIGINAL'
    assert story.children[0] is node and children[0] is node
    assert story.comment_metadata.author == 'Owned Author Ω'
    assert 'Owned Author' not in str(error.value)


def test_shared_comment_paragraph_without_cycle_is_not_rejected():
    model = dk.Document(BytesIO(owned_comment())).light_document_model
    paragraph = ldm.Paragraph()
    paragraph._children.append(ldm.Run(text='OWNED ALIAS'))
    model.source_stories[0].children = [paragraph, paragraph]
    with ZipFile(BytesIO(LdmDocxWriter().write_to_bytes(model))) as z:
        comment = ET.fromstring(z.read('word/comments.xml')).find(W + 'comment')
        assert len(comment.findall(W + 'p')) == 2
        assert ''.join(comment.itertext()) == 'OWNED ALIASOWNED ALIAS'


@pytest.mark.parametrize('flat_input', [False, True])
@pytest.mark.parametrize('via_json', [False, True])
@pytest.mark.parametrize('fmt', [dk.SaveFormat.DOCX, dk.SaveFormat.FLAT_OPC])
def test_comment_document_stream_save_retains_edits_and_output_metadata(flat_input, via_json, fmt):
    source = owned_comment()
    doc = dk.Document(BytesIO(encode(source) if flat_input else source))
    if via_json:
        doc._document = ldm.Document.model_validate_json(doc.light_document_model.model_dump_json())
    doc.light_document_model.source_stories[0].children[0].runs[0].text = 'EDITED STREAM COMMENT'
    for _ in range(2):
        stream = BytesIO()
        result = doc.save(stream, fmt)
        assert result.content_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        assert not stream.closed
        doc = dk.Document(BytesIO(stream.getvalue()))
        story = doc.light_document_model.source_stories[0]
        assert story.comment_metadata.author == 'Owned Author Ω'
        assert story.children[0].runs[0].text == 'EDITED STREAM COMMENT'


@pytest.mark.parametrize('fmt', [dk.SaveFormat.DOCX, dk.SaveFormat.FLAT_OPC])
def test_invalid_comment_stream_save_leaves_stream_unchanged(fmt):
    doc = dk.Document(BytesIO(owned_comment()))
    doc.light_document_model.source_stories.clear()
    stream = BytesIO(b'ORIGINAL');stream.seek(3)
    with pytest.raises(ValueError, match='Comment anchor has no matching story'):
        doc.save(stream, fmt)
    assert stream.getvalue() == b'ORIGINAL' and stream.tell() == 3 and not stream.closed


def owned_comment_hyperlink():
    with ZipFile(BytesIO(owned_comment())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    paragraph = root.find('.//' + W + 'p')
    anchor = next(r for r in paragraph.findall(W + 'r') if 'OWNED ANCHOR' in ''.join(r.itertext()))
    nodes = [paragraph.find(W + 'commentRangeStart'), anchor, paragraph.find(W + 'commentRangeEnd')]
    link = ET.Element(W + 'hyperlink', {'{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id': 'rIdOwnedLink'})
    position = list(paragraph).index(nodes[0])
    for node in nodes:
        paragraph.remove(node)
        link.append(node)
    paragraph.insert(position, link)
    parts['word/document.xml'] = ET.tostring(root)
    relationships = ET.fromstring(parts['word/_rels/document.xml.rels'])
    relationships.append(ET.Element('{http://schemas.openxmlformats.org/package/2006/relationships}Relationship', {
        'Id': 'rIdOwnedLink', 'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        'Target': 'https://owned.test/path', 'TargetMode': 'External'}))
    parts['word/_rels/document.xml.rels'] = ET.tostring(relationships)
    return package(parts)


@pytest.mark.parametrize('history', [None, '0', '1', 'false', 'true', 'off', 'on'])
@pytest.mark.parametrize('flat_input', [False, True])
@pytest.mark.parametrize('via_json', [False, True])
@pytest.mark.parametrize('flat_output', [False, True])
def test_comment_hyperlink_history_survives_json_and_save_cycles(history, flat_input, via_json, flat_output):
    with ZipFile(BytesIO(owned_comment_hyperlink())) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    if history is not None:
        root.find('.//' + W + 'hyperlink').set(W + 'history', history)
    parts['word/document.xml'] = ET.tostring(root)
    source = package(parts)
    model = dk.Document(BytesIO(encode(source) if flat_input else source)).light_document_model
    if via_json:
        model = ldm.Document.model_validate_json(model.model_dump_json())
    for _ in range(2):
        run = next(r for r in model.sections[0].body.paragraphs[0].runs if r.is_hyperlink)
        assert run.source_hyperlink_history == history
        raw = LdmDocxWriter().write_to_bytes(model)
        with ZipFile(BytesIO(raw)) as z:
            assert ET.fromstring(z.read('word/document.xml')).find('.//' + W + 'hyperlink').get(W + 'history') == history
        model = dk.Document(BytesIO(encode(raw) if flat_output else raw)).light_document_model


@pytest.mark.parametrize('history', ['', '2', 'invalid'])
@pytest.mark.parametrize('flat', [False, True])
def test_comment_hyperlink_invalid_history_is_validated_before_early_branch(history, flat):
    with ZipFile(BytesIO(owned_comment_hyperlink())) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    root = ET.fromstring(parts['word/document.xml'])
    root.find('.//' + W + 'hyperlink').set(W + 'history', history)
    parts['word/document.xml'] = ET.tostring(root)
    raw = package(parts)
    with pytest.raises(RuntimeError, match='Invalid hyperlink history'):
        dk.Document(BytesIO(encode(raw) if flat else raw))
