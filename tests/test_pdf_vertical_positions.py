"""Source superscript/subscript survive PDF measurement and rendering."""

import pymupdf
import pytest
from fpdf import FPDF
from fpdf.enums import CharVPos

from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.pdf_writer.font import apply_run_font, register_fonts
from aspose.words_foss.pdf_writer.run_renderer import RunRenderer
from aspose.words_foss.saving import PdfSaveOptions


def script_model(location="body", *, widow=True, size=20):
    runs = [ldm.Run(text="Q", font=ldm.Font(size=size)),
            ldm.Run(text="W", font=ldm.Font(size=size, superscript=True)),
            ldm.Run(text="E", font=ldm.Font(size=size, subscript=True)),
            ldm.Run(text="R", font=ldm.Font(size=size))]
    paragraph = ldm.Paragraph(children=runs, paragraph_format=ldm.ParagraphFormat(widow_control=widow))
    normal = ldm.Paragraph(children=[ldm.Run(text="T", font=ldm.Font(size=size))])
    section = ldm.Section(page_setup=ldm.PageSetup(page_width=320, page_height=220,
        left_margin=20, right_margin=20, top_margin=30, bottom_margin=30),
        body=ldm.Body(children=[paragraph, normal]))
    if location == "table":
        section.body.children = [ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(children=[paragraph])])]), normal]
    elif location in ("header", "footer"):
        section.body.children = [normal]
        section.headers_footers = [ldm.HeaderFooter(header_footer_type=0 if location == "header" else 1,
                                                   children=[paragraph])]
    elif location in ("heading", "quote", "code"):
        paragraph.paragraph_format.style_name = "Heading 1" if location == "heading" else "Code" if location == "code" else "Quote"
        if location == "heading":
            paragraph.paragraph_format.outline_level = 0
    elif location == "list":
        paragraph.list_format = ldm.ListFormat(is_list_item=True, list_id=1)
        paragraph.list_label = ldm.ListLabel(label_string="1.")
    return ldm.Document(sections=[section]), paragraph


def characters(page):
    return {c["c"]: (c, span["size"]) for block in page.get_text("rawdict")["blocks"]
            for line in block.get("lines", []) for span in line["spans"] for c in span["chars"]}


@pytest.mark.parametrize("location", ["body", "table", "header", "footer", "list", "heading", "quote", "code"])
@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("widow", [False, True])
def test_source_script_sizes_rises_and_following_normal_text(location, shaping, widow):
    model, _ = script_model(location, widow=widow)
    snapshot = model.model_dump()
    options = PdfSaveOptions()
    options.text_shaping = shaping
    with pymupdf.open(stream=LdmPdfWriter(options).write_to_bytes(model), filetype="pdf") as pdf:
        chars = characters(pdf[0])
        if location in ("body", "table", "list", "heading", "quote", "code"):
            assert chars["W"][0]["bbox"][1] >= 30 - 0.05
        assert set(chars) >= set("QWERT")
        assert [chars[c][1] for c in "QWERT"] == pytest.approx([20, 14, 14, 20, 20], abs=0.05)
        baseline = chars["Q"][0]["origin"][1]
        assert chars["W"][0]["origin"][1] == pytest.approx(baseline - 8, abs=0.05)
        assert chars["E"][0]["origin"][1] == pytest.approx(baseline + 3, abs=0.05)
        assert chars["R"][0]["origin"][1] == pytest.approx(baseline, abs=0.05)
    assert model.model_dump() == snapshot


@pytest.mark.parametrize("shaping", [False, True])
@pytest.mark.parametrize("location", ["body", "code"])
def test_script_links_and_decorations_match_reduced_raised_glyphs(shaping, location):
    model, paragraph = script_model(location)
    for run in paragraph.runs:
        run.text = f"[{run.text}](https://example.test/{run.text})"
        run.font.strike_through = True
        run.font.underline = True
        run.font.highlight_color = "FFFF00"
    options = PdfSaveOptions()
    options.text_shaping = shaping
    with pymupdf.open(stream=LdmPdfWriter(options).write_to_bytes(model), filetype="pdf") as pdf:
        page = pdf[0]
        chars = characters(page)
        links = page.get_links()
        assert len(links) == 4
        spans = [span for b in page.get_text("dict")["blocks"]
                 for line in b.get("lines", []) for span in line["spans"]]
        assert all(span["color"] == 238 for span in spans if any(c in span["text"] for c in "QWER"))
        for char in "QWER":
            c, size = chars[char]
            link = next(link for link in links if link["uri"].endswith(char))
            assert link["from"].height == pytest.approx(size, abs=0.05)
            assert link["from"].y0 == pytest.approx(c["origin"][1] - 0.8 * size, abs=0.05)
            strikes = [d["rect"] for d in page.get_drawings() if d["rect"].height < 2
                       and abs(d["rect"].x0 - c["origin"][0]) < 0.05]
            assert any(c["origin"][1] - size * 0.5 < rect.y0 < c["origin"][1] - size * 0.1
                       for rect in strikes)


@pytest.mark.parametrize("shaping", [False, True])
def test_width_measurement_restores_vpos_and_uses_scaled_width(shaping):
    pdf = FPDF()
    if shaping:
        pdf.set_text_shaping(True)
    register_fonts(pdf)
    pdf.add_page()
    font = ldm.Font(size=20, superscript=True)
    run = ldm.Run(text="WWWW", font=font)
    apply_run_font(pdf, font)
    width = pdf.get_string_width(run.text)
    pdf.char_vpos = CharVPos.LINE
    normal_width = pdf.get_string_width(run.text)
    assert width == pytest.approx(normal_width * 0.7)
    pdf.set_text_color(0, 0, 238)
    saved_color = pdf.text_color
    rows = RunRenderer.wrap_segments(pdf, [(run, run.text, width, 20, None)], width + 0.01)
    assert len(rows) == 1
    assert rows[0][0][2] == pytest.approx(width)
    assert pdf.char_vpos == CharVPos.LINE
    assert pdf.text_color == saved_color


def test_superscript_precedence_matches_docx_writer():
    pdf = FPDF()
    register_fonts(pdf)
    apply_run_font(pdf, ldm.Font(superscript=True, subscript=True))
    assert pdf.char_vpos == CharVPos.SUP


@pytest.mark.parametrize("location", ["body", "table", "header", "footer"])
@pytest.mark.parametrize("shaping", [False, True])
def test_public_docx_script_flags_reach_pdf(tmp_path, location, shaping):
    from docx import Document as Docx

    from aspose.words_foss import Document
    from aspose.words_foss.docx_writer import LdmDocxWriter

    model, _ = script_model(location)
    path = tmp_path / "scripts.docx"
    LdmDocxWriter().write(model, path)
    source = Docx(path)
    paragraph = (source.tables[0].cell(0, 0).paragraphs[0] if location == "table" else
                 source.sections[0].header.paragraphs[0] if location == "header" else
                 source.sections[0].footer.paragraphs[0] if location == "footer" else source.paragraphs[0])
    assert paragraph.runs[1].font.superscript
    assert paragraph.runs[2].font.subscript
    options = PdfSaveOptions()
    options.text_shaping = shaping
    with pymupdf.open(stream=Document(path).to_bytes(options), filetype="pdf") as pdf:
        chars = characters(pdf[0])
        assert [chars[c][1] for c in "QWERT"] == pytest.approx([20, 14, 14, 20, 20], abs=0.05)
        assert chars["W"][0]["origin"][1] < chars["Q"][0]["origin"][1] < chars["E"][0]["origin"][1]


@pytest.mark.parametrize("location", ["body", "table"])
@pytest.mark.parametrize("columns", [1, 2])
@pytest.mark.parametrize("shaping", [False, True])
def test_wrapped_scripts_keep_all_glyphs_inside_page_regions(location, columns, shaping):
    model, paragraph = script_model(location)
    section = model.sections[0]
    section.page_setup.page_width = 180
    section.page_setup.page_height = 180
    section.page_setup.text_columns = ldm.TextColumns(count=columns, spacing=10)
    paragraph.runs[1].text = "2 " * 150
    options = PdfSaveOptions()
    options.text_shaping = shaping
    with pymupdf.open(stream=LdmPdfWriter(options).write_to_bytes(model), filetype="pdf") as pdf:
        assert len(pdf) > 1
        assert sum(page.get_text().count("2") for page in pdf) == 150
        for page in pdf:
            digits = [(c, span["size"]) for b in page.get_text("rawdict")["blocks"]
                      for line in b.get("lines", []) for span in line["spans"] for c in span["chars"]
                      if c["c"] == "2"]
            for char, size in digits:
                assert size == pytest.approx(14, abs=0.05)
                assert char["bbox"][1] >= 30 - 0.05
                assert char["bbox"][3] <= page.rect.height - 30 + 0.05


def test_unimplemented_rotated_script_layout_is_reported():
    from aspose.words_foss.pdf_writer.diagnostics import PdfContentLossWarning

    model, _ = script_model("table")
    model.sections[0].body.children[0].rows[0].cells[0].cell_format.orientation = 1
    with pytest.warns(PdfContentLossWarning, match="superscript/subscript"):
        assert LdmPdfWriter().write_to_bytes(model).startswith(b"%PDF")
