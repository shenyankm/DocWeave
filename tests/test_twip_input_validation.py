"""Malformed dimensions fail before conversion and do not disclose source data."""

import os
import subprocess
import sys
import traceback
from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document as IndependentDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from aspose.words_foss import Document, DocxDocument

INVALID = ("PRIVATE POINT VALUE", "1.13", "1_13", "١١٣", "1e2", "", "9" * 5000, "9" * 400)
DIMENSIONS = {"left": "left_indent", "start": "left_indent", "right": "right_indent", "end": "right_indent",
              "firstLine": "first_line_indent", "hanging": "first_line_indent",
              "before": "space_before", "after": "space_after"}


def source(attribute, kind, value):
    document = IndependentDocument()
    document.add_paragraph("PRIVATE BODY")
    if kind == "body":
        properties = document.paragraphs[0]._p.get_or_add_pPr()
    elif kind == "table":
        properties = document.add_table(1, 1).cell(0, 0).paragraphs[0]._p.get_or_add_pPr()
    elif kind in {"header", "footer"}:
        properties = getattr(document.sections[0], kind).paragraphs[0]._p.get_or_add_pPr()
    elif kind == "style":
        properties = document.styles["Normal"].element.get_or_add_pPr()
    elif kind == "default":
        properties = document.styles.element.find(qn("w:docDefaults") + "/" + qn("w:pPrDefault") + "/" + qn("w:pPr"))
    elif kind in {"numbering", "numbering_override"}:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), "999")
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for name, item in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1.")):
            node = OxmlElement("w:" + name)
            node.set(qn("w:val"), item)
            level.append(node)
        properties = OxmlElement("w:pPr")
        level.append(properties)
        abstract.append(level)
        document.part.numbering_part.element.append(abstract)
        number = OxmlElement("w:num")
        number.set(qn("w:numId"), "999")
        reference = OxmlElement("w:abstractNumId")
        reference.set(qn("w:val"), "999")
        number.append(reference)
        if kind == "numbering_override":
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), "0")
            inline_level = deepcopy(level)
            properties = inline_level.find(qn("w:pPr"))
            override.append(inline_level)
            number.append(override)
        document.part.numbering_part.element.append(number)
    else:
        if kind == "table_style_margin":
            parent = OxmlElement("w:tblPr")
            document.styles.add_style("Margin", WD_STYLE_TYPE.TABLE).element.append(parent)
        else:
            table = document.add_table(1, 1)
            parent = table._tbl.tblPr if kind == "table_margin" else table.cell(0, 0)._tc.get_or_add_tcPr()
        properties = OxmlElement("w:tcMar" if kind == "cell_margin" else "w:tblCellMar")
        parent.append(properties)
        element = OxmlElement("w:left")
        element.set(qn("w:type"), "dxa")
        element.set(qn("w:w"), value)
        properties.append(element)
    if not kind.endswith("margin"):
        tag = "spacing" if attribute in {"before", "after", "line"} else "ind"
        element = properties.find(qn("w:" + tag))
        if element is None:
            element = OxmlElement("w:" + tag)
            properties.append(element)
        element.set(qn("w:" + attribute), value)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def assert_private_error(error, raw):
    text = "".join(traceback.format_exception(error))
    assert "PRIVATE BODY" not in text
    if raw:
        assert raw not in text
    assert error.__suppress_context__


@pytest.mark.parametrize("attribute", DIMENSIONS)
@pytest.mark.parametrize("kind", ["body", "table", "header", "footer", "style", "default"])
@pytest.mark.parametrize("raw", INVALID)
def test_invalid_paragraph_dimensions_fail_without_loss_or_disclosure(attribute, kind, raw):
    document = DocxDocument(BytesIO(source(attribute, kind, raw)))
    before = document.to_bytes()
    with pytest.raises(ValueError, match="integer twips") as caught:
        document.to_light_document()
    assert_private_error(caught.value, raw)
    assert document.to_bytes() == before


@pytest.mark.parametrize("kind", ["body", "table", "header", "footer", "style", "default"])
@pytest.mark.parametrize("raw,expected,rule", [
    (INVALID[0], 0, 2), (INVALID[1], .05, 2), (INVALID[2], 0, 2),
    (INVALID[3], 0, 2), (INVALID[4], 5, 2), (INVALID[5], 0, 2),
    (INVALID[6], .05, 1), (INVALID[7], .05, 1),
])
def test_line_measure_uses_native_recovery_without_mutating_source(kind, raw, expected, rule):
    document = DocxDocument(BytesIO(source("line", kind, raw)))
    before = document.to_bytes()
    section = document.to_light_document().sections[0]
    if kind == "table":
        paragraph = section.body.tables[0].rows[0].cells[0].paragraphs[0]
    elif kind in {"header", "footer"}:
        paragraph = section.headers_footers[0].paragraphs[0]
    else:
        paragraph = section.body.paragraphs[0]
    pf = paragraph.paragraph_format
    assert (pf.line_spacing, pf.line_spacing_rule) == (expected, rule)
    assert "line_spacing" in pf.model_fields_set
    assert document.to_bytes() == before


@pytest.mark.parametrize("attribute", DIMENSIONS)
@pytest.mark.parametrize("kind", ["body", "style", "default"])
@pytest.mark.parametrize("raw", INVALID)
def test_direct_and_effective_getters_fail_privately(attribute, kind, raw):
    document = DocxDocument(BytesIO(source(attribute, kind, raw)))
    before = document.to_bytes()
    prop = DIMENSIONS[attribute]
    fmt = (document.body.paragraphs[0].paragraph_format if kind == "body"
           else document.styles.get_by_name("Normal").paragraph_format)
    with pytest.raises(ValueError, match="integer twips") as caught:
        getattr(fmt, prop)
    assert_private_error(caught.value, raw)
    with pytest.raises(ValueError, match="integer twips") as caught:
        _ = document.body.paragraphs[0].effective_paragraph_format
    assert_private_error(caught.value, raw)
    assert document.to_bytes() == before


@pytest.mark.parametrize("attribute", ["left", "hanging", "firstLine"])
@pytest.mark.parametrize("kind", ["numbering", "numbering_override"])
@pytest.mark.parametrize("raw", INVALID)
def test_invalid_numbering_dimensions_fail_privately(attribute, kind, raw):
    document = DocxDocument(BytesIO(source(attribute, kind, raw)))
    before = document.to_bytes()
    with pytest.raises(ValueError, match="integer twips") as caught:
        document.to_light_document()
    assert_private_error(caught.value, raw)
    assert document.to_bytes() == before


@pytest.mark.parametrize("kind", ["table_margin", "cell_margin", "table_style_margin"])
@pytest.mark.parametrize("raw", INVALID)
def test_invalid_margin_dimensions_use_the_same_private_error(kind, raw):
    document = DocxDocument(BytesIO(source("left", kind, raw)))
    before = document.to_bytes()
    with pytest.raises(ValueError, match="integer twips") as caught:
        document.to_light_document()
    assert_private_error(caught.value, raw)
    assert document.to_bytes() == before


@pytest.mark.parametrize("raw,expected", [("0", 0), ("113", 5.65), ("-113", -5.65), (" +00113 ", 5.65)])
def test_valid_twip_lexical_forms_and_explicit_zero_keep_their_values(raw, expected):
    document = DocxDocument(BytesIO(source("left", "body", raw)))
    assert document.body.paragraphs[0].paragraph_format.left_indent == expected
    assert document.to_light_document().sections[0].body.paragraphs[0].paragraph_format.left_indent == expected


@pytest.mark.parametrize("kind", ["body", "style", "default", "numbering", "numbering_override", "table_margin"])
@pytest.mark.parametrize("format", ["docx", "flat_opc"])
def test_public_load_rejects_invalid_dimensions_and_preserves_source(kind, format):
    raw = source("left", kind, "PRIVATE POINT VALUE")
    document = DocxDocument(BytesIO(raw))
    data = document.to_flat_opc() if format == "flat_opc" else raw
    with pytest.raises(ValueError, match="integer twips") as caught:
        Document(BytesIO(data))
    assert_private_error(caught.value, "PRIVATE POINT VALUE")
    assert document.to_bytes() == raw


def test_cli_retains_existing_output_on_malformed_dimensions(tmp_path):
    path = tmp_path / "source.docx"
    path.write_bytes(source("left", "body", "PRIVATE POINT VALUE"))
    output = tmp_path / "output.pdf"
    output.write_bytes(b"existing output")
    result = subprocess.run([sys.executable, "-m", "aspose.words_foss.convert", str(path), str(output)],
                            capture_output=True, text=True, timeout=20, check=False)
    error = "requires POSIX process groups" if os.name == "nt" else "integer twips"
    assert result.returncode != 0 and error in result.stderr
    assert "PRIVATE POINT VALUE" not in result.stderr and "PRIVATE BODY" not in result.stderr
    assert output.read_bytes() == b"existing output"
    assert not list(tmp_path.glob(".conversion-*"))


@pytest.mark.parametrize("attribute,alias", [("left", "start"), ("right", "end"), ("firstLine", "hanging")])
def test_later_valid_alias_does_not_hide_malformed_twips(attribute, alias):
    document = IndependentDocument(BytesIO(source(attribute, "body", "PRIVATE POINT VALUE")))
    document.paragraphs[0]._p.get_or_add_pPr().find(qn("w:ind")).set(qn("w:" + alias), "113")
    stream = BytesIO()
    document.save(stream)
    editable = DocxDocument(BytesIO(stream.getvalue()))
    before = editable.to_bytes()
    with pytest.raises(ValueError, match="integer twips") as caught:
        getattr(editable.body.paragraphs[0].paragraph_format, DIMENSIONS[attribute])
    assert_private_error(caught.value, "PRIVATE POINT VALUE")
    with pytest.raises(ValueError, match="integer twips"):
        editable.to_light_document()
    assert editable.to_bytes() == before


def test_numbering_subtracts_twips_before_scaling_without_rounding_drift():
    document = IndependentDocument(BytesIO(source("left", "numbering", "244")))
    abstract = next(node for node in document.part.numbering_part.element if node.get(qn("w:abstractNumId")) == "999")
    abstract.find(qn("w:lvl") + "/" + qn("w:pPr") + "/" + qn("w:ind")).set(qn("w:hanging"), "1")
    stream = BytesIO()
    document.save(stream)
    model = DocxDocument(BytesIO(stream.getvalue())).to_light_document()
    level = next(item for item in model.lists if item.list_id == 999).list_levels[0]
    assert level.number_position == 12.15 and level.text_position == 12.2


def test_overflowing_numbering_difference_fails_without_disclosure():
    raw = "1" + "0" * 308
    document = IndependentDocument(BytesIO(source("left", "numbering", raw)))
    abstract = next(node for node in document.part.numbering_part.element if node.get(qn("w:abstractNumId")) == "999")
    abstract.find(qn("w:lvl") + "/" + qn("w:pPr") + "/" + qn("w:ind")).set(qn("w:hanging"), "-" + raw)
    stream = BytesIO()
    document.save(stream)
    editable = DocxDocument(BytesIO(stream.getvalue()))
    before = editable.to_bytes()
    with pytest.raises(ValueError, match="numeric range") as caught:
        editable.to_light_document()
    assert_private_error(caught.value, raw)
    assert editable.to_bytes() == before
