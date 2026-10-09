"""External Markdown images must not overwrite unrelated or earlier assets."""

from io import BytesIO

import pytest
from docx import Document as NativeDocument
from PIL import Image

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.md_writer import LdmMarkdownWriter
from aspose.words_foss.models import ConversionOptions


def picture(color, name="logo.png"):
    stream = BytesIO()
    Image.new("RGB", (8, 8), color).save(stream, format="PNG")
    return ldm.Shape(has_image=True, image_data=ldm.ImageData(
        image_bytes=stream.getvalue(), source_full_name=name))


@pytest.mark.parametrize("html", [False, True])
@pytest.mark.parametrize("direct", [False, True])
def test_conflicting_images_and_existing_files_are_preserved(tmp_path, html, direct):
    images = tmp_path / "assets"
    images.mkdir()
    (images / "logo.png").write_bytes(b"EXISTING")
    red, blue = picture("red"), picture("blue")
    doc = aw.Document(BytesIO(b"BODY"), aw.MarkdownLoadOptions())
    para = ldm.Paragraph(children=[red, blue, red])
    doc.sections[0].body.children = [ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[para])])])] if html else [para]
    if direct:
        opts = ConversionOptions(images_folder=str(images))
        if html:
            opts.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
        result = LdmMarkdownWriter(opts).write(doc.light_document_model, tmp_path / "out.md")
    else:
        opts = aw.saving.MarkdownSaveOptions()
        opts.images_folder = str(images)
        if html:
            opts.export_as_html = aw.saving.MarkdownExportAsHtml.TABLES
        doc.save(tmp_path / "out.md", opts)
        result = (tmp_path / "out.md").read_text()
    assert (images / "logo.png").read_bytes() == b"EXISTING"
    assets = [p for p in images.iterdir() if p.name != "logo.png"]
    assert len(assets) == 2
    assert {p.read_bytes() for p in assets} == {red.image_data.image_bytes, blue.image_data.image_bytes}
    for asset in assets:
        assert "assets/" + asset.name in result
    assert result.count("assets/" + next(p.name for p in assets if p.read_bytes() == red.image_data.image_bytes)) == 2


def test_image_destination_symlink_is_not_followed(tmp_path):
    images = tmp_path / "assets"
    images.mkdir()
    target = tmp_path / "private"
    target.write_bytes(b"PRIVATE")
    try:
        (images / "logo.png").symlink_to(target)
    except OSError:
        pytest.skip("Host does not permit symlink creation")
    opts = ConversionOptions(images_folder=str(images))
    writer = LdmMarkdownWriter(opts)
    doc = aw.Document(BytesIO(b"BODY"), aw.MarkdownLoadOptions())
    shape = picture("blue")
    doc.sections[0].body.children = [ldm.Paragraph(children=[shape])]
    result = writer.write(doc.light_document_model, tmp_path / "out.md")
    assert target.read_bytes() == b"PRIVATE"
    assert (images / "logo.png").is_symlink()
    assert "assets/logo.png" not in result
    assert any(p.is_file() and not p.is_symlink() and p.read_bytes() == shape.image_data.image_bytes for p in images.iterdir())


def test_image_cannot_replace_the_markdown_destination(tmp_path):
    opts = aw.saving.MarkdownSaveOptions()
    opts.images_folder = str(tmp_path / "sub" / "..")
    (tmp_path / "sub").mkdir()
    doc = aw.Document(BytesIO(b"BODY"), aw.MarkdownLoadOptions())
    image = picture("blue", "../out.md")
    doc.sections[0].body.children = [ldm.Paragraph(children=[image])]
    doc.save(tmp_path / "out.md", opts)
    result = (tmp_path / "out.md").read_text()
    assert "out_2.md" in result
    assert (tmp_path / "out_2.md").read_bytes() == image.image_data.image_bytes


def test_images_from_two_real_docx_packages_keep_distinct_contents(tmp_path):
    documents = []
    for color in ("red", "blue"):
        native = NativeDocument()
        native.add_picture(BytesIO(picture(color).image_data.image_bytes))
        stream = BytesIO()
        native.save(stream)
        documents.append(aw.Document(BytesIO(stream.getvalue())))
    doc = documents[0]
    doc.sections[0].body.children.extend(documents[1].sections[0].body.children)
    opts = aw.saving.MarkdownSaveOptions()
    opts.images_folder = str(tmp_path / "assets")
    doc.save(tmp_path / "out.md", opts)
    colors = set()
    for asset in (tmp_path / "assets").iterdir():
        with Image.open(asset) as image:
            colors.add(image.convert("RGB").getpixel((0, 0)))
    assert colors == {(255, 0, 0), (0, 0, 255)}
