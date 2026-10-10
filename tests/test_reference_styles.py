"""Reference styles change generated DOCX without importing reference content."""

from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor
import pytest

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.docx_writer.runs import color_to_hex
from aspose.words_foss.docx_writer.styles_part import apply_reference_styles, render_styles_xml

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def reference(tmp_path):
    doc = Document()
    doc.styles['Heading 1'].font.size = Pt(30)
    doc.styles['Heading 1'].font.color.rgb = RGBColor.from_string('008800')
    doc.styles['Normal'].font.name = 'Courier New'
    custom = doc.styles.add_style('Brand Title', WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = doc.styles['Heading 1']
    custom.paragraph_format.space_after = Pt(19)
    doc.add_paragraph('REFERENCE BODY MUST NOT APPEAR')
    path = tmp_path / 'reference.docx'
    doc.save(path)
    return path


def test_reference_styles_affect_markdown_heading_and_keep_source_model(tmp_path):
    path = reference(tmp_path)
    before = path.read_bytes()
    doc = aw.Document(BytesIO(b'# Report\n\nBody\n'), aw.MarkdownLoadOptions())
    original = doc.light_document_model.model_dump()
    opts = aw.saving.OoxmlSaveOptions()
    opts.reference_docx = path
    raw = doc.to_bytes(opts)
    independent = Document(BytesIO(raw))
    assert independent.paragraphs[0].style.font.size.pt == 30
    assert str(independent.paragraphs[0].style.font.color.rgb) == '008800'
    assert independent.styles['Normal'].font.name == 'Courier New'
    assert independent.styles['Brand Title'].base_style.name == 'Heading 1'
    assert 'REFERENCE BODY' not in '\n'.join(p.text for p in independent.paragraphs)
    assert independent.paragraphs[0].style.name == 'Heading 1'
    # Verify effective formatting with another read, not just a style row in styles.xml.
    reread = aw.Document(BytesIO(raw))
    assert reread.light_document_model.sections[0].body.paragraphs[0].runs[0].font.size == 30
    assert doc.light_document_model.model_dump() == original
    assert path.read_bytes() == before


def test_direct_formatting_still_overrides_reference_and_links_survive(tmp_path):
    source = Document()
    para = source.add_heading('Direct', 1)
    para.runs[0].font.size = Pt(21)
    source.add_paragraph('Body')
    stream = BytesIO()
    source.save(stream)
    retained = aw.DocxDocument(BytesIO(stream.getvalue()))
    retained.body.paragraphs[1].add_hyperlink('Link', 'https://example.com')
    doc = aw.Document(BytesIO(retained.to_bytes()))
    opts = aw.saving.OoxmlSaveOptions()
    opts.reference_docx = reference(tmp_path)
    raw = doc.to_bytes(opts)
    independent = Document(BytesIO(raw))
    assert independent.paragraphs[0].runs[0].font.size.pt == 21
    assert independent.paragraphs[0].style.font.size.pt == 30
    with ZipFile(BytesIO(raw)) as package:
        relationships = ET.fromstring(package.read('word/_rels/document.xml.rels'))
        assert any(r.get('Target') == 'https://example.com' for r in relationships)


def test_reference_remaps_colliding_ids_and_based_on(tmp_path):
    original = ldm.Document(styles=[ldm.Style(name='Brand Mark', type=1)])
    incoming = ldm.Document(styles=[
        ldm.Style(name='BrandMark', type=1),
        ldm.Style(name='Child', type=1, base_style_name='BrandMark')])
    result = ET.fromstring(apply_reference_styles(render_styles_xml(original), incoming))
    styles = {s.find(W+'name').get(W+'val'): s for s in result.findall(W+'style')}
    assert styles['Brand Mark'].get(W+'styleId') != styles['BrandMark'].get(W+'styleId')
    assert styles['Child'].find(W+'basedOn').get(W+'val') == styles['BrandMark'].get(W+'styleId')


def test_reference_type_conflict_is_rejected():
    source = ldm.Document(styles=[ldm.Style(name='Brand', type=1)])
    incoming = ldm.Document(styles=[ldm.Style(name='Brand', type=2)])
    with pytest.raises(ValueError, match='type differs'):
        apply_reference_styles(render_styles_xml(source), incoming)


def test_ambiguous_reference_names_are_rejected_before_normalization():
    source = ldm.Document()
    incoming = ldm.Document(styles=[ldm.Style(name='Brand', type=1),
                                   ldm.Style(name='Brand', type=2)])
    with pytest.raises(ValueError, match='ambiguous'):
        apply_reference_styles(render_styles_xml(source), incoming)


def test_invalid_reference_does_not_replace_output(tmp_path):
    target = tmp_path / 'existing.docx'
    target.write_bytes(b'unchanged')
    opts = aw.saving.OoxmlSaveOptions()
    opts.reference_docx = tmp_path / 'missing.docx'
    with pytest.raises(FileNotFoundError):
        aw.Document(BytesIO(b'Body')).save(target, opts)
    assert target.read_bytes() == b'unchanged'
    opts.reference_docx = b'not a path'
    with pytest.raises(ValueError, match='path'):
        LdmDocxWriter(opts).write_to_bytes(ldm.Document())


@pytest.mark.parametrize('location', ['body', 'table', 'header', 'footer'])
def test_reference_character_chain_is_not_flattened_to_direct_format(tmp_path, location):
    source = Document()
    base = source.styles.add_style('Brand Base', WD_STYLE_TYPE.CHARACTER)
    base.font.size = Pt(16)
    base.font.color.rgb = RGBColor.from_string('0000FF')
    child = source.styles.add_style('Brand Child', WD_STYLE_TYPE.CHARACTER)
    child.base_style = base
    if location == 'body':
        para = source.add_paragraph()
    elif location == 'table':
        para = source.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    else:
        para = getattr(source.sections[0], location).paragraphs[0]
    para.add_run('INHERITED').style = child
    direct = para.add_run('DIRECT')
    direct.style = child
    direct.font.size = Pt(21)
    stream = BytesIO()
    source.save(stream)
    base.font.size = Pt(24)
    base.font.color.rgb = RGBColor.from_string('FF0000')
    reference_path = tmp_path / 'character-reference.docx'
    source.save(reference_path)
    opts = aw.saving.OoxmlSaveOptions()
    opts.reference_docx = reference_path
    doc = aw.Document(BytesIO(stream.getvalue()))
    fonts = []
    def collect_fonts(node):
        if isinstance(node, ldm.Paragraph):
            for run in node.runs:
                collect_fonts(run)
            return
        if isinstance(node, ldm.Run) and node.text in ('INHERITED', 'DIRECT'):
            fonts.append(node.font)
        for name in getattr(type(node), 'model_fields', {}):
            value = getattr(node, name)
            for child in value if isinstance(value, list) else [value]:
                collect_fonts(child)
    for section in doc.light_document_model.sections:
        collect_fonts(section)
    assert [font.size for font in fonts] == [16, 21]
    assert all(color_to_hex(font.color) == '0000FF' for font in fonts)
    before = doc.light_document_model.model_dump()
    raw = doc.to_bytes(opts)
    result = Document(BytesIO(raw))
    if location == 'body':
        actual = result.paragraphs[0]
    elif location == 'table':
        actual = result.tables[0].cell(0, 0).paragraphs[0]
    else:
        actual = getattr(result.sections[0], location).paragraphs[0]
    assert actual.runs[0].font.size is None
    assert actual.runs[0].font.color.rgb is None
    assert actual.runs[1].font.size.pt == 21
    assert actual.runs[1].font.color.rgb is None
    assert actual.runs[0].style.base_style.font.size.pt == 24
    assert str(actual.runs[0].style.base_style.font.color.rgb) == 'FF0000'
    fonts.clear()
    for section in aw.Document(BytesIO(raw)).light_document_model.sections:
        collect_fonts(section)
    assert [font.size for font in fonts] == [24, 21]
    assert all(color_to_hex(font.color) == 'FF0000' for font in fonts)
    assert doc.light_document_model.model_dump() == before


def test_character_inheritance_reaches_pdf_and_direct_overrides_win():
    import fitz

    source = Document()
    base = source.styles.add_style('Brand Base', WD_STYLE_TYPE.CHARACTER)
    base.font.size = Pt(16)
    base.font.color.rgb = RGBColor.from_string('0000FF')
    child = source.styles.add_style('Brand Child', WD_STYLE_TYPE.CHARACTER)
    child.base_style = base
    source.add_paragraph().add_run('INHERITED').style = child
    direct = source.add_paragraph().add_run('DIRECT')
    direct.style = child
    direct.font.size = Pt(21)
    stream = BytesIO()
    source.save(stream)
    doc = aw.Document(BytesIO(stream.getvalue()))
    with fitz.open(stream=doc.to_bytes('pdf'), filetype='pdf') as pdf:
        spans = [span for page in pdf for block in page.get_text('dict')['blocks']
                 for line in block.get('lines', []) for span in line['spans']]
        actual = {span['text']: span for span in spans}
        assert actual['INHERITED']['size'] == pytest.approx(16)
        assert actual['DIRECT']['size'] == pytest.approx(21)
        assert actual['INHERITED']['color'] == 0x0000FF
        assert actual['DIRECT']['color'] == 0x0000FF



def test_reference_theme_colors_resolve_without_mutating_source_snapshots():
    font = ldm.Font(color='Color [Empty]', source_color=ldm.SourceColor(value='auto', theme_color='accent1'),
                    color_rendering='Color [A=255, R=18, G=52, B=86]')
    reference = ldm.Document(styles=[ldm.Style(name='Brand', font=font)])
    before = reference.model_dump()
    generated = render_styles_xml(ldm.Document())
    output = ET.fromstring(apply_reference_styles(generated, reference))
    brand = next(s for s in output.findall(W + 'style') if s.find(W + 'name').get(W + 'val') == 'Brand')
    color = brand.find(W + 'rPr/' + W + 'color')
    assert color.get(W + 'val') == '123456' and color.get(W + 'themeColor') is None
    assert reference.model_dump() == before
