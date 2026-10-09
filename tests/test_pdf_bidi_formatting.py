"""A format boundary must not change paragraph-level bidirectional ordering."""
from io import BytesIO

import pymupdf
from pypdf import PdfReader
import pytest

from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.saving import PdfSaveOptions
from .test_pdf_shaping import arabic_font


def glyphs(raw):
    with pymupdf.open(stream=raw, filetype='pdf') as pdf:
        return sorted((char['c'], *char['bbox']) for page in pdf
                      for block in page.get_text('rawdict')['blocks']
                      for line in block.get('lines', []) for span in line['spans']
                      for char in span['chars'] if not char['c'].isspace())


@pytest.mark.parametrize('parts', [
    ('سلام ', '(123)', ' عالم ', 'ABC', ' نهاية'),
    ('START ', 'سلام ', '(123) ', 'عالم ', 'END'),
    ('سلام ', ' - ', 'ABC', ' (123) ', 'عالم'),
])
@pytest.mark.parametrize('links', [False, True])
@pytest.mark.parametrize('wrapped', [False, True])
def test_format_boundaries_keep_global_bidi_positions(parts, links, wrapped):
    pytest.importorskip('uharfbuzz')
    options = PdfSaveOptions()
    options.text_shaping = True
    options.fallback_fonts = [str(arabic_font())]
    if wrapped:
        parts = (parts + (' ',)) * 3

    def render(fragmented):
        children = []
        for index, part in enumerate(parts if fragmented else (''.join(parts),)):
            link = f'https://example.org/{index}' if links and fragmented else None
            children.append(ldm.Run(text=f'[{part}]({link})' if link else part,
                is_hyperlink=bool(link), font=ldm.Font(size=22,
                color='Color [A=255, R=255, G=0, B=0]' if fragmented and index % 2 else '')))
        model = ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
            page_width=300 if wrapped else 600, page_height=300, left_margin=25, right_margin=25,
            top_margin=25, bottom_margin=25), body=ldm.Body(children=[ldm.Paragraph(children=children)]))])
        snapshot = model.model_dump()
        raw = LdmPdfWriter(options).write_to_bytes(model)
        assert model.model_dump() == snapshot
        return raw

    actual_raw, reference_raw = render(True), render(False)
    actual, expected = glyphs(actual_raw), glyphs(reference_raw)
    assert [item[0] for item in actual] == [item[0] for item in expected]
    for item, target in zip(actual, expected, strict=True):
        assert item[1:] == pytest.approx(target[1:], abs=0.05)
    if links:
        with pymupdf.open(stream=actual_raw, filetype='pdf') as pdf:
            targets = {item['uri'] for item in pdf[0].get_links()}
            assert {f'https://example.org/{index}' for index, part in enumerate(parts)
                    if part.strip()} <= targets
            assert targets <= {f'https://example.org/{index}' for index in range(len(parts))}
    else:
        with pymupdf.open(stream=actual_raw, filetype='pdf') as pdf:
            assert any(span['color'] == 0xff0000 for block in pdf[0].get_text('dict')['blocks']
                       for line in block.get('lines', []) for span in line['spans'])
    assert len(PdfReader(BytesIO(actual_raw)).pages) == 1


def test_bold_bidi_boundaries_match_native_paragraph():
    pytest.importorskip('uharfbuzz')
    from fpdf import FPDF
    from aspose.words_foss.pdf_writer.constants import PT_TO_MM
    from aspose.words_foss.pdf_writer.font import close_fonts, register_fonts

    options = PdfSaveOptions()
    options.text_shaping = True
    options.fallback_fonts = [str(arabic_font())]
    model = ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
        page_width=600, page_height=300, left_margin=25, right_margin=25,
        top_margin=25, bottom_margin=25), body=ldm.Body(children=[ldm.Paragraph(children=[
            ldm.Run(text=part, font=ldm.Font(size=22, bold=index % 2 == 1))
            for index, part in enumerate(('سلام ', 'ABC', ' عالم ', '123', ' نهاية'))])]))])
    actual = LdmPdfWriter(options).write_to_bytes(model)
    native = FPDF(format=(600 * PT_TO_MM, 300 * PT_TO_MM))
    try:
        native.set_margins(25 * PT_TO_MM, 25 * PT_TO_MM, 25 * PT_TO_MM)
        native.set_auto_page_break(True, margin=25 * PT_TO_MM)
        native.set_text_shaping(True)
        register_fonts(native, model, options.fallback_fonts)
        native.add_page()
        native.set_font('DocumentSansSC', size=22)
        native.multi_cell(native.epw, 22 * 1.4 * PT_TO_MM,
                          'سلام **ABC** عالم **123** نهاية', markdown=True)
        reference = bytes(native.output())
    finally:
        close_fonts(native)
    expected, received = glyphs(reference), glyphs(actual)
    assert [item[0] for item in received] == [item[0] for item in expected]
    for item, target in zip(received, expected, strict=True):
        assert item[1:] == pytest.approx(target[1:], abs=0.05)


def test_empty_wrapped_fragments_do_not_move_text_outside_page():
    pytest.importorskip('uharfbuzz')
    options = PdfSaveOptions()
    options.text_shaping = True
    options.fallback_fonts = [str(arabic_font())]
    # Joining the repetitions also exercises a known cross-format shaping limit.
    parts = ('سلام ', ' - ', 'ABC', ' (123) ', 'عالم') * 3

    def render(parts):
        model = ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
            page_width=300, page_height=300, left_margin=25, right_margin=25,
            top_margin=25, bottom_margin=25), body=ldm.Body(children=[ldm.Paragraph(children=[
                ldm.Run(text=part, font=ldm.Font(size=22,
                    color='Color [A=255, R=255, G=0, B=0]' if index % 2 else ''))
                for index, part in enumerate(parts)])]))])
        return LdmPdfWriter(options).write_to_bytes(model)

    raw = render(parts)
    actual, expected = glyphs(raw), glyphs(render((''.join(parts),)))
    assert [item[0] for item in actual] == [item[0] for item in expected]
    assert all(24 <= item[1] <= item[3] <= 276 for item in actual)
    with pymupdf.open(stream=raw, filetype='pdf') as pdf:
        assert pdf[0].get_text().count('ABC') == 3
        assert pdf[0].get_text().count('123') == 3


def test_public_docx_color_boundaries_preserve_bidi_layout():
    pytest.importorskip('uharfbuzz')
    from docx import Document
    from docx.shared import Pt, RGBColor
    import aspose.words_foss as aw

    options = PdfSaveOptions()
    options.text_shaping = True
    options.fallback_fonts = [str(arabic_font())]
    parts = ('سلام ', '(123)', ' عالم ', 'ABC', ' نهاية')
    outputs = []
    for fragmented in (False, True):
        doc = Document()
        para = doc.add_paragraph()
        for index, part in enumerate(parts if fragmented else (''.join(parts),)):
            run = para.add_run(part)
            run.font.size = Pt(22)
            if fragmented and index % 2:
                run.font.color.rgb = RGBColor(255, 0, 0)
        source = BytesIO()
        doc.save(source)
        source.seek(0)
        outputs.append(aw.Document(source).to_bytes(options))
    reference, actual = map(glyphs, outputs)
    assert [item[0] for item in actual] == [item[0] for item in reference]
    for item, target in zip(actual, reference, strict=True):
        assert item[1:] == pytest.approx(target[1:], abs=0.05)


@pytest.mark.parametrize('parts', [('س', 'لام'), ('عالم', 'سلام'), ('سس', 'سس')])
def test_cross_link_word_context_matches_whole_word_glyphs(parts):
    pytest.importorskip('uharfbuzz')
    options = PdfSaveOptions()
    options.text_shaping = True
    options.fallback_fonts = [str(arabic_font())]
    rendered = []
    for texts in ((''.join(parts),), parts):
        model = ldm.Document(sections=[ldm.Section(body=ldm.Body(children=[ldm.Paragraph(children=[
            ldm.Run(text=f'[{text}](https://example.org/{index})', is_hyperlink=True,
                    font=ldm.Font(size=28)) for index, text in enumerate(texts)])]))])
        raw = LdmPdfWriter(options).write_to_bytes(model)
        with pymupdf.open(stream=raw, filetype='pdf') as pdf:
            assert pdf[0].get_pixmap(matrix=pymupdf.Matrix(2, 2)).samples
            rendered.append(raw)
            assert len(pdf[0].get_links()) == len(texts)
    actual, reference = map(glyphs, rendered)
    assert [item[0] for item in actual] == [item[0] for item in reference]
    for item, target in zip(actual, reference, strict=True):
        assert item[1:] == pytest.approx(target[1:], abs=0.05)


@pytest.mark.parametrize('cut', [1, 2, 3])
def test_cross_color_glyphs_match_independent_harfbuzz_ranges(cut):
    hb = pytest.importorskip('uharfbuzz')
    from fontTools.ttLib import TTFont

    path = arabic_font()
    source = 'سلام'
    original = TTFont(path)
    font = hb.Font(hb.Face(path.read_bytes()))
    expected = set()
    for start, end in ((0, cut), (cut, len(source))):
        buffer = hb.Buffer()
        buffer.add_str(source, start, end - start)
        buffer.guess_segment_properties()
        hb.shape(font, buffer)
        for info in buffer.glyph_infos:
            coords = original['glyf'][original.getGlyphName(info.codepoint)].getCoordinates(original['glyf'])[0]
            expected.add(tuple(map(tuple, coords)))
    original.close()
    options = PdfSaveOptions()
    options.text_shaping = True
    options.fallback_fonts = [str(path)]
    model = ldm.Document(sections=[ldm.Section(body=ldm.Body(children=[ldm.Paragraph(children=[
        ldm.Run(text=source[:cut], font=ldm.Font(size=28)),
        ldm.Run(text=source[cut:], font=ldm.Font(size=28, color='Color [A=255, R=255, G=0, B=0]'))])]))])
    raw = LdmPdfWriter(options).write_to_bytes(model)
    observed = set()
    # RTL extraction across separate color spans is parser-dependent; verify character coverage.
    assert sorted(source) == sorted(PdfReader(BytesIO(raw)).pages[0].extract_text())
    for reference in PdfReader(BytesIO(raw)).pages[0]['/Resources']['/Font'].values():
        resource = reference.get_object()
        if 'Arial' not in str(resource['/BaseFont']) and 'DejaVu' not in str(resource['/BaseFont']):
            continue
        stream = resource['/DescendantFonts'][0].get_object()['/FontDescriptor']['/FontFile2'].get_data()
        with TTFont(BytesIO(stream)) as embedded:
            for name in embedded.getGlyphOrder()[1:]:
                coords = embedded['glyf'][name].getCoordinates(embedded['glyf'])[0]
                if coords:
                    observed.add(tuple(map(tuple, coords)))
    assert observed == expected
    with pymupdf.open(stream=raw, filetype='pdf') as pdf:
        assert any(span['color'] == 0xff0000 for block in pdf[0].get_text('dict')['blocks']
                   for line in block.get('lines', []) for span in line['spans'])


@pytest.mark.parametrize('in_table', [False, True])
def test_repeated_joining_letters_wrap_consistently_across_links(in_table):
    pytest.importorskip('uharfbuzz')
    options = PdfSaveOptions()
    options.text_shaping = True
    options.fallback_fonts = [str(arabic_font())]
    outputs = []
    for parts in (('سس' * 30,), ('سس',) * 30):
        paragraph = ldm.Paragraph(children=[ldm.Run(
            text=f'[{part}](https://example.org/{index})', is_hyperlink=True,
            font=ldm.Font(size=18)) for index, part in enumerate(parts)])
        child = ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[paragraph])])]) if in_table else paragraph
        model = ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
            page_width=220, page_height=400, left_margin=25, right_margin=25,
            top_margin=25, bottom_margin=25), body=ldm.Body(children=[child]))])
        snapshot = model.model_dump()
        raw = LdmPdfWriter(options).write_to_bytes(model)
        assert model.model_dump() == snapshot
        outputs.append(raw)
    reference, actual = map(glyphs, outputs)
    assert len(actual) == len(reference) == 60
    for item, target in zip(actual, reference, strict=True):
        assert item[1:] == pytest.approx(target[1:], abs=0.05)
    assert len({round(item[2], 1) for item in actual}) > 1


def test_contextual_paragraph_measurement_matches_drawn_baselines():
    pytest.importorskip('uharfbuzz')
    from aspose.words_foss.pdf_writer.constants import PT_TO_MM

    options = PdfSaveOptions()
    options.text_shaping = True
    options.fallback_fonts = [str(arabic_font())]
    paragraph = ldm.Paragraph(children=[ldm.Run(
        text=f'[سس](https://example.org/{index})', is_hyperlink=True,
        font=ldm.Font(size=18)) for index in range(30)])
    model = ldm.Document(sections=[ldm.Section(page_setup=ldm.PageSetup(
        page_width=220, page_height=400, left_margin=25, right_margin=25,
        top_margin=25, bottom_margin=25), body=ldm.Body(children=[paragraph]))])
    writer = LdmPdfWriter(options)
    raw = writer.write_to_bytes(model)
    estimate = writer._estimate_paragraph_height(paragraph, 170 * PT_TO_MM)
    with pymupdf.open(stream=raw, filetype='pdf') as pdf:
        baselines = {round(span['origin'][1], 1) for block in pdf[0].get_text('dict')['blocks']
                     for line in block.get('lines', []) for span in line['spans']}
    line_height = writer._paragraph_renderer.line_height_mm(18, paragraph.paragraph_format)
    assert estimate == pytest.approx(len(baselines) * line_height, abs=0.01)
