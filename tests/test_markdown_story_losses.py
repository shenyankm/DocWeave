"""Markdown must diagnose omitted page-band content without leaking hidden text."""

import os
import subprocess
import sys
import warnings
from io import BytesIO

import pytest
from docx import Document as NativeDocument
from PIL import Image

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.md_writer import LdmMarkdownWriter


def loaded_document(header_type, children):
    doc = aw.Document(BytesIO(b"BODY_SENTINEL"), aw.MarkdownLoadOptions())
    doc.sections[0].headers_footers = [ldm.HeaderFooter(header_footer_type=header_type, children=children)]
    return doc


@pytest.mark.parametrize("header_type", [0, 1, 2, 3])
@pytest.mark.parametrize("kind", ["text", "table"])
@pytest.mark.parametrize("output_mode", ["memory", "path", "writer"])
def test_omitted_page_band_content_warns(tmp_path, header_type, kind, output_mode):
    paragraph = ldm.Paragraph(children=[ldm.Run(text="BAND_SENTINEL")])
    child = paragraph if kind == "text" else ldm.Table(rows=[ldm.Row(cells=[ldm.Cell(paragraphs=[paragraph])])])
    doc = loaded_document(header_type, [child])
    before = doc.light_document_model.model_dump()
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        if output_mode == "memory":
            result = doc.to_bytes("md").decode()
        elif output_mode == "path":
            path = tmp_path / "output.md"
            doc.save(path)
            result = path.read_text()
        else:
            result = LdmMarkdownWriter().write(doc.light_document_model)
    assert any(issubclass(item.category, aw.ContentLossWarning) and "header/footer" in str(item.message)
               for item in caught)
    assert doc.light_document_model.model_dump() == before
    assert "BODY_SENTINEL" in result and "BAND_SENTINEL" not in result
    if output_mode != "writer":
        assert any(item.code == "markdown.header_footer_content_omitted" for item in doc.diagnostics)


@pytest.mark.parametrize("kind", ["hidden", "instruction", "empty"])
def test_invisible_page_band_text_does_not_warn(kind):
    children = [ldm.Run(text="SECRET", font=ldm.Font(hidden=True))] if kind == "hidden" else (
        [ldm.FieldStart(), ldm.Run(text="SECRET"), ldm.FieldEnd()] if kind == "instruction" else [])
    doc = loaded_document(0, [ldm.Paragraph(children=children)])
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        result = doc.to_bytes("md").decode()
    assert "SECRET" not in result
    assert not caught and not doc.diagnostics


def test_warning_error_preserves_output_and_records_diagnostic(tmp_path):
    doc = loaded_document(0, [ldm.Paragraph(children=[ldm.Run(text="HEADER")])])
    output = tmp_path / "original.md"
    output.write_bytes(b"ORIGINAL")
    with warnings.catch_warnings():
        warnings.simplefilter("error", aw.ContentLossWarning)
        with pytest.raises(aw.ContentLossWarning):
            doc.save(output)
    assert output.read_bytes() == b"ORIGINAL"
    assert list(tmp_path.iterdir()) == [output]
    assert doc.diagnostics[-1].code == "markdown.header_footer_content_omitted"


@pytest.fixture
def real_docx():
    native = NativeDocument()
    native.add_paragraph("BODY_SENTINEL")
    native.sections[0].header.paragraphs[0].text = "HEADER_SENTINEL"
    native.sections[0].footer.paragraphs[0].text = "FOOTER_SENTINEL"
    stream = BytesIO()
    native.save(stream)
    assert any(p.text == "HEADER_SENTINEL" for p in native.sections[0].header.paragraphs)
    return stream.getvalue()


def test_real_docx_omission_is_recorded_even_with_ignored_warnings(real_docx):
    doc = aw.Document(BytesIO(real_docx))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", aw.ContentLossWarning)
        result = doc.to_bytes("md").decode()
    assert result.strip() == "BODY_SENTINEL"
    assert doc.diagnostics[-1].code == "markdown.header_footer_content_omitted"


def test_strict_cli_refuses_page_band_loss_without_replacing_output(tmp_path, real_docx):
    source = tmp_path / "source.docx"
    source.write_bytes(real_docx)
    output = tmp_path / "output.md"
    output.write_bytes(b"ORIGINAL")
    result = subprocess.run([sys.executable, "-m", "aspose.words_foss.convert",
                             str(source), str(output), "--strict"],
                            capture_output=True, text=True, timeout=20, check=False)
    expected = "header/footer" if os.name == "posix" else "requires POSIX process groups"
    assert result.returncode != 0 and expected in result.stderr
    assert output.read_bytes() == b"ORIGINAL"


@pytest.mark.parametrize("hidden", [False, True])
def test_page_band_note_reference_is_diagnosed_only_when_visible(hidden):
    doc = loaded_document(0, [ldm.Paragraph(children=[
        ldm.NoteReference(kind="footnote", identifier="1", hidden=hidden)])])
    options = aw.saving.MarkdownSaveOptions()
    options.export_notes = True
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        doc.to_bytes(options)
    assert bool(caught) == (not hidden)
    assert bool(doc.diagnostics) == (not hidden)


@pytest.mark.parametrize("header_type", [0, 1, 2, 3])
def test_supported_page_band_images_still_export_without_loss_warning(header_type):
    stream = BytesIO()
    Image.new("RGB", (2, 2), "red").save(stream, format="PNG")
    shape = ldm.Shape(has_image=True, image_data=ldm.ImageData(
        image_type=ldm.ImageData.from_mime("image/png"), image_bytes=stream.getvalue()))
    doc = loaded_document(header_type, [ldm.Paragraph(children=[shape])])
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        result = doc.to_bytes("md").decode()
    assert ("data:image/png;base64," in result) == (header_type in (0, 1))
    assert bool(caught) == (header_type not in (0, 1))


@pytest.mark.parametrize("kind", ["text_box", "rule", "field_result"])
def test_other_visible_page_band_content_is_diagnosed(kind):
    paragraph = ldm.Paragraph(children=[ldm.Run(text="VISIBLE")])
    if kind == "text_box":
        paragraph = ldm.Paragraph(children=[ldm.Shape(text_box={"paragraphs": [paragraph]})])
    elif kind == "rule":
        paragraph = ldm.Paragraph(children=[ldm.Shape(shape_type=1)])
    else:
        paragraph = ldm.Paragraph(children=[ldm.FieldStart(), ldm.Run(text="SECRET_CODE"),
            ldm.FieldSeparator(), ldm.Run(text="VISIBLE"), ldm.FieldEnd()])
    doc = loaded_document(0, [paragraph])
    with pytest.warns(aw.ContentLossWarning, match="header/footer"):
        result = doc.to_bytes("md")
    assert b"VISIBLE" not in result and b"SECRET_CODE" not in result
    assert "SECRET_CODE" not in doc.diagnostics[-1].message


def test_multiple_stories_report_once_per_export():
    doc = loaded_document(0, [ldm.Paragraph(children=[ldm.Run(text="HEADER")])])
    doc.sections.append(ldm.Section(headers_footers=[ldm.HeaderFooter(header_footer_type=1,
        children=[ldm.Paragraph(children=[ldm.Run(text="FOOTER")])])]))
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        doc.to_bytes("md")
    assert len(caught) == 1 and len(doc.diagnostics) == 1
