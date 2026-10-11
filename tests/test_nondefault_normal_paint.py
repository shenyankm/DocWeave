"""Native 26.9 paint follows Normal fallback even with a disabled default marker."""
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
import pymupdf
import pytest

import aspose.words_foss as aw
from aspose.words_foss import light_document_model as ldm
from aspose.words_foss.docx_writer import LdmDocxWriter
from aspose.words_foss.pdf_writer import LdmPdfWriter
from aspose.words_foss.saving import OoxmlSaveOptions


def owned_model(flag, size, defaults):
    source = Document()
    source.styles['Normal'].element.set(qn('w:default'), flag)
    source.styles['Normal'].font.size = Pt(size)
    if defaults:
        props = source.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
        props.find(qn('w:sz')).set(qn('w:val'), '28')
    source.add_paragraph('OWNED_FONT')
    stream = BytesIO()
    source.save(stream)
    return aw.Document(BytesIO(stream.getvalue())).light_document_model


def target(model):
    return next(run.font for run in model.get_child_nodes(ldm.NodeType.RUN, True)
                if run.text == 'OWNED_FONT')


def paint(model, expected):
    with pymupdf.open(stream=LdmPdfWriter().write_to_bytes(model), filetype='pdf') as pdf:
        spans = [span for page in pdf for span in page.get_texttrace()
                 if 'OWNED_FONT' in ''.join(chr(char[0]) for char in span['chars'])]
        assert len(spans) == 1
        assert spans[0]['size'] == pytest.approx(expected or .5, abs=.011)


@pytest.mark.parametrize('flag', ['0', 'false', 'off'])
@pytest.mark.parametrize('size', [0, 16])
@pytest.mark.parametrize('defaults', [False, True])
@pytest.mark.parametrize('route', ['original', 'json', 'deep'])
def test_disabled_normal_fallback_reaches_actual_pdf_and_roundtrip(flag, size, defaults, route):
    model = owned_model(flag, size, defaults)
    if route == 'json':
        model = ldm.Document.model_validate_json(model.model_dump_json())
    elif route == 'deep':
        model = model.model_copy(deep=True)
    held = target(model)
    snapshot = model.model_dump_json()
    assert held.size == size
    assert held.size_explicit is False
    paint(model, size)
    assert held.size == size
    assert model.model_dump_json() == snapshot
    for fmt in (aw.SaveFormat.DOCX, aw.SaveFormat.FLAT_OPC):
        current = model.model_copy(deep=True)
        for _ in range(2):
            data = LdmDocxWriter(OoxmlSaveOptions(fmt)).write_to_bytes(current)
            current = aw.Document(BytesIO(data)).light_document_model
            assert target(current).size == size
            assert target(current).size_explicit is False
        paint(current, size)
    assert held.size == size
    assert model.model_dump_json() == snapshot


def test_live_normal_font_edits_update_paint_without_rebinding_shared_copy():
    model = owned_model('off', 16, False)
    held = target(model)
    shallow = model.model_copy()
    independent = model.model_copy(deep=True)
    normal = next(style for style in model.styles if style.name == 'Normal')
    normal.font.size = 20
    assert held is target(shallow)
    assert held.size == target(shallow).size == 20
    assert target(independent).size == 16
    paint(model, 20)
    paint(shallow, 20)
    paint(independent, 16)
    assert held.size_explicit is False
