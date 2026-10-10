"""Verify frozen inventory and corpus integrity, not behavioral compatibility."""

import argparse
import hashlib
import json
import runpy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


def digest(data):
    return hashlib.sha256(data).hexdigest()


def saved_pagination(data, prop):
    from docx import Document
    from docx.oxml.ns import qn

    tags = {"keep_with_next": "keepNext", "keep_together": "keepLines",
            "page_break_before": "pageBreakBefore", "widow_control": "widowControl"}
    tag = tags[prop]
    document = Document(BytesIO(data))
    default = document.styles.element
    for name in ("docDefaults", "pPrDefault", "pPr", tag):
        default = default.find(qn("w:" + name)) if default is not None else None
    initial = prop == "widow_control" if default is None else default.get(qn("w:val"), "1") in {"1", "true", "on"}
    result = {}
    for name in ("Base", "Derived"):
        style, chain = document.styles[name], []
        while style is not None:
            assert style.style_id not in {item.style_id for item in chain}
            chain.append(style)
            style = style.base_style
        value = initial
        for style in reversed(chain):
            direct = getattr(style.paragraph_format, prop)
            if direct is not None:
                value = direct
        result[name.lower()] = value
    paragraph = next(p for p in document.paragraphs if "IMPORT" in p.text)
    direct = getattr(paragraph.paragraph_format, prop)
    result["paragraph"] = result["derived"] if direct is None else direct
    return result


def verify_paragraph_pagination(root):
    report = json.loads((root / "paragraph-pagination-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["rendering_acceptance"] is False
    source, raw = ((root / report[key]).read_bytes() for key in ("corpus", "outputs"))
    assert digest(source) == report["corpus_sha256"] and digest(raw) == report["outputs_sha256"]
    generated = {name: (prop, data) for name, prop, data in runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/paragraph_pagination.py"))["inputs"]()}
    rows = report["records"]
    assert len(generated) == 324 and len(rows) == 1296
    errors = report["setter_errors"]
    assert len(errors) == 24 and all(row["error"] == "TypeError" for row in errors)
    assert {(r["property"], r["target"], r["value"]) for r in errors} == {
        (prop, target, value) for prop in {p for p, _ in generated.values()}
        for target in ("style", "paragraph") for value in (None, 1, "bad")}
    assert {(r["input"], r["target"], r["value"]) for r in rows} == {
        (name, target, value) for name in generated for target in ("style", "paragraph") for value in (False, True)}
    with ZipFile(BytesIO(source)) as inputs, ZipFile(BytesIO(raw)) as outputs:
        assert set(inputs.namelist()) == set(generated)
        assert set(outputs.namelist()) == {r["output"] for r in rows}
        for row in rows:
            data = inputs.read(row["input"])
            prop, expected = generated[row["input"]]
            assert prop == row["property"] and digest(data) == row["input_sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(expected)) as generated_package:
                assert {key: actual.read(key) for key in actual.namelist()} == {key: generated_package.read(key) for key in generated_package.namelist()}
            assert saved_pagination(data, prop) == row["before"]
            assert row["edited"]["derived" if row["target"] == "style" else "paragraph"] == row["value"]
            assert row["edited"] == row["saved_live"] == row["reopened"]
            saved = outputs.read(row["output"])
            assert digest(saved) == row["output_sha256"]
            assert saved_pagination(saved, prop) == row["reopened"]
    return len(rows)


def verify_first_paragraph_page_break(root):
    report = json.loads((root / "first-paragraph-page-break-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["status"].startswith("unresolved_") and report["rendering_acceptance"] is False
    rows = report["records"]
    assert len(rows) == 4
    assert {(r["target"], r["value"]) for r in rows} == {(t, v) for t in ("style", "paragraph") for v in (False, True)}
    source, raw = ((root / report[key]).read_bytes() for key in ("corpus", "outputs"))
    assert digest(source) == report["corpus_sha256"] and digest(raw) == report["outputs_sha256"]
    generated = {name: data for name, _, data in runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/paragraph_pagination.py"))["inputs"](first_paragraph=True)}
    with ZipFile(BytesIO(source)) as inputs, ZipFile(BytesIO(raw)) as outputs:
        assert set(inputs.namelist()) == {r["input"] for r in rows}
        assert set(outputs.namelist()) == {r["output"] for r in rows}
        for row in rows:
            data, saved = inputs.read(row["input"]), outputs.read(row["output"])
            assert digest(data) == row["input_sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(generated[row["input"]])) as expected:
                assert {key: actual.read(key) for key in actual.namelist()} == {key: expected.read(key) for key in expected.namelist()}
            assert digest(saved) == row["output_sha256"]
            assert row["property"] == "page_break_before"
            assert saved_pagination(data, row["property"]) == {"base": False, "derived": False, "paragraph": True}
            assert row["before"] == {"base": False, "derived": False, "paragraph": False}
            assert row["edited"] == row["saved_live"] == row["reopened"] == saved_pagination(saved, row["property"])
    return len(rows)


def trial_pagination_snapshot(data):
    from docx import Document

    paragraphs = Document(BytesIO(data)).paragraphs
    return {"owned": [p.paragraph_format.page_break_before or False for p in paragraphs if "IMPORT" in p.text],
            "trial_banner_count": sum("Created with an evaluation copy" in p.text for p in paragraphs)}


def verify_first_paragraph_trial(root):
    report = json.loads((root / "first-paragraph-trial-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["licensed_behavior_confirmed"] is report["rendering_acceptance"] is False
    assert report["status"].startswith("trial_confounded_")
    source, raw = ((root / report[key]).read_bytes() for key in ("corpus", "outputs"))
    assert digest(source) == report["corpus_sha256"] and digest(raw) == report["outputs_sha256"]
    generated = dict(runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/first_paragraph_trial.py"))["inputs"]())
    rows = report["records"]
    assert len(rows) == 18 and len(generated) == 15
    assert {(r["mode"], r["input"]) for r in rows} == {
        *(('load', name) for name in generated), *(('builder', name) for name in ("none", "empty", "text"))}
    with ZipFile(BytesIO(source)) as inputs, ZipFile(BytesIO(raw)) as outputs:
        assert set(inputs.namelist()) == set(generated)
        assert set(outputs.namelist()) == {r["output"] for r in rows}
        for row in rows:
            name = row["input"]
            retained = name.split("-")[0] not in {"none", "table"}
            if row["mode"] == "load":
                data = inputs.read(name)
                assert digest(data) == row["input_sha256"]
                with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(generated[name])) as expected:
                    assert {key: actual.read(key) for key in actual.namelist()} == {key: expected.read(key) for key in expected.namelist()}
                assert trial_pagination_snapshot(data) == {"owned": [True], "trial_banner_count": 0}
                assert row["before_save"] == {"owned": [retained], "trial_banner_count": 1}
            else:
                assert row["input_sha256"] is None
                assert row["before_save"] == {"owned": [True], "trial_banner_count": 0}
            assert row["saved_live"] == row["reopened"] == {"owned": [retained], "trial_banner_count": 1}
            saved = outputs.read(row["output"])
            assert digest(saved) == row["output_sha256"]
            assert trial_pagination_snapshot(saved) == row["reopened"]
    return len(rows)


def saved_dimensions(data):
    from copy import deepcopy

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.parfmt import ParagraphFormat

    props = ("left_indent", "right_indent", "first_line_indent", "space_before", "space_after")
    document = Document(BytesIO(data))
    element = OxmlElement("w:p")
    defaults = document.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        group = defaults.find(qn("w:pPrDefault"))
        if group is not None and group.find(qn("w:pPr")) is not None:
            element.append(deepcopy(group.find(qn("w:pPr"))))
    def point(fmt, prop):
        alias = {"left_indent": "start", "right_indent": "end", "first_line_indent": "hanging"}.get(prop)
        ppr = fmt._element.find(qn("w:pPr"))
        ind = ppr.find(qn("w:ind")) if ppr is not None else None
        if alias and ind is not None:
            physical = {"start": "left", "end": "right", "hanging": "firstLine"}[alias]
            values = [(attribute, value) for attribute, value in ind.attrib.items()
                      if attribute in {qn("w:" + alias), qn("w:" + physical)}]
            if values:
                attribute, value = values[-1]
                return int(value) / (-20 if attribute == qn("w:hanging") else 20)
        value = getattr(fmt, prop)
        return value.pt if value is not None else None

    initial = {prop: point(ParagraphFormat(element), prop) or 0.0 for prop in props}

    def resolved(style, direct=None):
        chain = []
        while style is not None:
            assert style.style_id not in {item.style_id for item in chain}
            chain.append(style)
            style = style.base_style
        result = dict(initial)
        for fmt in [item.paragraph_format for item in reversed(chain)] + ([direct] if direct is not None else []):
            for prop in props:
                value = point(fmt, prop)
                if value is not None:
                    result[prop] = value
        return result

    paragraph = next(p for p in document.paragraphs if "IMPORT" in p.text)
    return {"base": resolved(document.styles["Base"]), "derived": resolved(document.styles["Derived"]),
            "paragraph": resolved(paragraph.style, paragraph.paragraph_format)}


def verify_paragraph_character_indents(root):
    from docx import Document
    from docx.oxml.ns import qn

    probe = runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/paragraph_character_indents.py"))
    report = json.loads((root / "paragraph-character-indents-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is report["sdk_acceptance"] is False
    assert report["font_contexts"] == [list(context) for context in probe["FONT_CONTEXTS"]]
    generated = dict(probe["nonfirst_inputs"]())
    source, raw = ((root / report[key]).read_bytes() for key in ("corpus", "outputs"))
    assert digest(source) == report["corpus_sha256"] and digest(raw) == report["outputs_sha256"]
    rows = report["records"]
    expected = {(name, None, layout) for name in generated if name.startswith("fonts/") for layout in (False, True)}
    expected |= {(name, value, False) for name in generated if not name.startswith("fonts/") for value in probe["VALUES"]}
    assert len(rows) == len(expected) == 61
    assert {(row["input"], row["value"], row["layout"]) for row in rows} == expected
    properties = probe["PROPERTIES"]
    attributes = ("left", "right", "firstLine", "leftChars", "rightChars", "firstLineChars")

    def saved(data):
        document = Document(BytesIO(data))
        paragraph = next(p for p in document.paragraphs if "IMPORT" in p.text)
        ind = paragraph._p.find(qn("w:pPr") + "/" + qn("w:ind"))
        result = {}
        for prop, attr in zip(properties, attributes):
            alias = {"firstLine": "hanging", "firstLineChars": "hangingChars"}.get(attr)
            value = ind.get(qn("w:" + attr)) if ind is not None else None
            if value is None and alias and ind is not None:
                other = ind.get(qn("w:" + alias))
                value = str(-int(other)) if other is not None else None
            result[prop] = int(value or "0") / (100 if attr.endswith("Chars") else 20)
        return result

    with ZipFile(BytesIO(source)) as inputs, ZipFile(BytesIO(raw)) as outputs:
        assert len(inputs.namelist()) == len(generated) == 17 and set(inputs.namelist()) == set(generated)
        assert len(outputs.namelist()) == len(rows) and set(outputs.namelist()) == {row["output"] for row in rows}
        for row in rows:
            data = inputs.read(row["input"])
            assert digest(data) == row["input_sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(generated[row["input"]])) as original:
                assert {name: actual.read(name) for name in actual.namelist()} == {name: original.read(name) for name in original.namelist()}
            assert all(set(row[phase]) == set(properties) for phase in ("loaded", "after_edit", "before_save", "after_save_live", "after_reopen"))
            initial = dict.fromkeys(properties, 0.0)
            if row["property"] is None:
                assert row["input"].startswith("fonts/")
                index = int(Path(row["input"]).stem)
                default, style, _, run = probe["FONT_CONTEXTS"][index]
                initial.update(zip(properties, (10.0, 20.0, (run or style or default or 10) / 2, 1.0, 2.0, 0.5)))
                assert row["after_edit"] == row["loaded"]
            else:
                assert row["property"] == row["input"].split("/")[0] and row["property"] in properties[3:]
                index = properties.index(row["property"]) - 3
                initial[properties[index]] = 13.75 if index == 2 else 12.5
                initial[row["property"]] = 1.25
                # Frozen native quantization observations, not Python round(value * 100).
                quantized = dict(zip(probe["VALUES"], (-1.23, 0.0, 1.23, 1.22, 12.37)))
                assert row["after_edit"][row["property"]] == quantized[row["value"]]
                if row["value"] == 0:
                    assert row["after_edit"][properties[index]] == row["loaded"][properties[index]]
            assert row["loaded"] == initial
            assert row["before_save"] == (row["after_save_live"] if row["layout"] else row["after_edit"])
            output = outputs.read(row["output"])
            assert digest(output) == row["output_sha256"]
            assert saved(output) == row["after_reopen"] == row["after_save_live"]
    errors = report["setter_errors"]
    assert len(errors) == 27
    assert {(row["input"], repr(row["value"])) for row in errors} == {
        (name, repr(value)) for name in generated if not name.startswith("fonts/") for value in (None, True, "bad")}
    assert all(row["property"] == row["input"].split("/")[0] and row["error"] == "TypeError" for row in errors)
    return len(rows)


def verify_paragraph_character_reads(root):
    probe = runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/paragraph_character_reads.py"))
    report = json.loads((root / "paragraph-character-reads-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["python"].startswith("3.13.") and report["platform"].startswith("macOS-")
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is report["sdk_acceptance"] is False
    generated = dict(probe["inputs"]())
    rows = report["records"]
    raw = (root / report["corpus"]).read_bytes()
    assert digest(raw) == report["corpus_sha256"]
    assert len(rows) == len(generated) == 24 and {row["input"] for row in rows} == set(generated)
    with ZipFile(BytesIO(raw)) as corpus:
        assert len(corpus.namelist()) == 24 and set(corpus.namelist()) == set(generated)
        for row in rows:
            data = corpus.read(row["input"])
            assert digest(data) == row["input_sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(generated[row["input"]])) as original:
                assert {name: actual.read(name) for name in actual.namelist()} == {name: original.read(name) for name in original.namelist()}
            prop = row["input"].split("/")[0]
            layers = probe["LAYERS"][int(Path(row["input"]).stem)]
            assert set(row["values"]) == {"base", "derived", "paragraph"}
            for target, end in (("base", 2), ("derived", 3), ("paragraph", 4)):
                expected = dict.fromkeys(probe["PROPERTIES"], 0.0)
                expected[prop] = next((value / 100 for value in reversed(layers[:end]) if value is not None), 0.0)
                assert row["values"][target] == expected
    return len(rows)


def verify_paragraph_character_inheritance_edits(root):
    from docx import Document
    from docx.oxml.ns import qn

    probes = Path(__file__).parents[1] / "docs/probes"
    source_probe = runpy.run_path(str(probes / "paragraph_character_reads.py"))
    edits = runpy.run_path(str(probes / "paragraph_character_inheritance_edits.py"))
    properties = runpy.run_path(str(probes / "paragraph_character_indents.py"))["PROPERTIES"]
    report = json.loads((root / "paragraph-character-inheritance-edits-26.9.json").read_text())
    baseline = json.loads((root / "paragraph-character-reads-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["python"].startswith("3.13.") and report["platform"].startswith("macOS-")
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is report["sdk_acceptance"] is False
    assert report["corpus"] == baseline["corpus"] and report["corpus_sha256"] == baseline["corpus_sha256"]
    generated = dict(source_probe["inputs"]())
    rows = report["records"]
    expected = {(name, target, value) for name in generated for target in edits["TARGETS"] for value in edits["VALUES"]}
    assert len(rows) == len(expected) == 216
    assert {(row["input"], row["target"], row["value"]) for row in rows} == expected
    initial_chars = {row["input"]: row["values"] for row in baseline["records"]}
    attributes = {"left": properties[0], "start": properties[0], "right": properties[1], "end": properties[1],
                  "firstLine": properties[2], "hanging": properties[2], "leftChars": properties[3], "startChars": properties[3],
                  "rightChars": properties[4], "endChars": properties[4], "firstLineChars": properties[5], "hangingChars": properties[5]}
    word_namespace = qn("w:p")[:-1]

    def saved(layers):
        result = dict.fromkeys(properties, 0.0)
        for layer in layers:
            ind = layer.find(qn("w:ind")) if layer is not None else None
            if ind is not None:
                for key, raw in ind.attrib.items():
                    if key.startswith(word_namespace):
                        name = key.split("}")[1]
                        if name in attributes:
                            divisor = 100 if name.endswith("Chars") else 20
                            result[attributes[name]] = int(raw) / (-divisor if name in {"hanging", "hangingChars"} else divisor)
        return result

    source, raw = ((root / report[key]).read_bytes() for key in ("corpus", "outputs"))
    assert digest(source) == report["corpus_sha256"] and digest(raw) == report["outputs_sha256"]
    with ZipFile(BytesIO(source)) as inputs, ZipFile(BytesIO(raw)) as outputs:
        assert len(inputs.namelist()) == 24 and set(inputs.namelist()) == set(generated)
        assert len(outputs.namelist()) == 216 and set(outputs.namelist()) == {row["output"] for row in rows}
        for row in rows:
            data = inputs.read(row["input"])
            assert digest(data) == row["input_sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(generated[row["input"]])) as original:
                assert {name: actual.read(name) for name in actual.namelist()} == {name: original.read(name) for name in original.namelist()}
            for phase in ("loaded", "after_edit", "after_save_live", "after_reopen"):
                assert set(row[phase]) == set(edits["TARGETS"])
                assert all(set(values) == set(properties) for values in row[phase].values())
            prop = row["input"].split("/")[0]
            assert row["property"] == prop
            index = properties.index(prop) - 3
            point = properties[index]
            layers = list(source_probe["LAYERS"][int(Path(row["input"]).stem)])
            _, base, derived, direct = layers
            # Measured load/edit rules for these 24 fixtures; not a general font resolver.
            size = 11 if index == 2 else 5
            base_point = (base or 0) * size / 100
            derived_point = derived * size / 100 if derived is not None else base_point
            paragraph_point = direct * (11 if index == 2 else 10) / 100 if direct else derived_point
            quantized = {0.0: 0.0, 1.235: 1.23, -1.235: -1.23}[row["value"]]
            for target, value in zip(edits["TARGETS"], (base_point, derived_point, paragraph_point)):
                initial = dict.fromkeys(properties, 0.0)
                initial.update(initial_chars[row["input"]][target])
                initial[point] = value
                assert row["loaded"][target] == initial
            layers[edits["TARGETS"].index(row["target"]) + 1] = quantized * 100
            for target, end in (("base", 2), ("derived", 3), ("paragraph", 4)):
                expected_edit = dict(row["loaded"][target])
                expected_edit[prop] = next((value / 100 for value in reversed(layers[:end]) if value is not None), 0.0)
                if target == row["target"] == "paragraph":
                    expected_edit[point] = round(quantized * 11 * 20) / 20 if quantized else expected_edit[point]
                    if index == 2 and quantized < 0:
                        expected_edit["left_indent"] = -expected_edit[point]
                assert row["after_edit"][target] == expected_edit
            assert row["after_save_live"] == row["after_reopen"]
            if row["target"] == "paragraph":
                assert row["after_edit"] == row["after_save_live"]
            output = outputs.read(row["output"])
            assert digest(output) == row["output_sha256"]
            document = Document(BytesIO(output))
            assert document.styles["Base"].element.find(qn("w:basedOn")) is None
            reference = document.styles["Derived"].element.find(qn("w:basedOn"))
            assert reference is not None and reference.get(qn("w:val")) == "Base"
            default = document.styles.element.find(qn("w:docDefaults") + "/" + qn("w:pPrDefault") + "/" + qn("w:pPr"))
            base = document.styles["Base"].element.find(qn("w:pPr"))
            derived = document.styles["Derived"].element.find(qn("w:pPr"))
            paragraph = next(p for p in document.paragraphs if "IMPORT" in p.text)
            paragraph = paragraph._p.find(qn("w:pPr"))
            reference = paragraph.find(qn("w:pStyle"))
            assert reference is not None and reference.get(qn("w:val")) == "Derived"
            for target, layers in (("base", [default, base]), ("derived", [default, base, derived]),
                                   ("paragraph", [default, base, derived, paragraph])):
                assert saved(layers) == row["after_reopen"][target]
    errors = report["setter_errors"]
    assert len(errors) == 216
    assert {(row["input"], row["target"], repr(row["value"])) for row in errors} == {
        (name, target, repr(value)) for name in generated for target in edits["TARGETS"] for value in (None, True, "bad")}
    assert all(row["error"] == "TypeError" and row["property"] == row["input"].split("/")[0] for row in errors)
    return len(rows)


def verify_paragraph_character_setters(root):
    from docx import Document
    from docx.oxml.ns import qn

    probe = runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/paragraph_character_indents.py"))
    edits = runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/paragraph_character_setters.py"))
    report = json.loads((root / "paragraph-character-setters-26.9.json").read_text())
    baseline = json.loads((root / "paragraph-character-indents-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["python"].startswith("3.13.") and report["platform"].startswith("macOS-")
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is report["sdk_acceptance"] is False
    assert report["corpus"] == baseline["corpus"] and report["corpus_sha256"] == baseline["corpus_sha256"]
    generated = {name: raw for name, raw in probe["nonfirst_inputs"]() if name.startswith("fonts/")}
    properties = probe["PROPERTIES"]
    rows = report["records"]
    expected = {(name, target, prop, value) for name in generated for target in edits["TARGETS"]
                for prop in properties[3:] for value in edits["VALUES"]}
    assert len(rows) == len(expected) == 144
    assert {(row["input"], row["target"], row["property"], row["value"]) for row in rows} == expected
    initial = {row["input"]: row["loaded"] for row in baseline["records"] if row["property"] is None and not row["layout"]}
    attributes = ("left", "right", "firstLine", "leftChars", "rightChars", "firstLineChars")

    def saved(element):
        ind = element.find(qn("w:pPr") + "/" + qn("w:ind"))
        result = {}
        for prop, attr in zip(properties, attributes):
            value = ind.get(qn("w:" + attr)) if ind is not None else None
            if value is None and attr in {"firstLine", "firstLineChars"} and ind is not None:
                value = ind.get(qn("w:hanging" if attr == "firstLine" else "w:hangingChars"))
                value = str(-int(value)) if value is not None else None
            result[prop] = int(value or "0") / (100 if attr.endswith("Chars") else 20)
        return result

    source, raw = ((root / report[key]).read_bytes() for key in ("corpus", "outputs"))
    assert digest(source) == report["corpus_sha256"] and digest(raw) == report["outputs_sha256"]
    with ZipFile(BytesIO(source)) as inputs, ZipFile(BytesIO(raw)) as outputs:
        assert len(outputs.namelist()) == len(rows) and set(outputs.namelist()) == {row["output"] for row in rows}
        for row in rows:
            data = inputs.read(row["input"])
            assert digest(data) == row["input_sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(generated[row["input"]])) as original:
                assert {name: actual.read(name) for name in actual.namelist()} == {name: original.read(name) for name in original.namelist()}
            for phase in ("loaded", "after_edit", "after_save_live", "after_reopen"):
                assert set(row[phase]) == {"target", "paragraph"}
                assert all(set(values) == set(properties) for values in row[phase].values())
            paragraph_target = row["target"] == "paragraph"
            loaded = initial[row["input"]] if paragraph_target else dict.fromkeys(properties, 0.0)
            assert row["loaded"] == {"target": loaded, "paragraph": initial[row["input"]]}
            quantized = {0.0: 0.0, 1.235: 1.23, -1.235: -1.23}[row["value"]]
            for prop in properties[3:]:
                assert row["after_edit"]["target"][prop] == (quantized if prop == row["property"] else loaded[prop])
            # These assertions describe the eight owned fixtures, not a general font/layout algorithm.
            default, style, _, run = probe["FONT_CONTEXTS"][int(Path(row["input"]).stem)]
            default = default or 10
            point = properties[properties.index(row["property"]) - 3]
            if paragraph_target:
                assert row["after_edit"] == row["after_save_live"]
                assert row["after_edit"]["target"] == row["after_edit"]["paragraph"]
                expected_points = dict(loaded, left_indent=default, right_indent=default * 2)
                expected_points[row["property"]] = quantized
                size = (run or style or default) if point == "first_line_indent" else default
                expected_points[point] = round(quantized * size * 20) / 20 if quantized else loaded[point]
                if point == "first_line_indent" and quantized < 0:
                    expected_points["left_indent"] -= expected_points[point]
                assert row["after_edit"]["target"] == expected_points
            else:
                assert all(row["after_edit"]["target"][prop] == 0 for prop in properties[:3])
                assert row["after_edit"]["paragraph"] == row["loaded"]["paragraph"]
                expected_points = dict(row["after_edit"]["target"])
                size = (style or default) if point == "first_line_indent" else 5
                expected_points[point] = round(quantized * size * 20) / 20
                assert row["after_reopen"]["target"] == expected_points
            assert row["after_save_live"] == row["after_reopen"]
            output = outputs.read(row["output"])
            assert digest(output) == row["output_sha256"]
            document = Document(BytesIO(output))
            paragraph = next(p for p in document.paragraphs if "IMPORT" in p.text)
            assert saved(paragraph._p) == row["after_reopen"]["paragraph"]
            target = paragraph._p if paragraph_target else document.styles["P"].element
            assert saved(target) == row["after_reopen"]["target"]
    errors = report["setter_errors"]
    assert len(errors) == 144
    assert {(row["input"], row["target"], row["property"], repr(row["value"])) for row in errors} == {
        (name, target, prop, repr(value)) for name in generated for target in edits["TARGETS"]
        for prop in properties[3:] for value in (None, True, "bad")}
    assert all(row["error"] == "TypeError" for row in errors)
    return len(rows)


def verify_character_indent_roundtrips(root):
    from xml.etree import ElementTree as ET

    from docx import Document
    from docx.oxml.ns import qn

    report = json.loads((root / "paragraph-character-indents-current.json").read_text())
    assert report["baseline_version"] == "26.9.0"
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is False
    native_raw = (root / report["native_report"]).read_bytes()
    assert digest(native_raw.replace(b"\r\n", b"\n")) == report["native_report_sha256"]
    native = json.loads(native_raw)
    assert native["outputs_sha256"] == report["native_outputs_sha256"]
    originals = {row["output"]: row for row in native["records"]}
    rows = report["records"]
    assert len(rows) == 122
    assert {(row["native_output"], row["format"]) for row in rows} == {
        (name, fmt) for name in originals for fmt in ("DOCX", "FLAT_OPC")}
    raw = (root / report["outputs"]).read_bytes()
    assert digest(raw) == report["outputs_sha256"]
    fields = ("character_unit_left_indent", "character_unit_right_indent", "character_unit_first_line_indent")
    reread = report["official_reread"]
    assert reread["version"] == "26.9.0" and reread["licensed"] is False
    assert reread["outputs_sha256"] == report["outputs_sha256"]
    observed = {row["output"]: row for row in reread["records"]}
    assert len(observed) == len(reread["records"]) == len(rows)
    assert set(observed) == {row["output"] for row in rows}
    with ZipFile(BytesIO(raw)) as outputs:
        assert len(outputs.namelist()) == len(rows) and set(outputs.namelist()) == {row["output"] for row in rows}
        for row in rows:
            original = originals[row["native_output"]]
            assert row["native_output_sha256"] == original["output_sha256"]
            expected = {name: original["after_reopen"][name] for name in fields}
            assert row["character_units"] == expected
            assert observed[row["output"]]["character_units"] == expected
            assert observed[row["output"]]["output_sha256"] == row["output_sha256"]
            data = outputs.read(row["output"])
            assert digest(data) == row["output_sha256"]
            if row["format"] == "DOCX":
                document = Document(BytesIO(data))
                element = next(p._p for p in document.paragraphs if p.text == "IMPORT")
            else:
                package = "{http://schemas.microsoft.com/office/2006/xmlPackage}"
                parts = [part for part in ET.fromstring(data) if part.get(package + "name") == "/word/document.xml"]
                assert len(parts) == 1
                document = parts[0].find(package + "xmlData/" + qn("w:document"))
                element = next(p for p in document.iter(qn("w:p")) if "".join(t.text or "" for t in p.iter(qn("w:t"))) == "IMPORT")
            ind = element.find(qn("w:pPr") + "/" + qn("w:ind"))
            actual = {}
            for field, attribute in zip(fields, ("leftChars", "rightChars", "firstLineChars")):
                value = ind.get(qn("w:" + attribute)) if ind is not None else None
                if value is None and attribute == "firstLineChars" and ind is not None:
                    hanging = ind.get(qn("w:hangingChars"))
                    value = str(-int(hanging)) if hanging is not None else None
                actual[field] = int(value or "0") / 100
            assert actual == expected
    return len(rows)


def verify_paragraph_dimensions(root, filename="paragraph-dimensions-26.9.json"):
    assert filename in {"paragraph-dimensions-26.9.json", "paragraph-spacing-limits-26.9.json", "paragraph-indent-limits-26.9.json", "paragraph-logical-indents-26.9.json"}
    logical = filename == "paragraph-logical-indents-26.9.json"
    limits = filename == "paragraph-spacing-limits-26.9.json"
    indents = filename == "paragraph-indent-limits-26.9.json"
    report = json.loads((root / filename).read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    probe = runpy.run_path(str(Path(__file__).parents[1] / "docs/probes" / ("paragraph_logical_indents.py" if logical else "paragraph_indent_limits.py" if indents else "paragraph_spacing_limits.py" if limits else "paragraph_dimensions.py")))
    generated = {name: (prop, data) for name, prop, data in probe["inputs"]()}
    assert (len(generated), len(report["records"]), len(report["setter_errors"])) == ((42, 168, 252) if logical else (3, 54, 18) if indents else (2, 32, 12) if limits else (45, 360, 270))
    archives = {}
    for key in ("corpus", "outputs") + (("normalized_outputs",) if indents else ()):
        raw = (root / report[key]).read_bytes()
        assert digest(raw) == report[key + "_sha256"]
        with ZipFile(BytesIO(raw)) as archive:
            archives[key] = {name: archive.read(name) for name in archive.namelist()}
    assert set(archives["corpus"]) == set(generated)
    assert set(archives["outputs"]) == {row["output"] for row in report["records"] if row["error"] is None}
    if indents:
        affected = {row["output"] for row in report["records"] if row.get("native_xml_issue")}
        assert len(affected) == 10 and set(archives["normalized_outputs"]) == affected
    before = {}
    for name, (prop, expected) in generated.items():
        data = archives["corpus"][name]
        with ZipFile(BytesIO(data)) as source, ZipFile(BytesIO(expected)) as current:
            assert {n: source.read(n) for n in source.namelist()} == {n: current.read(n) for n in current.namelist()}
        before[name] = saved_dimensions(data)
    assert {(r["input"], r["target"], r["value"]) for r in report["records"]} == {(name, target, value) for name in generated for target in ("style", "paragraph") for value in probe["VALUES"]}
    for row in report["records"]:
        assert row["input_sha256"] == digest(archives["corpus"][row["input"]])
        assert row["property"] == generated[row["input"]][0] and row["before_edit"] == before[row["input"]]
        if row["property"].startswith("space_") and (row["value"] < 0 or row["value"] > 1584):
            assert row["error"] == "RuntimeError" and row["after_edit"] == row["before_edit"]
        else:
            assert row["error"] is None
            raw = archives["outputs"][row["output"]]
            assert digest(raw) == row["output_sha256"]
            if indents and row.get("native_xml_issue"):
                normalized = archives["normalized_outputs"][row["output"]]
                assert digest(normalized) == row["normalized_sha256"]
                assert row["normalized_reopen"] == row["after_reopen"]
                assert row["property"] == "first_line_indent"
                assert row["native_xml_issue"] == "negative unsigned hanging at signed twip minimum"
                with ZipFile(BytesIO(raw)) as native, ZipFile(BytesIO(normalized)) as valid:
                    assert set(native.namelist()) == set(valid.namelist())
                    modified = 0
                    for name in native.namelist():
                        original = native.read(name)
                        expected = original.replace(b'w:hanging="-2147483648"', b'w:hanging="2147483648"') if name in {"word/document.xml", "word/styles.xml"} else original
                        assert valid.read(name) == expected
                        modified += expected != original
                    assert modified == 1
                raw = normalized
            assert saved_dimensions(raw) == row["after_reopen"] == row["after_edit"] == row["after_save_live"]
    assert {(row["input"], row["target"], repr(row["value"])) for row in report["setter_errors"]} == {(name, target, repr(value)) for name in generated for target in ("style", "paragraph") for value in (None, True, "bad")}
    assert all(row["error"] == "TypeError" and row["property"] == generated[row["input"]][0] for row in report["setter_errors"])
    return len(report["records"])


def verify_pagination_rendering(root, filename="pagination-rendering-26.9.json"):
    dimensions = filename == "paragraph-dimensions-rendering-26.9.json"
    assert filename in {"pagination-rendering-26.9.json", "paragraph-dimensions-rendering-26.9.json"}
    report = json.loads((root / filename).read_text())
    assert report["commercial_environment"]["version"] == "26.9.0"
    assert report["commercial_environment"]["licensed"] is False
    assert report["full_rendering_acceptance"] is False
    assert report["tolerances"] == {"text_origin_pt": 0.02, "text_advance_pt": 0.02, "black_ink_difference_ratio": 0.01}
    probe = runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/pagination_rendering.py"))
    generator = runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/paragraph_dimensions_rendering.py")) if dimensions else probe
    generated = {name: (prop, data) for name, prop, data in generator["inputs"]()}
    rows = report["records"]
    assert len(rows) == 10 and len(generated) == 5
    assert {(r["input"], r["value"]) for r in rows} == {(name, value) for name in generated for value in ((0.0, 12.375) if dimensions else (False, True))}
    archives = {}
    keys = ("commercial_outputs", "current_outputs") + (() if dimensions else ("before_outputs",))
    for key in ("corpus", *keys):
        data = (root / report[key]).read_bytes()
        assert digest(data) == report[key + "_sha256"]
        with ZipFile(BytesIO(data)) as archive:
            archives[key] = {name: archive.read(name) for name in archive.namelist()}
    assert set(archives["corpus"]) == set(generated)
    for key in keys:
        assert set(archives[key]) == {r["output"] for r in rows}
    for row in rows:
        prop, expected = generated[row["input"]]
        data = archives["corpus"][row["input"]]
        assert row["property"] == prop and digest(data) == row["input_sha256"]
        with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(expected)) as source:
            assert {name: actual.read(name) for name in actual.namelist()} == {name: source.read(name) for name in source.namelist()}
        snapshots = (("commercial_outputs", "native_snapshot", "output_sha256"),
                     ("current_outputs", "current_snapshot", "current_output_sha256"))
        if not dimensions:
            snapshots += (("before_outputs", "before_snapshot", "before_output_sha256"),)
        for key, snapshot, sha in snapshots:
            raw = archives[key][row["output"]]
            assert digest(raw) == row[sha]
            actual, frozen = probe["pdf_snapshot"](raw), row[snapshot]
            assert actual["page_sizes"] == frozen["page_sizes"] and set(actual["lines"]) == set(frozen["lines"])
            for label, line in frozen["lines"].items():
                observed = actual["lines"][label]
                assert observed["page"] == line["page"] and observed["size"] == line["size"]
                assert max(abs(a - b) for a, b in zip(observed["origin"], line["origin"], strict=True)) <= report["tolerances"]["text_origin_pt"]
                # MuPDF's x86/ARM advances differ by a float ULP; use the declared geometry tolerance.
                assert abs(observed["advance"] - line["advance"]) <= report["tolerances"]["text_advance_pt"]
        native, current = row["native_snapshot"], row["current_snapshot"]
        assert native["page_sizes"] == current["page_sizes"]
        assert set(native["lines"]) == set(current["lines"])
        for label, line in native["lines"].items():
            own = current["lines"][label]
            assert line["page"] == own["page"]
            if label != "BEFORE":
                assert max(abs(a - b) for a, b in zip(line["origin"], own["origin"], strict=True)) <= 0.02
                assert abs(line["advance"] - own["advance"]) <= 0.02 and line["size"] == own["size"]
        pairs = (("current_outputs", "black_ink_difference"),)
        if not dimensions:
            pairs += (("before_outputs", "before_black_ink_difference"),)
        for key, field in pairs:
            ratios = probe["ink_difference"](archives["commercial_outputs"][row["output"]], archives[key][row["output"]])
            assert ratios == row[field]
        assert max(row["black_ink_difference"]) <= 0.01
        if not dimensions:
            assert max(row["before_black_ink_difference"]) > 0.01
    ledger = json.loads((root / "commercial-26.9-capabilities.json").read_text())
    field = "dimension_rendering_evidence" if dimensions else "pagination_rendering_evidence"
    linked = [record for record in ledger["records"] if field in record]
    assert {record["id"] for record in linked} == {"aspose.words.ParagraphFormat." + row["property"] for row in rows}
    for record in linked:
        evidence = record[field]
        assert evidence["file"] == filename and evidence["tolerances"] == report["tolerances"]
        assert evidence["observed_black_ink_difference_ratio"] == max(max(row["black_ink_difference"]) for row in rows)
    return len(rows)


def verify_format_outputs(report, archive_path):
    raw = archive_path.read_bytes()
    assert digest(raw) == report["outputs"]["sha256"], "format archive digest mismatch"
    checked = 0
    with ZipFile(BytesIO(raw)) as archive:
        for key, observed in report["reports"].items():
            for record in observed["records"]:
                if key.startswith("save_"):
                    for output in record["files"]:
                        value = archive.read(key + "/" + record["name"] + "/" + output["name"])
                        assert len(value) == output["size"] and digest(value) == output["sha256"]
                        checked += 1
                elif "saved_docx" in record:
                    output = record["saved_docx"]
                    assert digest(archive.read(key + "/" + output["name"])) == output["sha256"]
                    checked += 1
        from docx import Document

        for key, checks in report["independent_docx_checks"].items():
            for check in checks:
                doc = Document(BytesIO(archive.read(key + "/" + check["format"] + ".docx")))
                text = "\n".join([p.text for p in doc.paragraphs] +
                                 [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
                assert len(doc.paragraphs) == check["paragraph_count"] and len(doc.tables) == check["table_count"]
                assert {label: label in text for label in check["labels_present"]} == check["labels_present"]
                assert any(token in text for token in ("<pkg:package", "<html", "<w:wordDocument", "<w:document")) == check["literal_markup_present"]
    return checked


def verify_import_outputs(root):
    inspect = runpy.run_path(str(Path(__file__).parents[1] / "docs" / "probes" / "inspect_import_outputs.py"))["inspect_document"]
    baseline = json.loads((root / "import-node-26.9.json").read_text())
    current = json.loads((root / "import-node-current.json").read_text())
    checked, actual = 0, {}
    for report, phases in ((baseline, baseline["reports"]), (current, {"candidate": current["report"]})):
        raw = (root / report["archive"]).read_bytes()
        assert digest(raw) == report["archive_sha256"], "import archive digest mismatch"
        with ZipFile(BytesIO(raw)) as archive:
            for phase, observed in phases.items():
                checks = report["independent_checks"] if phase == "candidate" else report["independent_checks"][phase]
                recorded = {row["output"]: row for row in checks}
                assert len(recorded) == len(checks) == len(observed["records"])
                actual[phase] = {}
                for row in observed["records"]:
                    name = row["output"] if phase == "candidate" else phase + "/" + row["output"]
                    data = archive.read(name)
                    assert digest(data) == row["sha256"], "import output digest mismatch"
                    result = {"output": row["output"], **inspect(BytesIO(data))}
                    assert result == recorded[row["output"]], "import independent observation mismatch"
                    actual[phase][row["output"]] = result
                    checked += 1
    matched = {row["output"] for row in current["report"]["records"] if row["outcome"]["status"] == "returned"}
    assert matched == set(current["matched_observations"])
    for name in matched:
        assert actual["candidate"][name] == actual["commercial"][name]
    assert current["validation"]["full_import_acceptance"] is False
    return checked


def verify_style_toggles(root):
    report = json.loads((root / "style-toggles-26.9.json").read_text())
    raw = (root / "corpus" / "style-toggles-26.9.zip").read_bytes()
    assert report["version"] == "26.9.0" and report["module"] == "aspose.words"
    assert report["full_format_acceptance"] is False and report["licensed"] is False
    assert digest(raw) == report["corpus_sha256"], "style corpus digest mismatch"
    generated = runpy.run_path(str(Path(__file__).parents[1] / "docs" / "probes" / "style_toggles.py"))["inputs"]
    expected = {name + ".docx" for name, _ in generated()}
    assert len(expected) == len(report["records"]) == 745
    with ZipFile(BytesIO(raw)) as archive:
        assert set(archive.namelist()) == expected == {row["input"] for row in report["records"]}
        for row in report["records"]:
            assert digest(archive.read(row["input"])) == row["sha256"]
            assert isinstance(row["bold"], bool) and isinstance(row["italic"], bool)
    current = json.loads((root / "style-toggles-current.json").read_text())
    native = {row["input"]: row for row in report["records"]}
    mismatches = {}
    for phase in ("before", "candidate"):
        observed = current["reports"][phase]
        assert observed["corpus_sha256"] == report["corpus_sha256"]
        rows = {row["input"]: row for row in observed["records"]}
        assert len(rows) == len(observed["records"]) == 745 and set(rows) == expected
        assert all(row["sha256"] == native[name]["sha256"] for name, row in rows.items())
        assert all(isinstance(row["bold"], bool) and isinstance(row["italic"], bool) for row in rows.values())
        mismatches[phase] = sorted(name for name, row in rows.items() if row != native[name])
    assert mismatches["before"] == sorted(current["before_mismatches"])
    assert mismatches["candidate"] == current["candidate_mismatches"] == []
    assert current["validation"]["full_format_acceptance"] is False
    return len(report["records"])


def import_pixels_painted(pdf):
    import pymupdf

    return any(any(value < 200 for value in page.get_pixmap(
        clip=pymupdf.Rect(word[:4]), matrix=pymupdf.Matrix(2, 2), alpha=False).samples)
        for page in pdf for word in page.get_text('words') if word[4] == 'IMPORT')


def verify_font_boolean_contexts(root):
    report = json.loads((root / 'font-boolean-contexts-26.9.json').read_text())
    assert report['version'] == '26.9.0' and report['licensed'] is False
    assert report['full_format_acceptance'] is report['rendering_acceptance'] is False
    generated = {name: (field, list(values), raw) for name, field, values, raw in
                 runpy.run_path(str(Path(__file__).parents[1] / 'docs/probes/font_boolean_contexts.py'))['inputs']()}
    corpus = (root / report['corpus']).read_bytes()
    assert digest(corpus) == report['corpus_sha256']
    assert len(generated) == len(report['records']) == 891
    with ZipFile(BytesIO(corpus)) as archive:
        assert set(archive.namelist()) == set(generated) == {r['input'] for r in report['records']}
        for row in report['records']:
            field, values, raw = generated[row['input']]
            source = archive.read(row['input'])
            assert digest(source) == row['sha256']
            with ZipFile(BytesIO(source)) as actual, ZipFile(BytesIO(raw)) as rebuilt:
                assert {n: actual.read(n) for n in actual.namelist()} == {n: rebuilt.read(n) for n in rebuilt.namelist()}
            assert (row['field'], row['values']) == (field, values)
            default, paragraph, character, direct = values
            expected = (bool(default) or (paragraph ^ character) if paragraph is not None and character is not None
                        else paragraph if paragraph is not None else character if character is not None else bool(default))
            assert row['value'] is (expected if direct is None else direct)
    before = []
    for row in report['records']:
        default, paragraph, character, direct = row['values']
        previous = (direct if direct is not None else character if character is not None
                    else paragraph if paragraph is not None else bool(default))
        if row['value'] is not previous:
            before.append(row['input'])
    comparison = report['docweave_getter_comparison']
    assert len(before) == 66 and sorted(before) == sorted(comparison['before_mismatches'])
    assert comparison['after_mismatches'] == []
    pdf_report = json.loads((root / 'font-hidden-rendering-26.9.json').read_text())
    assert pdf_report['version'] == '26.9.0' and pdf_report['licensed'] is False
    assert pdf_report['rendering_acceptance'] is False
    outputs = (root / pdf_report['outputs']).read_bytes()
    assert digest(outputs) == pdf_report['outputs_sha256']
    hidden = {row['input']: row for row in report['records'] if row['field'] == 'hidden'}
    assert len(pdf_report['records']) == len(hidden) == 81
    import pymupdf
    differences = []
    with ZipFile(BytesIO(outputs)) as archive:
        assert {row['input'] for row in pdf_report['records']} == set(hidden)
        assert set(archive.namelist()) == {row['output'] for row in pdf_report['records']}
        for row in pdf_report['records']:
            raw = archive.read(row['output'])
            assert digest(raw) == row['sha256']
            with pymupdf.open(stream=raw, filetype='pdf') as pdf:
                assert row['import_visible'] is any('IMPORT' in page.get_text() for page in pdf)
                if row['import_visible']:
                    assert import_pixels_painted(pdf)
            if row['import_visible'] is hidden[row['input']]['value']:
                differences.append(row['input'])
    assert sorted(differences) == ['hidden-1-0-1-n.docx', 'hidden-1-1-0-n.docx']
    return len(generated)


def verify_hidden_style_contexts(root):
    report = json.loads((root / 'hidden-style-contexts-26.9.json').read_text())
    assert report['version'] == '26.9.0' and report['licensed'] is False
    assert report['rendering_acceptance'] is False
    generated = dict(runpy.run_path(str(Path(__file__).parents[1] / 'docs/probes/hidden_style_contexts.py'))['inputs']())
    source = (root / report['corpus']).read_bytes()
    outputs = (root / report['outputs']).read_bytes()
    assert digest(source) == report['corpus_sha256'] and digest(outputs) == report['outputs_sha256']
    assert len(generated) == len(report['records']) == 32
    import pymupdf
    with ZipFile(BytesIO(source)) as inputs, ZipFile(BytesIO(outputs)) as pdfs:
        assert set(inputs.namelist()) == set(generated) == {r['input'] for r in report['records']}
        assert set(pdfs.namelist()) == {r['output'] for r in report['records']}
        for row in report['records']:
            raw = inputs.read(row['input'])
            assert digest(raw) == row['source_sha256']
            with ZipFile(BytesIO(raw)) as actual, ZipFile(BytesIO(generated[row['input']])) as rebuilt:
                assert {n: actual.read(n) for n in actual.namelist()} == {n: rebuilt.read(n) for n in rebuilt.namelist()}
            prefix, _, default, paragraph, character, reference = row['input'][:-5].split('-')
            default, paragraph, character, reference = [v == '1' for v in (default, paragraph, character, reference)]
            expected = (default or paragraph ^ character) if reference else paragraph
            rendered = (default ^ paragraph ^ character if prefix == 'explicit' else character) if reference else paragraph
            assert row['hidden'] is expected and row['visible'] is not rendered
            raw = pdfs.read(row['output'])
            assert digest(raw) == row['output_sha256']
            with pymupdf.open(stream=raw, filetype='pdf') as pdf:
                assert row['visible'] is any('IMPORT' in page.get_text() for page in pdf)
                if row['visible']:
                    assert import_pixels_painted(pdf)
    return len(generated)


def font_size_parts(raw):
    from xml.etree import ElementTree as ET

    if raw.startswith(b"PK"):
        with ZipFile(BytesIO(raw)) as package:
            return tuple(ET.fromstring(package.read("word/" + name + ".xml")) for name in ("document", "styles"))
    package = "{http://schemas.microsoft.com/office/2006/xmlPackage}"
    parts = {part.get(package + "name"): part for part in ET.fromstring(raw)}
    return tuple(parts["/word/" + name + ".xml"].find(package + "xmlData")[0] for name in ("document", "styles"))


def saved_hidden_state(raw, tag="vanish"):
    """Toggle getter/declarations; the visibility component is only valid for vanish."""
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    document, styles = font_size_parts(raw)
    paragraph = next(p for p in document.iter(w + 'p') if ''.join(t.text or '' for t in p.iter(w + 't')) == 'IMPORT')
    run = next(r for r in paragraph.iter(w + 'r') if ''.join(t.text or '' for t in r.iter(w + 't')) == 'IMPORT')
    by_id = {s.get(w + 'styleId'): s for s in styles.findall(w + 'style')}
    def value(element):
        return None if element is None else element.get(w + 'val', '1') in {'1', 'true', 'on'}
    def nearest(identifier):
        seen = set()
        while identifier in by_id:
            assert identifier not in seen
            seen.add(identifier)
            style = by_id[identifier]
            declared = value(style.find(w + 'rPr/' + w + tag))
            if declared is not None:
                return declared
            parent = style.find(w + 'basedOn')
            identifier = parent.get(w + 'val') if parent is not None else None
        return None
    reference = paragraph.find(w + 'pPr/' + w + 'pStyle')
    identifier = reference.get(w + 'val') if reference is not None else next(
        (s.get(w + 'styleId') for s in reversed(styles.findall(w + 'style'))
         if s.get(w + 'type') == 'paragraph' and s.get(w + 'default') in {'1', 'true', 'on'}), None)
    character = run.find(w + 'rPr/' + w + 'rStyle')
    p = nearest(identifier)
    c = nearest(character.get(w + 'val')) if character is not None else None
    default = bool(value(styles.find(w + 'docDefaults/' + w + 'rPrDefault/' + w + 'rPr/' + w + tag)))
    getter = default or p ^ c if p is not None and c is not None else p if p is not None else c if c is not None else default
    rendered = (c if c is not None else p if p is not None else default) if reference is None else (
        default ^ p ^ c if p is not None and c is not None else getter)
    direct = value(run.find(w + 'rPr/' + w + tag))
    if direct is not None:
        getter = rendered = direct
    return getter, not rendered, direct, reference.get(w + 'val') if reference is not None else None


def verify_dom_font_boolean_errors(root):
    report = json.loads((root / 'dom-font-boolean-errors-26.9.json').read_text())
    assert report['version'] == '26.9.0' and report['licensed'] is False
    assert report['full_Font_acceptance'] is report['rendering_acceptance'] is False
    fields = {r['field'] for r in json.loads((root / report['source']).read_text())['records']}
    values = (None, 0, 1, -1, 1.5, 'false', '', [], {})
    expected = {(field, target, json.dumps(value)) for field in fields
                for target in ('run', 'paragraph', 'character') for value in values}
    native, before, after = (report[key] for key in ('native', 'current_before_error_fix', 'current_after_error_fix'))
    assert len(native) == len(before) == len(after) == len(expected) == 297
    assert {(r['field'], r['target'], json.dumps(r['input'])) for r in native} == expected
    type_differences = none_differences = 0
    for actual, old, current in zip(native, before, after):
        identity = ('field', 'target', 'input')
        assert all(actual[k] == old[k] == current[k] for k in identity)
        assert actual['before'] is actual['after'] is old['before'] is old['after'] is current['before'] is current['after'] is True
        assert actual['error'] == 'TypeError'
        cleared = actual['input'] is None and actual['target'] == 'run'
        assert current['error'] == (None if cleared else 'TypeError')
        assert old['error'] == (None if cleared else 'TypeError' if actual['input'] is None else 'ValueError')
        none_differences += cleared
        type_differences += actual['input'] is not None
    assert type_differences == report['non_bool_error_class_differences'] == 264
    assert none_differences == report['direct_None_extension_differences'] == 11
    assert report['non_bool_error_class_differences_after_fix'] == 0
    return len(native)


def verify_dom_font_booleans(root):
    report = json.loads((root / 'dom-font-booleans-26.9.json').read_text())
    assert report['full_Font_acceptance'] is report['rendering_acceptance'] is False
    assert report['version'] == '26.9.0' and report['licensed'] is False
    source = json.loads((root / report['source']).read_text())
    sources = {r['input']: r for r in source['records']}
    data = (root / report['observations']).read_bytes()
    assert digest(data) == report['observations_sha256']
    tags = runpy.run_path(str(Path(__file__).parents[1] / 'docs/probes/font_boolean_contexts.py'))['FIELDS']
    expected = {(name, target, value) for name in sources for target in ('run', 'paragraph', 'character')
                for value in (False, True)}
    with ZipFile(BytesIO(data)) as archive:
        live = json.loads(archive.read('live.json'))
        saves = json.loads(archive.read('saves.json'))
        cold = json.loads(archive.read('cold.json'))
        assert live['version'] == cold['version'] == '26.9.0'
        assert live['licensed'] is cold['licensed'] is False
        assert live['full_Font_acceptance'] is live['rendering_acceptance'] is False
        assert len(live['records']) == len(saves['records']) == len(expected) == report['live_edits'] == 5346
        assert {(r['input'], r['target'], r['value']) for r in live['records']} == expected
        assert len(cold['records']) == report['saved_outputs'] == 10692
        used, index = {'live.json', 'saves.json', 'cold.json'}, 0
        for native, current in zip(live['records'], saves['records']):
            assert native == {k: v for k, v in current.items() if k != 'outputs'}
            original = sources[native['input']]
            assert native['field'] == original['field'] and native['source_sha256'] == original['sha256']
            default, paragraph, character, direct = original['values']
            assert native['before'] == {'run': original['value'],
                                       'paragraph': bool(default if paragraph is None else paragraph),
                                       'character': bool(default if character is None else character)}
            assert native['after'][native['target']] is native['value']
            assert {r['format'] for r in current['outputs']} == {'docx', 'flat_opc'}
            assert len(current['outputs']) == 2
            for saved in current['outputs']:
                observed = cold['records'][index]
                index += 1
                assert saved == {k: v for k, v in observed.items() if k != 'observed'}
                raw = archive.read(saved['storage'])
                used.add(saved['storage'])
                assert digest(raw) == saved['sha256']
                assert raw.startswith(b'PK') is (saved['format'] == 'docx')
                assert saved['current_cold'] == observed['observed'] == native['after']
                tag = tags[native['field']]
                state = saved_hidden_state(raw, tag)
                assert state[0] is observed['observed']['run']
                assert state[2] is (native['value'] if native['target'] == 'run' else direct)
                assert state[3] == 'P'
                _, styles = font_size_parts(raw)
                w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
                for target, identifier, declaration in (('paragraph', 'P', paragraph), ('character', 'C', character)):
                    style = next(s for s in styles.findall(w + 'style') if s.get(w + 'styleId') == identifier)
                    element = style.find(w + 'rPr/' + w + tag)
                    actual = None if element is None else element.get(w + 'val', '1') in {'1', 'true', 'on'}
                    assert actual is (native['value'] if native['target'] == target else declaration)
                    assert observed['observed'][target] is bool(default if actual is None else actual)
        assert set(archive.namelist()) == used
    assert len(sources) == report['input_count'] == 891 and report['native_cold_getter_differences'] == 0
    return index


def verify_hidden_font_roundtrip(root):
    report = json.loads((root / 'hidden-font-roundtrip-26.9.json').read_text())
    inputs = {name: (raw, hidden, visible) for name, raw, hidden, visible in
              runpy.run_path(str(Path(__file__).parents[1] / 'docs/probes/hidden_font_roundtrip.py'))['inputs'](root)}
    assert len(inputs) == 113 and report['full_format_acceptance'] is report['rendering_acceptance'] is False
    expected = {(name, phase, fmt) for name in inputs for phase in ('direct', 'json') for fmt in ('docx', 'flat_opc')}
    differences = {}
    for phase, current in (('before', report['before_origin_fix']), ('after', report)):
        assert len(current['records']) == 452
        assert {(r['input'], r['phase'], r['format']) for r in current['records']} == expected
        outputs = (root / current['outputs']).read_bytes()
        assert digest(outputs) == current['outputs_sha256']
        native = current['native_reread']
        assert native['version'] == '26.9.0' and native['licensed'] is False
        assert native['visible_target_pixels_checked'] is True
        observed = {r['output']: r for r in native['records']}
        assert len(observed) == len(native['records']) == 452
        getter_errors = visible_errors = 0
        with ZipFile(BytesIO(outputs)) as archive:
            assert set(archive.namelist()) == set(observed) == {r['output'] for r in current['records']}
            for row in current['records']:
                source, hidden, visible = inputs[row['input']]
                assert digest(source) == row['source_sha256']
                assert row['expected_hidden'] is hidden and row['expected_visible'] is visible
                raw = archive.read(row['output'])
                assert digest(raw) == row['output_sha256']
                actual = saved_hidden_state(raw)
                recorded = observed[row['output']]
                assert recorded['hidden'] is actual[0] and recorded['visible'] is actual[1]
                getter_errors += actual[0] is not hidden
                visible_errors += actual[1] is not visible
                if phase == 'after':
                    assert actual[2:] == saved_hidden_state(source)[2:]
        differences[phase] = (getter_errors, visible_errors)
    assert differences == {'before': (20, 52), 'after': (0, 0)}
    return len(report['records'])


def verify_font_boolean_roundtrip(root):
    report = json.loads((root / 'font-boolean-roundtrip-26.9.json').read_text())
    rows = json.loads((root / 'font-boolean-contexts-26.9.json').read_text())
    sources = {row['input']: row for row in rows['records'] if row['field'] != 'hidden'}
    tags = runpy.run_path(str(Path(__file__).parents[1] / 'docs/probes/font_boolean_contexts.py'))['FIELDS']
    assert len(sources) == 810 and report['full_format_acceptance'] is report['rendering_acceptance'] is False
    expected = {(name, phase, fmt) for name in sources for phase in ('direct', 'json') for fmt in ('docx', 'flat_opc')}
    differences = {}
    with ZipFile(root / rows['corpus']) as inputs:
        for phase, current in (('before', report['before_origin_fix']), ('after', report)):
            assert len(current['records']) == 3240
            assert {(r['input'], r['phase'], r['format']) for r in current['records']} == expected
            raw = (root / current['outputs']).read_bytes()
            assert digest(raw) == current['outputs_sha256']
            native = current['native_reread']
            assert native['version'] == '26.9.0' and native['licensed'] is False
            observed = {r['output']: r['value'] for r in native['records']}
            assert len(observed) == len(native['records']) == 3240
            errors = []
            with ZipFile(BytesIO(raw)) as outputs:
                assert set(observed) == {r['output'] for r in current['records']}
                assert set(outputs.namelist()) == {r.get('storage', r['output']) for r in current['records']}
                for row in current['records']:
                    source = sources[row['input']]
                    assert row['field'] == source['field'] and row['expected'] is source['value']
                    original = inputs.read(row['input'])
                    assert digest(original) == row['source_sha256'] == source['sha256']
                    saved = outputs.read(row.get('storage', row['output']))
                    assert digest(saved) == row['output_sha256']
                    state = saved_hidden_state(saved, tags[row['field']])
                    assert state[0] is observed[row['output']]
                    if state[0] is not row['expected']:
                        errors.append(row['field'])
                    if phase == 'after':
                        assert state[2:] == saved_hidden_state(original, tags[row['field']])[2:]
            differences[phase] = {field: errors.count(field) for field in tags if field != 'hidden'}
    assert differences == {'before': {f: 20 for f in tags if f != 'hidden'},
                           'after': {f: 0 for f in tags if f != 'hidden'}}
    optimization = report['storage_optimization']
    assert len(optimization['records']) == 2
    for metric, (phase, current) in zip(optimization['records'], (('before', report['before_origin_fix']), ('after', report))):
        assert metric['phase'] == phase and metric['logical_outputs'] == len(current['records']) == 3240
        assert metric['all_logical_output_bytes_unchanged'] is True
        assert metric['after_archive_sha256'] == current['outputs_sha256']
        assert metric['after_archive_bytes'] == (root / current['outputs']).stat().st_size < metric['before_archive_bytes']
        assert metric['stored_outputs'] == len({r['storage'] for r in current['records']}) < 3240
    return len(report['records'])


def saved_font_sizes(raw):
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    document, styles = font_size_parts(raw)
    paragraph = next(p for p in document.iter(w + "p") if "".join(t.text or "" for t in p.iter(w + "t")) == "IMPORT")
    run = next(r for r in paragraph.iter(w + "r") if "".join(t.text or "" for t in r.iter(w + "t")) == "IMPORT")
    def size(element):
        return None if element is None else int(element.get(w + "val")) / 2
    resolved = size(styles.find(w + "docDefaults/" + w + "rPrDefault/" + w + "rPr/" + w + "sz"))
    if resolved is None:
        resolved = 10 if styles.find(w + "docDefaults/" + w + "rPrDefault") is not None else 11
    chain, identifier = [], "Derived"
    by_id = {style.get(w + "styleId"): style for style in styles.findall(w + "style")}
    while identifier:
        assert identifier not in chain and identifier in by_id
        chain.append(identifier)
        base = by_id[identifier].find(w + "basedOn")
        identifier = base.get(w + "val") if base is not None else None
    for identifier in reversed(chain):
        direct = size(by_id[identifier].find(w + "rPr/" + w + "sz"))
        if direct is not None:
            resolved = direct
    style_size = resolved
    direct = size(run.find(w + "rPr/" + w + "sz"))
    if direct is not None:
        resolved = direct
    return style_size, resolved


def saved_font_toggles(raw):
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    document, styles = font_size_parts(raw)
    paragraph = next(p for p in document.iter(w + "p") if "".join(t.text or "" for t in p.iter(w + "t")) == "IMPORT")
    run = next(r for r in paragraph.iter(w + "r") if "".join(t.text or "" for t in r.iter(w + "t")) == "IMPORT")
    by_id = {s.get(w + "styleId"): s for s in styles.findall(w + "style")}
    def onoff(element):
        return None if element is None else element.get(w + "val", "1") not in ("0", "false", "off")
    def nearest(identifier, tag):
        visited = set()
        while identifier:
            assert identifier not in visited and identifier in by_id
            visited.add(identifier)
            style = by_id[identifier]
            value = onoff(style.find(w + "rPr/" + w + tag))
            if value is not None:
                return value
            parent = style.find(w + "basedOn")
            identifier = parent.get(w + "val") if parent is not None else None
        return None
    pstyle = paragraph.find(w + "pPr/" + w + "pStyle")
    rstyle = run.find(w + "rPr/" + w + "rStyle")
    paragraph_id = pstyle.get(w + "val") if pstyle is not None else next(
        (s.get(w + "styleId") for s in reversed(list(by_id.values()))
         if s.get(w + "type") == "paragraph" and s.get(w + "default") == "1"), None)
    result = {}
    for tag, field in (("b", "bold"), ("i", "italic")):
        default = onoff(styles.find(w + "docDefaults/" + w + "rPrDefault/" + w + "rPr/" + w + tag)) or False
        para = nearest(paragraph_id, tag)
        char = nearest(rstyle.get(w + "val") if rstyle is not None else None, tag)
        value = default or (para != char) if para is not None and char is not None else (
            para if para is not None else char if char is not None else default)
        direct = onoff(run.find(w + "rPr/" + w + tag))
        result[field] = value if direct is None else direct
    return result


def font_toggle_declarations(raw):
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    document, styles = font_size_parts(raw)
    run = next(r for r in document.iter(w + "r") if "".join(t.text or "" for t in r.iter(w + "t")) == "IMPORT")
    def flags(element):
        return tuple(None if (value := element.find(w + "rPr/" + w + tag)) is None
                     else value.get(w + "val", "1") not in ("0", "false", "off") for tag in ("b", "i"))
    return flags(run), {s.get(w + "styleId"): flags(s) for s in styles.findall(w + "style")
                        if s.get(w + "type") == "character"}


def verify_font_json_origin(root):
    report = json.loads((root / "font-json-origin-26.9.json").read_text())
    source = json.loads((root / report["source_report"]).read_text())
    assert report["source_corpus_sha256"] == source["corpus_sha256"]
    original = {r["input"]: r for r in source["records"]}
    assert len(original) == len(source["records"]) == 745
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is False
    source_raw = (root / "corpus/style-toggles-26.9.zip").read_bytes()
    assert digest(source_raw) == source["corpus_sha256"]
    with ZipFile(BytesIO(source_raw)) as corpus:
        declarations = {}
        for name, row in original.items():
            data = corpus.read(name)
            assert digest(data) == row["sha256"]
            declarations[name] = font_toggle_declarations(data)
    for phase in (report["before_cascade_fix"], report):
        rows = {(r["input"], r["format"]): r for r in phase["records"]}
        assert len(rows) == len(phase["records"]) == 1490
        assert set(rows) == {(name, fmt) for name in original for fmt in ("docx", "flat_opc")}
        native = phase["native_reread"]
        assert native["version"] == "26.9.0" and native["licensed"] is False
        observed = {r["output"]: r for r in native["records"]}
        assert len(observed) == len(native["records"]) == 1490
        raw = (root / phase["outputs"]).read_bytes()
        assert digest(raw) == phase["outputs_sha256"]
        mismatches = []
        with ZipFile(BytesIO(raw)) as archive:
            assert len(archive.namelist()) == 1490 and set(archive.namelist()) == set(observed) == {r["output"] for r in rows.values()}
            for (name, fmt), row in rows.items():
                assert row["source_sha256"] == original[name]["sha256"]
                assert row["output"] == name.removesuffix(".docx") + "/" + fmt + (".docx" if fmt == "docx" else ".xml")
                data = archive.read(row["output"])
                assert digest(data) == row["output_sha256"] and data.startswith(b"PK") is (fmt == "docx")
                actual = saved_font_toggles(data)
                if phase is report:
                    assert font_toggle_declarations(data) == declarations[name]
                assert all(type(actual[field]) is bool and actual[field] == observed[row["output"]][field]
                           for field in ("bold", "italic"))
                assert all(row["expected_" + field] is original[name][field] for field in ("bold", "italic"))
                if any(actual[field] is not original[name][field] for field in ("bold", "italic")):
                    mismatches.append(row["output"])
        if phase is report:
            assert mismatches == []
        else:
            assert len(mismatches) == 194 and sorted(mismatches) == sorted(phase["mismatched_outputs"])
    return len(rows)


def verify_font_default_presence(root):
    report = json.loads((root / "font-default-presence-26.9.json").read_text())
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is False
    sources = ("font-defaults-26.9.json", "font-default-matrix-26.9.json", "font-size-loading-26.9.json")
    assert report["sources"] == list(sources)
    originals = {}
    for filename in sources:
        source = json.loads((root / filename).read_text())
        raw = (root / source["corpus"]).read_bytes()
        assert digest(raw) == source["corpus_sha256"]
        with ZipFile(BytesIO(raw)) as archive:
            for row in source["records"]:
                expected = row.get("loaded", row)
                if "run_size" in expected and 0 < expected["run_size"] < 1000:
                    data = archive.read(row["input"])
                    assert digest(data) == row["sha256"]
                    originals[filename, row["input"]] = row, expected, data
    assert len(originals) == 87
    rows = {(r["source_report"], r["input"], r["format"], r["json_roundtrip"]): r for r in report["records"]}
    assert len(rows) == len(report["records"]) == 348
    assert set(rows) == {(filename, name, fmt, use_json) for filename, name in originals
                         for fmt in ("docx", "flat_opc") for use_json in (False, True)}
    native = report["native_reread"]
    assert native["version"] == "26.9.0" and native["licensed"] is False
    observed = {r["output"]: r for r in native["records"]}
    assert len(observed) == len(native["records"]) == 348
    raw = (root / report["outputs"]).read_bytes()
    assert digest(raw) == report["outputs_sha256"]
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    def declarations(data):
        document, styles = font_size_parts(data)
        paragraph = next(p for p in document.iter(w + "p") if "".join(t.text or "" for t in p.iter(w + "t")) == "IMPORT")
        run = next(r for r in paragraph.iter(w + "r") if "".join(t.text or "" for t in r.iter(w + "t")) == "IMPORT")
        derived = next(s for s in styles.findall(w + "style") if s.get(w + "styleId") == "Derived")
        return (styles.find(w + "docDefaults/" + w + "rPrDefault") is not None,
                run.find(w + "rPr/" + w + "sz") is not None, derived.find(w + "rPr/" + w + "sz") is not None)
    with ZipFile(BytesIO(raw)) as archive:
        assert len(archive.namelist()) == 348 and set(archive.namelist()) == set(observed) == {r["output"] for r in rows.values()}
        for key, row in rows.items():
            source, expected, original = originals[key[:2]]
            filename, name, fmt, use_json = key
            assert type(use_json) is bool
            assert row["output"] == (filename.removesuffix(".json") + "/" + name.removesuffix(".docx") +
                                      f"/{fmt}-{int(use_json)}." + ("docx" if fmt == "docx" else "xml"))
            assert row["source_sha256"] == source["sha256"]
            data = archive.read(row["output"])
            assert data.startswith(b"PK") is (fmt == "docx")
            assert digest(data) == row["output_sha256"]
            state = declarations(original)
            assert declarations(data) == state and row["rpr_default_present"] is state[0]
            style_size, run_size = saved_font_sizes(data)
            assert style_size == expected["style_size"] == row["expected_style_size"] == observed[row["output"]]["style_size"]
            assert run_size == expected["run_size"] == row["expected_run_size"] == observed[row["output"]]["run_size"]
    return len(rows)


def verify_font_size_loading(root):
    report = json.loads((root / "font-size-loading-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is False
    generated = {name: (scope, value, data) for name, scope, value, data in
                 runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/font_size_loading.py"))["inputs"]()}
    source, outputs = ((root / report[key]).read_bytes() for key in ("corpus", "outputs"))
    assert digest(source) == report["corpus_sha256"] and digest(outputs) == report["outputs_sha256"]
    rows = {row["input"]: row for row in report["records"]}
    assert len(rows) == len(report["records"]) == len(generated) == 126 and set(rows) == set(generated)
    measured = (12, 12, 12.5, 12, None, 50, 12, 6, 36, 72, 72, 12, 12, None, None,
                0, -1, 0, 0, 0, 0, 0, 0, -.5, -.5, 12, None, .5, 11.5, 6, 6, 0,
                None, -1, 12, None, 12, 12.5, 11, 0, 108, None)
    with ZipFile(BytesIO(source)) as corpus, ZipFile(BytesIO(outputs)) as saved:
        assert len(corpus.namelist()) == 126 and set(corpus.namelist()) == set(rows)
        expected_outputs = {row["output"] for row in rows.values() if "output" in row}
        assert len(saved.namelist()) == len(expected_outputs) == 105 and set(saved.namelist()) == expected_outputs
        for name, row in rows.items():
            scope, value, expected_source = generated[name]
            data = corpus.read(name)
            assert row["scope"] == scope and row["value"] == value and digest(data) == row["sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(expected_source)) as expected:
                assert {key: actual.read(key) for key in actual.namelist()} == {key: expected.read(key) for key in expected.namelist()}
            expected_size = measured[int(name[-7:-5])]
            if expected_size is None:
                assert row["load_error"] == "RuntimeError" and "output" not in row and "loaded" not in row
                continue
            assert "load_error" not in row and "save_error" not in row and row["output"] == name
            raw = saved.read(name)
            assert digest(raw) == row["output_sha256"]
            style_size, resolved = saved_font_sizes(raw)
            assert resolved == expected_size == row["loaded"]["run_size"] == row["after_save"]["run_size"] == row["cold"]["run_size"]
            assert style_size == row["loaded"]["style_size"] == row["after_save"]["style_size"] == row["cold"]["style_size"]
    current = report["docweave_roundtrip"]
    assert current["full_acceptance"] is False and current["native_reread"]["version"] == "26.9.0"
    assert current["native_reread"]["licensed"] is False
    raw = (root / current["outputs"]).read_bytes()
    assert digest(raw) == current["outputs_sha256"]
    positive = {name for name, row in rows.items() if "loaded" in row and 0 < row["loaded"]["run_size"] < 1000}
    records = {row["input"]: row for row in current["records"]}
    native = {row["input"]: row for row in current["native_reread"]["records"]}
    assert len(records) == len(current["records"]) == len(native) == len(current["native_reread"]["records"]) == 66
    assert set(records) == set(native) == positive
    with ZipFile(BytesIO(raw)) as archive:
        assert len(archive.namelist()) == 66 and set(archive.namelist()) == positive
        for name in positive:
            data = archive.read(name)
            assert digest(data) == records[name]["output_sha256"]
            style_size, run_size = saved_font_sizes(data)
            assert run_size == native[name]["run_size"] == rows[name]["loaded"]["run_size"]
            assert style_size == native[name]["style_size"]
    return len(rows)


def verify_font_defaults(root):
    report = json.loads((root / "font-defaults-26.9.json").read_text())
    raw = (root / report["corpus"]).read_bytes()
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["full_format_acceptance"] is False and digest(raw) == report["corpus_sha256"]
    generated = dict(runpy.run_path(str(Path(__file__).parents[1] / "docs" / "probes" / "font_defaults.py"))["inputs"]())
    rows = {row["input"]: row for row in report["records"]}
    assert len(rows) == len(report["records"]) == len(generated) == 5 and set(rows) == set(generated)
    with ZipFile(BytesIO(raw)) as archive:
        assert len(archive.namelist()) == 5 and set(archive.namelist()) == set(rows)
        for name, row in rows.items():
            data = archive.read(name)
            assert digest(data) == row["sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(generated[name])) as expected:
                assert {key: actual.read(key) for key in actual.namelist()} == {key: expected.read(key) for key in expected.namelist()}
            assert row["run_size"] == row["style_size"] == (10.0 if name.startswith("rpr_") else 11.0)
    return len(rows)


def verify_font_default_matrix(root):
    report = json.loads((root / "font-default-matrix-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is False
    raw = (root / report["corpus"]).read_bytes()
    assert digest(raw) == report["corpus_sha256"]
    generated = dict(runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/font_default_matrix.py"))["inputs"]())
    rows = {row["input"]: row for row in report["records"]}
    assert len(rows) == len(report["records"]) == len(generated) == 16
    with ZipFile(BytesIO(raw)) as archive:
        assert len(archive.namelist()) == 16 and set(archive.namelist()) == set(rows) == set(generated)
        for name, row in rows.items():
            data = archive.read(name)
            assert digest(data) == row["sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(generated[name])) as expected:
                assert {key: actual.read(key) for key in actual.namelist()} == {key: expected.read(key) for key in expected.namelist()}
            size = 12.0 if name.endswith("-run_size.docx") else 10.0 if "-run_" in name else 11.0
            assert row["run_size"] == row["style_size"] == size
    return len(rows)


def verify_style_imports(root, name="style-import-conflicts-26.9.json"):
    report = json.loads((root / name).read_text())
    assert set(report["reports"]) in ({"commercial", "before", "candidate"}, {"candidate"})
    assert set(report["official_candidate_rereads"]) == set(report["reports"]) - {"commercial"}
    corpus = (root / report["corpus"]["archive"]).read_bytes()
    raw_outputs = (root / report["outputs"]["archive"]).read_bytes()
    assert digest(corpus) == report["corpus"]["sha256"], "style import corpus digest mismatch"
    assert digest(raw_outputs) == report["outputs"]["sha256"], "style import outputs digest mismatch"
    probes = Path(__file__).parents[1] / "docs" / "probes"
    generator = report["corpus"].get("generator", "import_style_conflicts.py")
    counts = {"import_style_conflicts.py": 247, "paragraph_style_defaults.py": 132}
    assert generator in counts
    generate = runpy.run_path(str(probes / generator))["inputs"]
    expected = {key + "/" + phase + ".docx" for key, phase, _ in generate()}
    keys = {name.split("/")[0] for name in expected}
    inspect = runpy.run_path(str(probes / "inspect_style_imports.py"))["inspect_document"]
    assert len(expected) == 2 * len(keys) and len(keys) == report["corpus"]["input_pairs"] == counts[generator]
    checked, phases, formats = 0, {}, {}
    with ZipFile(BytesIO(corpus)) as inputs, ZipFile(BytesIO(raw_outputs)) as outputs:
        assert len(inputs.namelist()) == len(expected) and set(inputs.namelist()) == expected
        assert len(outputs.namelist()) == report["outputs"]["files"] == len(keys) * len(report["reports"])
        for phase in report["reports"]:
            observed = report["reports"][phase]
            assert observed["corpus_sha256"] == report["corpus"]["sha256"]
            rows = {row["case"]: row for row in observed["records"]}
            assert len(rows) == len(observed["records"]) == len(keys) and set(rows) == keys
            phases[phase] = rows
            recorded = {row["case"]: row for row in report["independent_checks"][phase]}
            assert len(recorded) == len(keys) and set(recorded) == keys
            formats[phase] = {}
            for key, row in rows.items():
                for source in ("source", "destination"):
                    assert digest(inputs.read(key + "/" + source + ".docx")) == row["inputs"][source]
                data = outputs.read(phase + "/" + row["output"])
                assert digest(data) == row["output_sha256"]
                actual = {"case": key, "output": row["output"], **inspect(BytesIO(data))}
                assert actual == recorded[key], "style import independent observation mismatch"
                if "story_rereads" in report:
                    formats[phase][key] = saved_story_formats(data)
                if row["outcome"] == "raised":
                    assert phase != "commercial" and row["destination_unchanged"] is True
                else:
                    assert row["outcome"] == "returned" and "reopened_format" in row
                checked += 1
    if "commercial" in phases:
        native = phases["commercial"]
    else:
        baseline = (root / report["baseline"]["file"]).read_bytes()
        assert digest(baseline.replace(b"\r\n", b"\n")) == report["baseline"]["sha256"]
        native = {row["case"]: row for row in json.loads(baseline)["reports"]["commercial"]["records"]}
    for phase in report["official_candidate_rereads"]:
        read = report["official_candidate_rereads"][phase]
        assert read["version"] == "26.9.0" and read["licensed"] is False
        observed = {row["case"]: row for row in read["records"]}
        returned = {key for key, row in phases[phase].items() if row["outcome"] == "returned"}
        assert len(observed) == len(read["records"]) and set(observed) == returned
        assert all(row["output_sha256"] == phases[phase][key]["output_sha256"] for key, row in observed.items())
        mismatches = sorted(key for key, row in observed.items() if row["format"] != native[key]["reopened_format"])
        assert mismatches == sorted(report["reread_mismatches"][phase])
        if phase == "candidate":
            assert not mismatches
    assert sorted(report["native_source_changes"]) == sorted(key for key, row in native.items() if row["source_format"] != row["imported_format"])
    assert sorted(report["native_roundtrip_changes"]) == sorted(key for key, row in native.items() if row["imported_format"] != row["reopened_format"])
    repeated = {row["case"]: row for row in report["repeat_without_source_getter"]}
    assert len(repeated) == len(report["repeat_without_source_getter"]) and set(repeated) == set(report["native_roundtrip_changes"])
    for key, row in repeated.items():
        assert row["no_source_getter_before_import"] is True
        if "imported_format" in row:
            assert row["imported_format"] == native[key]["imported_format"]
            assert row["reopened_format"] == native[key]["reopened_format"]
        else:
            assert row["imported_bold"] == native[key]["imported_format"]["bold"]
            assert row["reopened_bold"] == native[key]["reopened_format"]["bold"]
    if "repeat_outputs" in report:
        raw = (root / report["repeat_outputs"]["archive"]).read_bytes()
        assert digest(raw) == report["repeat_outputs"]["sha256"]
        with ZipFile(BytesIO(raw)) as archive:
            assert len(archive.namelist()) == report["repeat_outputs"]["files"] == len(repeated)
            for row in repeated.values():
                data = archive.read(row["output"])
                assert digest(data) == row["output_sha256"]
                assert saved_story_formats(data)["IMPORT"] == row["reopened_format"]
    if "story_rereads" in report:
        previous = phases.get("before")
        if previous is None:
            path = root / report["before_baseline"]["file"]
            raw = path.read_bytes()
            assert digest(raw.replace(b"\r\n", b"\n")) == report["before_baseline"]["sha256"]
            previous = {row["case"]: row for row in json.loads(raw)["reports"]["candidate"]["records"]}
        source_rows = {"commercial": native, "before": previous, "candidate": phases["candidate"]}
        observations = {}
        assert set(report["story_rereads"]) == set(source_rows)
        for phase, observed in report["story_rereads"].items():
            assert observed["version"] == "26.9.0" and observed["licensed"] is False
            rows = {row["case"]: row for row in observed["records"]}
            assert len(rows) == len(observed["records"]) == len(keys) and set(rows) == keys
            observations[phase] = rows
            for key, row in rows.items():
                original = source_rows[phase][key]
                assert row["output_sha256"] == original["output_sha256"]
                assert set(row["formats"]) == ({"IMPORT", "DESTINATION"} if original["outcome"] == "returned" else {"DESTINATION"})
                if phase in formats:
                    assert row["formats"] == formats[phase][key], "saved story observation mismatch"
        for phase in ("before", "candidate"):
            mismatches = sorted(key for key, row in source_rows[phase].items() if row["outcome"] == "returned"
                                and observations[phase][key]["formats"] != observations["commercial"][key]["formats"])
            assert mismatches == sorted(report["story_mismatches"][phase])
        assert not report["story_mismatches"]["candidate"]
    assert report["validation"]["full_import_acceptance"] is False and report["validation"]["rendering_acceptance"] is False
    if "style_state_baseline" in report:
        raw = (root / report["style_state_baseline"]["file"]).read_bytes()
        assert digest(raw.replace(b"\r\n", b"\n")) == report["style_state_baseline"]["sha256"]
        native_styles = {row["case"]: row["after_reopen"]["styles"] for row in json.loads(raw)["reports"][0]["records"]}
        with ZipFile(BytesIO(raw_outputs)) as outputs:
            for key, row in phases["candidate"].items():
                if row["outcome"] == "returned":
                    assert saved_style_fonts(outputs.read("candidate/" + row["output"])) == native_styles[key], "saved style font mismatch"
    return checked


def saved_story_formats(data):
    from aspose.words_foss import DocxDocument
    from aspose.words_foss.dom.styles import _style_stories

    document = DocxDocument(BytesIO(data))
    paragraphs = []
    for root in _style_stories(document._package):
        part = next(name for name, tree in document._package._trees.items() if tree.documentElement is root)
        paragraphs.extend((document.body if part == "word/document.xml" else document.story(part)).paragraphs)
    result = {}
    for paragraph in paragraphs:
        if paragraph.text not in {"IMPORT", "DESTINATION"}:
            continue
        assert paragraph.text not in result
        font = paragraph.runs[0].effective_font
        alignment = paragraph.effective_paragraph_format.alignment
        result[paragraph.text] = {"bold": font.bold, "italic": font.italic, "size": font.size,
                                  "alignment": {"both": "JUSTIFY"}.get(alignment, (alignment or "left").upper())}
    return result


def verify_style_roundtrips(root):
    report = json.loads((root / "style-save-roundtrips-26.9.json").read_text())
    observed = report["report"]
    assert observed["version"] == "26.9.0" and observed["licensed"] is False
    corpus = (root / report["corpus"]["archive"]).read_bytes()
    raw = (root / report["outputs"]["archive"]).read_bytes()
    assert digest(corpus) == report["corpus"]["sha256"] == observed["corpus_sha256"]
    assert digest(raw) == report["outputs"]["sha256"]
    rows = {(row["case"], row["phase"]): row for row in observed["records"]}
    with ZipFile(BytesIO(corpus)) as inputs, ZipFile(BytesIO(raw)) as outputs:
        assert len(rows) == len(observed["records"]) == len(inputs.namelist()) == len(outputs.namelist()) == 494
        assert {key + "/" + phase + ".docx" for key, phase in rows} == set(inputs.namelist())
        for (key, phase), row in rows.items():
            data = inputs.read(key + "/" + phase + ".docx")
            output = outputs.read(row["output"])
            label = "IMPORT" if phase == "source" else "DESTINATION"
            assert digest(data) == row["input_sha256"] and digest(output) == row["output_sha256"]
            assert saved_story_formats(data) == {label: row["before"]}
            assert saved_story_formats(output) == {label: row["after"]}
    assert report["changes"] == sorted(key + "/" + phase for (key, phase), row in rows.items() if row["before"] != row["after"])
    assert report["full_format_acceptance"] is False and report["rendering_acceptance"] is False
    return len(rows)


def saved_style_fonts(data):
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(BytesIO(data))
    defaults = document.styles.element.find(qn("w:docDefaults"))
    properties = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    values = {}
    for name, tag in (("bold", "b"), ("italic", "i"), ("size", "sz")):
        element = properties.find(qn("w:" + tag))
        raw = None if element is None else element.get(qn("w:val"))
        values[name] = float(raw) / 2 if name == "size" else element is not None and raw not in {"0", "false", "off"}
    result = {}
    for name in ("Ancestor", "Base", "Derived", "P", "C"):
        if name not in document.styles:
            continue
        style, chain = document.styles[name], []
        while style is not None:
            assert style.style_id not in {item.style_id for item in chain}, "cyclic owned style"
            chain.append(style)
            style = style.base_style
        font = values.copy()
        for style in reversed(chain):
            for attribute in font:
                value = getattr(style.font, attribute)
                if value is not None:
                    font[attribute] = value.pt if attribute == "size" else value
        result[name] = font
    return result


def verify_style_save_state(root, name="style-save-state-26.9.json"):
    counts = {"style-save-state-26.9.json": (247, 132), "character-style-save-state-26.9.json": (108,),
              "style-normalization-contexts-26.9.json": (36,)}
    report = json.loads((root / name).read_text())
    raw = (root / report["outputs"]["archive"]).read_bytes()
    assert digest(raw) == report["outputs"]["sha256"]
    changes, members = [], set()
    with ZipFile(BytesIO(raw)) as outputs:
        for expected, observed in zip(counts[name], report["reports"], strict=True):
            assert observed["version"] == "26.9.0" and observed["licensed"] is False
            corpus = (root / observed["corpus"]).read_bytes()
            assert digest(corpus) == observed["corpus_sha256"]
            rows = {row["case"]: row for row in observed["records"]}
            assert len(rows) == len(observed["records"]) == expected
            with ZipFile(BytesIO(corpus)) as inputs:
                assert set(inputs.namelist()) == {key + "/" + phase + ".docx" for key in rows for phase in ("source", "destination")}
                for key, row in rows.items():
                    for phase in ("source", "destination"):
                        assert digest(inputs.read(key + "/" + phase + ".docx")) == row["inputs"][phase]
                    assert row["before_save"] == row["after_save_live"], "live save observation changed"
                    data = outputs.read(row["output"])
                    assert digest(data) == row["output_sha256"]
                    assert saved_story_formats(data) == row["after_reopen"]["paragraphs"], "reopened paragraph observation"
                    assert saved_style_fonts(data) == row["after_reopen"]["styles"], "reopened style observation"
                    assert row["before_save"]["styles"] == row["after_reopen"]["styles"]
                    if row["before_save"]["paragraphs"] != row["after_reopen"]["paragraphs"]:
                        changes.append(str(expected) + "/" + key)
                    assert row["output"] not in members
                    members.add(row["output"])
        assert members == set(outputs.namelist())
    assert sorted(changes) == report["reopen_paragraph_changes"]
    assert report["live_changes"] == report["reopen_style_changes"] == []
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is False
    if "candidate" in report:
        native = {row["case"]: row for row in report["reports"][0]["records"]}
        candidate = {row["case"]: row for row in report["candidate"]["records"]}
        reread = report["candidate_official_rereads"]
        assert reread["version"] == "26.9.0" and reread["licensed"] is False
        actual = {row["case"]: row for row in reread["records"]}
        assert len(actual) == len(reread["records"]) == len(candidate) == len(report["candidate"]["records"]) == 108
        assert set(native) == set(candidate) == set(actual)
        assert report["candidate"]["corpus_sha256"] == report["reports"][0]["corpus_sha256"]
        for key, row in candidate.items():
            assert row["inputs"] == native[key]["inputs"]
            assert row["output_sha256"] == actual[key]["output_sha256"]
            if row["outcome"] == "returned":
                assert actual[key]["formats"] == native[key]["after_reopen"]["paragraphs"]
            else:
                assert row["outcome"] == "raised" and row["destination_unchanged"] is True
    return len(members)


def saved_style_alignments(data):
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(BytesIO(data))
    default = document.styles.element
    for tag in ("docDefaults", "pPrDefault", "pPr", "jc"):
        default = default.find(qn("w:" + tag)) if default is not None else None
    value = "left" if default is None else default.get(qn("w:val"))
    result = {}
    for name in ("Base", "Derived"):
        style, chain = document.styles[name], []
        while style is not None:
            assert style.style_id not in {item.style_id for item in chain}
            chain.append(style)
            style = style.base_style
        alignment = {"left": "LEFT", "center": "CENTER", "right": "RIGHT", "both": "JUSTIFY"}[value]
        for style in reversed(chain):
            if style.paragraph_format.alignment is not None:
                alignment = style.paragraph_format.alignment.name
        result[name] = alignment
    return result


def verify_style_paragraph_formats(root):
    report = json.loads((root / "style-paragraph-format-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is False
    source = (root / report["corpus"]).read_bytes()
    raw = (root / report["outputs"]).read_bytes()
    assert digest(source) == report["corpus_sha256"] and digest(raw) == report["outputs_sha256"]
    generated = dict(runpy.run_path(str(Path(__file__).parents[1] / "docs/probes/style_paragraph_formats.py"))["inputs"]())
    rows = report["records"]
    assert len(rows) == 108 and len(generated) == 27
    assert {(row["input"], row["alignment"]) for row in rows} == {
        (name, alignment) for name in generated for alignment in ("LEFT", "CENTER", "RIGHT", "JUSTIFY")}
    with ZipFile(BytesIO(source)) as inputs, ZipFile(BytesIO(raw)) as outputs:
        assert set(inputs.namelist()) == set(generated)
        assert set(outputs.namelist()) == {row["output"] for row in rows}
        for row in rows:
            data = inputs.read(row["input"])
            assert digest(data) == row["input_sha256"]
            with ZipFile(BytesIO(data)) as actual, ZipFile(BytesIO(generated[row["input"]])) as expected:
                assert {key: actual.read(key) for key in actual.namelist()} == {key: expected.read(key) for key in expected.namelist()}
            assert saved_style_alignments(data) == row["before_edit"]["style_alignment"]
            assert row["after_edit"]["style_alignment"]["Derived"] == row["alignment"]
            assert row["after_edit"] == row["after_save_live"] == row["after_reopen"]
            saved = outputs.read(row["output"])
            assert digest(saved) == row["output_sha256"]
            assert saved_style_alignments(saved) == row["after_reopen"]["style_alignment"]
            assert saved_style_fonts(saved) == row["after_reopen"]["styles"]
            assert saved_story_formats(saved) == row["after_reopen"]["paragraphs"]
    return len(rows)


def verify_style_font_edits(root):
    report = json.loads((root / "style-font-edits-26.9.json").read_text())
    assert report["version"] == "26.9.0" and report["licensed"] is False
    assert report["full_format_acceptance"] is report["rendering_acceptance"] is False
    source = (root / "corpus" / report["corpus"]).read_bytes()
    raw = (root / "corpus/style-font-edits-26.9.zip").read_bytes()
    assert digest(source) == report["corpus_sha256"]
    assert digest(raw) == report["outputs_sha256"]
    rows = report["records"]
    assert len(rows) == 1482
    with ZipFile(BytesIO(source)) as inputs, ZipFile(BytesIO(raw)) as outputs:
        expected = {(name, prop, value) for name in inputs.namelist() if name.endswith("/source.docx")
                    for prop, values in (("bold", (False, True)), ("italic", (False, True)), ("size", (10, 17.5)))
                    for value in values}
        assert {(row["input"], row["property"], row["value"]) for row in rows} == expected
        assert set(outputs.namelist()) == {row["output"] for row in rows}
        for row in rows:
            assert row["style"] == "Derived"
            assert digest(inputs.read(row["input"])) == row["input_sha256"]
            assert row["before_edit"]["styles"] == saved_style_fonts(inputs.read(row["input"]))
            assert row["after_edit"]["styles"]["Derived"][row["property"]] == row["value"]
            assert row["after_edit"] == row["after_save_live"]
            data = outputs.read(row["output"])
            assert digest(data) == row["output_sha256"]
            assert saved_story_formats(data) == row["after_reopen"]["paragraphs"]
            assert saved_style_fonts(data) == row["after_reopen"]["styles"]
            assert row["after_edit"]["styles"] == row["after_reopen"]["styles"]
    assert sum(row["after_edit"] != row["after_reopen"] for row in rows) == 36
    return len(rows)


def verify_style_projections(root):
    report = json.loads((root / "style-import-projections.json").read_text())
    assert report["baseline_version"] == "26.9.0"
    raw = (root / report["outputs"]["archive"]).read_bytes()
    assert digest(raw) == report["outputs"]["sha256"]
    inspect = runpy.run_path(str(Path(__file__).parents[1] / "docs" / "probes" / "inspect_style_imports.py"))["inspect_document"]
    members = set()
    with ZipFile(BytesIO(raw)) as outputs:
        for count, matrix in zip((247, 132, 108, 36), report["reports"], strict=True):
            assert matrix["count"] == count
            baseline = (root / matrix["native"]["file"]).read_bytes()
            assert digest(baseline.replace(b"\r\n", b"\n")) == matrix["native"]["sha256"]
            native = {row["case"]: row for row in json.loads(baseline)["reports"][matrix["native"]["report_index"]]["records"]}
            corpus = (root / matrix["corpus"]["archive"]).read_bytes()
            assert digest(corpus) == matrix["corpus"]["sha256"] == matrix["candidate"]["corpus_sha256"]
            rows = {row["case"]: row for row in matrix["candidate"]["records"]}
            reread = matrix["official_reread"]
            assert reread["version"] == "26.9.0" and reread["licensed"] is False
            observed = {row["case"]: row for row in reread["records"]}
            checks = {row["case"]: row for row in matrix["independent_checks"]}
            assert len(rows) == len(matrix["candidate"]["records"]) == len(native) == len(observed) == len(reread["records"]) == len(checks) == len(matrix["independent_checks"]) == count
            assert set(rows) == set(native) == set(observed) == set(checks)
            with ZipFile(BytesIO(corpus)) as inputs:
                assert set(inputs.namelist()) == {key + "/" + phase + ".docx" for key in rows for phase in ("source", "destination")}
                for key, row in rows.items():
                    assert row["outcome"] == "returned"
                    assert all(digest(inputs.read(key + "/" + phase + ".docx")) == row["inputs"][phase] == native[key]["inputs"][phase]
                               for phase in ("source", "destination"))
                    member = str(count) + "/" + row["output"]
                    data = outputs.read(member)
                    assert digest(data) == row["output_sha256"] == observed[key]["output_sha256"]
                    assert saved_story_formats(data) == observed[key]["formats"] == native[key]["after_reopen"]["paragraphs"]
                    assert saved_style_fonts(data) == native[key]["after_reopen"]["styles"]
                    warm = row["imported_format"]
                    warm = {**warm, "alignment": {"both": "JUSTIFY"}.get(warm["alignment"], (warm["alignment"] or "left").upper())}
                    assert warm == native[key]["before_save"]["paragraphs"]["IMPORT"]
                    assert {"case": key, **inspect(BytesIO(data))} == checks[key]
                    assert member not in members
                    members.add(member)
        assert set(outputs.namelist()) == members and len(members) == report["outputs"]["files"] == 523
    assert report["validation"]["full_import_acceptance"] is report["validation"]["full_format_acceptance"] is report["validation"]["rendering_acceptance"] is False
    return len(members)


def verify(root):
    def read(name):
        return json.loads((root / name).read_text())

    declarations = read("commercial-26.9-api.json")
    ledger = read("commercial-26.9-capabilities.json")
    enums = read("commercial-26.9-all-enums.json")
    baseline = read("commercial-26.9-baseline.json")
    registry = baseline["documentation_registry"]
    raw_registry = (root / registry["path"]).read_bytes()
    assert digest(raw_registry.replace(b"\r\n", b"\n")) == registry["sha256"]
    assert [(row["url"], row["sha256"]) for row in json.loads(raw_registry)["documents"]] == [
        (row["url"], row["sha256"]) for row in baseline["documentation"]]
    symbols = [item["id"] for module in declarations["modules"] for item in module["symbols"]]
    assert declarations["version"] == ledger["baseline_version"] == enums["version"] == "26.9.0"
    assert declarations["module_count"] == len(declarations["modules"])
    assert declarations["symbol_count"] == len(symbols) == len(set(symbols))
    assert ledger["capability_count"] == len(ledger["records"]) == len({row["id"] for row in ledger["records"]})
    assert enums["enum_count"] == len(enums["enums"])
    assert enums["value_count"] == sum(len(item["values"]) for item in enums["enums"])
    assert enums["alias_count"] == sum(v["name"] != v["canonical_name"] for item in enums["enums"] for v in item["values"])
    formats = read("commercial-26.9-format-behavior.json")
    assert digest((root / formats["source"]).read_bytes()) == formats["source_sha256"]
    checked = verify_format_outputs(formats, root / formats["outputs"]["archive"])
    assert len(ledger["format_capabilities"]) == sum(len(formats["reports"][key]["records"])
                                                   for key in ("save_commercial", "load_commercial"))
    literal = read("commercial-26.9-literal-text.json")
    raw = (root / literal["fixtures"]["archive"]).read_bytes()
    assert digest(raw) == literal["fixtures"]["sha256"]
    outputs = (root / literal["raw_outputs"]["archive"]).read_bytes()
    assert digest(outputs) == literal["raw_outputs"]["sha256"]
    with ZipFile(BytesIO(raw)) as archive, ZipFile(BytesIO(outputs)) as generated:
        for backend in ("commercial", "docweave"):
            for row in literal[backend]:
                assert digest(archive.read(row["source"])) == row["source_sha256"]
                data = generated.read(row["output_member"])
                assert digest(data) == row["output_sha256"]
                assert data.decode().replace("\r\n", "\n") == row["markdown"]
    imports = verify_import_outputs(root)
    toggles = verify_style_toggles(root)
    style_imports = verify_style_imports(root)
    style_imports += verify_style_imports(root, "style-import-translated.json")
    style_imports += verify_style_imports(root, "style-import-default-on.json")
    style_imports += verify_style_imports(root, "style-import-character-default-on.json")
    style_imports += verify_style_imports(root, "paragraph-style-defaults-26.9.json")
    roundtrips = verify_style_roundtrips(root)
    save_states = verify_style_save_state(root)
    save_states += verify_style_save_state(root, "character-style-save-state-26.9.json")
    save_states += verify_style_save_state(root, "style-normalization-contexts-26.9.json")
    projections = verify_style_projections(root)
    edits = verify_style_font_edits(root)
    paragraph_edits = verify_style_paragraph_formats(root)
    pagination = verify_paragraph_pagination(root)
    first_paragraph = verify_first_paragraph_page_break(root)
    first_trial = verify_first_paragraph_trial(root)
    pagination_rendering = verify_pagination_rendering(root)
    dimensions = verify_paragraph_dimensions(root)
    spacing_limits = verify_paragraph_dimensions(root, "paragraph-spacing-limits-26.9.json")
    indent_limits = verify_paragraph_dimensions(root, "paragraph-indent-limits-26.9.json")
    logical_indents = verify_paragraph_dimensions(root, "paragraph-logical-indents-26.9.json")
    character_indents = verify_paragraph_character_indents(root)
    character_setters = verify_paragraph_character_setters(root)
    character_reads = verify_paragraph_character_reads(root)
    character_inheritance_edits = verify_paragraph_character_inheritance_edits(root)
    character_roundtrips = verify_character_indent_roundtrips(root)
    dimension_rendering = verify_pagination_rendering(root, "paragraph-dimensions-rendering-26.9.json")
    font_sizes = verify_font_size_loading(root)
    default_presence = verify_font_default_presence(root)
    font_origins = verify_font_json_origin(root)
    font_booleans = verify_font_boolean_contexts(root)
    hidden_style_contexts = verify_hidden_style_contexts(root)
    hidden_roundtrips = verify_hidden_font_roundtrip(root)
    font_boolean_roundtrips = verify_font_boolean_roundtrip(root)
    dom_font_booleans = verify_dom_font_booleans(root)
    dom_font_boolean_errors = verify_dom_font_boolean_errors(root)
    defaults = verify_font_defaults(root)
    default_matrix = verify_font_default_matrix(root)
    return {"declared_symbols": len(symbols), "capability_rows": ledger["capability_count"],
            "checked_import_outputs": imports,
            "checked_style_inputs": toggles,
            "checked_style_import_outputs": style_imports,
            "checked_font_size_loading_inputs": font_sizes,
            "checked_font_default_presence_outputs": default_presence,
            "checked_font_json_origin_outputs": font_origins,
            "checked_font_boolean_inputs": font_booleans,
            "checked_hidden_style_contexts": hidden_style_contexts,
            "checked_hidden_font_roundtrips": hidden_roundtrips,
            "checked_font_boolean_roundtrips": font_boolean_roundtrips,
            "checked_dom_font_boolean_outputs": dom_font_booleans,
            "checked_dom_font_boolean_errors": dom_font_boolean_errors,
            "checked_font_default_inputs": defaults,
            "checked_font_default_matrix_inputs": default_matrix,
            "checked_style_roundtrip_outputs": roundtrips,
            "checked_style_save_states": save_states,
            "checked_style_projection_outputs": projections,
            "checked_style_font_edit_outputs": edits,
            "checked_style_paragraph_edit_outputs": paragraph_edits,
            "checked_paragraph_pagination_outputs": pagination,
            "checked_unresolved_first_paragraph_outputs": first_paragraph,
            "checked_trial_first_paragraph_outputs": first_trial,
            "checked_pagination_rendering_pairs": pagination_rendering,
            "checked_paragraph_dimension_edits": dimensions,
            "checked_paragraph_spacing_limits": spacing_limits,
            "checked_paragraph_indent_limits": indent_limits,
            "checked_paragraph_logical_indents": logical_indents,
            "checked_paragraph_character_indents": character_indents,
            "checked_paragraph_character_setters": character_setters,
            "checked_paragraph_character_reads": character_reads,
            "checked_paragraph_character_inheritance_edits": character_inheritance_edits,
            "checked_character_indent_roundtrips": character_roundtrips,
            "checked_paragraph_dimension_rendering_pairs": dimension_rendering,
            "checked_format_outputs": checked, "behavioral_acceptance": False}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("root", nargs="?", type=Path,
                        default=Path(__file__).parents[1] / "docs" / "benchmarks")
    args = parser.parse_args()
    print(json.dumps(verify(args.root)))
