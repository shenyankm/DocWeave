"""Compose docxtpl with this converter; install docxtpl separately.

Usage: python ApiExamples/template_report.py template.docx context.json report.docx report.pdf
Templates must be trusted: Jinja templates are code, not a secure upload format.
"""

import argparse
import json
from pathlib import Path

import aspose.words_foss as aw
from aspose.words_foss._io import atomic_output


def render_report(template_path, context, docx_path, pdf_path):
    """Render with a dict or a context factory receiving the actual DocxTemplate."""
    from docxtpl import DocxTemplate
    from jinja2 import Environment, StrictUndefined

    if len({Path(path).resolve() for path in (template_path, docx_path, pdf_path)}) != 3:
        raise ValueError("Template, DOCX output and PDF output paths must be distinct")
    template = DocxTemplate(template_path)
    if callable(context):
        context = context(template)
    if not isinstance(context, dict):
        raise ValueError("Template context must be a dict")
    template.render(context, jinja_env=Environment(undefined=StrictUndefined), autoescape=True)
    with atomic_output(docx_path) as temporary:
        template.save(temporary)
    document = aw.Document(docx_path)
    document.save(pdf_path, aw.SaveFormat.PDF)
    return document.diagnostics


def test_template_report(tmp_path):
    import pytest
    pytest.importorskip("docxtpl")
    from docx import Document
    from pypdf import PdfReader

    template = tmp_path / "template.docx"
    source = Document()
    source.add_paragraph("客户：{{ company }}")
    source.save(template)
    docx, pdf = tmp_path / "report.docx", tmp_path / "report.pdf"
    render_report(template, {"company": "中文 <公司> & 合作方"}, docx, pdf)
    assert "中文 <公司> & 合作方" in aw.Document(docx).get_text()
    assert "中文" in "".join(page.extract_text() for page in PdfReader(pdf).pages)
    before = docx.read_bytes()
    for docx_output, pdf_output in [(template, pdf), (docx, docx)]:
        with pytest.raises(ValueError, match="distinct"):
            render_report(template, {"company": "test"}, docx_output, pdf_output)
    assert docx.read_bytes() == before
    raw_pdf = tmp_path / "report.bin"
    render_report(template, {"company": "test"}, docx, raw_pdf)
    assert raw_pdf.read_bytes().startswith(b"%PDF-")
    before = docx.read_bytes()
    from jinja2 import UndefinedError
    with pytest.raises(UndefinedError):
        render_report(template, {}, docx, pdf)
    assert docx.read_bytes() == before


def test_template_report_loop_and_embedded_picture(tmp_path):
    import pytest
    pytest.importorskip("docxtpl")
    from docx import Document
    from docx.shared import Inches
    from PIL import Image
    from pypdf import PdfReader

    logo = tmp_path / "logo.png"
    Image.new("RGB", (40, 20), "navy").save(logo)
    source = Document()
    source.add_picture(str(logo), width=Inches(1))
    source.add_paragraph("{%p for item in items %}")
    source.add_paragraph("{{ item.name }}: {{ item.value }}")
    source.add_paragraph("{%p endfor %}")
    template = tmp_path / "loop.docx"
    source.save(template)
    docx, pdf = tmp_path / "report.docx", tmp_path / "report.pdf"
    render_report(template, {"items": [{"name": "Alpha <&>", "value": 3},
                                     {"name": "Beta", "value": 7}]}, docx, pdf)
    independent = Document(docx)
    assert [p.text for p in independent.paragraphs if p.text] == ["Alpha <&>: 3", "Beta: 7"]
    assert len(independent.inline_shapes) == 1
    rendered = PdfReader(pdf)
    text = "".join(page.extract_text() for page in rendered.pages)
    assert "Alpha <&>: 3" in text and "Beta: 7" in text
    assert sum(len(page.images) for page in rendered.pages) == 1


def test_template_report_dynamic_picture_context(tmp_path):
    import pytest
    pytest.importorskip('docxtpl')
    from docxtpl import InlineImage
    from docx import Document
    from docx.shared import Mm
    from PIL import Image
    import fitz
    from io import BytesIO
    from zipfile import ZipFile

    source = Document()
    source.add_paragraph('客户：{{ company }}')
    source.add_paragraph('{{ logo }}')
    template = tmp_path / 'dynamic.docx'
    source.save(template)
    original = template.read_bytes()
    logo = tmp_path / 'logo.png'
    docx, pdf = tmp_path / 'report.docx', tmp_path / 'report.pdf'
    calls = []

    def context(tpl):
        calls.append(tpl)
        return {'company': '图片 <&>', 'logo': InlineImage(tpl, str(logo), width=Mm(25))}

    for index, color in enumerate(['navy', 'red'], 1):
        Image.new('RGB', (40, 20), color).save(logo)
        render_report(template, context, docx, pdf)
        assert len(calls) == index
        independent = Document(docx)
        assert independent.paragraphs[0].text == '客户：图片 <&>'
        assert len(independent.inline_shapes) == 1
        assert independent.inline_shapes[0].width.mm == pytest.approx(25)
        assert independent.inline_shapes[0].height.mm == pytest.approx(12.5)
        with ZipFile(docx) as package:
            media = [name for name in package.namelist() if name.startswith('word/media/')]
            assert len(media) == 1 and package.read(media[0]) == logo.read_bytes()
        with fitz.open(pdf) as rendered:
            assert '图片 <&>' in rendered[0].get_text()
            images = rendered[0].get_images()
            assert len(images) == 1
            rect = rendered[0].get_image_rects(images[0][0])[0]
            assert rect.width == pytest.approx(25 * 72 / 25.4, abs=0.01)
            assert rect.height == pytest.approx(12.5 * 72 / 25.4, abs=0.01)
            with Image.open(BytesIO(rendered.extract_image(images[0][0])['image'])) as pixels:
                assert pixels.getpixel((0, 0)) == Image.new('RGB', (1, 1), color).getpixel((0, 0))
    assert template.read_bytes() == original


def test_dynamic_context_failure_preserves_both_outputs(tmp_path):
    import pytest
    pytest.importorskip('docxtpl')
    from docxtpl import InlineImage
    from docx import Document

    source = Document()
    source.add_paragraph('{{ logo }}')
    template = tmp_path / 'dynamic.docx'
    source.save(template)
    docx, pdf = tmp_path / 'report.docx', tmp_path / 'report.pdf'
    docx.write_bytes(b'KEEP DOCX')
    pdf.write_bytes(b'KEEP PDF')
    with pytest.raises(FileNotFoundError):
        render_report(template, lambda tpl: {'logo': InlineImage(tpl, str(tmp_path / 'missing.png'))}, docx, pdf)
    assert docx.read_bytes() == b'KEEP DOCX' and pdf.read_bytes() == b'KEEP PDF'
    with pytest.raises(ValueError, match='context'):
        render_report(template, lambda tpl: [], docx, pdf)
    assert docx.read_bytes() == b'KEEP DOCX' and pdf.read_bytes() == b'KEEP PDF'


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", type=Path)
    parser.add_argument("context", type=Path)
    parser.add_argument("docx", type=Path)
    parser.add_argument("pdf", type=Path)
    args = parser.parse_args()
    diagnostics = render_report(args.template, json.loads(args.context.read_text(encoding="utf-8")), args.docx, args.pdf)
    for diagnostic in diagnostics:
        print(f"{diagnostic.code}: {diagnostic.message}")


if __name__ == "__main__":
    main()
